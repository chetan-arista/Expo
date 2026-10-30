import argparse
import getpass
import sys
import os
import time

from crawler.config import AuthConfig, CrawlConfig, USER_AGENTS, MAX_CONCURRENCY
from crawler.engine import WebsiteCrawler, ExternalLink
from crawler.exporter import ExcelExporter, EXPORT_CATEGORIES
from crawler.issues import CATEGORY_LABELS, ISSUE_CATEGORIES
from crawler.url_utils import get_domain

MAX_SCREENSHOT_ISSUES = 25
PASSWORD_ENV = "CRAWLER_PASSWORD"


def _seo_index(crawler: WebsiteCrawler):
    """(Site, TabIndex) of the finished crawl, or (None, None) when the SEO analysis fails."""
    try:
        return crawler.site_index(force=True)
    except Exception as e:
        print(f"SEO analysis failed: {e}")
        return None, None


def _mask(text: str, secrets) -> str:
    """`text` with every secret replaced by ***."""
    text = str(text)
    for s in secrets:
        if s and s in text:
            text = text.replace(s, "***")
    return text


def _password(args) -> str:
    """--password, or the CRAWLER_PASSWORD environment variable when the value is "-" (or left out with
    --username); asks on the terminal as a last resort. Never printed."""
    pw = args.password
    if pw == "-" or (pw is None and args.username):
        pw = os.environ.get(PASSWORD_ENV, "")
        if not pw and args.password is None and sys.stdin.isatty():
            pw = getpass.getpass("Password: ")
    return pw or ""


def write_pdf_report(crawler: WebsiteCrawler, excel_path: str, screenshots: bool = True,
                     validation=None, intel=None, content=None, issues=None, checks=None) -> str:
    """Builds <excel name>_issues.pdf: one picture pass (the broken links' parent pages with the link boxed and
    their error pages, plus a pinpointed picture per SEO issue) when Chrome is available, plus the Validation,
    Fix plan, Content and Checks sections when given."""
    from crawler.pdf_report import build_pdf_report

    base = os.path.splitext(excel_path)[0]
    shots_dir = base + "_evidence"
    evidence, seo_evidence, note = {}, {}, ""
    site, index = _seo_index(crawler)
    if not screenshots:
        note = "Screenshots were disabled for this crawl."
    elif index is not None and index.issues:
        try:
            from crawler.browser import chrome_status
            available, why = chrome_status()
        except Exception as e:
            available, why = False, str(e)
        if not available:
            note = f"Screenshots unavailable: {why or 'Chrome is not available.'} Link positions are shown as page sketches."
            print(note)
        from crawler.seo_evidence import capture_issue_evidence, link_evidence
        print("Capturing pictures of the issues (broken links included)...")
        seo_evidence = capture_issue_evidence(crawler, (site, index), shots_dir, logger=print)
        evidence = link_evidence(seo_evidence) or {}
        broken = sum(1 for r in crawler.results_list if r.category in ISSUE_CATEGORIES)
        if available and broken > MAX_SCREENSHOT_ISSUES:
            note = f"Screenshots captured for the {MAX_SCREENSHOT_ISSUES} most severe of {broken} broken URLs."

    pdf_path = base + "_issues.pdf"
    seo_issues = issues if issues is not None else (list(index.issues) if index is not None else None)
    integrations = None
    if site is not None and getattr(crawler, "integrations", None):
        try:
            from crawler.integrations import report_summary
            integrations = report_summary(site)
        except Exception:
            integrations = None
    build_pdf_report(crawler.results_list, pdf_path, crawler.get_stats(), evidence, shots_dir, note,
                     seo_issues=seo_issues, site=site, external_links=list(crawler.external_links_list),
                     seo_evidence=seo_evidence or None, validation=validation, intel=intel, content=content,
                     seo_evidence_dir=shots_dir, checks=checks, integrations=integrations, seo_index=index)
    return pdf_path


def _read_list(path: str):
    """List mode: the URL lines of a .txt / .csv (first column, or a URL / Address column) / .xlsx file, or of
    stdin when path is "-"."""
    if path == "-":
        return [x.strip() for x in sys.stdin.read().splitlines() if x.strip()]
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
    else:
        import csv
        with open(path, "r", encoding="utf-8-sig", errors="replace") as fh:
            rows = list(csv.reader(fh)) if path.lower().endswith(".csv") else [[x] for x in fh.read().splitlines()]
    if not rows:
        return []
    head = [str(c or "").strip().lower() for c in rows[0]]
    col = next((i for i, h in enumerate(head) if h in ("url", "urls", "address", "loc")), None)
    body = rows[1:] if col is not None else rows
    col = col or 0
    return [str(r[col]).strip() for r in body if len(r) > col and r[col] is not None and str(r[col]).strip()]


def _headers(values):
    """--header "Name: value" pairs as [{"name", "value"}] (values are secrets: never printed)."""
    out = []
    for raw in values or []:
        name, sep, value = str(raw).partition(":")
        if not sep or not name.strip():
            print('Error: --header must look like "Name: value" (got a value without a colon).')
            sys.exit(1)
        out.append({"name": name.strip(), "value": value.strip()})
    return out


def run_validation(crawler: WebsiteCrawler) -> dict:
    """Re-checks every reported error live and prints the verdicts."""
    from crawler.validate import Validator

    print("\nRe-checking the reported errors live (3 attempts each)...")
    last = [0.0]

    def progress(done: int, total: int):
        if time.time() - last[0] > 2 or done == total:
            last[0] = time.time()
            print(f"  re-checked {done} of {total}", flush=True)

    out = Validator(crawler, progress=progress).run(scope="errors")
    s = out.get("summary") or {}
    print(f" Validation: {s.get('confirmed', 0)} confirmed, {s.get('not_reproduced', 0)} not reproduced, "
          f"{s.get('intermittent', 0)} intermittent"
          + (f", {s.get('not_checked')} not checked" if s.get("not_checked") else ""))
    exposed = [g for g in out.get("login_guard") or [] if g.get("exposed")]
    if out.get("login_guard"):
        print(f" Login check: {len(out['login_guard'])} member page(s) requested without the login, "
              f"{len(exposed)} exposed")
        for g in exposed[:10]:
            print(f"   EXPOSED: {g.get('url')}")
    return out


def run_intel(crawler: WebsiteCrawler, site, index, validation=None, use_ai: bool = True) -> dict:
    """Root-cause groups and content review (AI-enriched when a free AI key is configured)."""
    from crawler.intel import run_intel as intel_run

    print("\nAnalysing root causes and page content...")
    out = intel_run(site, index, crawler, progress=lambda pct, msg: print(f"  [{pct:>3}%] {msg}", flush=True),
                    use_ai=use_ai, fetch=True, validation=validation)
    errors = out.get("errors") or {}
    groups = sorted(errors.get("groups") or [], key=lambda g: int(g.get("rank") or 10 ** 6))
    print(f" Fix plan: {len(groups)} group(s)" + (f" ({errors.get('ai_note')})" if errors.get("ai_note") else ""))
    for g in groups[:5]:
        print(f"   {g.get('rank')}. {g.get('title')} - {g.get('urls_count')} URL(s), impact {g.get('impact')}, "
              f"effort {g.get('effort')}")
    content = out.get("content") or {}
    counts = content.get("counts") or {}
    print(f" Content: {counts.get('mismatches', 0)} wrong page type, {counts.get('thin', 0)} thin, "
          f"{counts.get('off_topic', 0)} off topic")
    return out


def run_business(url: str, crawler: WebsiteCrawler, base_path: str) -> str:
    """Business insights report as <base>_business.pdf."""
    from crawler.ai.business import build_business_report
    from crawler.business_pdf import build_business_pdf

    print("\nResearching the business behind the site...")
    report = build_business_report(url, crawler=crawler,
                                   progress=lambda pct, msg: print(f"  [{pct:>3}%] {msg}", flush=True))
    company = report.get("company") or {}
    print(f" Company: {company.get('name') or 'not found'}"
          + (f" - {company.get('what_they_do')}" if company.get("what_they_do") else ""))
    if report.get("note"):
        print(f" Note: {report['note']}")
    path = base_path + "_business.pdf"
    build_business_pdf(report, path)
    return path


MAX_PASTE_CHARS = 2_000_000


def _read_paste(path: str) -> str:
    """The cookie paste in `path` ("-" = stdin), at most MAX_PASTE_CHARS characters."""
    if path == "-":
        return sys.stdin.read(MAX_PASTE_CHARS)
    with open(path, encoding="utf-8-sig", errors="replace") as fh:
        return fh.read(MAX_PASTE_CHARS)


def _test_cookies(raw: str, url: str, page: str, ua_preset: str, custom_ua: str, verify_ssl: bool,
                  secrets) -> int:
    """--test-cookies: one page that needs sign-in, opened with the cookies (and once without). Exit code 0 when
    signed in, 1 otherwise. Cookie values are never printed."""
    from crawler.auth import test_cookies
    ua = custom_ua if ua_preset == "custom" else USER_AGENTS.get(ua_preset, "")
    print("Testing the pasted cookies (two or three requests; a sign-out address is never opened)...")
    res = test_cookies(raw, url, test_url=page or None, user_agent=ua or None, verify_ssl=verify_ssl)
    verdict = {True: "Signed in", False: "Not signed in"}.get(res.get("signed_in"), "Couldn't tell")
    print(f"\n Cookie test : {verdict}")
    if res.get("test_url"):
        print(_mask(f" Page tested : {res['test_url']}", secrets))
    if res.get("status") is not None:
        print(_mask(f" Answer      : HTTP {res['status']} at {res.get('final_url') or res.get('test_url')}", secrets))
    print(_mask(f" {res.get('reason') or ''}", secrets))
    for step in (res.get("challenge") or {}).get("owner_steps") or []:
        print(f"   - {step}")
    return 0 if res.get("signed_in") is True else 1


def find_subdomains(config: CrawlConfig, typed) -> int:
    """--find-subdomains: the hosts of the site's domain found by every source, each checked with one request.
    Hosts of other domains are never probed. Exit code 0 when the search ran."""
    from crawler import subdomains
    from crawler.engine import proxy_url
    from crawler.url_utils import registrable_domain
    p = proxy_url(getattr(config, "proxy", None))
    seen = set()

    def progress(snap):
        msg = str(snap.get("message") or "")
        if msg and msg not in seen and not snap.get("total"):
            seen.add(msg)
            print(f"  {msg}", flush=True)

    print(f"\nFinding the subdomains of {registrable_domain(config.start_url)} ...")
    res = subdomains.discover(config.start_url, extra_hosts=typed, user_agent=config.user_agent,
                              proxies={"http": p, "https": p} if p else None, verify_ssl=config.verify_ssl,
                              on_progress=progress)
    labels = res.get("source_labels") or {}
    hosts = res.get("hosts") or []
    print(f"\n {len(hosts)} subdomain(s) of {res.get('domain') or '?'} found"
          + (f" (first {len(hosts)} of {res.get('found')})" if res.get("capped") else "") + ":")
    if hosts:
        print(f"  {'Host':<40} {'Found in':<28} {'DNS':<4} {'Status':<7} Title / note")
    for h in hosts:
        found_in = ", ".join(labels.get(s, s) for s in h.get("sources") or [])
        dns = {True: "yes", False: "no"}.get(h.get("resolves"), "-")
        status = str(h.get("status")) if h.get("status") is not None else "-"
        text = " - ".join(x for x in (h.get("title") or "", h.get("note") or "") if x)
        print(f"  {h.get('host', '')[:40]:<40} {found_in[:28]:<28} {dns:<4} {status:<7} {text[:80]}")
    for source, message in (res.get("errors") or {}).items():
        print(f" {labels.get(source, source)}: {message}")
    if hosts:
        print("\n To crawl some of them too, add --subdomain HOST (repeatable) to the crawl; for a host that answers "
              "on http only, give its home URL (e.g. --subdomain http://HOST/).")
    return 0 if res.get("state") == "done" else 1


def print_bot_checks_and_subdomains(crawler: WebsiteCrawler, secrets) -> None:
    """After the crawl: pages a bot check answered instead (their status codes stay as they are) with the routes
    that work, and every host of the site's domain seen, crawled or not, with the reason."""
    try:
        from crawler.auth import crawl_challenge_status
        ch = crawl_challenge_status(crawler)
    except Exception:
        ch = {}
    if ch.get("seen"):
        print(_mask(f"\n BOT CHECK: {ch['count']} page(s) were answered by a bot check ({ch['provider']}) instead of "
                    f"the page, the first at {ch['first_url']}. Their status codes are kept as they are.", secrets))
        print(_mask(f" {ch.get('advice') or ''}", secrets))
        for step in ch.get("owner_steps") or []:
            print(f"   - {step}")
    try:
        from crawler.subdomains import subdomain_report
        rep = subdomain_report(crawler)
    except Exception:
        return
    hosts = rep.get("hosts") or []
    if len(hosts) <= 1:
        return
    print(f"\n Hosts of {rep.get('domain')} seen: {len(hosts)}")
    for h in hosts[:40]:
        print(f"  {str(h.get('host'))[:40]:<40} {int(h.get('crawled') or 0):>7,} URL(s)  {h.get('reason') or ''}")
    if len(hosts) > 40:
        print(f"  ... and {len(hosts) - 40} more")


def require_all_checks(where: str) -> None:
    """Stops with exit code 2 when an SEO check module failed to import (names the module and the error);
    ARACHNIX_ALLOW_BROKEN_TABS=1 turns that into a loud warning (reports are then flagged "checks incomplete")."""
    from crawler.seo import registry
    try:
        registry.require_all_tabs(where)
    except registry.TabLoadError as e:
        print(f"\nERROR: {e}\n", file=sys.stderr, flush=True)
        sys.exit(2)


def _build4_options(args) -> dict:
    """W9: the Build 4 crawl options of --profile, --config and the single flags, checked with the crawl's own rules
    (crawler.config.validate_options): {option: value}. Exits with a plain message on a bad value."""
    import json
    opts = {}
    if args.profile:
        from crawler.settings_store import load_profile
        prof = load_profile(args.profile)
        if prof is None:
            print(f"Error: no saved profile called {args.profile!r} (Configuration > Profiles in the dashboard).")
            sys.exit(1)
        opts.update(prof.get("options") or {})
    if args.config:
        try:
            with open(args.config, encoding="utf-8") as f:
                data = json.load(f)
        except (OSError, ValueError) as e:
            print(f"Error: --config {args.config}: {e}")
            sys.exit(1)
        opts.update(data.get("options", data) if isinstance(data, dict) else {})
    if args.a11y_audit:
        opts["a11y_audit"] = True
    if args.max_url_length:
        opts["max_url_length"] = args.max_url_length
    if args.retries is not None:
        opts["max_retries"] = args.retries
    if args.max_redirects:
        opts["max_redirects"] = args.max_redirects
    from crawler.config import validate_options
    errors = validate_options(opts)
    if errors:
        print("Error: " + "; ".join(f"{k.replace('_', ' ')}: {v}" for k, v in list(errors.items())[:4]))
        sys.exit(1)
    return opts


def main():
    if len(sys.argv) > 1 and sys.argv[1] == "mcp":
        # W9: "cli.py mcp" runs the stdio bridge for AI agents (ARACHNIX_URL + ARACHNIX_MCP_TOKEN; no access key)
        from crawler.mcp.bridge import main as bridge_main
        sys.exit(bridge_main(sys.argv[2:]))
    parser = argparse.ArgumentParser(
        description="ARACHNIX: website crawler, SEO audit and link checker with Excel / CSV export"
    )
    parser.add_argument("url", nargs="?", default="", help="Website URL to crawl (e.g. https://aristasystems.in/)")
    parser.add_argument("-o", "--output", help="Path to output Excel (.xlsx) file")
    parser.add_argument("-m", "--max-urls", type=int, default=1000, help="Maximum internal URLs to crawl (0 = unlimited, default: 1000)")
    parser.add_argument("-d", "--max-depth", type=int, default=0,
                        help="Maximum link depth (default: 0 = no limit, as in the dashboard)")
    parser.add_argument("-t", "--threads", type=int, default=8,
                        help=f"Concurrent worker threads, 1-{MAX_CONCURRENCY} (default: 8)")
    parser.add_argument("--rate", type=float, default=0.0,
                        help="Maximum URLs requested per second from the site (default: 0 = no limit)")
    parser.add_argument("--user-agent", default="chrome", metavar="PRESET|STRING",
                        help=f"User-agent preset ({', '.join(USER_AGENTS)}) or a full custom User-Agent string")
    parser.add_argument("--include", action="append", default=[], metavar="REGEX",
                        help="Only crawl internal URLs matching this regex (repeatable; the start URL is always crawled)")
    parser.add_argument("--exclude-regex", action="append", default=[], metavar="REGEX",
                        help="Never crawl internal URLs matching this regex (repeatable)")
    parser.add_argument("--no-resources", action="store_true",
                        help="Do not check internal images, scripts, stylesheets and iframes")
    parser.add_argument("--no-external-check", action="store_true",
                        help="Record external links without checking their status codes")
    parser.add_argument("--no-sitemap", action="store_true", help="Disable sitemap discovery")
    parser.add_argument("--no-external", action="store_true", help="Disable collecting external outbound links")
    parser.add_argument("--dynamic", action="store_true", help="Enable Selenium Chrome for JavaScript dynamic links")
    parser.add_argument("--delay", type=float, default=0.0, help="Polite delay between requests in seconds")
    parser.add_argument("--subdomains", action="store_true", help="Include subdomains as internal links")
    parser.add_argument("--verify-ssl", action="store_true", help="Enforce strict SSL certificate verification")
    parser.add_argument("--ignore-robots", action="store_true", help="Crawl URLs that robots.txt disallows")
    parser.add_argument("--category", default="all", choices=sorted(EXPORT_CATEGORIES),
                        help="Excel scope: all (default), 2xx, 3xx, 4xx, 5xx, error, issues or external")
    parser.add_argument("--pdf", action="store_true", help="Also create a visual PDF issue report")
    parser.add_argument("--no-screenshots", action="store_true", help="PDF without screenshots (no Chrome needed)")
    # Login (form login over HTTP, or in Chrome when the login page needs JavaScript)
    parser.add_argument("--username", help="Log in with this username / email before crawling")
    parser.add_argument("--password", metavar="PASSWORD|-",
                        help=f'Password for --username. "-" (or leaving it out) reads the {PASSWORD_ENV} '
                             "environment variable, which keeps the password out of the shell history")
    parser.add_argument("--login-url", help="Login page URL (found automatically when left out)")
    # Sign in with your own browser's session: ARACHNIX never tries to pass a bot check / CAPTCHA
    parser.add_argument("--cookies", metavar="FILE|-",
                        help="Crawl signed in with cookies copied from your browser (- = stdin): a Cookie header, "
                             "a JSON export, a Netscape cookies.txt or a 'Copy as cURL' command. Values are never "
                             "printed")
    parser.add_argument("--use-pasted-user-agent", action="store_true",
                        help="Send the User-Agent that came with --cookies (e.g. from 'Copy as cURL') for the whole "
                             "crawl: sessions are often tied to the browser that made them")
    parser.add_argument("--test-cookies", nargs="?", const="", default=None, metavar="PAGE",
                        help="Only test the --cookies: open one page that needs sign-in (PAGE, e.g. /dashboard; "
                             "found automatically when left out) and say whether they are signed in, then exit")
    # Subdomains: other hosts of the site's domain
    parser.add_argument("--subdomain", action="append", default=[], metavar="HOST",
                        help="Also crawl this subdomain of the site, from its home page (repeatable; Spider mode)")
    parser.add_argument("--find-subdomains", action="store_true",
                        help="List the site's subdomains (links on the home page, sitemaps, robots.txt, public "
                             "certificate logs, Search Console when connected, and --subdomain hosts), checking "
                             "each with one request, then exit")
    # Follow-up work after the crawl
    parser.add_argument("--validate", action="store_true",
                        help="Re-check every reported error live (confirmed / not reproduced / intermittent) "
                             "and, on logged-in crawls, check that member pages are not open to everyone")
    parser.add_argument("--intel", action="store_true",
                        help="Group issues by root cause with fix steps and review page content "
                             "(AI-enriched when GEMINI_API_KEY or GROQ_API_KEY is set)")
    parser.add_argument("--business", action="store_true",
                        help="Also write a business insights PDF (company, market, contacts, tech stack)")
    parser.add_argument("--bot-check", action="store_true",
                        help="Also fetch a sample of pages as search and AI crawlers (about a minute)")
    # Build 3: scope, List mode, rendering, headers, custom rules, more exports
    parser.add_argument("--scope", default="", choices=["", "subdomain", "subfolder", "all_subdomains", "exact_url"],
                        help="Crawl scope: subdomain (default), subfolder, all_subdomains or exact_url")
    parser.add_argument("--list", metavar="FILE|-",
                        help="List mode: crawl exactly the URLs in this .txt / .csv / .xlsx file (- = stdin); "
                             "their links are recorded, not followed")
    parser.add_argument("--render-all", action="store_true",
                        help="Render every internal page in Chrome (JavaScript tab, console log, rendered checks)")
    parser.add_argument("--header", action="append", default=[], metavar='"NAME: VALUE"',
                        help="Extra HTTP header sent to the crawled site only (repeatable; values are never printed)")
    parser.add_argument("--search", action="append", default=[], metavar="TEXT",
                        help="Custom Search: count pages containing this text (repeatable, at most 10)")
    parser.add_argument("--extract-css", action="append", default=[], metavar="NAME=SELECTOR",
                        help="Custom Extraction: the text of the elements this CSS selector matches (repeatable)")
    parser.add_argument("--format", default="xlsx", choices=["xlsx", "csv"],
                        help="File format of --report exports (default xlsx; csv gives .csv or a .zip of CSVs)")
    parser.add_argument("--report", action="append", default=[], metavar="ID",
                        help="Also write this report (e.g. crawl_overview, redirect_chains, orphan_pages; repeatable; "
                             "--list-reports shows them all)")
    parser.add_argument("--list-reports", action="store_true", help="Print the report and bulk export ids and exit")
    parser.add_argument("--sitemap", action="store_true", help="Also write an XML sitemap of the indexable pages")
    parser.add_argument("--save-crawl", action="store_true",
                        help="Save the crawl so the dashboard can open it again (File > Crawls)")
    parser.add_argument("--no-source", action="store_true", help="Do not keep each page's HTML (View Source)")
    # ---- ARACHNIX 2.0 (W9) ----
    parser.add_argument("--profile", help="Use the crawl options of this saved profile (Configuration > Profiles)")
    parser.add_argument("--config", metavar="FILE", help="Crawl options from a JSON file (a profile export, or "
                        "{option: value}: preferences, follow, crawl_types, limits, carbon, custom_robots_txt...)")
    parser.add_argument("--a11y-audit", action="store_true", help="Run the axe-core accessibility audit on every "
                        "page in Chrome (without rendering every page)")
    parser.add_argument("--max-url-length", type=int, default=0, help="Skip URLs longer than this (crawl trap guard)")
    parser.add_argument("--retries", type=int, default=None, help="Retries of a 500 / 502 / 504 answer (default 2)")
    parser.add_argument("--max-redirects", type=int, default=0, help="Redirect hops followed in one chain (default 10)")
    parser.add_argument("--search-where", default="html", choices=["html", "text", "head", "body", "text_no_anchors"],
                        help="Where --search looks (default: the HTML source)")

    args = parser.parse_args()
    # Every SEO check module must load, or the crawl's report would silently miss its checks
    require_all_checks("ARACHNIX command line")

    if args.list_reports:
        from crawler import reports
        print("Reports:")
        for r in reports.REPORTS:
            print(f"  {r['id']:<40} {' > '.join(r.get('menu_path') or [r['name']])}")
        print("Bulk exports:")
        for r in reports.BULK_EXPORTS:
            print(f"  {r['id']:<40} {' > '.join(r.get('menu_path') or [r['name']])}")
        return

    seed_urls = []
    if args.list:
        try:
            seed_urls = _read_list(args.list)
        except OSError as e:
            print(f"Error: the list file could not be read: {e}")
            sys.exit(1)
        if not seed_urls:
            print("Error: the list holds no URLs.")
            sys.exit(1)
        if not args.url:
            first = next((u for u in seed_urls if u.startswith(("http://", "https://"))), "")
            if not first:
                print("Error: the list holds no web addresses (each line must start with http:// or https://).")
                sys.exit(1)
            args.url = first

    if not args.url:
        print("\n=== ARACHNIX - website crawler and SEO audit ===")
        url_input = input("Enter Website URL (e.g. https://aristasystems.in/): ").strip()
        if not url_input:
            print("Error: No URL provided.")
            sys.exit(1)
        args.url = url_input

    if not args.url.startswith("http://") and not args.url.startswith("https://"):
        args.url = "https://" + args.url

    auth = AuthConfig()
    secrets = []
    if args.username or args.password:
        password = _password(args)
        if not args.username or not password:
            print(f"Error: a login needs --username and a password (--password, or set {PASSWORD_ENV} "
                  "and pass --password -).")
            sys.exit(1)
        auth = AuthConfig(enabled=True, auth_type="form", login_url=args.login_url or None,
                          username=args.username, password=password)
        secrets = [password]
    cookie_paste, cookie_preview = "", {}
    if args.cookies:
        if args.username or args.password:
            print("Error: sign in either with --cookies or with --username / --password, not both.")
            sys.exit(1)
        try:
            cookie_paste = _read_paste(args.cookies)
        except OSError as e:
            print(f"Error: the cookies file could not be read: {e.strerror or type(e).__name__}")
            sys.exit(1)
        from crawler.auth import preview_cookies, secret_values
        auth = AuthConfig(enabled=True, auth_type="cookies", cookies=cookie_paste)
        secrets = secret_values(auth, args.url)
        cookie_preview = preview_cookies(cookie_paste, args.url)
        for w in cookie_preview.get("warnings") or []:
            print(_mask(f"Note (cookies): {w}", secrets))
        if not cookie_preview.get("usable") and not cookie_preview.get("storage_origins"):
            why = cookie_preview.get("error") or "none of the pasted cookies belongs to this site or can still be used"
            print(_mask(f"Error: no cookies for this site could be read from {args.cookies}: {why}", secrets))
            sys.exit(1)
        if args.use_pasted_user_agent:
            if cookie_preview.get("user_agent"):
                auth.user_agent = cookie_preview["user_agent"]   # the sign-in and the whole crawl send it
            else:
                print("Note: the paste carried no User-Agent, so --user-agent is used.")
    elif args.use_pasted_user_agent or args.test_cookies is not None:
        print("Error: --use-pasted-user-agent and --test-cookies go with --cookies.")
        sys.exit(1)

    from crawler.url_utils import split_url_credentials
    domain = get_domain(split_url_credentials(args.url)[0]) or "website"   # never the user name or password
    timestamp = time.strftime("%Y%m%d_%H%M%S")
    output_path = args.output or f"crawl_{domain}_{timestamp}.xlsx"

    threads = max(1, min(MAX_CONCURRENCY, args.threads))
    if threads != args.threads:
        print(f"Note: --threads must be between 1 and {MAX_CONCURRENCY}; using {threads}.")
    ua_preset = args.user_agent.strip().lower()
    custom_ua = ""
    if ua_preset not in USER_AGENTS:
        ua_preset, custom_ua = "custom", args.user_agent.strip()
    if getattr(auth, "user_agent", None):
        ua_preset, custom_ua = "custom", auth.user_agent

    if args.test_cookies is not None:
        sys.exit(_test_cookies(cookie_paste, args.url, args.test_cookies, ua_preset, custom_ua,
                               args.verify_ssl, secrets))

    config = CrawlConfig(
        start_url=args.url,
        max_urls=args.max_urls,
        max_depth=args.max_depth,
        concurrency=threads,
        max_urls_per_second=max(0.0, args.rate),
        user_agent_preset=ua_preset,
        user_agent=custom_ua,
        include_patterns=args.include,
        exclude_regex=args.exclude_regex,
        crawl_resources=not args.no_resources,
        check_external_links=not args.no_external_check,
        use_sitemaps=not args.no_sitemap,
        collect_external_links=not args.no_external,
        use_dynamic_js=args.dynamic,
        delay=args.delay,
        include_subdomains=args.subdomains,
        verify_ssl=args.verify_ssl,
        respect_robots=not args.ignore_robots,
        auth=auth,
        output_file=output_path,
        post_crawl_checks=False,   # run below with postcrawl.run(..., {"bot_check": args.bot_check})
    )
    # Build 3 options, set after construction (and checked here, so a bad rule stops before the crawl)
    request_headers = _headers(args.header)
    secrets += [h["value"] for h in request_headers if len(h["value"]) >= 4]
    searches = [{"name": f"Search {i}: {t[:40]}", "mode": "contains", "text": t, "regex": False,
                 "where": args.search_where}
                for i, t in enumerate(args.search[:10], start=1)]
    extractors = []
    for i, raw in enumerate(args.extract_css[:10], start=1):
        name, sep, sel = raw.partition("=")
        if not sep:
            name, sel = f"Extractor {i}", raw
        extractors.append({"name": name.strip() or f"Extractor {i}", "type": "css", "expr": sel.strip(),
                           "extract": "text"})
    try:
        from crawler.engine import clean_request_headers
        from crawler.seo.extract import compile_custom_extractors, compile_custom_searches
        problems = clean_request_headers(request_headers)[1] + compile_custom_searches(searches)[1] \
            + compile_custom_extractors(extractors)[1]
    except ImportError:
        problems = []
    if problems:
        print("Error: " + "; ".join(str(p) for p in problems[:5]))
        sys.exit(1)
    try:
        from crawler.settings_store import proxy_config
        config.proxy = proxy_config()     # Settings > Proxy of this computer's reports folder
        if config.proxy.get("password"):
            secrets.append(config.proxy["password"])
    except Exception:
        pass
    config.request_headers = request_headers
    config.custom_searches, config.custom_extractors = searches, extractors
    b4 = _build4_options(args)
    if b4:
        from crawler.config import normalize_options
        for k, v in b4.items():
            if hasattr(config, k) and k not in ("start_url", "auth", "output_file"):
                setattr(config, k, v)
        normalize_options(config)
    config.render_all = bool(args.render_all)
    config.store_source = not args.no_source
    if args.scope:
        config.crawl_scope = args.scope
    if seed_urls:
        config.mode, config.seed_urls = "list", seed_urls
        config.use_sitemaps = False
        config.respect_robots = False     # List mode checks exactly the listed URLs (Screaming Frog's default)
    extra_hosts = [h.strip() for h in args.subdomain if h.strip()][:200]
    if extra_hosts and seed_urls:
        print("Note: --subdomain is ignored in List mode (it crawls exactly the listed URLs).")
    elif extra_hosts:
        config.extra_hosts = extra_hosts
    from crawler.config import sync_scope
    sync_scope(config)
    from crawler.url_utils import move_url_credentials
    secrets += move_url_credentials(config)   # a password in the URL: a site-only Basic auth header, masked
    if args.find_subdomains:
        sys.exit(find_subdomains(config, extra_hosts))

    print("\n========================================================")
    print(f" Target Website    : {config.start_url}")
    print(f" Output File       : {os.path.abspath(output_path)}")
    print(f" Concurrency       : {config.concurrency} threads"
          f"{f', max {config.max_urls_per_second:g} URLs/s' if config.max_urls_per_second else ''}")
    print(f" User-Agent        : {ua_preset} ({config.user_agent[:60]}{'...' if len(config.user_agent) > 60 else ''})")
    if config.include_patterns or config.exclude_regex:
        print(f" Include / Exclude : {config.include_patterns or '-'} / {config.exclude_regex or '-'}")
    print(f" Resources         : {'Images, scripts, CSS and iframes checked' if config.crawl_resources else 'Not checked'}")
    print(f" Sitemaps Probe    : {'Enabled' if config.use_sitemaps else 'Disabled'}")
    print(f" External Links    : {('Collect and check status (not crawled)' if config.check_external_links else 'Collect without crawling') if config.collect_external_links else 'Disabled'}")
    print(f" Dynamic JS DOM    : {'Enabled (Selenium)' if config.use_dynamic_js else 'Disabled'}")
    print(f" Max URLs Limit    : {'Unlimited' if config.max_urls == 0 else config.max_urls}")
    print(f" robots.txt        : {'Respected (as ' + config.robots_user_agent + ')' if config.respect_robots else 'Ignored'}")
    if auth.enabled and auth.auth_type == "cookies":
        n = cookie_preview.get("usable", 0)
        print(f" Login             : {n} pasted cookie{'s' if n != 1 else ''} "
              f"({cookie_preview.get('format') or 'header'} format; values ***)"
              f"{', with the pasted browser User-Agent' if getattr(auth, 'user_agent', None) else ''}")
    elif auth.enabled:
        print(f" Login             : {args.username} (password ***)"
              f"{' at ' + args.login_url if args.login_url else ', login page found automatically'}")
    if getattr(config, "extra_hosts", None):
        shown = ", ".join(config.extra_hosts[:5]) + (f" and {len(config.extra_hosts) - 5} more"
                                                     if len(config.extra_hosts) > 5 else "")
        print(f" Subdomains too    : {shown}")
    extras = [name for name, on in (("validation", args.validate), ("root causes", args.intel),
                                    ("business insights", args.business), ("bot check", args.bot_check)) if on]
    if extras:
        print(f" After the crawl   : {', '.join(extras)}")
    print("========================================================\n")

    def on_result(result):
        code_str = str(result.status_code) if result.status_code else "ERR"
        title_disp = (result.page_title[:42] + '..') if len(result.page_title) > 44 else result.page_title
        print(f"[{code_str:>3}] (D:{result.depth}) [Internal] {result.url} | {title_disp}", flush=True)

    def on_external_link(ext: ExternalLink):
        # Concise indicator when an external URL is found and recorded
        pass

    def on_log(msg):
        print(_mask(msg, secrets), flush=True)

    crawler = WebsiteCrawler(
        config=config,
        on_result=on_result,
        on_external_link=on_external_link,
        on_log=on_log
    )
    render = {}
    if config.render_all:
        # every internal 2xx HTML page rendered in Chrome (the pool starts after the login, on the first page)
        crawler.evidence_dir = os.path.splitext(output_path)[0] + "_evidence"

        def render_hook(url, seo):
            if "pool" not in render:
                from crawler.render_capture import RenderPool
                render["pool"] = RenderPool.from_crawler(crawler, evidence_dir=crawler.evidence_dir)
                crawler.render_pool = render["pool"]
            return render["pool"].render(url, raw=seo, load_url=str((seo or {}).get("final_url") or url))
        crawler.render_hook = render_hook
    elif getattr(config, "a11y_audit", False):
        from crawler.render_capture import AuditHook
        render["audit"] = AuditHook(crawler)
        crawler.audit_hook = render["audit"]

    try:
        results = crawler.run()
        if render.get("pool") is not None:
            render["pool"].close()
        if render.get("audit") is not None:
            render["audit"].close()
            rs = render["pool"].stats()
            print(f" Rendered in Chrome: {rs.get('rendered', 0)} page(s), {rs.get('failed', 0)} could not be rendered"
                  + (f"; {rs['promoted']} use their rendered HTML for the tabs" if rs.get("promoted") else ""))
        stats = crawler.get_stats()
        if crawler.status == "auth_failed":
            print(_mask(f"\n[LOGIN FAILED] Nothing was crawled: {crawler.auth_error}", secrets))
            sys.exit(2)

        print("\n" + "="*58)
        print(" Crawl Summary:")
        print(f" Total Internal Discovered : {stats['total_discovered']}")
        print(f" Total Internal URLs       : {stats['total_crawled']}   (the status rows below add up to it)")
        try:   # the dashboard's "Total URLs Crawled": internal + external, minus the URLs never requested
            from crawler.exporter import crawl_totals, crawled_note
            totals = crawl_totals(crawler.results_list, crawler.external_links_list)
            print(f" Total URLs Crawled        : {totals['crawled']}   ({crawled_note(totals)})")
        except Exception:
            pass
        print(f" {CATEGORY_LABELS['2xx']:<26}: {stats['status_2xx']}")
        print(f" {CATEGORY_LABELS['3xx']:<26}: {stats['status_3xx']}")
        print(f" {CATEGORY_LABELS['4xx']:<26}: {stats['status_4xx']}")
        print(f" {CATEGORY_LABELS['5xx']:<26}: {stats['status_5xx']}")
        print(f" {CATEGORY_LABELS['error']:<26}: {stats['status_err']}")
        print(f" {CATEGORY_LABELS['blocked']:<26}: {stats['status_blocked']}")
        print(f" External Links Collected  : {stats.get('total_external', 0)} ({stats.get('unique_external_domains', 0)} unique domains)")
        if crawler.check_external and config.collect_external_links:
            print(f" (External links checked for their status: {stats.get('external_checked', 0)}; "
                  "the other sites' pages are not crawled)")
        else:
            print(" (Note: external links are recorded only; their status was not checked)")
        for note in crawler.limit_notes():
            print(f" NOTE: {note}")
        print(f" Elapsed Time              : {stats['elapsed_seconds']}s")
        print("="*58)
        print_bot_checks_and_subdomains(crawler, secrets)

        # Post-crawl checks: robots.txt for search / AI crawlers, random-URL probe, optional bot check
        try:
            from crawler.seo import postcrawl
            postcrawl.run(crawler, {"bot_check": args.bot_check})
        except Exception as e:
            print(f"Post-crawl checks skipped: {e}")

        site, index = _seo_index(crawler)
        validation = run_validation(crawler) if args.validate else None
        if args.intel and index is not None:
            intel = run_intel(crawler, site, index, validation=validation)
            fix_plan, content = intel.get("errors"), intel.get("content")
        elif index is not None:   # rule-based, instant
            from crawler.intel import content as intel_content, root_cause
            fix_plan = root_cause.analyse(site, index, crawler, validation=validation)
            content = intel_content.analyse(site)
        else:
            fix_plan = content = None
        checks = None
        if index is not None:
            from crawler.intel.checks import check_statuses
            checks = check_statuses(index, crawler, site)

        # Export to Excel
        print(f"\nExporting to Excel spreadsheet: {output_path} ...")
        saved = ExcelExporter.export(   # another name when the file is open in Excel (Windows locks it)
            results=results,
            filepath=output_path,
            stats=crawler.get_stats(),
            external_links=crawler.external_links_list,
            category=args.category,
            site=site,
            seo_index=index,
            checks=checks,
            validation=validation,
            fix_plan=fix_plan,
            content=content,
        )
        print(f"[SUCCESS] Excel file saved successfully at: {os.path.abspath(saved)}")
        if args.pdf:
            pdf_path = write_pdf_report(crawler, output_path, screenshots=not args.no_screenshots,
                                        validation=validation, intel=fix_plan, content=content, checks=checks)
            print(f"[SUCCESS] PDF issue report saved at: {os.path.abspath(pdf_path)}")
        base_path = os.path.splitext(output_path)[0]
        if args.report and index is not None:
            from crawler import reports
            for rid in args.report:
                try:
                    name, data, _mime = reports.build_report(rid, site, index, crawler, fmt=args.format,
                                                             secrets=secrets, stats=crawler.get_stats())
                except KeyError:
                    print(f"Unknown report id: {rid} (see --list-reports)")
                    continue
                path = f"{base_path}_{name}"
                with open(path, "wb") as fh:
                    fh.write(data)
                print(f"[SUCCESS] Report saved at: {os.path.abspath(path)}")
        if args.sitemap and site is not None:
            from crawler import sitemap_gen
            try:
                name, data, _mime = sitemap_gen.build_sitemap(site, crawler, "xml")
                path = f"{base_path}_{name}"
                with open(path, "wb") as fh:
                    fh.write(data)
                print(f"[SUCCESS] XML sitemap saved at: {os.path.abspath(path)}")
            except sitemap_gen.SitemapEmpty as e:
                print(f"No sitemap written: {e}")
        if args.save_crawl:
            from crawler import store
            try:
                meta = store.save_crawl(crawler, index=(site, index) if index is not None else None,
                                        validation=validation, secrets=secrets)
                print(f"[SUCCESS] Crawl saved: {meta.get('name')} ({meta.get('id')}); open it in the dashboard under "
                      "File > Crawls")
            except store.StoreError as e:
                print(f"The crawl could not be saved: {e}")
        if args.business:
            biz_path = run_business(config.start_url, crawler, os.path.splitext(output_path)[0])
            print(f"[SUCCESS] Business insights PDF saved at: {os.path.abspath(biz_path)}")

    except KeyboardInterrupt:
        print("\n[Interrupt] Stopping crawler gracefully...")
        crawler.stop()
        results = crawler.results_list
        if results or crawler.external_links_list:
            saved = ExcelExporter.export(
                results=results,
                filepath=output_path,
                stats=crawler.get_stats(),
                external_links=crawler.external_links_list
            )
            print(f"[Saved partial results to {saved}]")
        sys.exit(0)
    finally:
        if render.get("pool") is not None:
            render["pool"].close()
        if render.get("audit") is not None:
            render["audit"].close()
        try:
            crawler.close_browser()
        except Exception:
            pass

if __name__ == "__main__":
    main()
