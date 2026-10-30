import base64
import contextlib
import copy
import hashlib
import hmac
import html as html_lib
import io
import ipaddress
import json
import logging
import math
import os
import platform as platform_lib
import re
import shutil
import subprocess
import sys
import time
import tempfile
import threading
import asyncio
import uuid
from typing import Any, Callable, Dict, List, Optional, Tuple, Union
from collections import Counter, OrderedDict, deque
from urllib.parse import parse_qsl, quote, unquote_plus, urlparse, urljoin
import requests
from bs4 import BeautifulSoup
from fastapi import FastAPI, Header, HTTPException, Query, Request
from fastapi.responses import HTMLResponse, FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from web import agents as _agents          # Build 4 W1: /mcp, agent rules, export links (imports nothing of app.py)
from web import connections as _connections   # Build 4 W2: the Connections hub endpoints
from web import ai_routes as _ai_routes       # Build 4 W3 / W4: Ask ARACHNIX, AI tab, spelling, n-grams
from web import extension_routes as _extension   # Build 4 W6: browser extension sign-in hand-off
from web import v2_misc as _v2_misc             # Build 4 W10 / W11: trend history, small endpoints
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, ConfigDict, Field

from crawler.config import CrawlConfig, AuthConfig, DEFAULT_USER_AGENT
try:
    from crawler.config import CRAWL_MODES, CRAWL_SCOPES, MAX_CUSTOM_RULES, default_analysis, default_url_rewrite
except ImportError:   # an older config module: the newer crawl options are refused with a clear message
    CRAWL_MODES, CRAWL_SCOPES, MAX_CUSTOM_RULES = ("spider",), ("subdomain", "all_subdomains"), 10

    def default_analysis() -> Dict[str, bool]:
        return {}

    def default_url_rewrite() -> Dict[str, Any]:
        return {}
from crawler.engine import WebsiteCrawler
from crawler.exporter import (ExcelExporter, EXPORT_CATEGORIES, ISSUE_URL_HEADERS, SOURCE_HEADERS,
                              EXTERNAL_FIX_ERROR, EXTERNAL_FIX_NO_RESPONSE, filter_source_rows, found_on,
                              inlink_rows, is_broken_external, issue_url_rows, json_value, resource_type_label,
                              row_category, row_status_code, sources_of, tab_grid, text_value)
from crawler.evidence import EvidenceCollector
from crawler.issues import ISSUE_CATEGORIES, issue_sort_key
from crawler.url_utils import get_domain, is_login_url, normalize_url

try:
    from crawler.config import USER_AGENTS
except ImportError:  # older config without presets
    USER_AGENTS = {}

try:
    from crawler.seo.text_metrics import pixel_width
except ImportError:
    def pixel_width(text: str, kind: str = "title") -> int:
        return 0

# --------------------------------------------------------------------------
# Team access key (public deployments): with CRAWLER_ACCESS_KEY set, every /api/* request needs it.
# Unset or empty (a PC) = no key, everything open as before. DEPLOY.md, 2.3.
# --------------------------------------------------------------------------
# Always open: the health checks (Render, uptime monitors), whether the dashboard has to ask for the key, and
# the page Google sends the browser back to after "Connect" (it cannot carry the key; the single-use state of
# the sign-in, made by a request that had the key, is what it is checked against)
GOOGLE_CALLBACK_PATH = "/api/integrations/google/callback"
ACCESS_OPEN_PATHS = ("/api/health", "/api/access", GOOGLE_CALLBACK_PATH)
# GET addresses the browser opens itself (<a href> downloads, <img src> screenshots, the live window's frame),
# which cannot send a header: these also take ?key=. The access log masks it (_RedactSessionIds).
KEY_IN_URL_RE = re.compile(r"^/api/(crawl/(report/[\w-]+|export(/[\w-]+){0,2}|download/[^/]+|evidence/[^/]+)"
                           r"|crawls/[^/]+/[^/]+/export|profiles/[^/]+/export|schedules/export|debug"
                           r"|business/pdf|export/(fixes|tickets)|live/frame|mcp/bridge|extension/download)$")
ACCESS_MAX_FAILS = 10           # different wrong keys from one client within ACCESS_WINDOW_S seconds...
ACCESS_WINDOW_S = 600
ACCESS_BLOCK_S = 600            # ...lock that client out (429) for this many seconds
ACCESS_MAX_CLIENTS = 10_000     # clients with recent wrong keys (and locked-out clients) kept in memory
# Both tables are kept in time order (oldest first), so a full one drops its oldest entry without a scan
_access_lock = threading.Lock()
_access_fails: "OrderedDict[str, Dict[str, float]]" = OrderedDict()   # client -> {fingerprint of a wrong key: last tried}
_access_blocked: "OrderedDict[str, float]" = OrderedDict()            # client -> locked out until (time.time())
# A pass: what /api/access hands out for the right key (see access_pass). Made from a secret created at start-up
ACCESS_PASS_RE = re.compile(r"^pass\.[0-9a-f]{64}$")
_ACCESS_PASS_SECRET = os.urandom(32)


def access_key() -> str:
    """The team access key (CRAWLER_ACCESS_KEY), '' when none is set."""
    return (os.environ.get("CRAWLER_ACCESS_KEY") or "").strip()


def access_key_matches(given: str, key: str) -> bool:
    """Constant-time comparison, so the answer time does not tell how much of a guess was right."""
    return bool(given) and bool(key) and hmac.compare_digest(given.encode("utf-8"), key.encode("utf-8"))


def access_pass(key: str) -> str:
    """The pass for `key`: 'pass.' + an HMAC of the key with a secret made when the server starts. It opens the
    API like the key, also while the visitor's address is locked out after someone else's wrong keys: nobody can
    guess it or work it out from a guessed key, so trying passes tells nothing about the key (and is not
    counted). It stops working when the key changes or the server restarts; the dashboard then trades the key
    for a new one."""
    return "pass." + hmac.new(_ACCESS_PASS_SECRET, key.encode("utf-8"), hashlib.sha256).hexdigest()


def access_pass_matches(given: str, key: str) -> bool:
    return bool(key) and bool(ACCESS_PASS_RE.match(given)) and hmac.compare_digest(given, access_pass(key))


def _trust_proxy() -> bool:
    """CRAWLER_TRUST_PROXY=1: a proxy in front appends the visitor's address to X-Forwarded-For."""
    return os.environ.get("CRAWLER_TRUST_PROXY", "").strip().lower() in ("1", "true", "yes")


def _header_values(scope, name: bytes) -> List[str]:
    """Every value of one request header, in order (a client can send a header several times)."""
    return [v.decode("latin-1") for k, v in scope.get("headers") or [] if k == name]


def _address_group(value: str) -> str:
    """An address as it is counted: IPv4 as is, IPv6 as its /64 network (one subscriber usually holds a whole
    /64, so its addresses count as one), '::ffff:1.2.3.4' as 1.2.3.4, a port removed. Anything else (an odd
    proxy value) is cut to 64 characters."""
    m = re.match(r"^\[([^\]]+)\](?::\d+)?$|^(\d{1,3}(?:\.\d{1,3}){3}):\d+$", value)
    try:
        ip = ipaddress.ip_address((m.group(1) or m.group(2)) if m else value)
        if ip.version == 6 and ip.ipv4_mapped:
            return str(ip.ipv4_mapped)
        return str(ipaddress.IPv6Network((int(ip) >> 64 << 64, 64))) if ip.version == 6 else str(ip)
    except ValueError:
        return value[:64] or "unknown"


def _client_address(scope) -> str:
    """The address wrong keys are counted for. Behind a proxy every request comes from the proxy itself, so the
    visitor's address is read from a header that only the proxy sets, never from one a visitor can choose
    (True-Client-IP, the first X-Forwarded-For entry):
    - CRAWLER_CLIENT_IP_HEADER (a header name, e.g. X-Real-IP): that header, as the proxy in front sets it;
    - on Render (RENDER is set): CF-Connecting-IP, which Cloudflare in front of Render always overwrites;
    - CRAWLER_TRUST_PROXY=1, or on Render without CF-Connecting-IP: the last X-Forwarded-For entry, the one the
      proxy appended (a visitor can only put entries in front of it).
    Otherwise the direct peer."""
    named = (os.environ.get("CRAWLER_CLIENT_IP_HEADER") or "").strip().lower()
    render = bool(os.environ.get("RENDER"))
    if named:
        found = _header_values(scope, named.encode("latin-1"))
    else:
        found = _header_values(scope, b"cf-connecting-ip") if render else []
        if not "".join(found).strip() and (render or _trust_proxy()):
            found = _header_values(scope, b"x-forwarded-for")
    value = ",".join(found).split(",")[-1].strip()
    return _address_group(value or (scope.get("client") or ("",))[0] or "unknown")


def _access_wait(client: str, now: float) -> float:
    """Seconds this client is still locked out (0 = not locked out)."""
    with _access_lock:
        until = _access_blocked.get(client, 0.0)
        if until > now:
            return until - now
        _access_blocked.pop(client, None)
        return 0.0


def _access_failed(client: str, given: str, now: float) -> int:
    """Counts a wrong key; returns how many more different wrong keys this client may send (0: it is locked out
    from now on). The same wrong key sent again (a dashboard's parallel requests with an old key) counts once:
    only different guesses add up. Constant time: an update moves the client to the end of _access_fails, so
    expired entries and, when the table is full, the oldest one are taken from the front."""
    mark = hashlib.sha256(given.encode("utf-8", "replace")).hexdigest()[:16]
    with _access_lock:
        fails = {m: t for m, t in _access_fails.pop(client, {}).items() if now - t < ACCESS_WINDOW_S}
        fails[mark] = now
        left = max(0, ACCESS_MAX_FAILS - len(fails))
        if not left:
            _access_blocked.pop(client, None)
            while _access_blocked and (len(_access_blocked) >= ACCESS_MAX_CLIENTS
                                       or next(iter(_access_blocked.values())) <= now):
                _access_blocked.popitem(last=False)
            _access_blocked[client] = now + ACCESS_BLOCK_S
            return 0
        while _access_fails and (len(_access_fails) >= ACCESS_MAX_CLIENTS
                                 or now - max(next(iter(_access_fails.values())).values(), default=0) >= ACCESS_WINDOW_S):
            _access_fails.popitem(last=False)
        _access_fails[client] = fails
        return left


def _access_denied(wait: float = 0.0, left: Optional[int] = None) -> JSONResponse:
    """401 (no or wrong key), or 429 while the client is locked out. X-Access-Required tells the dashboard the
    answer is about the team key (a site login can fail with 401 too); X-Access-Tries-Left, after a wrong key,
    how many more different wrong keys lock this address out."""
    headers = {"X-Access-Required": "1", "Cache-Control": "no-store"}
    if wait <= 0:
        if left is not None:
            headers["X-Access-Tries-Left"] = str(left)
        return JSONResponse({"detail": "Access key required"}, status_code=401, headers=headers)
    minutes = max(1, math.ceil(wait / 60))
    headers["Retry-After"] = str(math.ceil(wait))
    return JSONResponse({"detail": f"Too many wrong access keys. Try again in {minutes} minute{'s' if minutes > 1 else ''}."},
                        status_code=429, headers=headers)


def _given_key(scope, method: str, path: str) -> str:
    """The key a request carries: the X-Access-Key header, or ?key= on the GET addresses in KEY_IN_URL_RE.
    Sent several times, the non-empty values are joined with ', ' (as a proxy joins them) and checked, and
    counted, as one guess."""
    given = ", ".join(v for v in (h.strip() for h in _header_values(scope, b"x-access-key")) if v)
    if not given and method in ("GET", "HEAD") and KEY_IN_URL_RE.match(path):
        query = parse_qsl(scope.get("query_string", b"").decode("latin-1"), keep_blank_values=True)
        given = ", ".join(v for v in (v.strip() for k, v in query if k == "key") if v)
    return given


class AccessKeyMiddleware:
    """With CRAWLER_ACCESS_KEY set, /api/* answers 401 without the key or its pass (header X-Access-Key; ?key=
    only on the GET addresses in KEY_IN_URL_RE) and 429 for ACCESS_BLOCK_S seconds once one client sent
    ACCESS_MAX_FAILS different wrong keys within ACCESS_WINDOW_S; a pass still works then. The dashboard page,
    its static files, /healthz, /api/health and /api/access (without a key) stay open. The value decided here
    is handed to /api/access in request.state.access_key, so no route reads the header a second way.
    CORSMiddleware wraps this one, so a dashboard on another origin can read the 401 and 429 answers."""

    def __init__(self, app):
        self.app = app

    async def __call__(self, scope, receive, send):
        internal = _agents.internal_state(scope) if scope["type"] == "http" else ""
        if internal == "internal":       # an in-process tool call (MCP, Ask ARACHNIX): never needs the key
            return await self.app(scope, receive, send)
        key = access_key() if scope["type"] == "http" else ""
        path, method = scope.get("path", ""), scope.get("method", "")
        if not key or method == "OPTIONS" or not (path == "/api" or path.startswith("/api/")):
            return await self.app(scope, receive, send)
        if internal == "forged":         # the internal header from the network counts as a wrong key
            client, now = _client_address(scope), time.time()
            wait = _access_wait(client, now)
            left = None if wait else _access_failed(client, "internal:" + _agents._header(
                scope, b"x-slickup-internal"), now)
            return await _access_denied(wait if wait else (0.0 if left else ACCESS_BLOCK_S), left)(scope, receive, send)
        if _agents.export_link_ok(path, method):    # a single-purpose download link (W1), valid and not used up
            return await self.app(scope, receive, send)
        given = _given_key(scope, method, path)
        scope.setdefault("state", {})["access_key"] = given
        if method in ("GET", "HEAD") and path in ACCESS_OPEN_PATHS and not (given and path == "/api/access"):
            return await self.app(scope, receive, send)
        if access_pass_matches(given, key):
            return await self.app(scope, receive, send)
        client, now = _client_address(scope), time.time()
        wait, left = _access_wait(client, now), None
        # a request without any key is not a guess, nor is a pass (it is never compared with the key)
        if not wait and given and not ACCESS_PASS_RE.match(given):
            if access_key_matches(given, key):
                return await self.app(scope, receive, send)
            left = _access_failed(client, given, now)
            wait = 0.0 if left else ACCESS_BLOCK_S
        await _access_denied(wait, left)(scope, receive, send)


class DownloadHeadersMiddleware:
    """Every file download (Content-Disposition: attachment) also says X-Content-Type-Options: nosniff, so the
    browser keeps the declared type (application/pdf, the Excel type) instead of guessing one. It does not stop
    Chrome's "insecure download" warning on plain HTTP from another computer's address: only HTTPS does (run.py
    --ssl-certfile / --ssl-keyfile, DEPLOY.md)."""

    def __init__(self, app):
        self.app = app

    async def __call__(self, scope, receive, send):
        if scope["type"] != "http":
            return await self.app(scope, receive, send)

        async def send_wrapped(message):
            if message.get("type") == "http.response.start":
                headers = list(message.get("headers") or [])
                names = {k.lower() for k, _ in headers}
                disposition = next((v for k, v in headers if k.lower() == b"content-disposition"), b"")
                if disposition.lower().startswith(b"attachment") and b"x-content-type-options" not in names:
                    headers.append((b"x-content-type-options", b"nosniff"))
                    message = dict(message, headers=headers)
            await send(message)
        return await self.app(scope, receive, send_wrapped)


def cors_origins() -> List[str]:
    """CRAWLER_CORS_ORIGINS (comma-separated, e.g. the Vercel dashboard's address); unset = every origin."""
    raw = os.environ.get("CRAWLER_CORS_ORIGINS") or ""
    return [o.strip().rstrip("/") for o in raw.split(",") if o.strip()] or ["*"]


@contextlib.asynccontextmanager
async def _lifespan(_app):
    """Server start / stop: the scheduler's tick thread and the integrations listener (see _on_startup).
    Refuses to start when an SEO check module failed to import (crawls would silently miss its checks), unless
    ARACHNIX_ALLOW_BROKEN_TABS=1 (then a loud warning, /healthz and the dashboard's red banner report it)."""
    from crawler.seo import registry as _seo_registry
    _seo_registry.require_all_tabs("ARACHNIX server")
    _on_startup()
    try:
        async with _agents.mcp_app().lifespan():      # /mcp (W1): created whether MCP is on or off
            yield
    finally:
        _agents.agent_log.flush()
        _on_shutdown()


app = FastAPI(title="ARACHNIX API", lifespan=_lifespan)

# Added first, so it runs inside CORSMiddleware (the last one added is the outermost)
app.add_middleware(DownloadHeadersMiddleware)
app.add_middleware(AccessKeyMiddleware)
app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_origins(),
    # The dashboard sends no cookies (the team key travels in the X-Access-Key header), and browsers refuse
    # credentials with "*" anyway
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],          # Content-Type, X-Access-Key, X-Live-Session
    # the hosted dashboard (another origin) reads the download name, whether fix drafts came from the AI,
    # the address the live login window shows, whether a 401 / 429 is about the team key and how many wrong keys
    # are left before a lock-out
    expose_headers=["Content-Disposition", "X-Fixes-AI", "X-Live-Url", "X-Access-Required", "X-Access-Tries-Left", "Retry-After"],
)

MAX_SCREENSHOT_ISSUES = int(os.environ.get("CRAWLER_MAX_SCREENSHOTS", "25"))
REPORT_RETENTION_HOURS = float(os.environ.get("CRAWLER_REPORT_RETENTION_HOURS", "72"))
REPORT_FILE_RE = re.compile(r"^crawl_[\w.-]+\.(xlsx|pdf|csv|zip)$")
GENERATED_RE = re.compile(r"^crawl_[\w.-]+(\.(xlsx|pdf|csv|zip)|_evidence)$")
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FILE_MEDIA = {".xlsx": XLSX_MEDIA, ".pdf": "application/pdf", ".csv": "text/csv; charset=utf-8",
              ".zip": "application/zip"}
MAX_CONCURRENCY = 50
MAX_PAGE_SIZE = 5000          # rows per /api/crawl/tab page
LINKED_FROM_SHOWN = 50        # linking pages listed per broken URL in /api/crawl/issues
PDF_ISSUE_EXAMPLES = 3        # example URLs per SEO issue in the PDF
UA_PRESETS = ("chrome", "googlebot_desktop", "googlebot_smartphone", "bingbot", "custom")
MAX_LIST_URLS = int(os.environ.get("CRAWLER_MAX_LIST_URLS", "10000"))   # List mode: URLs per list (see _list_cap)
MAX_UPLOAD_BYTES = int(float(os.environ.get("CRAWLER_MAX_UPLOAD_MB", "20")) * 1024 * 1024)   # list / SERP uploads
MAX_SOURCE_CHARS = 3_000_000  # View Source: characters of HTML returned per page
MAX_CUSTOM_JS_CHARS = 100_000  # one Custom JavaScript snippet
APP_NAME = "ARACHNIX"
DETAIL_LIST_KEYS = ("links", "resources", "images")   # PageData lists served as outlinks/resources/images
DETAIL_HEAVY_KEYS = DETAIL_LIST_KEYS + ("rendered", "ids", "cookies", "original_html")   # served by their own
# detail endpoints (original_html: the raw-HTML values of a page whose tabs read its rendered HTML, in the
# Rendered Page summary as "html")
PRIORITY_ORDER = {"High": 0, "Medium": 1, "Low": 2}
TYPE_ORDER = {"Issue": 0, "Warning": 1, "Opportunity": 2}
# Crawl snapshots (REPORTS_DIR/snapshots/<host>/) and dismissed issues (REPORTS_DIR/dismissed/) live here
REPORTS_DIR = os.environ.get("CRAWLER_REPORTS_DIR") or tempfile.gettempdir()
# Server-wide URL cap per crawl (0 = none): a hosted instance has a fixed amount of memory (see DEPLOY.md 2.1)
SERVER_MAX_URLS = max(0, int(os.environ.get("CRAWLER_MAX_URLS", "0") or 0))
PDF_PINPOINT_WAIT = float(os.environ.get("CRAWLER_PDF_PINPOINT_WAIT", "240"))  # s a PDF request waits for pictures
FIXES_AI_WAIT = float(os.environ.get("CRAWLER_FIXES_AI_WAIT", "90"))          # s a fix-file request waits for AI drafts
BOT_CHECK_SAMPLE = int(os.environ.get("CRAWLER_BOT_CHECK_SAMPLE", "10"))       # pages fetched as each crawler
MAX_RECHECK_URLS = 500
ROBOTS_TEXT_SHOWN = 20_000          # characters of robots.txt returned by /api/crawl/robots-audit
LOGIN_TEST_DIR = os.path.join(tempfile.gettempdir(), "crawl_logintest_evidence")
# background jobs that belong to one crawl ("analysis" = Crawl Analysis > Start, "export" = Multi-Export)
JOB_KINDS = ("validate", "intel", "postcrawl", "pinpoint", "fixes", "analysis", "export")
VERDICT_KEYS = ("confirmed", "not_reproduced", "intermittent", "fixed")


CUT_SECRET_MIN = 10     # an error message that quotes a secret cut short still shows at least this much of it


def _mask_cut_secret(text: str, secret: str) -> str:
    """`text` with every run of CUT_SECRET_MIN+ characters that matches the start of `secret` replaced by ***
    (a library error quoting the first 160 characters of a long token)."""
    head, out, i = secret[:CUT_SECRET_MIN], [], 0
    while True:
        j = text.find(head, i)
        if j < 0:
            break
        k = CUT_SECRET_MIN
        while k < len(secret) and j + k < len(text) and text[j + k] == secret[k]:
            k += 1
        out.append(text[i:j] + "***")
        i = j + k
    out.append(text[i:])
    return "".join(out)


def mask_secrets(text: Any, secrets: List[str]) -> str:
    """`text` with every secret (password, header value, cookie value) replaced by ***, also when only the
    start of a long secret is quoted."""
    text = str(text or "")
    for s in secrets or ():
        if not s:
            continue
        if s in text:
            text = text.replace(s, "***")
        if len(s) >= 16 and s[:CUT_SECRET_MIN] in text:
            text = _mask_cut_secret(text, s)
    return text


# Response headers whose value can be a working login (the site's refreshed session cookie, a token)
_SECRET_HEADER = re.compile(r"cookie|authorization|token|secret|session|api-?key|signature|csrf|xsrf|passw", re.I)
_COOKIE_PAIR = re.compile(r"(^|[;,]\s*)([^=;,\s]+)=([^;,]*)")
_COOKIE_ATTRS = {"path", "domain", "expires", "max-age", "samesite", "priority", "partitioned", "secure", "httponly"}


def _masked_set_cookie(value: str) -> str:
    """A Set-Cookie header with every cookie value replaced by *** (names and attributes stay readable)."""
    found = []

    def sub(m):
        if m.group(2).lower() in _COOKIE_ATTRS:
            return m.group(0)
        found.append(1)
        return f"{m.group(1)}{m.group(2)}=***"
    masked = _COOKIE_PAIR.sub(sub, value)
    return masked if found else "***"


def safe_headers(headers: Optional[Dict[str, Any]], secrets: Optional[List[str]] = None) -> Dict[str, str]:
    """Response headers fit for an API answer: cookie values, tokens and the crawl's own secrets are masked.
    On a logged-in crawl a site's Set-Cookie carries a refreshed session, so its values never leave the server."""
    out = {}
    for k, v in (headers or {}).items():
        v = ", ".join(map(str, v)) if isinstance(v, (list, tuple)) else str(v)
        name = str(k).lower()
        if name == "set-cookie":
            v = _masked_set_cookie(v)
        elif _SECRET_HEADER.search(name):
            v = "***"
        out[str(k)] = mask_secrets(v, secrets or [])
    return out


class _RedactSessionIds(logging.Filter):
    """Keeps live-login session ids and the team access key out of the server's access log: after
    /api/live/finish the id unlocks the kept login for hours, and download and picture addresses carry the
    key as ?key= (KEY_IN_URL_RE). Every query parameter whose name, decoded the way the server decodes it (or
    more often), is key / session_id / live_session_id is masked, also after ';' or as '%6Bey'; so is any value
    that is the current key or its pass, and the key itself wherever else it shows up."""
    _NAMES = ("key", "session_id", "live_session_id", "handoff", "handoff_id")
    _QUERY = re.compile(r"\?(\S*)")
    _LINK = re.compile(r"(/api/export/link/)[A-Za-z0-9_\-]{8,}")

    @classmethod
    def _secret_name(cls, name: str) -> bool:
        for _ in range(3):          # %256Bey decodes to %6Bey, then to key
            name = unquote_plus(name)
        return name.strip().lower() in cls._NAMES

    def _redact(self, text: str, values: Tuple[str, ...]) -> str:
        def query(m):
            parts = re.split(r"([&;])", m.group(1))
            for i in range(0, len(parts), 2):
                name, eq, value = parts[i].partition("=")
                if eq and value and (self._secret_name(name) or unquote_plus(value).strip() in values):
                    parts[i] = name + "=***"
            return "?" + "".join(parts)
        if "?" in text:
            text = self._QUERY.sub(query, text)
        if "/api/export/link/" in text:
            text = self._LINK.sub(r"\1***", text)
        for v in values:
            if len(v) >= 8 and v in text:
                text = text.replace(v, "***")
        return text

    def filter(self, record: logging.LogRecord) -> bool:
        key = access_key()
        values = (key, quote(key, safe=""), access_pass(key)) if key else ()
        args = record.args
        if isinstance(args, tuple) and any(isinstance(a, str) for a in args):
            record.args = tuple(self._redact(a, values) if isinstance(a, str) else a for a in args)
        elif isinstance(record.msg, str):
            record.msg = self._redact(record.msg, values)
        return True


def _install_log_redaction():
    access = logging.getLogger("uvicorn.access")
    if not any(isinstance(f, _RedactSessionIds) for f in access.filters):
        access.addFilter(_RedactSessionIds())


_install_log_redaction()


def _auth_secrets(auth: Optional[AuthConfig], start_url: str = "") -> List[str]:
    """Values that must never show up in a log line or an API answer, longest first: crawler.auth.secret_values
    (password, header value and token, every pasted cookie value however it was pasted, escaped forms)."""
    if not auth or not getattr(auth, "enabled", False):
        return []
    try:
        from crawler.auth import secret_values
        return secret_values(auth, start_url or auth.login_url or "https://example.com/")
    except Exception:   # never lose the masking: fall back to the simple split below
        pass
    found = [auth.password or "", auth.custom_header_value or ""]
    found += (auth.custom_header_value or "").split()[1:]          # the token of "Bearer <token>"
    raw = (auth.cookies or "").strip()
    try:
        data = json.loads(raw) if raw[:1] in "[{" else None
    except ValueError:
        data = None
    if data is not None:
        items = data.get("cookies") if isinstance(data, dict) else data
        found += [str(c.get("value") or "") for c in items or [] if isinstance(c, dict)]
    else:
        found += [part.partition("=")[2].strip() for part in raw.split(";")]
    return sorted({s for s in found if len(s) >= 4}, key=len, reverse=True)


class JobBusy(Exception):
    """A job of this kind is already running."""


class JobClosed(JobBusy):
    """The crawl this runner belongs to was replaced: none of its jobs start any more."""


def _ai_scope(stop_event: threading.Event):
    """crawler.ai.client.job_scope(stop_event), or a no-op when the AI client cannot be imported."""
    try:
        from crawler.ai import client as ai_client
        return ai_client.job_scope(stop_event)
    except Exception:
        import contextlib
        return contextlib.nullcontext()


class JobRunner:
    """One kind of background job (business, validate, intel, ...): a worker thread, its state and a stop flag.
    One job of a kind runs at a time; the last result stays until the next job of that kind starts."""

    def __init__(self, kind: str, mask: Optional[Callable[[str], str]] = None):
        self.kind = kind
        self.mask = mask or (lambda s: s)
        self.job_id = ""
        self.state = "idle"             # idle -> running -> done | error
        self.progress = 0
        self.message = ""
        self.result: Any = None
        self.meta: Dict[str, Any] = {}
        self.started_at = 0.0
        self.finished_at = 0.0
        self.stop_event = threading.Event()
        self.thread: Optional[threading.Thread] = None
        self.closed = False             # close(): stopped for good (its crawl was replaced)
        self._lock = threading.Lock()

    @property
    def busy(self) -> bool:
        return bool(self.thread and self.thread.is_alive())

    def start(self, work: Callable[["JobRunner"], Any], message: str = "Starting...", **meta) -> str:
        """Runs work(runner) in a new thread; its return value becomes the result. Raises JobBusy while
        a job of this kind is still running."""
        with self._lock:
            if self.closed:
                raise JobClosed(f"The {self.kind} job belongs to a crawl that was replaced.")
            if self.busy:
                raise JobBusy(f"A {self.kind} job is already running.")
            self.job_id = uuid.uuid4().hex[:12]
            self.state, self.progress, self.message = "running", 0, message
            self.result, self.meta = None, dict(meta)
            self.stop_event = threading.Event()
            self.started_at, self.finished_at = time.time(), 0.0
            self.thread = threading.Thread(target=self._run, args=(work,), daemon=True, name=f"job-{self.kind}")
            self.thread.start()
            return self.job_id

    def _run(self, work: Callable[["JobRunner"], Any]):
        try:
            with _ai_scope(self.stop_event):   # its AI calls end when the job stops or a new crawl starts
                result = work(self)
            self.result = result          # the result first: a poller that sees "done" also sees it
            self.progress = 100
            self.state = "done"
        except Exception as e:            # a job never takes the server down; the message says why
            self.message = self.mask(f"The {self.kind} job failed: {type(e).__name__}: {str(e)[:300]}")
            self.state = "error"
        finally:
            self.finished_at = time.time()

    def set_done(self, result: Any, message: str = "", **meta) -> str:
        """Shows an earlier result (a saved business report) as this runner's finished job, without running
        anything. Raises JobBusy while a job of this kind is running."""
        with self._lock:
            if self.closed:
                raise JobClosed(f"The {self.kind} job belongs to a crawl that was replaced.")
            if self.busy:
                raise JobBusy(f"A {self.kind} job is already running.")
            self.job_id = uuid.uuid4().hex[:12]
            self.result, self.meta = result, dict(meta)
            self.state, self.progress, self.message = "done", 100, self.mask(message)
            self.stop_event = threading.Event()
            self.started_at = self.finished_at = time.time()
            return self.job_id

    def update(self, progress: Optional[int] = None, message: Optional[str] = None):
        if progress is not None:
            self.progress = max(0, min(99, int(progress)))
        if message:
            self.message = self.mask(message)

    def stop(self):
        self.stop_event.set()

    def close(self):
        """Stops the running job and keeps any later start() from running (a new stop_event per start would
        otherwise forget this stop)."""
        with self._lock:
            self.closed = True
            self.stop_event.set()

    def wait(self, timeout: float) -> bool:
        """Waits up to `timeout` seconds; True when no job of this kind is running any more."""
        thread = self.thread
        if thread is not None and thread.is_alive():
            thread.join(max(0.0, timeout))
        return not self.busy

    def status(self) -> dict:
        return {"job_id": self.job_id, "state": self.state, "progress": self.progress, "message": self.message}


class ViewCache:
    """Grids and per-URL lookups derived from one TabIndex; replaced when the index is rebuilt."""

    def __init__(self, index):
        self.index = index
        self.grids: Dict[Tuple[str, str], tuple] = {}   # (tab, filter) -> (columns, rows, search texts)
        self.url_issues: Optional[Dict[str, List[dict]]] = None
        self.ext_by_url: Optional[Dict[str, Any]] = None
        self.res_inlinks: Optional[Dict[str, List[dict]]] = None   # url -> pages embedding it (img, css, js)


class CrawlJob:
    """Everything belonging to one crawl run. A new crawl gets a new job, so a finishing
    background thread can never overwrite the state of the crawl that replaced it."""

    def __init__(self, crawler: Optional[WebsiteCrawler], base_name: str, screenshots: bool = True,
                 secrets: Optional[List[str]] = None):
        self.crawler = crawler
        self.base_name = base_name  # crawl_<domain>_<timestamp>
        self.thread: Optional[threading.Thread] = None
        self.logs = deque(maxlen=100)
        self.phase = "crawling"  # crawling -> analyzing -> done
        self.export_file: Optional[str] = None
        self.export_stale = False   # follow-up data changed: the saved workbook is rebuilt on its next download
        self.pdf_file: Optional[str] = None
        self.pdf_key: Optional[tuple] = None   # inputs the saved PDF was built from (see _pdf_key)
        self.pdf_lock = threading.Lock()
        self.excel_lock = threading.RLock()   # one workbook build at a time: they share file names and ~400 MB
        self.retired = False                   # replaced by a newer crawl (stop_jobs)
        self.evidence: Dict[str, dict] = {}
        self.evidence_dir = os.path.join(get_export_dir(), f"{base_name}_evidence")
        self.evidence_progress = {"done": 0, "total": 0}
        self.screenshot_note = ""
        self.screenshots = screenshots
        self.seo_evidence: dict = {}           # pinpointed issue pictures {(tab, filter): [item]}
        self.cancel_event = threading.Event()
        self.view: Optional[ViewCache] = None
        self.view_lock = threading.Lock()
        self.index_error = ""
        self.secrets = list(secrets or [])
        self.runners: Dict[str, JobRunner] = {k: JobRunner(k, mask=self.mask) for k in JOB_KINDS}
        self.validation: Optional[dict] = None   # every live re-check of this crawl merged, latest verdict per URL
        self.validation_version = 0
        self.results_version = 0    # bumped when re-checks change crawl rows: the dashboard reloads them
        self.snapshot_id = ""
        self.memo: Dict[tuple, Any] = {}
        self.memo_lock = threading.Lock()
        self.run_ref = uuid.uuid4().hex[:16]   # what the scheduler knows this crawl by
        self.schedule_ref = ""                 # set when a scheduled task started this crawl
        self.origin = "manual"                 # manual | schedule | opened (a saved crawl) | direct
        self.render: Dict[str, Any] = {}       # render_all: {"pool": RenderPool, "html": SourceStore}
        self.notes: List[str] = []             # settings that were changed or ignored when the crawl started
        self.saved_meta: Optional[dict] = None   # meta of the saved crawl (File > Crawls) this job was saved as
        self.save_name = ""                    # name / project for the saved crawl (crawl request)
        self.save_project = ""
        self.after_done = False                # the follow-up chain (pictures, bot check, email...) has started
        self.saved_deleted = False             # its saved copy was deleted (File > Crawls): not saved again by itself
        self.post_since = 0.0                  # when a scheduled run's after-crawl steps started (0: none running)
        self.challenge_cache: Optional[tuple] = None   # (key, challenge object, made at, cost) - see _crawl_challenge

    def mask(self, text: Any) -> str:
        return mask_secrets(text, self.secrets)

    def log(self, msg: str):
        self.logs.append(self.mask(msg))

    def stop_jobs(self):
        """The job was replaced (new crawl, reset): stops its follow-up jobs and keeps them from starting later,
        e.g. the bot check that run_crawler_task starts once the reports are built."""
        self.retired = True
        for runner in self.runners.values():
            runner.close()
        try:   # its Search Console / GA4 / PageSpeed fetches stop too (they would keep using the Google quota)
            from crawler.integrations import stop_for_crawler
            stop_for_crawler(self.crawler)
        except Exception:
            pass

    @property
    def busy(self) -> bool:
        return bool(self.thread and self.thread.is_alive())


class CrawlSession:
    def __init__(self):
        self.job: Optional[CrawlJob] = None
        self.business = JobRunner("business")   # works with or without a crawl


session = CrawlSession()


class AuthRequest(BaseModel):
    enabled: bool = False
    auth_type: str = "none" # "none", "form", "cookies", "header", "live"
    login_url: Optional[str] = None
    username: Optional[str] = None
    password: Optional[str] = None
    # "a=1; b=2", or JSON: a list of full cookie objects, or {"cookies": [...], "local": {...}, "session": {...}}
    cookies: Optional[Union[str, list, dict]] = None
    custom_header_name: Optional[str] = None
    custom_header_value: Optional[str] = None
    live_session_id: Optional[str] = None      # auth_type "live": from POST /api/live/finish
    handoff_id: Optional[str] = None           # W6: a browser-extension hand-off (consumed at crawl start)
    username_selector: Optional[str] = None
    password_selector: Optional[str] = None
    submit_selector: Optional[str] = None
    # The browser User-Agent the person chose to keep from a paste (e.g. a "Copy as cURL" command): sessions are
    # often tied to the browser that made them, so the sign-in and the whole crawl send it (see _signin_user_agent)
    user_agent: Optional[str] = None

class CrawlRequest(BaseModel):
    url: Optional[str] = ""             # the start URL (List mode: optional, the first listed URL is used)
    auth: Optional[AuthRequest] = None
    # Smart defaults for automated crawling:
    max_urls: Optional[int] = 0        # 0 = Unlimited (crawls until queue is exhausted)
    max_depth: Optional[int] = 0       # 0 = no limit (Screaming Frog's default; W7 / P2)
    # ---- Build 4 options (PLAN4_build.md 10.7; None = CrawlConfig's default; checked by config.validate_options)
    max_redirects: Optional[int] = None
    custom_robots_txt: Optional[str] = None
    sitemap_urls: Optional[Union[List[str], str]] = None
    max_sitemap_files: Optional[int] = None
    max_sitemap_only_urls: Optional[int] = None
    max_sitemap_index_urls: Optional[int] = None
    cdn_hosts: Optional[Union[List[str], str]] = None
    near_dup_threshold: Optional[float] = None
    content_area: Optional[dict] = None
    preferences: Optional[dict] = None
    soft_404_phrases: Optional[Union[List[str], str]] = None
    non_descriptive_anchors: Optional[Union[List[str], str]] = None
    link_positions: Optional[List[dict]] = None
    crawl_types: Optional[dict] = None
    follow: Optional[dict] = None
    follow_internal_nofollow: Optional[bool] = None
    list_follow_canonicals: Optional[bool] = None
    cookie_storage: Optional[str] = None
    assume_html: Optional[bool] = None
    ignore_non_indexable_issues: Optional[bool] = None
    limits: Optional[dict] = None
    block_addresses: Optional[str] = None
    carbon: Optional[dict] = None
    spelling: Optional[dict] = None
    a11y_audit: Optional[bool] = None
    a11y_audit_max_pages: Optional[int] = None
    a11y_audit_max_minutes: Optional[int] = None
    render_device: Optional[str] = None
    render_timeout: Optional[float] = None
    flatten_shadow_dom: Optional[bool] = None
    flatten_iframes: Optional[bool] = None
    concurrency: Optional[int] = 6
    timeout: Optional[int] = 15
    use_sitemaps: Optional[bool] = True
    collect_external: Optional[bool] = True
    use_dynamic_js: Optional[bool] = True
    enable_spa_mode: Optional[bool] = False
    delay: Optional[float] = 0.0
    subdomains: Optional[bool] = False
    verify_ssl: Optional[bool] = False
    respect_robots: Optional[bool] = None   # None = on in Spider mode, off in List mode (Screaming Frog's default)
    screenshots: Optional[bool] = True
    force: bool = False
    # Crawl settings (Screaming Frog's Speed / User-Agent / Include / Exclude / Spider configuration)
    max_urls_per_second: Optional[float] = 0.0          # 0 = unlimited
    user_agent: Optional[str] = "chrome"                 # one of UA_PRESETS
    custom_user_agent: Optional[str] = None              # used when user_agent == "custom"
    include_patterns: Optional[Union[List[str], str]] = None   # regexes; a string is split per line
    exclude_regex: Optional[Union[List[str], str]] = None
    crawl_resources: Optional[bool] = True
    check_external_links: Optional[bool] = True
    # Per-page limits (0 = no limit); defaults come from CrawlConfig
    max_page_size_kb: Optional[int] = None
    max_links_per_page: Optional[int] = None
    # After the crawl: fetch a sample of pages as search / AI crawlers (background, about a minute)
    bot_check: Optional[bool] = True
    # Crawl scope (top bar): subdomain | subfolder | all_subdomains | exact_url (None: from `subdomains`)
    crawl_scope: Optional[str] = None
    # Mode > List: crawl exactly these URLs (a list, or one per line); links are recorded, not followed
    mode: Optional[str] = "spider"
    urls: Optional[Union[List[str], str]] = None
    list_max_depth: Optional[int] = 0
    # Configuration > URL Rewriting ({"remove_params", "remove_all_params", "regex_replace", "lowercase",
    # "sort_params", "percent_encoding"}; missing keys keep today's normalisation)
    url_rewrite: Optional[dict] = None
    # Configuration > HTTP Header: [{"name", "value"}] sent to the crawled site only (values are secrets)
    request_headers: Optional[List[dict]] = None
    # Configuration > Custom > Search / Extraction / JavaScript (at most MAX_CUSTOM_RULES each)
    custom_searches: Optional[List[dict]] = None
    custom_extractors: Optional[List[dict]] = None
    custom_js: Optional[List[dict]] = None
    render_all: Optional[bool] = False       # render every internal HTML page in Chrome (JavaScript tab etc.)
    store_source: Optional[bool] = True      # keep each page's HTML for View Source
    analysis: Optional[dict] = None          # Crawl Analysis steps (None: the server's Crawl Analysis defaults)
    loop_guard: Optional[bool] = True
    # Advanced spider settings (None: CrawlConfig's defaults)
    max_retries: Optional[int] = None
    adaptive_throttle: Optional[bool] = None
    skip_crawl_traps: Optional[bool] = None
    max_url_length: Optional[int] = None
    robots_user_agent: Optional[str] = None
    exclude_patterns: Optional[Union[List[str], str]] = None   # the built-in skip list (plain text parts)
    bot_check_sample: Optional[int] = None
    # Name and project of the saved crawl (File > Crawls), when crawls are saved automatically
    name: Optional[str] = None
    project: Optional[str] = None
    # Find subdomains: other hosts of the site's domain to crawl too (Spider mode; each host's home page is an
    # extra start URL and in scope). Bare names ("blog.example.com", "host:port") or home URLs (a discovery
    # entry's "url" keeps an http-only host on http); a string is split on spaces, commas and lines
    extra_hosts: Optional[Union[List[str], str]] = None

def get_export_dir() -> str:
    """Return a writable directory across both local and serverless/Vercel environments."""
    return tempfile.gettempdir()

def _header_field(value: Optional[str], what: str) -> Optional[str]:
    """A pasted header name / value without surrounding spaces or line breaks. A line break inside is refused:
    the HTTP library would reject the header with an error message that quotes the secret value."""
    if value is None:
        return None
    value = value.strip()
    if any(ch in value for ch in "\r\n"):
        raise HTTPException(status_code=400, detail=f"The header {what} has a line break in it. "
                                                    "Paste it again on a single line.")
    return value

_CTRL_CHARS = re.compile(r"[\x00-\x1f\x7f]")
MAX_SIGNIN_UA_CHARS = 512

def _signin_user_agent(auth: Optional[AuthRequest], force: bool = False) -> str:
    """auth.user_agent checked ('' when none or when the sign-in settings are off): one line of at most
    MAX_SIGNIN_UA_CHARS characters, else a 400."""
    if not auth or not (auth.enabled or force):
        return ""
    ua = (auth.user_agent or "").strip()
    if not ua:
        return ""
    if len(ua) > MAX_SIGNIN_UA_CHARS or _CTRL_CHARS.search(ua):
        _bad(f"The User-Agent in the sign-in settings must be one line of at most {MAX_SIGNIN_UA_CHARS} "
             "characters. Paste it again, or keep the default User-Agent.")
    return ua

def _paste_text(value: Any) -> str:
    """A cookie paste as text: a JSON export that arrived as a list or an object is written back as JSON."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value)

def _auth_config(auth: Optional[AuthRequest], force: bool = False) -> AuthConfig:
    """AuthConfig from the request's "auth" object (force: use it even when "enabled" is false). The chosen
    browser User-Agent rides along as AuthConfig.user_agent (read by crawler.auth.auth_user_agent)."""
    if not auth or not (auth.enabled or force):
        return AuthConfig()
    config = AuthConfig(
        enabled=True,
        auth_type=auth.auth_type,
        login_url=auth.login_url,
        username=auth.username,
        password=auth.password,
        cookies=auth.cookies if isinstance(auth.cookies, (str, type(None))) else json.dumps(auth.cookies),
        custom_header_name=_header_field(auth.custom_header_name, "name"),
        custom_header_value=_header_field(auth.custom_header_value, "value"),
        live_session_id=auth.live_session_id,
        username_selector=auth.username_selector,
        password_selector=auth.password_selector,
        submit_selector=auth.submit_selector,
    )
    # set after construction: an AuthConfig without the field still carries it (auth_user_agent uses getattr)
    config.user_agent = _signin_user_agent(auth, force=True) or None
    return config

def _build_auth(req: CrawlRequest) -> AuthConfig:
    return _auth_config(req.auth)

def _patterns(raw) -> List[str]:
    """Regex list from the request (a list, or one pattern per line); invalid ones are a 400."""
    if not raw:
        return []
    items = raw.splitlines() if isinstance(raw, str) else list(raw)
    out = [str(p).strip() for p in items if str(p or "").strip()]
    for p in out:
        try:
            re.compile(p)
        except re.error as e:
            raise HTTPException(status_code=400, detail=f"Invalid regular expression '{p}': {e}")
    return out

def _ua_presets() -> Tuple[str, ...]:
    """Every User-Agent preset a request may name (crawler/config.py USER_AGENTS plus "custom")."""
    return tuple(dict.fromkeys(list(USER_AGENTS) + [p for p in UA_PRESETS if p != "custom"] + ["custom"]))

def _user_agent(req: CrawlRequest) -> Tuple[str, str]:
    """(preset, User-Agent string) for the request; unknown presets fall back to Chrome."""
    preset = (req.user_agent or "chrome").strip().lower()
    if preset not in _ua_presets():
        preset = "chrome"
    if preset == "custom":
        custom = (req.custom_user_agent or "").strip()
        if not custom:
            raise HTTPException(status_code=400, detail="Enter a custom User-Agent string or pick a preset.")
        return preset, custom
    return preset, USER_AGENTS.get(preset, DEFAULT_USER_AGENT)

def _is_serverless() -> bool:
    return bool(os.environ.get("VERCEL") or os.environ.get("AWS_LAMBDA_FUNCTION_NAME"))

def _bad(detail: str, status: int = 400):
    raise HTTPException(status_code=status, detail=detail)

def _list_cap() -> int:
    """URLs one List mode crawl may hold: CRAWLER_MAX_LIST_URLS (10,000), lowered to the server's URL cap."""
    return min(MAX_LIST_URLS, SERVER_MAX_URLS) if SERVER_MAX_URLS else MAX_LIST_URLS

def _lines(raw: Any) -> List[str]:
    """A list of strings from a list or one item per line, blanks dropped."""
    if raw is None:
        return []
    items = raw.splitlines() if isinstance(raw, str) else list(raw)
    return [str(x).strip() for x in items if str(x or "").strip()]

def _as_web_address(value: str) -> str:
    """An http(s) address from a list line ('' when the line is not one): "example.com/x" gets https://."""
    v = (value or "").strip().strip('"').strip("'").strip()
    if not v or any(c.isspace() for c in v) or len(v) > 4000:
        return ""
    if re.match(r"^https?://", v, re.I):
        return v if urlparse(v).netloc else ""
    if re.match(r"^[a-z][a-z0-9+.-]*:", v, re.I):     # mailto:, ftp:, javascript: ...
        return ""
    host = v.split("/")[0]
    return "https://" + v if re.match(r"^[A-Za-z0-9-]+(\.[A-Za-z0-9-]+)+(:\d+)?$", host) else ""

def _bool_map(given: Any, defaults: Dict[str, Any], what: str) -> Dict[str, Any]:
    """`defaults` with the known keys of `given` (true / false only) applied; other keys are refused."""
    out = dict(defaults)
    if given is None:
        return out
    if not isinstance(given, dict):
        _bad(f"{what} must be an object of on / off switches.")
    for k, v in given.items():
        if k not in defaults:
            _bad(f"Unknown {what} switch: {k}")
        if isinstance(defaults[k], bool):
            if not isinstance(v, bool):
                _bad(f"{what} '{k}' must be true or false.")
            out[k] = v
        else:
            out[k] = v
    return out

def _analysis_config(req: CrawlRequest) -> Dict[str, Any]:
    """CrawlConfig.analysis: the server's Crawl Analysis defaults (Settings), overridden by the request."""
    base = dict(default_analysis())
    try:
        from crawler import settings_store
        base.update({k: v for k, v in settings_store.analysis_defaults(REPORTS_DIR).items() if k in base})
    except Exception:
        pass
    return _bool_map(req.analysis, base, "Crawl Analysis")

def _url_rewrite_config(given: Optional[dict]) -> Dict[str, Any]:
    """CrawlConfig.url_rewrite from the request (missing keys keep the defaults); invalid rules are a 400."""
    rw = default_url_rewrite()
    if given is None:
        return rw
    if not isinstance(given, dict):
        _bad("URL Rewriting must be an object.")
    unknown = [k for k in given if k not in rw and k != "strip_tracking_params"]
    if unknown:
        _bad(f"Unknown URL Rewriting option: {', '.join(map(str, unknown))}")
    rw.update({k: v for k, v in given.items() if k in rw})
    if given.get("strip_tracking_params") is False:   # the dialog's "remove tracking parameters" switch
        rw["remove_params"] = []
    try:
        from crawler.url_utils import rules_from_config
        _rules, problems = rules_from_config(rw)
    except Exception as e:
        _bad(f"URL Rewriting could not be read: {type(e).__name__}: {str(e)[:160]}")
    if problems:
        _bad("URL Rewriting: " + "; ".join(map(str, problems[:5])))
    return rw

def _request_headers(given: Optional[List[dict]]) -> List[Dict[str, str]]:
    """Configuration > HTTP Header rows, checked (names are HTTP tokens, no line breaks); never echoed."""
    if not given:
        return []
    if not isinstance(given, list) or len(given) > 50:
        _bad("HTTP headers: send at most 50 rows of {name, value}.")
    rows = [{"name": str(h.get("name") or "").strip(), "value": str(h.get("value") or "")}
            for h in given if isinstance(h, dict) and str(h.get("name") or "").strip()]
    try:
        from crawler.engine import clean_request_headers
        _clean, problems = clean_request_headers(rows)
    except ImportError:
        problems = []
    if problems:
        _bad("HTTP headers: " + "; ".join(map(str, problems[:5])))
    return rows

def _custom_rules(req: CrawlRequest) -> Tuple[List[dict], List[dict], List[dict]]:
    """Custom Search / Extraction / JavaScript rules, compiled once here so a bad rule is a 400, not a crawl
    that silently finds nothing."""
    searches, extractors, snippets = list(req.custom_searches or []), list(req.custom_extractors or []), \
        list(req.custom_js or [])
    for name, items in (("Custom Search", searches), ("Custom Extraction", extractors), ("Custom JavaScript", snippets)):
        if len(items) > MAX_CUSTOM_RULES:
            _bad(f"{name}: at most {MAX_CUSTOM_RULES} rules.")
        if any(not isinstance(x, dict) for x in items):
            _bad(f"{name}: every rule must be an object.")
    try:
        from crawler.seo.extract import compile_custom_extractors, compile_custom_searches
        problems = (compile_custom_searches(searches)[1] if searches else []) + \
            (compile_custom_extractors(extractors)[1] if extractors else [])
    except ImportError:
        problems = []
    if problems:
        _bad("; ".join(str(p) for p in problems[:6]))
    names = set()
    for i, s in enumerate(snippets, start=1):
        code = str(s.get("code") or "")
        kind = str(s.get("type") or "extraction").strip().lower()
        label = str(s.get("name") or "").strip() or f"Snippet {i}"
        if not code.strip():
            _bad(f"Custom JavaScript '{label}': the snippet is empty.")
        if len(code) > MAX_CUSTOM_JS_CHARS:
            _bad(f"Custom JavaScript '{label}': the snippet is longer than {MAX_CUSTOM_JS_CHARS:,} characters.")
        if kind not in ("extraction", "action"):
            _bad(f"Custom JavaScript '{label}': the type must be extraction or action.")
        if label.lower() in names:
            _bad(f"Custom JavaScript: two snippets are named '{label}'.")
        names.add(label.lower())
        snippets[i - 1] = {"name": label[:80], "code": code, "type": kind}
    return searches, extractors, snippets

MAX_EXTRA_HOSTS = 200          # subdomains one crawl may add (crawler/subdomains.py MAX_HOSTS)
MAX_HOST_CHARS = 300           # one extra host entry (a name, or a home URL with its port)

def _host_values(raw: Any, what: str) -> List[str]:
    """Host names / home URLs from a list or a text (split on spaces, commas, semicolons and lines), blanks
    and repeats dropped; more than MAX_EXTRA_HOSTS, or an entry longer than MAX_HOST_CHARS, is a 400."""
    if raw is None:
        return []
    if isinstance(raw, str):
        items = re.split(r"[\s,;]+", raw)
    elif isinstance(raw, list):
        items = raw
    else:
        _bad(f"{what} must be a list of host names.")
    out = list(dict.fromkeys(str(x).strip() for x in items if x is not None and str(x).strip()))
    if len(out) > MAX_EXTRA_HOSTS:
        _bad(f"{what}: at most {MAX_EXTRA_HOSTS} hosts at a time ({len(out):,} were sent).")
    if any(len(x) > MAX_HOST_CHARS for x in out):
        _bad(f"{what}: an entry is longer than {MAX_HOST_CHARS} characters; enter host names such as "
             "blog.example.com.")
    return out

def _extra_hosts(req: CrawlRequest, mode: str, notes: List[str]) -> List[str]:
    """CrawlConfig.extra_hosts from the request. The engine keeps only hosts of the site's own domain that
    resolve to a public address (extra_hosts_skipped says why the others were left out); List mode crawls
    exactly its list, so they are ignored there (with a note)."""
    hosts = _host_values(req.extra_hosts, "Subdomains to crawl")
    if hosts and mode == "list":
        notes.append("The subdomains chosen under Find subdomains are not crawled in List mode: it crawls exactly "
                     "the listed URLs.")
        return []
    return hosts

def _browser_enabled() -> bool:
    """Settings > Embedded Browser (server-wide): off = no Chrome for crawls (rendering, screenshots)."""
    try:
        from crawler import settings_store
        return bool(settings_store.effective("browser", REPORTS_DIR).get("enabled", True))
    except Exception:
        return True

def _proxy_config() -> Optional[Dict[str, Any]]:
    try:
        from crawler import settings_store
        return settings_store.proxy_config(REPORTS_DIR)
    except Exception:
        return None

def _extra_config(req: CrawlRequest) -> Dict[str, Any]:
    """Advanced spider settings the request sets (others keep CrawlConfig's defaults), range-checked."""
    out: Dict[str, Any] = {}
    for name, lo, hi in (("max_retries", 0, 10), ("max_url_length", 200, 100_000), ("bot_check_sample", 1, 50)):
        v = getattr(req, name, None)
        if v is not None:
            if not isinstance(v, int) or not lo <= v <= hi:
                _bad(f"{name.replace('_', ' ').capitalize()} must be a whole number from {lo} to {hi:,}.")
            out[name] = v
    for name in ("adaptive_throttle", "skip_crawl_traps"):
        v = getattr(req, name, None)
        if v is not None:
            out[name] = bool(v)
    if req.robots_user_agent is not None:
        token = req.robots_user_agent.strip()
        if not re.fullmatch(r"[A-Za-z0-9_.*-]{1,100}", token):
            _bad("The robots.txt user-agent must be one word such as Googlebot.")
        out["robots_user_agent"] = token
    if req.exclude_patterns is not None:
        pats = _lines(req.exclude_patterns)
        if len(pats) > 200 or any(len(p) > 500 for p in pats):
            _bad("The skip list holds at most 200 entries of 500 characters.")
        out["exclude_patterns"] = pats
    return out

BUILD4_OPTIONS = ("max_redirects", "custom_robots_txt", "sitemap_urls", "max_sitemap_files", "max_sitemap_only_urls",
                  "max_sitemap_index_urls", "cdn_hosts", "near_dup_threshold", "content_area", "preferences",
                  "soft_404_phrases", "non_descriptive_anchors", "link_positions", "crawl_types", "follow",
                  "follow_internal_nofollow", "list_follow_canonicals", "cookie_storage", "assume_html",
                  "ignore_non_indexable_issues", "limits", "block_addresses", "carbon", "spelling", "a11y_audit",
                  "a11y_audit_max_pages", "a11y_audit_max_minutes", "render_device", "render_timeout",
                  "flatten_shadow_dom", "flatten_iframes")


def _build_config(req: CrawlRequest, raw_url: str, export_path: str, serverless: bool = False,
                  notes: Optional[List[str]] = None) -> CrawlConfig:
    """CrawlConfig from a crawl request, every option checked (a bad one is a 400 before anything stops or
    starts). Serverless runs are capped so they finish inside 60 s. `notes` collects the options this server
    changed (Embedded Browser off, serverless limits)."""
    notes = notes if notes is not None else []
    preset, ua = _user_agent(req)
    signin_ua = _signin_user_agent(req.auth)
    if signin_ua:
        # the browser the pasted session belongs to: the sign-in and every request of the crawl send it
        if signin_ua != ua:
            notes.append("The crawl sends the browser User-Agent kept with the sign-in settings (the browser the "
                         "cookies came from), instead of the User-Agent chosen under Configuration.")
        preset, ua = "custom", signin_ua
    mode = (req.mode or "spider").strip().lower()
    if mode not in CRAWL_MODES:
        _bad(f"Unknown mode. Use one of: {', '.join(CRAWL_MODES)}")
    scope = (req.crawl_scope or ("all_subdomains" if req.subdomains else "subdomain")).strip().lower()
    if scope not in CRAWL_SCOPES:
        _bad(f"Unknown crawl scope. Use one of: {', '.join(CRAWL_SCOPES)}")
    seeds: List[str] = []
    if mode == "list":
        seeds = _lines(req.urls)
        if not seeds:
            _bad("List mode needs at least one URL (upload or paste the list).")
        if len(seeds) > _list_cap():
            _bad(f"The list has {len(seeds):,} lines; this server crawls at most {_list_cap():,} URLs per list.")
    browser_on = _browser_enabled()
    if serverless:
        max_urls = min(req.max_urls, 100) if req.max_urls else 100
        concurrency = min(req.concurrency or 6, 8)
        timeout = min(req.timeout or 15, 12)
        dynamic_js = False
    else:
        max_urls = req.max_urls if req.max_urls is not None else 0  # 0 = unlimited
        if SERVER_MAX_URLS:   # this server's cap (CRAWLER_MAX_URLS): "no limit" means the cap
            max_urls = min(max_urls, SERVER_MAX_URLS) if max_urls else SERVER_MAX_URLS
        concurrency = req.concurrency or 6
        timeout = req.timeout or 15
        dynamic_js = True if req.use_dynamic_js is None else req.use_dynamic_js
    if dynamic_js and not browser_on:
        dynamic_js = False
        notes.append("JavaScript link discovery is off: the embedded browser is switched off in Settings.")
    render_all = bool(req.render_all)
    if render_all and (serverless or not browser_on):
        render_all = False
        notes.append("Rendering every page in Chrome is off: " + ("serverless servers have no Chrome."
                     if serverless else "the embedded browser is switched off in Settings."))
    elif render_all and not _chrome_status().get("available"):
        _bad("Rendering every page needs Chrome, which is not available on this server: "
             + (_chrome_status().get("error") or "Chrome was not found."))
    searches, extractors, snippets = _custom_rules(req)
    if snippets and not render_all:
        notes.append("Custom JavaScript runs only when every page is rendered in Chrome; its tab shows "
                     "Not checked for this crawl.")
    respect = req.respect_robots if req.respect_robots is not None else mode != "list"
    list_depth = max(0, min(int(req.list_max_depth or 0), 10))
    store_source = req.store_source is not False and not serverless
    # Build 4 options (10.7): plain 400 messages, then set on the config (normalize_options cleans them)
    b4 = {k: getattr(req, k) for k in BUILD4_OPTIONS if getattr(req, k, None) is not None}
    for k in ("sitemap_urls", "cdn_hosts", "soft_404_phrases", "non_descriptive_anchors"):
        if isinstance(b4.get(k), str):
            b4[k] = _lines(b4[k])
    from crawler.config import normalize_options, validate_options
    b4_errors = dict(validate_options(dict(b4, custom_searches=req.custom_searches or [],
                                           custom_js=req.custom_js or [])))
    if b4_errors:
        _bad("; ".join(f"{k.replace('_', ' ')}: {v}" for k, v in list(b4_errors.items())[:4]))
    config = CrawlConfig(
        start_url=raw_url,
        max_urls=max(0, int(max_urls or 0)),
        max_depth=max(0, int(req.max_depth or 0)),
        concurrency=max(1, min(int(concurrency), MAX_CONCURRENCY)),
        max_urls_per_second=max(0.0, float(req.max_urls_per_second or 0.0)),
        timeout=timeout,
        use_sitemaps=(True if req.use_sitemaps is None else req.use_sitemaps) and mode != "list",
        collect_external_links=True if req.collect_external is None else req.collect_external,
        use_dynamic_js=dynamic_js,
        enable_spa_mode=bool(req.enable_spa_mode),
        delay=req.delay or 0.0,
        include_subdomains=scope == "all_subdomains",
        verify_ssl=bool(req.verify_ssl),
        respect_robots=bool(respect),
        include_patterns=_patterns(req.include_patterns),
        exclude_regex=_patterns(req.exclude_regex),
        crawl_resources=req.crawl_resources is not False,
        check_external_links=req.check_external_links is not False,
        user_agent_preset=preset,
        user_agent=ua,
        auth=_build_auth(req),
        output_file=export_path,
        # This app runs the post-crawl checks (_quick_postcrawl, background bot check) and saves the snapshot
        # (_save_snapshot) itself, so the engine must not do them again inside crawler.run()
        post_crawl_checks=False,
        save_snapshot=False,
        **_page_limits(req),
        **_extra_config(req),
    )
    for k, v in b4.items():
        setattr(config, k, v)
    if b4:
        normalize_options(config)
    # the newer options are set after construction, so an older CrawlConfig without them still starts
    newer = {"crawl_scope": scope, "mode": mode, "seed_urls": seeds, "list_max_depth": list_depth,
             "url_rewrite": _url_rewrite_config(req.url_rewrite), "request_headers": _request_headers(req.request_headers),
             "custom_searches": searches, "custom_extractors": extractors, "custom_js": snippets,
             "render_all": render_all, "store_source": store_source, "analysis": _analysis_config(req),
             "loop_guard": req.loop_guard is not False, "extra_hosts": _extra_hosts(req, mode, notes)}
    proxy = _proxy_config()
    if proxy is not None:
        newer["proxy"] = proxy
    for k, v in newer.items():
        setattr(config, k, v)
    try:
        from crawler.config import sync_scope
        sync_scope(config)
    except ImportError:
        pass
    return config

def _config_secrets(config: CrawlConfig) -> List[str]:
    """Values of a crawl's configuration that must never be logged or returned: the login's secrets, every
    custom HTTP header value, the proxy password and a password written in the start URL (moved into a Basic
    authentication header by url_utils.move_url_credentials)."""
    from crawler.url_utils import move_url_credentials, split_url_credentials
    found = move_url_credentials(config)
    for u in [config.start_url] + list(getattr(config, "seed_urls", None) or [])[:1]:
        pw = split_url_credentials(str(u or ""))[2]
        if len(pw) >= 4:
            found.append(pw)
    found += _auth_secrets(config.auth, config.start_url)
    for h in getattr(config, "request_headers", None) or []:
        v = str(h.get("value") or "").strip()
        if len(v) >= 4:
            found.append(v)
            found += [p for p in v.split() if len(p) >= 8]          # the token of "Bearer <token>"
    pw = str((getattr(config, "proxy", None) or {}).get("password") or "")
    if len(pw) >= 4:
        found.append(pw)
    return sorted(set(found), key=len, reverse=True)

def _page_limits(req: CrawlRequest) -> Dict[str, int]:
    """max_page_size_kb / max_links_per_page from the request, only when given (else CrawlConfig's defaults)."""
    out = {}
    for name in ("max_page_size_kb", "max_links_per_page"):
        value = getattr(req, name, None)
        if value is not None:
            out[name] = max(0, int(value))
    return out

def _normalize_target(raw_url: str) -> str:
    raw_url = raw_url.strip()
    if not raw_url.startswith("http://") and not raw_url.startswith("https://"):
        raw_url = "https://" + raw_url
    return raw_url

def _base_name(raw_url: str) -> str:
    from crawler.url_utils import split_url_credentials
    domain = get_domain(split_url_credentials(raw_url)[0]) or "website"   # never the user name or password
    return f"crawl_{domain}_{time.strftime('%Y%m%d_%H%M%S')}"

def cleanup_old_reports(keep: Optional[str] = None):
    """Deletes reports and screenshot folders this app generated more than
    REPORT_RETENTION_HOURS ago. Nothing else in the temp directory is touched."""
    import shutil
    cutoff = time.time() - REPORT_RETENTION_HOURS * 3600
    root = get_export_dir()
    try:
        names = os.listdir(root)
    except OSError:
        return
    for name in names:
        if not GENERATED_RE.match(name) or (keep and name.startswith(keep)):
            continue
        path = os.path.join(root, name)
        try:
            if os.path.getmtime(path) >= cutoff:
                continue
            if os.path.isdir(path):
                shutil.rmtree(path, ignore_errors=True)
            else:
                os.remove(path)
        except OSError:
            pass

def _issues_for_evidence(crawler: WebsiteCrawler) -> list:
    issues = [r for r in crawler.results_list if r.category in ISSUE_CATEGORIES]
    return sorted(issues, key=issue_sort_key)[:MAX_SCREENSHOT_ISSUES]

def _site_index(job: Optional[CrawlJob], force: bool = False):
    """(Site, TabIndex) of the job's crawl, or (None, None) when there is no crawl or the SEO
    analysis failed. Never call this from /api/crawl/status: it can take seconds on big sites."""
    if not job or not job.crawler or not hasattr(job.crawler, "site_index"):
        return None, None
    try:
        return job.crawler.site_index(force=force)
    except Exception as e:
        msg = f"SEO analysis failed: {e}"
        if msg != job.index_error:  # log each distinct failure once, not on every poll
            job.index_error = msg
            job.crawler.log(msg)
        return None, None

def _issues_with_examples(index, issues: Optional[List[dict]] = None,
                          dismissed_urls: Optional[Dict[str, List[str]]] = None) -> Optional[List[dict]]:
    """Issues overview rows (default index.issues) plus a few example URLs each (for the PDF); URLs
    dismissed by the team are not used as examples."""
    if index is None:
        return None
    out = []
    for issue in (issues if issues is not None else index.issues):
        tr = index.tabs.get(issue.get("tab"))
        hits = tr.filter_rows.get(issue.get("filter"), []) if tr is not None else []
        hidden = set((dismissed_urls or {}).get(f"{issue.get('tab')}:{issue.get('filter')}") or ())
        examples = [getattr(r, "url", "") for r in hits if getattr(r, "url", "") not in hidden]
        out.append(dict(issue, examples=examples[:PDF_ISSUE_EXAMPLES]))
    return out

def _memo(job: CrawlJob, key: tuple, build: Callable[[], Any]) -> Any:
    """build() once per key for this crawl (keys carry the index identity, so a rebuilt index misses)."""
    with job.memo_lock:
        if key in job.memo:
            return job.memo[key]
    value = build()
    with job.memo_lock:
        if len(job.memo) > 32:
            job.memo.clear()
        job.memo[key] = value
    return value

def _safe(build: Callable[[], Any], default: Any, job: Optional[CrawlJob] = None, what: str = "") -> Any:
    """build(), or `default` when it fails (logged once per call): one broken add-on never breaks a report."""
    try:
        return build()
    except Exception as e:
        if job is not None and job.crawler is not None and what:
            job.crawler.log(f"{what} skipped: {type(e).__name__}: {str(e)[:160]}")
        return default

def _dismiss_store(job: Optional[CrawlJob], site_url: str = ""):
    """The dismissed-issues store of `site_url` (default: the current crawl's site), or None."""
    from crawler.dismiss import store_for
    url = (site_url or "").strip() or (job.crawler.start_url if job and job.crawler else "")
    return store_for(REPORTS_DIR, url) if url else None

def _empty_dismissals(index) -> dict:
    return {"issues": list(index.issues) if index is not None else [], "dismissed": [], "dismissed_urls": {},
            "fully_dismissed_issues": [], "entries": [], "version": "",
            "summary": {"dismissals": 0, "issues_hidden": 0, "urls_hidden": 0}}

def _dismissals(job: Optional[CrawlJob], index) -> dict:
    """Issues overview with dismissed URLs taken out of the counts, plus every dismissal and how many URLs it
    covered ("dismissed"), "dismissed_urls" {"tab:filter": [urls]}, the stored "entries" and a "version"."""
    if job is None or index is None:
        return _empty_dismissals(index)
    store = _safe(lambda: _dismiss_store(job), None)
    entries = store.list() if store is not None else []
    if not entries:
        return _empty_dismissals(index)
    version = hashlib.md5(json.dumps(entries, sort_keys=True).encode("utf-8")).hexdigest()[:12]

    def build():
        from crawler.dismiss import apply_dismissals
        out = apply_dismissals(index, entries)
        out.update(entries=entries, version=version)
        return out
    return _memo(job, ("dismiss", id(index), version), build)

def _validation_of(job: Optional[CrawlJob]) -> Optional[dict]:
    return job.validation if job is not None else None

def _merge_validation(job: CrawlJob, out: dict):
    """Adds one validation run to the crawl's merged validation (the latest verdict per URL wins)."""
    old = job.validation or {}
    # a URL the newest run skipped (e.g. a sign-out link on a logged-in crawl) loses its older verdict: the
    # latest outcome wins, and "Not checked" is an outcome
    newly_skipped = {s.get("url") for s in (out.get("skipped") or []) if s.get("url")}
    results = {u: e for u, e in (old.get("results") or {}).items() if u not in newly_skipped}
    results.update(out.get("results") or {})
    skipped = {s.get("url"): s for s in (old.get("skipped") or []) + (out.get("skipped") or [])
               if s.get("url") and s.get("url") not in results}
    guard_ran = out.get("scope") != "urls"
    summary = {k: 0 for k in VERDICT_KEYS}
    for e in results.values():
        summary[e.get("verdict")] = summary.get(e.get("verdict"), 0) + 1
    summary["link_removed"] = sum(1 for e in results.values() if e.get("link_on_parent") is False)
    summary["not_checked"] = len(skipped)
    job.validation = {
        "state": "done", "done": len(results) + len(skipped), "total": len(results) + len(skipped),
        "scope": out.get("scope") or "", "summary": summary, "results": results, "skipped": list(skipped.values()),
        "login_guard": (out.get("login_guard") or []) if guard_ran else (old.get("login_guard") or []),
        "login_guard_summary": (out.get("login_guard_summary") or {}) if guard_ran
        else (old.get("login_guard_summary") or {}),
        "started_at": old.get("started_at") or out.get("started_at") or "",
        "finished_at": out.get("finished_at") or "",
    }
    job.validation_version += 1

def _intel_errors(job: Optional[CrawlJob], site, index) -> dict:
    """GET /api/intel/errors: the AI-enriched result of the last intel job for this index, else the instant
    rule-based grouping (with validation verdicts and without dismissed URLs)."""
    from crawler import intel
    if job is None or index is None:
        return intel.empty_errors()
    d = _dismissals(job, index)
    runner = job.runners["intel"]
    if (runner.state == "done" and isinstance(runner.result, dict) and runner.meta.get("index") is index
            and runner.meta.get("dismiss_version") == d["version"]):
        return runner.result["errors"]
    validation = _validation_of(job)
    # the AI job ran on an older index (rebuilt since, e.g. when the bot check finished): keep its AI text on
    # the groups that did not change instead of falling back to "AI off"
    earlier = runner.result.get("errors") if runner.state == "done" and isinstance(runner.result, dict) else None

    def build():
        from crawler.intel import root_cause
        out = root_cause.analyse(site, index, job.crawler, validation=validation,
                                 issues=d["issues"] if d["entries"] else None,
                                 dismissed=d["dismissed_urls"] or None)
        return root_cause.carry_ai_text(out, earlier)
    return _memo(job, ("intel_errors", id(index), job.validation_version, d["version"], id(earlier)), build)

def _intel_content(job: Optional[CrawlJob], site, index) -> dict:
    """GET /api/intel/content: the AI-reviewed result of the last intel job for this index, else rule-based."""
    from crawler import intel
    if job is None or index is None:
        return intel.empty_content()
    runner = job.runners["intel"]
    if runner.state == "done" and isinstance(runner.result, dict) and (
            runner.meta.get("index") is index
            or runner.meta.get("results_version", -1) == job.results_version):   # same pages, index rebuilt
        return runner.result["content"]
    from crawler.intel import content
    return content.analyse(site)

def _checks(job: Optional[CrawlJob], site, index) -> List[dict]:
    """Passed / Failed / Not checked per check. A check the team dismissed completely keeps its status and carries
    "dismissed" (counted apart from failed, see checks.summary); a partly dismissed one counts only the URLs left
    and says how many were dismissed ("dismissed_urls")."""
    if index is None:
        return []
    from crawler.intel.checks import check_statuses
    out = check_statuses(index, job.crawler if job else None, site)
    d = _dismissals(job, index)
    gone = {f"{i.get('tab')}:{i.get('filter')}" for i in d["fully_dismissed_issues"]}
    for c in out:
        key = f"{c.get('tab')}:{c.get('filter')}"
        hidden = len(d["dismissed_urls"].get(key) or ())
        if key in gone:
            c["dismissed"] = True
            c["dismissed_urls"] = hidden or int(c.get("urls") or 0)
        elif hidden and c.get("status") == "failed":
            c["urls"] = max(0, int(c.get("urls") or 0) - hidden)
            c["dismissed_urls"] = hidden
            c["reason"] = ((c.get("reason") or "") + f" ({hidden:,} dismissed URL{'s' if hidden != 1 else ''} "
                           "not counted.)").strip()
    return out

def _report_extras(job: CrawlJob, site, index) -> dict:
    """Follow-up data for the PDF and the full workbook: rule-based results when a job has not run yet."""
    if index is None:
        return {"issues": None, "dismissed": [], "dismissed_urls": {}, "checks": [],
                "validation": _validation_of(job), "intel": None, "content": None}
    d = _safe(lambda: _dismissals(job, index), _empty_dismissals(index), job, "Dismissed issues")
    return {
        "issues": d["issues"], "dismissed": d["dismissed"], "dismissed_urls": d["dismissed_urls"],
        "checks": _safe(lambda: _checks(job, site, index), [], job, "Checks"),
        "validation": _validation_of(job),
        "intel": _safe(lambda: _intel_errors(job, site, index), None, job, "Fix plan"),
        "content": _safe(lambda: _intel_content(job, site, index), None, job, "Content intelligence"),
    }

def _public_stats(job: CrawlJob) -> dict:
    """crawler.get_stats() with the login messages masked: a login error can quote a pasted header value."""
    stats = job.crawler.get_stats()
    for key in ("auth_status", "auth_error"):
        if stats.get(key):
            stats[key] = job.mask(stats[key])
    if stats.get("limitations"):
        stats["limitations"] = [job.mask(x) for x in stats["limitations"]]
    return stats

def _write_excel(job: CrawlJob, path: str, category: str = "all") -> str:
    """The Excel report for `category`; the full workbook also gets the follow-up sheets. Builds run one at a
    time per crawl (job.excel_lock). Returns the path written: on Windows a report that is still being downloaded
    cannot be replaced, so the new one then gets its own name."""
    with job.excel_lock:
        return _build_excel(job, path, category)

def _build_excel(job: CrawlJob, path: str, category: str) -> str:
    crawler = job.crawler
    site, index = _site_index(job)   # linking pages for every broken URL; tab sheets for "all"
    extras = _report_extras(job, site, index) if category == "all" else {}
    return ExcelExporter.export(
        results=list(crawler.results_list),
        filepath=path,
        stats=_public_stats(job),
        external_links=list(crawler.external_links_list),
        category=category,
        site=site,
        seo_index=index,
        issues=extras.get("issues"),
        checks=extras.get("checks"),
        validation=extras.get("validation"),
        fix_plan=extras.get("intel"),
        content=extras.get("content"),
        dismissed=extras.get("dismissed"),
        dismissed_urls=extras.get("dismissed_urls"),
    )

def _excel_file(job: CrawlJob, category: str = "all") -> str:
    """The Excel report to download. Builds wait for each other (job.excel_lock). Once the crawl is done, the full
    workbook saved after it is reused while nothing it shows changed (job.export_stale), so a second click, or a
    download that arrives during a rebuild, gets that file instead of building another one."""
    suffix = "" if category == "all" else f"_{category}"
    path = os.path.join(get_export_dir(), f"{job.base_name}{suffix}.xlsx")
    with job.excel_lock:
        final = category == "all" and job.phase == "done"
        if final and job.export_file and not job.export_stale and os.path.exists(job.export_file):
            return job.export_file
        if final:
            job.export_stale = False   # cleared first: a change during the build marks the new file stale again
        out = _write_excel(job, path, category)
        if final:
            job.export_file = out
        return out

def _pdf_key(job: CrawlJob, index) -> tuple:
    """What the PDF shows: a saved PDF built from other inputs is rebuilt on its next request."""
    pin, intel = job.runners["pinpoint"], job.runners["intel"]
    d = _safe(lambda: _dismissals(job, index), _empty_dismissals(index))
    integ = getattr(job.crawler, "integrations", None) if job.crawler is not None else None
    fetched = tuple(sorted((s, str((v or {}).get("fetched_at") or ""), str((v or {}).get("status") or ""))
                           for s, v in (integ or {}).items() if isinstance(v, dict))) if isinstance(integ, dict) else ()
    return (id(index), job.validation_version, intel.job_id if intel.state == "done" else "",
            pin.job_id if pin.state == "done" else "", d["version"], job.screenshot_note, len(job.evidence), fetched)

def _pdf_extras(job: CrawlJob, site, index, extras: dict) -> Dict[str, Any]:
    """The compact PDF's optional inputs (Checks tally, connected-data tables, per-page issue counts), passed
    only when this pdf_report accepts them."""
    import inspect
    from crawler.pdf_report import build_pdf_report
    params = inspect.signature(build_pdf_report).parameters
    out: Dict[str, Any] = {}
    if "checks" in params:
        out["checks"] = extras.get("checks") or None
    if "seo_index" in params:
        out["seo_index"] = index
    if "integrations" in params and site is not None and getattr(job.crawler, "integrations", None):
        out["integrations"] = _safe(lambda: _integrations().report_summary(site), None, job, "Connected data in the PDF")
    return out

def build_pdf(job: CrawlJob) -> str:
    from crawler.pdf_report import build_pdf_report
    path = os.path.join(get_export_dir(), f"{job.base_name}_issues.pdf")
    with job.pdf_lock:
        site, index = _site_index(job)
        key = _pdf_key(job, index)
        extras = _report_extras(job, site, index)
        partial = path + ".partial"
        build_pdf_report(
            results=list(job.crawler.results_list),
            filepath=partial,
            stats=_public_stats(job),
            evidence=job.evidence,
            evidence_dir=job.evidence_dir,
            screenshot_note=job.screenshot_note,
            seo_issues=_issues_with_examples(index, extras["issues"], extras["dismissed_urls"]),
            site=site,
            external_links=list(job.crawler.external_links_list),
            seo_evidence=job.seo_evidence or None,
            validation=extras["validation"],
            intel=extras["intel"],
            content=extras["content"],
            seo_evidence_dir=job.evidence_dir,
            dismissed_urls=extras["dismissed_urls"] or None,
            **_pdf_extras(job, site, index, extras),
        )
        try:
            os.replace(partial, path)
        except OSError:   # the old PDF is still open (being downloaded on Windows): use a new name
            path = os.path.join(get_export_dir(), f"{job.base_name}_issues_{int(time.time())}.pdf")
            os.replace(partial, path)
        job.pdf_file, job.pdf_key = path, key
    return path

def _save_snapshot(job: CrawlJob) -> str:
    """Saves (or replaces) the crawl's snapshot for GET /api/crawls and comparisons. Never raises."""
    crawler = job.crawler
    if crawler is None or crawler.status == "auth_failed" or not crawler.results_list:
        return ""
    site, index = _site_index(job)
    if index is None:
        return ""
    from crawler.seo import compare
    sid = _safe(lambda: compare.save_snapshot(crawler, index, REPORTS_DIR, site=site), "", job, "Crawl snapshot")
    if sid:
        job.snapshot_id = sid
    return sid

def _reports_changed(job: CrawlJob, resave: bool = True):
    """Follow-up data changed (validation, intel, bot check, re-check): the saved workbook is rebuilt on its
    next download, the PDF on its next request (see _pdf_key), and the snapshot is saved again. A re-check that
    changed an answer goes to its own "<id>-r" history entry (compare.save_snapshot), so the crawl's own
    entry keeps what the crawl found."""
    job.export_stale = True
    if resave:
        _save_snapshot(job)
        _autosave(job)

def _quick_postcrawl(job: CrawlJob, bot_check: bool):
    """robots.txt audit and the random-URL probe (a few requests) before the reports are built. The bot
    check runs afterwards in the background (_start_bot_check); its report says "running" until then."""
    from crawler.seo import postcrawl
    crawler = job.crawler
    analysis = getattr(crawler.config, "analysis", None) or {}
    # the bot check runs later in the background (_start_bot_check): not logged as "turned off" here
    opts: Dict[str, Any] = {"bot_check": "deferred" if bot_check else False}
    if analysis.get("post_crawl_checks") is False:   # Crawl Analysis > Configure: robots.txt audit off
        opts.update(robots_audit=False, trap_probe=False)
    _safe(lambda: postcrawl.run(crawler, opts), None, job, "Post-crawl checks")
    if not bot_check:
        if analysis.get("bot_check") is False and crawler.results_list:
            crawler.bot_report = postcrawl.status_report("skipped", "turned off for this crawl (Crawl Analysis > "
                                                                    "Configure)")
            crawler._index_key = None
        return
    if crawler.status == "stopped" or crawler._stop_event.is_set():
        crawler.bot_report = postcrawl.status_report("skipped", "the crawl was stopped, so the bot check did not start")
    elif crawler.results_list:
        crawler.bot_report = postcrawl.status_report("running", "starts when the reports are ready")
    crawler._index_key = None

def _start_bot_check(job: CrawlJob):
    rep = getattr(job.crawler, "bot_report", None)
    if not (isinstance(rep, dict) and rep.get("status") == "running"):
        return
    post_stop = getattr(job.crawler, "_post_stop", None)
    if job.retired or job.cancel_event.is_set() or (post_stop is not None and post_stop.is_set()):
        # replaced or stopped while the reports were built: no Chrome for a crawl nobody is looking at
        from crawler.seo import postcrawl
        job.crawler.bot_report = postcrawl.status_report("stopped", "the crawl was stopped or replaced before "
                                                                    "the bot check started")
        job.crawler._index_key = None
        return
    try:
        job.runners["postcrawl"].start(lambda r: _bot_check_work(job, r), "Checking pages as search and AI crawlers...")
    except JobBusy:
        pass

def _bot_check_work(job: CrawlJob, runner: JobRunner) -> dict:
    """Pages fetched as a browser and as search / AI crawlers (postcrawl's bot_check step, in the background)."""
    from crawler.seo import bot_check, postcrawl
    crawler = job.crawler
    started = time.time()
    sample = BOT_CHECK_SAMPLE if os.environ.get("CRAWLER_BOT_CHECK_SAMPLE") else \
        int(getattr(crawler.config, "bot_check_sample", 0) or BOT_CHECK_SAMPLE)
    try:
        report = bot_check.run(crawler, sample=sample, render=_browser_enabled(), stop_event=runner.stop_event,
                               on_progress=lambda msg: runner.update(None, msg))
    except Exception as e:   # a check must never break the crawl
        report = postcrawl.status_report("error", f"{type(e).__name__}: {str(e)[:160]}")
    if not isinstance(report, dict):
        report = postcrawl.status_report("error", "the check returned nothing")
    if report.get("status") == "running":
        report = postcrawl.status_report("stopped", "the bot check was stopped")
    report.setdefault("seconds", round(time.time() - started, 1))
    crawler.bot_report = report
    crawler._index_key = None   # the AI Crawlers tab reads the new report on the next index build
    crawler.log(f"[post-crawl] bot check: {report.get('status')}"
                + (f" ({report.get('reason')})" if report.get("reason") else ""))
    runner.update(None, f"Bot check {report.get('status')}.")
    _reports_changed(job)
    return {"status": report.get("status"), "pages_checked": report.get("pages_checked", 0)}

def capture_evidence(job: CrawlJob):
    """Screenshots of broken pages and of the parent page with the broken link highlighted."""
    crawler = job.crawler
    issues = _issues_for_evidence(crawler)
    if not issues:
        return
    if job.cancel_event.is_set():
        job.screenshot_note = "Screenshots skipped by user."
        return
    collector = EvidenceCollector(
        out_dir=job.evidence_dir,
        user_agent=crawler.config.user_agent,
        timeout=crawler.config.timeout,
        cookies=crawler.auth_cookies,
        logger=crawler.log,
        cancel_event=job.cancel_event,
        crawler=crawler,  # screenshots in the crawl's own logged-in browser
    )
    if not collector.start():
        job.screenshot_note = f"Screenshots unavailable: {collector.unavailable_reason} Link positions are shown as page sketches."
        crawler.log(f"[Screenshots] {job.screenshot_note}")
        return
    try:
        total_issues = sum(1 for r in crawler.results_list if r.category in ISSUE_CATEGORIES)
        crawler.log(f"[Screenshots] Capturing visual evidence for {len(issues)} issue(s)...")
        job.evidence_progress = {"done": 0, "total": 1}

        def progress(done, total):
            job.evidence_progress = {"done": done, "total": total}

        job.evidence = collector.capture(issues, on_progress=progress)
        if total_issues > len(issues):
            job.screenshot_note = f"Screenshots captured for the {len(issues)} most severe of {total_issues} issues."
        if collector.login_lost:
            job.screenshot_note = ("The screenshot browser could not use the login session, so pages behind "
                                   "the login were not captured. Check that the session cookies are complete.")
        if job.cancel_event.is_set():
            job.screenshot_note = "Screenshot capture was stopped early; some issues have no screenshot."
    finally:
        collector.close()

def _attach_render_pool(job: CrawlJob):
    """render_all: renders every internal 2xx HTML page in Chrome through the engine's render_hook. The pool
    starts on the first page (after the login, so it carries the session) and closes when the crawl ends
    (_close_render_pool); the rendered HTML is kept next to the raw HTML for View Source."""
    crawler = job.crawler
    lock = threading.Lock()

    def pool():
        with lock:
            p = job.render.get("pool")
            if p is None:
                from crawler.render_capture import RenderPool
                p = RenderPool.from_crawler(crawler, evidence_dir=job.evidence_dir)
                job.render["pool"] = p
                crawler.render_pool = p        # render_one() falls back to it when no other Chrome may start
                store = getattr(crawler, "source_store", None)
                if store is not None:
                    try:
                        from crawler.source_store import SourceStore
                        job.render["html"] = SourceStore(os.path.join(store.dir, "rendered"))
                    except Exception:
                        job.render["html"] = None
            return p

    def on_html(url: str, html: str):
        store = job.render.get("html")
        if store is not None and html:
            store.put(url, html)

    def hook(url: str, seo: dict) -> Optional[dict]:
        if job.retired or job.cancel_event.is_set():
            return None
        seen: Dict[str, str] = {}

        def keep_html(page: str, html: str):
            seen["html"] = html
            on_html(page, html)

        rd = pool().render(url, raw=seo, load_url=str((seo or {}).get("final_url") or url), on_html=keep_html)
        _custom_rules_on_render(crawler, rd, seen.get("html") or "")
        return rd

    crawler.render_hook = hook

def _custom_rules_on_render(crawler, rd: Optional[dict], html: str):
    """Custom Search / Custom Extraction on the rendered HTML (Screaming Frog's JavaScript rendering mode): the
    results go in the rendered dict and the engine uses them in place of the raw-HTML ones. Not for a render
    that ended on another page (a meta refresh or script redirect): those would be the destination's values."""
    opts = getattr(crawler, "extract_options", None)
    if not rd or not html or opts is None or not (opts.searches or opts.extractors):
        return
    from crawler.seo.model import rendered_left_page
    if rendered_left_page(rd):
        return
    try:
        from crawler.seo.extract import custom_rules_on_html
        rd["custom_search"], rd["custom_extraction"] = custom_rules_on_html(html, opts)
    except Exception:
        pass   # the raw-HTML values stay

def _browser_proxy(crawler, config: CrawlConfig):
    """Settings > Proxy also for the crawl's shared Chrome (login, JavaScript links, screenshots, bot check):
    SharedBrowser reads its proxy when Chrome starts."""
    proxy = getattr(config, "proxy", None)
    browser = getattr(crawler, "browser", None)
    if browser is None or not isinstance(proxy, dict) or not proxy.get("enabled"):
        return
    try:
        from crawler.browser import proxy_server_arg
        if proxy_server_arg(proxy):
            browser.proxy = proxy
    except Exception:
        pass

def _attach_audit_hook(job: CrawlJob):
    """The Accessibility audit (a11y_audit on, render_all off): crawler.audit_hook runs axe-core only, in its own
    Chrome pool that starts on the first page and closes with the crawl (_close_render_pool)."""
    from crawler.render_capture import AuditHook
    hook = AuditHook(job.crawler)
    job.render["audit"] = hook
    job.crawler.audit_hook = hook


def _close_render_pool(job: CrawlJob):
    """Quits the render Chromes once the crawl ends; their counts go to the log and the coverage notes."""
    audit = job.render.pop("audit", None)
    if audit is not None:
        try:
            s = audit.close() or {}
            job.crawler.log(f"[Accessibility] axe-core ran on {s.get('rendered', s.get('audited', 0))} page(s)"
                            + (f", {s['failed']} could not be checked" if s.get("failed") else "") + ".")
        except Exception:
            pass
    p = job.render.get("pool")
    if p is None or getattr(p, "_app_closed", False):
        return
    p._app_closed = True
    try:
        p.close()
        s = p.stats()
        job.crawler.log(f"[Render] {s.get('rendered', 0)} page(s) rendered in Chrome, {s.get('failed', 0)} could "
                        f"not be rendered" + (f" (last reason: {job.mask(s.get('last_error'))})"
                                              if s.get("failed") and s.get("last_error") else "") + ".")
        if s.get("promoted"):
            job.crawler.log(f"[Render] {s['promoted']} page(s) use their rendered HTML for the tabs (titles, "
                            "headings, words, canonical, directives), as in JavaScript rendering mode.")
        cov = getattr(job.crawler, "coverage", None)
        if s.get("failed") and isinstance(cov, dict) and isinstance(cov.get("notes"), list):
            cov["notes"].append(f"{s['failed']} page(s) could not be rendered in Chrome; their JavaScript "
                                "checks are Not checked.")
    except Exception as e:
        job.crawler.log(f"[Render] closing the render browsers failed: {type(e).__name__}")

def _crawl_outcome(job: CrawlJob) -> Tuple[str, str]:
    """(status, error) of a finished crawl for the scheduler's history."""
    crawler = job.crawler
    _settle_status(crawler)
    status = str(getattr(crawler, "status", "") or "")
    if status == "auth_failed":
        return "auth_failed", job.mask(getattr(crawler, "auth_error", "") or "The login failed.")
    if job.retired or getattr(crawler, "superseded", False):
        return "stopped", "Replaced by another crawl before it finished."
    return ("completed" if status == "completed" else "stopped"), ""

def run_crawler_task(job: CrawlJob, export_path: str, screenshots: bool = True, bot_check: bool = True):
    """The crawl, then what the dashboard needs at "done": the post-crawl checks, the SEO index and the
    snapshot. Everything slower runs after "done" in the background (_after_done): the Excel workbook, the
    pinpointed pictures (broken-link screenshots included), then the bot check; the PDF is built when it is
    first requested."""
    crawler = job.crawler
    finished_ok = False
    try:
        crawler.run()
        _close_render_pool(job)
        if crawler.status == "auth_failed":
            pass  # nothing was crawled; the status carries the login error to the UI
        elif getattr(crawler, "superseded", False) or job.retired:
            crawler.log("Reports skipped: a new crawl replaced this one.")
        else:
            job.phase = "analyzing"
            analysis = getattr(crawler.config, "analysis", None) or {}
            if analysis.get("auto") is False:
                crawler.log("Crawl Analysis is set to run on request: start it from Crawl Analysis > Start.")
            else:
                _quick_postcrawl(job, bot_check and analysis.get("bot_check") is not False)
            site, index = _site_index(job, force=True)
            if index is not None:
                crawler.log(f"SEO analysis done: {len(index.issues)} issue type(s) flagged across {len(index.tabs)} tabs.")
            if not screenshots:
                job.screenshot_note = "Screenshots were disabled for this crawl."
            if _save_snapshot(job):
                crawler.log(f"Crawl snapshot saved for later comparison: {job.snapshot_id}")
            _autosave(job, log=True)
            finished_ok = True
    except Exception as e:
        crawler.log(f"Crawl failed with error: {job.mask(e)}")
    finally:
        _close_render_pool(job)
        job.phase = "done"
    if finished_ok:
        _safe(lambda: _after_done(job, screenshots, bot_check), None, job, "Follow-up work")
    _schedule_finished(job)   # always: a scheduled run never stays "started" after its crawl ended

def _after_done(job: CrawlJob, screenshots: bool, bot_check: bool):
    """Follow-up work once the crawl is "done", none of it on the done path: the full workbook is prepared,
    connected data (Search Console, GA4, PageSpeed) is fetched, the crawl-complete email goes out, and the
    pinpointed pictures run first with the bot check after them (two Chrome jobs never run at once)."""
    if job.after_done:
        return
    job.after_done = True
    threading.Thread(target=_prebuild_excel, args=(job,), daemon=True, name="excel-prebuild").start()
    _safe(lambda: _auto_fetch_integrations(job), None, job, "Connected data")
    if not job.schedule_ref:
        threading.Thread(target=_email_on_complete, args=(job,), daemon=True, name="email").start()
    _safe(lambda: _connections.after_crawl(job), None, job, "Connections")   # W2: Slack / Teams / webhooks, alerts

    def chain():
        if screenshots and job.screenshots:
            _ensure_pinpoint(job)
            runner = job.runners["pinpoint"]
            while runner.busy and not job.retired:
                runner.wait(2.0)
        if bot_check:
            _start_bot_check(job)
    threading.Thread(target=chain, daemon=True, name="after-done").start()

def _prebuild_excel(job: CrawlJob):
    """The full workbook, built once after "done" so the first download is quick (a click during the build
    waits for it; job.excel_lock). A file briefly locked by Windows (an antivirus scan of a temporary file) is
    tried again."""
    for attempt in range(3):
        if job.retired or job.crawler is None:
            return
        try:
            _excel_file(job, "all")
            job.crawler.log(f"Excel report generated: {os.path.basename(job.export_file or '')}")
            return
        except OSError as e:
            if attempt == 2:
                job.crawler.log(f"Excel report could not be built: {job.mask(e)}")
            time.sleep(1.0 + attempt)
        except Exception as e:
            job.crawler.log(f"Excel report could not be built: {job.mask(e)}")
            return

def _patch_openpyxl_cleanup():
    """openpyxl's write-only sheets go through temporary files that it deletes as soon as the workbook is
    written. On Windows a just-written file can be held open for a moment (an antivirus scan), and the failed
    delete used to abort the whole workbook. The delete is retried, then left to openpyxl's exit clean-up."""
    try:
        from openpyxl.worksheet import _writer as ws_writer
    except Exception:
        return
    original = ws_writer.WorksheetWriter.cleanup
    if getattr(original, "_site_auditor", False):
        return

    def cleanup(self):
        for i in range(6):
            try:
                return original(self)
            except PermissionError:
                time.sleep(0.05 * (i + 1))
            except FileNotFoundError:
                return None
        return None   # still locked: the file stays listed for openpyxl's clean-up at exit

    cleanup._site_auditor = True
    ws_writer.WorksheetWriter.cleanup = cleanup

    def tidy_at_exit():
        """Runs before openpyxl's own exit clean-up: a temporary file that is still locked is left behind
        quietly (the system temp folder) instead of printing a traceback at shutdown."""
        for path in list(getattr(ws_writer, "ALL_TEMP_FILES", [])):
            try:
                os.remove(path)
            except OSError:
                pass
            try:
                ws_writer.ALL_TEMP_FILES.remove(path)
            except ValueError:
                pass
    import atexit
    atexit.register(tidy_at_exit)

_patch_openpyxl_cleanup()

def _pinpoint_work(job: CrawlJob, runner: JobRunner) -> dict:
    """Pinpointed pictures for the flagged SEO issues (boxed screenshots, source panels, search previews).
    Broken links are pictured here too (up to MAX_SCREENSHOT_ISSUES): their parent page with the link boxed
    and the error page go into job.evidence, which the Broken Links view, the URL details and tickets read."""
    from crawler.seo_evidence import capture_issue_evidence
    crawler = job.crawler
    site, index = _site_index(job)
    if index is None or not index.issues:
        runner.update(None, "No SEO issues to picture.")
        return {"pictures": 0, "issues": 0}
    d = _safe(lambda: _dismissals(job, index), _empty_dismissals(index))
    skip = {(i.get("tab"), i.get("filter")) for i in d["fully_dismissed_issues"]}
    chrome = _chrome_status()
    if not chrome.get("available"):
        job.screenshot_note = (f"Screenshots unavailable: {chrome.get('error') or 'Chrome is not available on this server.'} "
                               "Link positions are shown as page sketches.")

    def progress(done: int, total: int):
        runner.meta.update(done=done, total=total)
        job.evidence_progress = {"done": done, "total": total}
        runner.update(done * 100 // max(1, total), f"Pinpointed pictures: {done} of {total}")

    evidence = capture_issue_evidence(crawler, (site, index), job.evidence_dir, logger=crawler.log,
                                      cancel_event=runner.stop_event, on_progress=progress, skip=skip or None)
    job.seo_evidence = evidence
    links = _safe(lambda: _link_evidence(evidence), {}, job, "Broken-link screenshots")
    if links:
        job.evidence = links
    total_broken = sum(1 for r in crawler.results_list if r.category in ISSUE_CATEGORIES)
    if runner.stop_event.is_set() or job.cancel_event.is_set():
        job.screenshot_note = "Screenshot capture was stopped early; some issues have no screenshot."
    elif chrome.get("available") and total_broken > MAX_SCREENSHOT_ISSUES:
        job.screenshot_note = (f"Screenshots captured for the {MAX_SCREENSHOT_ISSUES} most severe of "
                               f"{total_broken} broken URLs.")
    pictures = sum(len(v) for v in evidence.values())
    runner.update(None, f"{pictures} pinpointed picture(s) for {len(evidence)} issue type(s).")
    _autosave(job)
    return {"pictures": pictures, "issues": len(evidence)}

def _link_evidence(evidence: dict) -> Dict[str, dict]:
    try:
        from crawler.seo_evidence import link_evidence
    except ImportError:
        return {}
    return link_evidence(evidence) or {}

def _ensure_pinpoint(job: CrawlJob, wait: float = 0.0):
    """Starts the pinpointed issue pictures once per crawl (only when screenshots are on) and waits up to
    `wait` seconds for them. Capture runs in its own thread, never inside the request."""
    if not job.screenshots or job.crawler is None or job.phase != "done" or job.origin == "opened":
        return
    runner = job.runners["pinpoint"]
    if runner.state == "idle":
        try:
            runner.start(lambda r: _pinpoint_work(job, r), "Preparing pinpointed pictures...")
        except JobBusy:
            pass
    if wait > 0 and runner.busy:
        runner.wait(wait)

@app.get("/", response_class=HTMLResponse)
def serve_index():
    static_file = os.path.join(os.path.dirname(__file__), "static", "index.html")
    with open(static_file, "r", encoding="utf-8") as f:
        return f.read()

@app.get("/config.js")
def serve_config():
    # The dashboard loads config.js on any host it does not recognise (e.g. a LAN address like
    # http://10.10.6.56:8000). Served by the backend itself, the API is this same origin.
    return Response("window.CRAWLER_API_BASE_URL = window.location.origin;\n",
                    media_type="application/javascript", headers={"Cache-Control": "no-cache"})

@app.get("/healthz")
@app.get("/api/health")
async def health_check():
    """Health check endpoint for Render and uptime monitoring. "checks" lists the SEO check modules that failed
    to load ("failed": {module: error}); status is then "degraded" and the dashboard shows a red banner."""
    try:
        from crawler.seo import registry
        checks = registry.checks_status()
        checks = {"complete": checks["complete"], "loaded": len(checks["loaded"]), "failed": checks["failed"]}
    except Exception as e:
        checks = {"complete": False, "loaded": 0, "failed": {"registry": f"{type(e).__name__}: {str(e)[:200]}"}}
    return {
        "status": "ok" if checks["complete"] else "degraded",
        "service": "website-crawler-backend",
        "timestamp": time.time(),
        "checks": checks,
    }

@app.get("/api/access")
def get_access(request: Request):
    """Whether this server asks for the team access key (CRAWLER_ACCESS_KEY), so the dashboard knows to ask.
    Open without a key. Sent with X-Access-Key, the key is checked as on every other route first (wrong: 401,
    counted; locked out: 429), so "valid" is true here only for the right key or its pass. The value checked is
    the one AccessKeyMiddleware decided on (request.state), never a second reading of the header. The right key
    also gets its "pass" (access_pass), which the dashboard sends from then on."""
    key = access_key()
    if not key:
        return {"required": False}
    given = getattr(request.state, "access_key", "")
    if access_key_matches(given, key) or access_pass_matches(given, key):
        return JSONResponse({"required": True, "valid": True, "pass": access_pass(key)}, headers={"Cache-Control": "no-store"})
    return {"required": True, "valid": False}

CHROME_PROBE_TTL = 300          # s /api/environment reuses its Chrome check
_chrome_probe: Dict[str, Any] = {"at": 0.0, "value": None}
_chrome_probe_lock = threading.Lock()

def _chrome_status() -> Dict[str, Any]:
    """{"available", "error"} from crawler.browser.chrome_status(): Selenium and the Chrome binary every
    browser here starts (no browser is started), plus the last start failure. Checked at most every 5 minutes."""
    with _chrome_probe_lock:
        if _chrome_probe["value"] is not None and time.time() - _chrome_probe["at"] < CHROME_PROBE_TTL:
            return dict(_chrome_probe["value"])
        try:
            from crawler.browser import chrome_status
            available, error = chrome_status()
        except Exception as e:
            available, error = False, f"Could not check for Chrome: {str(e).splitlines()[0][:200] if str(e) else type(e).__name__}"
        value = {"available": bool(available), "error": str(error or "")}
        _chrome_probe.update(at=time.time(), value=value)
        return dict(value)

@app.get("/api/environment")
def get_environment():
    """Detects whether running in serverless (e.g. Vercel) or persistent container/local (Render)."""
    is_vercel = bool(os.environ.get("VERCEL") or os.environ.get("AWS_LAMBDA_FUNCTION_NAME"))
    is_render = bool(os.environ.get("RENDER"))
    platform = "Vercel Serverless" if is_vercel else ("Render Persistent Container" if is_render else "Persistent Server / Local")
    try:
        from crawler.browser import chrome_available
        chrome = chrome_available()   # Selenium installed + a Chrome/Chromium binary found (no browser started)
    except Exception:
        chrome = False
    return {
        "is_serverless": is_vercel,
        "platform": platform,
        "is_render": is_render,
        "max_duration_seconds": 60 if is_vercel else None,
        "chrome_available": chrome,
        # whether Chrome can be used here and, if not, why (a missing binary, a CHROME_BIN that points nowhere,
        # the last failed start); "error" can also hold a warning while Chrome works
        "chrome": _chrome_status(),
        # URLs one crawl may queue on this server (0 = no cap); Max URLs above it is lowered to it
        "max_urls_cap": 100 if is_vercel else SERVER_MAX_URLS,
        "list_max_urls": _list_cap(),
        "app": APP_NAME,
        "version": app_version(),
        "ua_presets": list(_ua_presets()),
        "crawl_scopes": list(CRAWL_SCOPES),
        "crawl_modes": list(CRAWL_MODES),
        "max_custom_rules": MAX_CUSTOM_RULES,
        # Find subdomains: the places it can look ({code: plain label}) and how many hosts one crawl may add
        "subdomain_sources": _safe(lambda: {s: _subdomains().SOURCE_LABELS.get(s, s) for s in _subdomains().SOURCES},
                                   {}),
        "max_extra_hosts": MAX_EXTRA_HOSTS,
        "browser_enabled": _browser_enabled(),
        "notify_configured": _safe(lambda: _notify().configured(reports_dir=REPORTS_DIR), False),
        "psi_connected": _safe(lambda: bool(_integrations().connection("psi").get("connected")), False),
        "scheduler": _scheduler_state(),
        "reports_dir_set": bool(os.environ.get("CRAWLER_REPORTS_DIR")),
        "access_key_set": bool(access_key()),
    }

class InspectRequest(BaseModel):
    url: str

@app.post("/api/auth/inspect")
@app.get("/api/auth/inspect")
def inspect_auth_endpoint(url: Optional[str] = None, req: Optional[InspectRequest] = None):
    """
    Fast pre-flight inspection of target URL to determine if login/credentials are required or available.
    Detects HTTP 401/403, redirects to login, password input fields, or login portal links.
    """
    raw_url = (req.url if req and req.url else url) or ""
    raw_url = raw_url.strip()
    if not raw_url:
        return {"requires_auth": False, "login_url": None, "reason": "No URL provided", "challenge": _no_challenge()}

    if not raw_url.startswith("http://") and not raw_url.startswith("https://"):
        raw_url = "https://" + raw_url

    try:
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8"
        }
        r = requests.get(raw_url, headers=headers, timeout=6, allow_redirects=True, verify=False)
        final_url = str(r.url)
        status = r.status_code
        # A bot check (a firewall challenge instead of the page, or a CAPTCHA in the login form): ARACHNIX can't
        # complete it, so the answer points at Paste cookies (or, for the owner, an allow rule)
        try:
            from crawler.auth import challenge_status, response_bot_check
            bot = response_bot_check(r, "sign-in", login_page=True)
        except Exception:
            bot = None
        challenge = challenge_status(bot["provider"], 1, final_url, "sign-in") if bot else _no_challenge()

        # 0. The site answers with a bot check instead of the page
        if bot and bot.get("kind") == "firewall":
            return {
                "requires_auth": True,
                "reason": bot["advice"],
                "login_url": final_url,
                "auth_type": "cookies",
                "has_captcha": True,
                "challenge": challenge,
            }

        # 1. HTTP 401 Unauthorized or 403 Forbidden
        if status in (401, 403):
            return {
                "requires_auth": True,
                "reason": f"HTTP {status} Access Restricted",
                "login_url": final_url,
                "auth_type": "form",
                "challenge": challenge,
            }

        p_orig = urlparse(raw_url).path.rstrip('/')
        p_final = urlparse(final_url).path.rstrip('/')

        # 2. Redirected to login/signin/auth page (whole path segments: /author/ or /authority-guide do not count)
        if p_orig != p_final and is_login_url(final_url):
            return {
                "requires_auth": True,
                "reason": bot["advice"] if bot else "Redirected to login portal",
                "login_url": final_url,
                "auth_type": "cookies" if bot else "form",
                "has_captcha": bool(bot),
                "challenge": challenge,
            }

        # 3. Direct login URL provided
        if is_login_url(final_url):
            return {
                "requires_auth": True,
                "reason": bot["advice"] if bot else "Direct login portal URL",
                "login_url": final_url,
                "auth_type": "cookies" if bot else "form",
                "has_captcha": bool(bot),
                "challenge": challenge,
            }

        # 4. Password input form detected on page (e.g. login form)
        soup = BeautifulSoup(r.text, 'html.parser')
        has_captcha = bool(bot) or any(k in r.text.lower() for k in ['turnstile', 'challenges.cloudflare.com', 'hcaptcha', 'cf-challenge', 'datadome'])
        pwd_input = soup.find('input', {'type': 'password'})
        if pwd_input:
            return {
                "requires_auth": True,
                "reason": bot["advice"] if bot else ("Login portal with Cloudflare Turnstile / CAPTCHA detected (Session Cookies recommended)" if has_captcha else "Password login form detected on page"),
                "login_url": final_url,
                "auth_type": "cookies" if has_captcha else "form",
                "has_captcha": has_captcha,
                "challenge": challenge,
            }

        # 5. A login link on a public page: a hint only (shops and blogs have one on every page), never "needs a
        # login" - the crawl starts without asking and the login page is pre-filled in the login options
        for a in soup.find_all('a', href=True):
            href = str(a.get('href') or '')
            text = a.get_text().strip().lower()
            full_login_url = urljoin(final_url, href)
            if is_login_url(full_login_url) or text in ('log in', 'login', 'sign in', 'signin'):
                return {
                    "requires_auth": False,
                    "login_available": True,
                    "reason": "Public site with a login link (log in only to crawl the members' pages)",
                    "login_url": full_login_url,
                    "auth_type": "form",
                    "has_captcha": False,
                    "challenge": challenge,
                }

        return {"requires_auth": False, "login_url": None, "reason": "Public access (No login detected)",
                "has_captcha": False, "challenge": challenge}
    except Exception as e:
        return {"requires_auth": False, "login_url": None, "reason": str(e), "challenge": _no_challenge()}

# --------------------------------------------------------------------------
# Direct Mode for Serverless / Vercel (Synchronous execution within timeout)
# --------------------------------------------------------------------------
@app.post("/api/crawl/direct")
def direct_crawl(req: CrawlRequest):
    """
    Synchronous crawl endpoint specifically optimized for Serverless platforms (Vercel).
    Executes within the 60s function limit and returns completed results & download filename.
    It is the server's one shared crawl like any other: refused (409) while another crawl runs, on screen and
    counted as running while it works (other starts wait, or replace it with force), and answered 409 when
    another crawl replaced it before it finished.
    """
    raw_url = _start_target(req)
    base_name = _base_name(raw_url)
    export_filename = f"{base_name}.xlsx"
    export_path = os.path.join(get_export_dir(), export_filename)

    is_vercel = _is_serverless()

    # On Vercel, cap at 100 URLs to guarantee completion within the 60s lambda limit
    config = _build_config(req, raw_url, export_path, serverless=is_vercel)
    config.render_all = False   # a synchronous crawl never opens a Chrome per page
    if not is_vercel:
        config.concurrency = min(config.concurrency, 8)

    outcome: Dict[str, Any] = {}
    with _start_lock:
        old = session.job
        if old is not None and _crawl_busy(old):
            raise HTTPException(status_code=409, detail=_busy_detail(old))
        try:
            from crawler.ai import client as ai_client
            ai_client.reset_usage()
        except Exception:
            pass
        job = CrawlJob(None, base_name, screenshots=False, secrets=_config_secrets(config))
        job.origin = "direct"
        job.screenshot_note = "Screenshots are not available in serverless mode."
        job.crawler = WebsiteCrawler(config=config, on_log=job.log)
        job.crawler.status = "running"
        job.crawler.start_time = time.time()
        job.thread = threading.Thread(target=_direct_work, args=(job, export_path, outcome), daemon=True,
                                      name="direct-crawl")
        if old is not None:
            old.stop_jobs()
        session.job = job
        job.thread.start()
    job.thread.join()

    crawler = job.crawler
    if crawler.status == "auth_failed":
        raise HTTPException(status_code=401, detail=job.mask(f"Login failed, nothing was crawled: {crawler.auth_error}"))
    if outcome.get("error"):
        raise HTTPException(status_code=500, detail=outcome["error"])
    if not outcome.get("finished"):
        raise HTTPException(status_code=409, detail="Another crawl replaced this one before it finished.")
    results = crawler.results_list
    return {
        "status": "completed",
        "stats": _public_stats(job),
        "results_count": len(results),
        "external_count": len(crawler.external_links_list),
        "recent_results": [r.to_dict() for r in results],
        "recent_external": [e.to_dict() for e in crawler.external_links_list],
        "export_file": os.path.basename(job.export_file or export_path),
        "challenge": _crawl_challenge(job),
    }

def _direct_work(job: CrawlJob, export_path: str, outcome: Dict[str, Any]):
    """The direct crawl in its own thread (job.busy while it runs): the crawl, the quick post-crawl checks,
    the SEO index, the workbook and the snapshot. A crawl replaced meanwhile stops without its reports."""
    crawler = job.crawler
    try:
        crawler.run()
        if crawler.status == "auth_failed" or getattr(crawler, "superseded", False) or job.retired:
            return
        job.phase = "analyzing"
        _quick_postcrawl(job, bot_check=False)   # robots.txt audit + random-URL probe only (60 s limit)
        _site_index(job, force=True)
        job.export_file = _write_excel(job, export_path)
        job.phase = "done"
        _save_snapshot(job)
        outcome["finished"] = True
    except Exception as e:
        outcome["error"] = job.mask(f"The crawl failed: {type(e).__name__}: {str(e)[:300]}")
        crawler.log(outcome["error"])
    finally:
        job.phase = "done"

# --------------------------------------------------------------------------
# Background Async Mode (For Local Machine, Docker, Render, Railway, VPS)
# --------------------------------------------------------------------------
_start_lock = threading.RLock()   # one crawl start (dashboard, scheduler, opening a saved crawl) at a time
SCHEDULE_POST_BUSY_MAX = 1800     # s a scheduled run's after-crawl steps count as busy at most (a hung step)
SETTLE_GRACE = 2.0                # s after end_time by which the engine has written a crawl's final status

def _final_status(crawler) -> Optional[str]:
    """The status a crawl must show once its run is over (the engine sets end_time when run() ends, after a
    failed login or an error, and writes its final status right after), or None while it still crawls, runs its
    post-crawl checks or is about to write that status itself (SETTLE_GRACE)."""
    end = getattr(crawler, "end_time", 0) or 0
    if not end or time.time() - end < SETTLE_GRACE:
        return None
    done, stop = getattr(crawler, "_crawl_done", None), getattr(crawler, "_stop_event", None)
    if done is not None and not done.is_set():   # ended before the crawl loop finished: a login failure or an error
        return "stopped" if getattr(crawler, "_error_stop", False) else "auth_failed"
    return "stopped" if stop is not None and stop.is_set() else "completed"

def _settle_status(crawler) -> None:
    """A Pause or Resume that reached a crawl just as it ended left it "paused" / "running" for ever, and the
    server busy (Stop, Clear, Close, Open and scheduled runs all refused): such a crawl gets its final status."""
    if crawler is None or str(getattr(crawler, "status", "")) not in ("running", "paused"):
        return
    final = _final_status(crawler)
    if final is None:
        return
    pause = getattr(crawler, "_pause_event", None)
    if pause is not None:
        pause.set()
    crawler.status = final

def _crawl_loop_running(crawler) -> bool:
    """The crawl itself (not its post-crawl checks) is still under way: Pause and Resume only act then."""
    done = getattr(crawler, "_crawl_done", None)
    return (str(getattr(crawler, "status", "")) in ("running", "paused") and not getattr(crawler, "end_time", 0)
            and not (done is not None and done.is_set()))

def _post_busy(job: Optional[CrawlJob]) -> bool:
    """A scheduled run's after-crawl steps (exports, PDF, saved crawl, comparison, email) are still working on
    this crawl. They count as busy, so the next task, Open, Close and Clear wait instead of replacing it."""
    since = getattr(job, "post_since", 0.0) if job is not None else 0.0
    return bool(since) and time.time() - since < SCHEDULE_POST_BUSY_MAX

def _crawl_running(job: Optional[CrawlJob]) -> bool:
    """The crawl itself or its report building is under way (not only a Crawl Analysis re-run or a scheduled
    run's exports)."""
    return bool(job is not None and job.crawler is not None
                and (job.busy or str(getattr(job.crawler, "status", "")) in ("running", "paused")))

def _crawl_busy(job: Optional[CrawlJob] = None) -> bool:
    """One shared crawl per server: True while a crawl runs, is paused or builds its reports, while a Crawl
    Analysis re-run works on it, or while a scheduled run's after-crawl steps build its exports. Opening a
    saved crawl, List mode, scheduled runs and analysis re-runs refuse (or, for scheduled runs, wait) while
    this is True."""
    job = job if job is not None else session.job
    if job is None or job.crawler is None:
        return False
    _settle_status(job.crawler)
    return _crawl_running(job) or job.runners["analysis"].busy or _post_busy(job)

def _busy_detail(job: Optional[CrawlJob] = None) -> str:
    job = job if job is not None else session.job
    url = job.crawler.config.start_url if job and job.crawler else ""
    if job is not None and job.runners["analysis"].busy and not job.busy:
        return f"Crawl Analysis is running for {url}. Wait for it to finish or stop it first."
    if job is not None and not job.busy and _post_busy(job):
        return (f"A scheduled task is still building the exports of its crawl of {url} (Excel, PDF, saved crawl, "
                "email). Wait until it has finished.")
    return f"A crawl is already running for {url}."

def _start_target(req: CrawlRequest) -> str:
    """The start URL: the request's url, or in List mode the first listed web address."""
    if (req.mode or "spider").strip().lower() == "list":
        if not _lines(req.urls):
            _bad("List mode needs at least one URL (upload or paste the list).")
        first = next((a for a in (_as_web_address(x) for x in _lines(req.urls)) if a), "")
        if not first:
            _bad("The list holds no web addresses (each line must be a URL such as https://example.com/page).")
        return first
    if not (req.url or "").strip():
        _bad("Enter the website address to crawl.")
    return _normalize_target(req.url)

def _launch_crawl(req: CrawlRequest, origin: str = "manual", schedule_ref: str = "", busy_status: int = 0,
                  block_addresses: str = "", extra_notes: Optional[List[str]] = None) -> dict:
    """Starts a crawl in the background (POST /api/crawl/start and scheduled tasks share this). Every option is
    checked before anything stops; a running crawl is refused unless `force` (then it is replaced). busy_status:
    the HTTP status of that refusal (default 400 for a manual start, 409 otherwise); block_addresses: refused
    addresses for this crawl (P24; on Render at least 'metadata')."""
    raw_url = _start_target(req)
    base_name = _base_name(raw_url)
    export_filename = f"{base_name}.xlsx"
    export_path = os.path.join(get_export_dir(), export_filename)
    notes: List[str] = []
    config = _build_config(req, raw_url, export_path, notes=notes)  # validates first: a bad regex stops nothing
    notes.extend(extra_notes or ())
    if block_addresses:
        config.block_addresses = block_addresses
    elif os.environ.get("RENDER") and not getattr(config, "block_addresses", ""):
        config.block_addresses = "metadata"       # a hosted server never opens its cloud's metadata service

    with _start_lock:
        job = session.job
        if job and _crawl_busy(job):
            if req.force:
                job.cancel_event.set()
                job.runners["analysis"].stop()
                # no post-crawl checks, reports or snapshot for a crawl still under way; a finished one (Crawl
                # Analysis or a scheduled run's exports still busy) stays complete, so its task can still save it
                if _crawl_running(job):
                    job.crawler.stop(superseded=True)
            else:
                raise HTTPException(status_code=busy_status or (400 if origin == "manual" else 409),
                                    detail=_busy_detail(job))

        cleanup_old_reports(keep=job.base_name if job else None)
        if job is not None:
            job.stop_jobs()   # validation, intel, bot check... of the previous crawl
        try:
            from crawler.ai import client as ai_client
            ai_client.reset_usage()     # a new per-crawl AI call budget
        except Exception:
            pass

        screenshots = req.screenshots is not False and _browser_enabled()
        job = CrawlJob(None, base_name, screenshots=screenshots, secrets=_config_secrets(config))
        job.origin, job.schedule_ref, job.notes = origin, schedule_ref, notes
        if schedule_ref:
            job.run_ref = schedule_ref
            _scheduled_jobs[schedule_ref] = job
            while len(_scheduled_jobs) > 50:
                _scheduled_jobs.pop(next(iter(_scheduled_jobs)))
        job.save_name, job.save_project = (req.name or "").strip()[:120], (req.project or "").strip()[:60]
        job.crawler = WebsiteCrawler(config=config, on_log=job.log)
        job.crawler.evidence_dir = job.evidence_dir   # render_all screenshots land with the other pictures
        _browser_proxy(job.crawler, config)
        if getattr(config, "render_all", False):
            _attach_render_pool(job)
        elif getattr(config, "a11y_audit", False) and _browser_enabled():
            _attach_audit_hook(job)       # W7 / A2: axe-core on pages that are not rendered anyway
        job.crawler.status = "running"
        job.crawler.start_time = time.time()
        session.job = job
        for note in notes:
            job.log(note)

        job.thread = threading.Thread(
            target=run_crawler_task,
            args=(job, export_path, screenshots, req.bot_check is not False),
            daemon=True
        )
        job.thread.start()
    out = {"status": "started", "target": config.start_url or raw_url, "output_file": export_filename,
           "mode": config.mode,   # target: the start URL without a user name and password (move_url_credentials)
           "crawl_scope": getattr(config, "crawl_scope", "subdomain"), "notes": notes, "run_ref": job.run_ref}
    if config.mode == "list":
        out["list"] = {"lines": len(config.seed_urls)}
    if getattr(config, "extra_hosts", None):
        out.update(_extra_hosts_answer(config))
    return out

def _extra_hosts_answer(config: CrawlConfig) -> Dict[str, Any]:
    """{"extra_hosts": [hosts that will be tried], "extra_hosts_skipped": {value: reason code}} for a crawl start
    (worked out offline; hosts without DNS or on a private address are dropped once the crawl starts and listed
    in the crawl status' stats.extra_hosts_skipped)."""
    try:
        from crawler.subdomains import extra_host_urls
        homes, skipped = extra_host_urls(config.extra_hosts, config.start_url)
    except Exception:
        return {"extra_hosts": list(config.extra_hosts), "extra_hosts_skipped": {}}
    return {"extra_hosts": [urlparse(h).netloc for h in homes], "extra_hosts_skipped": skipped}

@app.post("/api/crawl/start")
def start_crawl(req: CrawlRequest, request: Request):
    """Starts a crawl (Spider or List mode) with every crawl option of the dashboard; 400 for a bad option or
    while another crawl runs (unless force). A crawl an AI agent starts (W1) never replaces a running one (409),
    may not start on a metadata / private address and has refused addresses set (block_addresses)."""
    if _agents.is_agent(request):
        if req.auth is not None and getattr(req.auth, "handoff_id", None):
            raise HTTPException(status_code=400, detail="AI agents cannot use a browser sign-in hand-off.")
        mode, notes = _agents.agent_crawl_rules(req, request)
        return _launch_crawl(req, busy_status=409, block_addresses=mode, extra_notes=notes)
    _extension.apply_to_crawl(req)          # W6: auth.handoff_id -> the hand-off's cookies (used once)
    return _launch_crawl(req)

@app.post("/api/crawl/stop")
async def stop_crawl():
    """Stops crawling. If the crawl already finished and screenshots are being captured,
    skips the remaining screenshots instead. Reports are still generated."""
    job = session.job
    if not job or not job.crawler:
        return {"status": "idle"}
    _settle_status(job.crawler)
    if job.phase == "crawling":
        job.crawler.stop()
    else:
        job.cancel_event.set()
        job.runners["pinpoint"].stop()
        job.runners["analysis"].stop()
    _settle_status(job.crawler)
    return {
        "status": "stopped",
        "export_ready": True,
        "export_file": f"{job.base_name}.xlsx"
    }

@app.post("/api/crawl/pause")
async def pause_crawl():
    """Pauses the crawl while it crawls. A click that arrives once the crawl itself has finished (a stale page,
    or a click racing the end of the crawl) changes nothing: it used to leave the finished crawl "paused" and
    the server busy for good."""
    job = session.job
    crawler = job.crawler if job else None
    if crawler is None:
        return {"status": "idle"}
    if not _crawl_loop_running(crawler):
        _settle_status(crawler)
        return {"status": _public_status(job), "detail": "The crawl is not crawling any more: nothing to pause."}
    if str(getattr(crawler, "status", "")) == "paused":
        return {"status": "paused", "detail": "The crawl is already paused."}
    crawler.pause()
    if not _crawl_loop_running(crawler):
        # it finished between the check and the pause: its post-crawl checks go on and it keeps its final status
        crawler._pause_event.set()
        if crawler.status == "paused":
            crawler.status = "running"
        _settle_status(crawler)
        return {"status": _public_status(job), "detail": "The crawl had just finished: nothing to pause."}
    return {"status": "paused"}

@app.post("/api/crawl/resume")
async def resume_crawl():
    job = session.job
    crawler = job.crawler if job else None
    if crawler is None:
        return {"status": "idle"}
    if str(getattr(crawler, "status", "")) != "paused":
        _settle_status(crawler)
        return {"status": _public_status(job), "detail": "The crawl is not paused."}
    crawler.resume()
    _settle_status(crawler)   # a crawl left "paused" after it ended gets its final status instead of "running"
    return {"status": "resumed" if crawler.status == "running" else _public_status(job)}

@app.post("/api/crawl/reset")
async def reset_crawl():
    job = session.job
    if job and _crawl_busy(job):
        return {"status": "running", "detail": "A crawl is currently in progress." if _crawl_running(job)
                else _busy_detail(job)}
    if job is not None:
        job.stop_jobs()
    session.business.stop()
    session.business.wait(1)   # its AI calls stop within ~0.3 s (job scope); never replace it mid-write
    session.job = None
    session.business = JobRunner("business")
    return {"status": "idle"}

def _jobs_status(job: Optional[CrawlJob]) -> Dict[str, dict]:
    """State of every background job: {kind: {job_id, state, progress, message}}."""
    out = {k: (job.runners[k].status() if job is not None else JobRunner(k).status()) for k in JOB_KINDS}
    out["business"] = session.business.status()
    meta = job.runners["pinpoint"].meta if job is not None else {}
    out["pinpoint"].update(done=int(meta.get("done") or 0), total=int(meta.get("total") or 0))
    return out

CHALLENGE_REUSE_S = 2.0   # a running crawl's bot-check summary is rebuilt at most this often (or 4x its cost)

def _no_challenge() -> Dict[str, Any]:
    try:
        from crawler.auth import challenge_status
        return challenge_status()
    except Exception:
        return {"seen": False, "provider": "", "count": 0, "first_url": "", "advice": "", "owner_steps": []}

def _masked_challenge(found: Any, mask: Callable[[Any], str]) -> Dict[str, Any]:
    """A challenge object fit for an API answer: the contract's keys always present, texts masked."""
    out = _no_challenge()
    if isinstance(found, dict) and found.get("seen"):
        out.update({"seen": True, "provider": mask(found.get("provider") or ""),
                    "count": int(found.get("count") or 0), "first_url": mask(found.get("first_url") or ""),
                    "advice": mask(found.get("advice") or ""),
                    "owner_steps": [mask(s) for s in found.get("owner_steps") or []]})
    return out

def _crawl_challenge(job: CrawlJob) -> Dict[str, Any]:
    """The "challenge" object of GET /api/crawl/status: pages a bot check / firewall answered instead of the page
    (their status codes and categories stay as they are), else a bot check the sign-in met. Rebuilt when the
    results change, but while a crawl runs at most every CHALLENGE_REUSE_S seconds (it reads every result)."""
    crawler = job.crawler
    handler = getattr(crawler, "auth_handler", None)
    key = (len(getattr(crawler, "results_list", None) or ()), job.results_version,
           bool(getattr(handler, "challenge", None)), str(getattr(crawler, "status", "")))
    cached = job.challenge_cache
    if cached is not None:
        old_key, value, made_at, cost = cached
        running = str(getattr(crawler, "status", "")) in ("running", "paused")
        if old_key == key or (running and time.time() - made_at < max(CHALLENGE_REUSE_S, 4 * cost)):
            return value
    began = time.time()
    try:
        from crawler.auth import crawl_challenge_status
        value = _masked_challenge(crawl_challenge_status(crawler), job.mask)
    except Exception:
        value = _no_challenge()
    job.challenge_cache = (key, value, time.time(), time.time() - began)
    return value

def _public_status(job: CrawlJob) -> str:
    """running/paused while crawling, 'analyzing' while reports & screenshots are built."""
    _settle_status(job.crawler)
    status = job.crawler.status
    if status in ("completed", "stopped") and job.busy:
        return "analyzing"
    return status

@app.get("/api/crawl/status")
async def get_status(since: int = 0, ext_since: int = 0):
    """
    Live crawl status. `since` / `ext_since` are cursors into the full result lists so the
    client receives every row exactly once (up to 1000 per poll).
    """
    job = session.job
    if not job or not job.crawler:
        return {
            "status": "idle",
            "phase": "idle",
            "stats": {
                "status": "idle",
                "elapsed_seconds": 0,
                "total_discovered": 0,
                "total_crawled": 0,
                "queue_size": 0,
                "status_2xx": 0,
                "status_3xx": 0,
                "status_4xx": 0,
                "status_5xx": 0,
                "status_err": 0,
                "failed_count": 0,
                "auth_required_count": 0,
                "total_external": 0,
                "unique_external_domains": 0,
                "start_url": "",
                "auth_used": False,
                "auth_status": "Not required",
                "auth_error": None,
                "limitations": []
            },
            "results_count": 0,
            "external_count": 0,
            "new_results": [],
            "new_external": [],
            "next_since": 0,
            "next_ext_since": 0,
            "recent_results": [],
            "recent_external": [],
            "logs": [],
            "export_ready": False,
            "report_ready": False,
            "pdf_ready": False,
            "results_version": 0,
            "snapshot_id": "",
            "jobs": _jobs_status(None),
            "origin": "",
            "notes": [],
            "saved_crawl": None,
            "busy": False,
            "challenge": _no_challenge(),
        }

    crawler = job.crawler
    stats = _public_stats(job)
    stats["status"] = _public_status(job)
    results = crawler.results_list
    externals = crawler.external_links_list
    total_results, total_external = len(results), len(externals)
    since = max(0, min(since, total_results))
    ext_since = max(0, min(ext_since, total_external))
    new_results = [r.to_dict() for r in results[since:since + 1000]]
    new_external = [e.to_dict() for e in externals[ext_since:ext_since + 1000]]
    # Excel and PDF are built on request (the workbook is also prepared in the background after "done")
    report_ready = (job.phase == "done" and crawler.status != "auth_failed" and not job.retired
                    and not getattr(crawler, "superseded", False))
    export_name = os.path.basename(job.export_file) if job.export_file else (
        f"{job.base_name}.xlsx" if report_ready else None)

    return {
        "status": stats["status"],
        "phase": job.phase,
        "stats": stats,
        "results_count": total_results,
        "external_count": total_external,
        "new_results": new_results,
        "new_external": new_external,
        "next_since": since + len(new_results),
        "next_ext_since": ext_since + len(new_external),
        "recent_results": [r.to_dict() for r in results[-50:]],
        "recent_external": [e.to_dict() for e in externals[-50:]],
        "logs": list(job.logs)[-30:],
        "evidence_progress": job.evidence_progress,
        "screenshot_note": job.screenshot_note,
        "export_ready": report_ready,
        "report_ready": report_ready,
        "pdf_ready": report_ready,
        "pdf_built": bool(job.pdf_file),
        "export_file": export_name,
        "auth_used": stats.get("auth_used", False),
        "auth_status": stats.get("auth_status", "Not required"),
        "auth_error": stats.get("auth_error"),
        "limitations": stats.get("limitations", []),
        # bumped when a re-check rewrote crawl rows: reload the results from since=0
        "results_version": job.results_version,
        "snapshot_id": job.snapshot_id,
        "jobs": _jobs_status(job),
        "origin": job.origin,
        "notes": list(job.notes),
        "saved_crawl": _public_saved_meta(job),
        "busy": _crawl_busy(job),
        # bot checks met (the banner's text is "advice"); the pages keep their status codes
        "challenge": _crawl_challenge(job),
        "private": _crawl_private(job),     # a signed-in crawl (W1: AI agents get its content only when allowed)
    }


def _crawl_private(job: CrawlJob) -> bool:
    try:
        from crawler.fixfiles import _private
        crawler = job.crawler
        return bool(_private(None, crawler))
    except Exception:
        return False

def _external_issue(e, site) -> dict:
    """A broken link to another website, shaped like an internal issue row for the Broken Links tab."""
    first, pages = found_on(site, e) if site is not None else (e.parent_url, [])
    code = getattr(e, "status_code", None)
    item = e.to_dict()
    item.update({
        "external": True, "category": row_category(e), "parent_url": first or e.parent_url,
        "page_title": "", "depth": None, "resource_type": "external", "found_in_code": False,
        "linked_from": pages[:LINKED_FROM_SHOWN], "linked_from_count": len(pages),
        "issue": {
            "category": row_category(e), "code_label": e.http_status or "No response",
            "title": "External link: no response" if code is None else f"External link: {code}",
            "summary": "A link to another website that gave no answer when checked." if code is None
                       else f"A link to another website that answered {e.http_status or code}.",
            "causes": ["The other page was moved or deleted", "The other site blocks crawlers or is down"],
            "fix": EXTERNAL_FIX_NO_RESPONSE if code is None else EXTERNAL_FIX_ERROR,
            "severity": "low",
        },
    })
    return item

@app.get("/api/crawl/issues")
def get_issues():
    """All internal 4xx / 5xx / no-response URLs with diagnosis, link position, screenshot URLs and every
    page that links to them ("linked_from"); "external_issues" lists broken links to other websites."""
    job = session.job
    if not job or not job.crawler:
        return {"issues": [], "external_issues": [], "phase": "idle", "screenshot_note": ""}
    issues = sorted([r for r in job.crawler.results_list if r.category in ISSUE_CATEGORIES], key=issue_sort_key)
    site, _ = _site_index(job) if job.phase != "crawling" else (None, None)
    items = []
    for r in issues:
        ev = job.evidence.get(r.url) or {}
        item = r.to_dict()
        item["position"] = ev.get("position")
        item["parent_shot"] = f"/api/crawl/evidence/{ev['parent_shot']}" if ev.get("parent_shot") else None
        item["error_shot"] = f"/api/crawl/evidence/{ev['error_shot']}" if ev.get("error_shot") else None
        if site is not None:
            first, pages = found_on(site, r)
            item["found_on"] = first
            item["linked_from"] = pages[:LINKED_FROM_SHOWN]
            item["linked_from_count"] = len(pages)
        items.append(item)
    external = [_external_issue(e, site) for e in list(job.crawler.external_links_list) if is_broken_external(e)]
    return {"issues": items, "external_issues": external, "phase": job.phase,
            "screenshot_note": job.screenshot_note}

@app.get("/api/crawl/evidence/{name}")
async def get_evidence(name: str):
    """A screenshot file: broken-link evidence, pinpointed issue pictures (pin-*) or a login-test picture."""
    job = session.job
    safe = os.path.basename(name)
    if not re.fullmatch(r"[\w.-]+\.(jpg|png)", safe):
        raise HTTPException(status_code=404, detail="Screenshot not found.")
    folders = [job.evidence_dir] if job else []
    if safe.startswith("login_test_"):
        folders.append(LOGIN_TEST_DIR)
    path = next((os.path.join(d, safe) for d in folders if os.path.exists(os.path.join(d, safe))), None)
    if not path:
        raise HTTPException(status_code=404, detail="Screenshot not found.")
    return FileResponse(path, media_type="image/jpeg" if safe.endswith(".jpg") else "image/png")

@app.get("/api/crawl/report/excel")
def download_category_excel(category: str = "all"):
    """Excel report for one category: all, 2xx, 3xx, 4xx, 5xx, error, issues or external. The full workbook
    ("all") also has the Fix Plan, Validation, Checks, Content Intelligence and Dismissed Issues sheets."""
    job = session.job
    if not job or not job.crawler:
        raise HTTPException(status_code=404, detail="No crawl data available.")
    if category not in EXPORT_CATEGORIES:
        raise HTTPException(status_code=400, detail=f"Unknown category. Use one of: {', '.join(EXPORT_CATEGORIES)}")
    suffix = "" if category == "all" else f"_{category}"
    path = _excel_file(job, category)
    return FileResponse(path=path, filename=f"{job.base_name}{suffix}.xlsx", media_type=XLSX_MEDIA)

@app.get("/api/crawl/report/pdf")
def download_pdf_report():
    """Visual PDF issue report (summary + one card per issue with screenshots). On the first request after a
    crawl with screenshots, the pinpointed issue pictures are captured in the background and the request waits
    for them (up to CRAWLER_PDF_PINPOINT_WAIT seconds; later requests include them once ready). The PDF is
    rebuilt when validation, intel, pictures or dismissals changed since it was made."""
    job = session.job
    if not job or not job.crawler:
        raise HTTPException(status_code=404, detail="No crawl data available.")
    if job.phase == "analyzing":
        raise HTTPException(status_code=409, detail="The PDF report is still being generated.")
    _ensure_pinpoint(job, wait=PDF_PINPOINT_WAIT)
    _, index = _site_index(job)
    fresh = job.pdf_file and os.path.exists(job.pdf_file) and job.pdf_key == _pdf_key(job, index)
    path = job.pdf_file if fresh else build_pdf(job)
    return FileResponse(path=path, filename=os.path.basename(path), media_type="application/pdf")

@app.get("/api/crawl/pinpoint")
def get_pinpoint():
    """Pinpointed issue pictures: job state and the pictures per issue ("tab/filter" keys; each item's
    "shot" is served by /api/crawl/evidence/<shot>)."""
    job = session.job
    if not job or not job.crawler:
        return {"state": "idle", "progress": 0, "message": "", "done": 0, "total": 0, "evidence": {}}
    from crawler.seo_evidence import to_json
    out = _jobs_status(job)["pinpoint"]
    out["evidence"] = to_json(job.seo_evidence)
    if out.get("state") == "idle" and not job.screenshots and out["evidence"]:
        # an opened saved crawl: its pictures came with it (no new pictures are taken for it)
        out.update(state="done", progress=100, message="Pictures of the saved crawl.")
    return out

@app.post("/api/crawl/pinpoint")
def start_pinpoint():
    """Starts the pinpointed issue pictures now (normally the first PDF request does)."""
    job = session.job
    if not job or not job.crawler:
        raise HTTPException(status_code=404, detail="No crawl data available.")
    if job.phase != "done":
        raise HTTPException(status_code=409, detail="Wait until the crawl and its reports are finished.")
    if not job.screenshots:
        raise HTTPException(status_code=409, detail="Screenshots were disabled for this crawl.")
    _ensure_pinpoint(job)
    return {"job_id": job.runners["pinpoint"].job_id}

# --------------------------------------------------------------------------
# SEO tabs (Screaming Frog parity): grids, URL details, issues overview, exports.
# All are plain `def` so FastAPI runs them in its thread pool: building the
# index can take seconds on large crawls and must not block the event loop.
# --------------------------------------------------------------------------
def _view(job: CrawlJob, index) -> ViewCache:
    with job.view_lock:
        if job.view is None or job.view.index is not index:
            job.view = ViewCache(index)
        return job.view

def _grid(job: CrawlJob, site, index, tab_id: str, filter_id: str):
    """(columns, rows, lowercase search text per row) for a tab grid, cached per index.
    rows is None when the tab or filter does not exist."""
    view = _view(job, index)
    key = (tab_id, filter_id or "all")
    hit = view.grids.get(key)
    if hit is not None:
        return hit
    columns, rows = tab_grid(site, index, tab_id, filter_id)
    texts = None
    if rows is not None:
        texts = ["\n".join([row["url"]] + [text_value(v) for v in row["cells"].values()]).lower() for row in rows]
    hit = (columns, rows, texts)
    view.grids[key] = hit
    return hit

# --------------------------------------------------------------------------
# Several filters at once: ?filter=a,b,c&match=any|all
# --------------------------------------------------------------------------
FILTER_MATCHES = ("any", "all")
MAX_FILTERS = 64              # filter ids one request may combine

def _filter_ids(filter_param: str) -> List[str]:
    """The ids of ?filter=a,b,c in order, blanks and repeats dropped; ["all"] when there is none."""
    out: List[str] = []
    for part in str(filter_param or "").split(","):
        p = part.strip()
        if p and p not in out:
            out.append(p)
    return out or ["all"]

def _filter_match(match: str) -> str:
    m = (match or "any").strip().lower()
    if m not in FILTER_MATCHES:
        _bad("Unknown match. Use any (rows in at least one filter) or all (rows in every filter).")
    return m

def _filter_sel(tab_id: str, filter_param: str, match: str, index) -> Tuple[List[str], str]:
    """The validated filter selection of a grid request: (ids, match). One id is passed through as before
    (an unknown one is a 404 where the grid is built); in a list every id must exist (400 otherwise).
    "all" in a list is every row with match=any and is dropped with match=all (it limits nothing)."""
    ids, m = _filter_ids(filter_param), _filter_match(match)
    if len(ids) == 1:
        return ids, m
    if len(ids) > MAX_FILTERS:
        _bad(f"Too many filters: at most {MAX_FILTERS} at once.")
    tab = _tab_def(tab_id, index)
    if tab is None:
        raise HTTPException(status_code=404, detail=f"Unknown tab: {tab_id}")
    known = {f.id for f in tab.filters} | {"all"}
    unknown = [i for i in ids if i not in known]
    if unknown:
        _bad(f"Unknown filter{'s' if len(unknown) > 1 else ''} for {tab.name}: {', '.join(unknown[:10])}")
    if "all" in ids:
        ids = ["all"] if m == "any" else ([i for i in ids if i != "all"] or ["all"])
    return ids, m

def _sel_key(ids: List[str], match: str) -> str:
    """Cache key / echo of a selection: the id itself for one filter, "any:a,b" (sorted) for several."""
    return ids[0] if len(ids) == 1 else f"{match}:{','.join(sorted(ids))}"

def _sel_title(tab_def, ids: List[str], match: str) -> str:
    """"Missing + Duplicate (any)" style label of a selection for export titles."""
    from crawler.exporter import filter_name
    if len(ids) == 1:
        return filter_name(tab_def, ids[0])
    names = [filter_name(tab_def, i) for i in ids]
    return f"{' + '.join(names)} ({'in every filter' if match == 'all' else 'in any filter'})"

def _sel_slug(ids: List[str], match: str) -> str:
    """File-name part of a selection: the filter ids joined with "-", or "multiple-filters" when long."""
    if len(ids) == 1:
        return _safe_id(ids[0])
    parts = [re.sub(r"[^\w]+", "_", i).strip("_") or "f" for i in ids]
    slug = "-".join(parts)
    if len(parts) > 3 or len(slug) > 50:
        slug = "multiple-filters"
    return slug + ("_match-all" if match == "all" else "")

def _grid_sel(job: CrawlJob, site, index, tab_id: str, ids: List[str], match: str = "any"):
    """_grid for a selection of filters: the rows of the tab (in its own order) that are in any (match=any) or
    every (match=all) selected filter, one row per distinct URL. Built from the cached "All" grid; cached too."""
    if len(ids) == 1:
        return _grid(job, site, index, tab_id, ids[0])
    view = _view(job, index)
    key = (tab_id, _sel_key(ids, match))
    hit = view.grids.get(key)
    if hit is not None:
        return hit
    columns, all_rows, all_texts = _grid(job, site, index, tab_id, "all")
    tr = index.tabs.get(tab_id) if index is not None else None
    if all_rows is None or tr is None:
        return columns, None, None
    sets = []
    for fid in ids:
        items = tr.filter_rows.get(fid)
        if items is None:
            return columns, None, None
        sets.append({id(r) for r in items})
    keep = set.union(*sets) if match == "any" else set.intersection(*sets)
    rows, texts, seen = [], [], set()
    for i, r in enumerate(tr.rows):
        if id(r) not in keep or i >= len(all_rows):
            continue
        url = all_rows[i]["url"]
        if url in seen:            # distinct URLs: a URL listed twice in the tab counts once
            continue
        seen.add(url)
        rows.append(all_rows[i])
        texts.append(all_texts[i])
    hit = (columns, rows, texts)
    view.grids[key] = hit
    return hit

def _empty_tab_summary() -> list:
    """Every registered tab with zero counts, so the UI can draw its tab bar before a crawl."""
    try:
        from crawler.seo import registry
        idx = registry.TabIndex(tabs={t.id: registry.TabResult(tab=t, rows=[]) for t in registry.tabs()}, issues=[])
        return registry.tab_summary(idx)
    except Exception:
        return []

def _tab_def(tab_id: str, index=None):
    """The tab as this crawl shows it: the per-crawl copy for tabs built from the crawl's rules (Custom Search /
    Extraction / JavaScript names and columns), else the registered tab; None when unknown."""
    tr = index.tabs.get(tab_id) if index is not None else None
    if tr is not None and getattr(tr, "tab", None) is not None:
        return tr.tab
    try:
        from crawler.seo import registry
        return registry.get_tab(tab_id)
    except Exception:
        return None

def _registry_columns(tab_id: str, index=None) -> list:
    from crawler.exporter import user_label
    tab = _tab_def(tab_id, index)
    return [{"id": c.id, "label": user_label(c.label), "kind": c.kind} for c in tab.columns] if tab else []

def _is_blank(v) -> bool:
    return v is None or v == "" or (isinstance(v, (list, tuple)) and not v)

def _sort_key(v):
    """Numbers sort numerically (before text); text sorts case-insensitively."""
    if isinstance(v, (bool, int, float)):
        return (0, float(v), "")
    return (1, 0.0, text_value(v).strip().lower())

@app.get("/api/crawl/tabs")
def get_seo_tabs():
    """Tab bar + filter tree: every tab with its filter counts and the number of flagged issues."""
    job = session.job
    phase = job.phase if job and job.crawler else "idle"
    site, index = _site_index(job)
    if index is None:
        return {"phase": phase, "tabs": _empty_tab_summary(), "issues_total": 0, "dismissed_total": 0}
    from crawler.seo.registry import tab_summary
    d = _safe(lambda: _dismissals(job, index), _empty_dismissals(index))
    return {"phase": phase, "tabs": tab_summary(index), "issues_total": len(d["issues"]),
            "dismissed_total": len(d["fully_dismissed_issues"]),
            "summary": _safe(lambda: _indexable_summary(site), None),
            "checks_incomplete": bool(getattr(index, "checks_failed", None)),
            "checks_failed": dict(getattr(index, "checks_failed", None) or {})}

def _indexable_summary(site) -> Optional[dict]:
    """Overview summary rows: internal URLs that can / cannot appear in Google (they add up to the internal URLs,
    the same totals as Site Structure)."""
    if site is None:
        return None
    from crawler.seo.structure import _internal
    indexable = getattr(site, "indexable", None) or {}
    rows = _internal(site)
    n = sum(1 for r in rows if indexable.get(getattr(r, "url", "")))
    return {"internal": len(rows), "indexable": n, "non_indexable": len(rows) - n}

def _grid_view(job: CrawlJob, site, index, tab_id: str, filter_id: Union[str, List[str]], q: str = "",
               sort: str = "", sort_dir: str = "asc", segment: str = "",
               match: str = "any") -> Tuple[List[dict], List[dict], List[int]]:
    """(columns, rows, order): a tab grid searched (q), limited to one segment and sorted, as row positions.
    filter_id is one filter id or a selection (a list of ids, see _filter_sel) combined with `match`.
    404 for an unknown tab or filter."""
    ids = list(filter_id) if isinstance(filter_id, (list, tuple)) else [filter_id or "all"]
    columns, rows, texts = _grid_sel(job, site, index, tab_id, ids, match)
    if rows is None:
        what = "tab" if not columns else "filter"
        raise HTTPException(status_code=404, detail=f"Unknown {what}: {tab_id if what == 'tab' else ','.join(ids)}")
    order = list(range(len(rows)))
    needle = (q or "").strip().lower()
    if needle:
        order = [i for i in order if needle in texts[i]]
    if segment:
        _segs, assignment = _segment_assignment(job, site, index)
        want = "" if segment == SEGMENT_NONE else segment
        order = [i for i in order if (assignment.get(rows[i]["url"]) or "") == want]
    col_ids = {c["id"] for c in columns}
    if sort in col_ids or sort == "url":
        def value(i):
            return rows[i]["url"] if sort == "url" else rows[i]["cells"].get(sort)
        filled = [i for i in order if not _is_blank(value(i))]
        blank = [i for i in order if _is_blank(value(i))]
        filled.sort(key=lambda i: _sort_key(value(i)), reverse=(sort_dir or "").lower() == "desc")
        order = filled + blank  # empty cells always last, like Screaming Frog
    return columns, rows, order

def _with_segments(job: CrawlJob, site, index, columns: List[dict], page: List[dict]) -> Tuple[List[dict], List[dict]]:
    """Adds a Segment column (when the site has segments) to one page of grid rows, without touching the
    cached rows."""
    segs, assignment = _segment_assignment(job, site, index)
    if not segs:
        return columns, page
    cols = list(columns) + [{"id": "segment", "label": "Segment", "kind": "text"}]
    out = []
    for row in page:
        r = dict(row)
        r["cells"] = dict(row.get("cells") or {}, segment=assignment.get(row["url"]) or SEGMENT_NONE)
        out.append(r)
    return cols, out

@app.get("/api/crawl/tab/{tab_id}")
def get_seo_tab(tab_id: str, filter_id: str = Query("all", alias="filter"), q: str = "", sort: str = "",
                sort_dir: str = Query("asc", alias="dir"), offset: int = 0, limit: int = 200, view: str = "list",
                path: str = "", segment: str = "", format: str = "", match: str = "any"):
    """One page of a tab grid, filtered, searched (q: substring of the URL or any cell), limited to a segment
    (segment=<name>, or "(No segment)") and sorted server-side. filter= takes one filter id or several
    ("a,b,c"): match=any (default) lists the URLs in at least one of them, match=all the URLs in every one,
    each URL once. view=tree groups the rows by URL folder: the folders directly under `path` with their row
    counts, and the rows in that folder. format=csv|xlsx downloads every row of the view (search, segment and
    sort applied) instead. "filter_total" is the selection's URL count before search and segment."""
    job = session.job
    site, index = _site_index(job)
    ids, m = _filter_sel(tab_id, filter_id, match, index)
    filter_id = ids[0] if len(ids) == 1 else ",".join(ids)
    multi = {"filters": ids, "match": m} if len(ids) > 1 else {}
    fmt = (format or "").strip().lower()
    if fmt:
        return _tab_download(job, site, index, tab_id, ids, fmt, q, sort, sort_dir, segment, m)
    if index is None:
        return dict({"tab": tab_id, "filter": filter_id, "columns": _registry_columns(tab_id), "total": 0,
                     "rows": [], "view": view, "filter_total": 0}, **multi)
    columns, rows, order = _grid_view(job, site, index, tab_id, ids, q, sort, sort_dir, segment, m)
    offset = max(0, int(offset or 0))
    limit = max(1, min(int(limit or 200), MAX_PAGE_SIZE))
    tr = index.tabs.get(tab_id)
    nc_map = (getattr(tr, "not_checked", None) or {}) if tr is not None else {}
    not_checked = nc_map.get(ids[0], "") if len(ids) == 1 else ""
    if len(ids) > 1:
        # a selected filter that could not run is named, so "0 rows" is never read as "none found"
        multi["not_checked_filters"] = {i: nc_map[i] for i in ids if nc_map.get(i)}
    if (view or "").strip().lower() == "tree":
        from crawler.seo.structure import tree_view
        tree = tree_view([rows[i] for i in order], path or "")
        page = list(tree.get("rows") or [])
        cols, page = _with_segments(job, site, index, columns, page[offset:offset + limit])
        return dict({"tab": tab_id, "filter": filter_id, "columns": cols, "view": "tree", "path": tree.get("path", ""),
                     "parent": tree.get("parent"), "folders": tree.get("folders") or [], "total": tree.get("total", 0),
                     "rows_total": len(tree.get("rows") or []), "rows": page, "not_checked": not_checked,
                     "filter_total": len(rows)}, **multi)
    cols, page = _with_segments(job, site, index, columns, [rows[i] for i in order[offset:offset + limit]])
    return dict({"tab": tab_id, "filter": filter_id, "columns": cols, "total": len(order), "rows": page,
                 "view": "list", "not_checked": not_checked, "filter_total": len(rows)}, **multi)

def _tab_download(job, site, index, tab_id: str, filter_id: Union[str, List[str]], fmt: str, q: str = "",
                  sort: str = "", sort_dir: str = "asc", segment: str = "", match: str = "any") -> Response:
    """CSV / XLSX of one tab grid view (the tab's columns, every row of the filter or filter selection)."""
    from crawler.exporter import check_format, grid_table, tables_bytes, safe_name
    try:
        check_format(fmt)
    except ValueError:
        _bad("Unknown format. Use csv or xlsx.")
    ids = list(filter_id) if isinstance(filter_id, (list, tuple)) else [filter_id or "all"]
    tab_def = _tab_def(tab_id, index)
    if tab_def is None:
        raise HTTPException(status_code=404, detail=f"Unknown tab: {tab_id}")
    for fid in ids:
        if fid != "all" and not any(f.id == fid for f in tab_def.filters):
            raise HTTPException(status_code=404, detail=f"Unknown filter: {fid}")
    columns, rows = _registry_columns(tab_id, index), []
    if index is not None:
        columns, all_rows, order = _grid_view(job, site, index, tab_id, ids, q, sort, sort_dir, segment, match)
        columns, rows = _with_segments(job, site, index, columns, [all_rows[i] for i in order])
    title = f"{tab_def.name} - {_sel_title(tab_def, ids, match)}" + (f" - {segment}" if segment else "")
    stem = f"{safe_name(tab_id)}_{safe_name(ids[0]) if len(ids) == 1 else _sel_slug(ids, match)}"
    name, data, mime = tables_bytes([grid_table(title[:200], columns, rows)], fmt, stem)
    return _download(name, data, mime)

@app.get("/api/crawl/issues-overview")
def get_issues_overview():
    """Screaming Frog's Issues tab: one row per flagged filter, most important first. URLs the team dismissed
    drop out of the counts (each row says how many in "dismissed_urls"); issues dismissed completely move
    to "dismissed", which lists every dismissal with the number of URLs it covered."""
    job = session.job
    _, index = _site_index(job)
    if index is None:
        return {"issues": [], "dismissed": []}
    d = _safe(lambda: _dismissals(job, index), _empty_dismissals(index))
    from crawler.seo.registry import checks_warning
    failed = dict(getattr(index, "checks_failed", None) or {})
    return {"issues": list(d["issues"]), "dismissed": list(d["dismissed"]),
            "checks_incomplete": bool(failed), "checks_failed": failed, "checks_warning": checks_warning(failed)}

def _url_issues(view: ViewCache) -> Dict[str, List[dict]]:
    """url -> every issue filter that lists it (built once per index)."""
    if view.url_issues is not None:
        return view.url_issues
    out: Dict[str, List[dict]] = {}
    for tab_id, tr in view.index.tabs.items():
        tab = tr.tab
        for f in tab.filters:
            if not f.issue_type:
                continue
            entry = {"tab": tab.id, "tab_name": tab.name, "filter": f.id, "name": f"{tab.name}: {f.name}",
                     "issue_type": f.issue_type, "priority": f.priority or "Low",
                     "description": f.description, "fix": f.fix}
            for r in tr.filter_rows.get(f.id, []):
                url = getattr(r, "url", "")
                bucket = out.setdefault(url, [])
                if not any(e["tab"] == tab.id and e["filter"] == f.id for e in bucket):
                    bucket.append(entry)
    for bucket in out.values():
        bucket.sort(key=lambda e: (PRIORITY_ORDER.get(e["priority"], 3), TYPE_ORDER.get(e["issue_type"], 3)))
    view.url_issues = out
    return out

def _externals(view: ViewCache, site) -> Dict[str, Any]:
    if view.ext_by_url is None:
        ext: Dict[str, Any] = {}
        for e in site.externals or []:
            ext.setdefault(getattr(e, "url", ""), e)
        view.ext_by_url = ext
    return view.ext_by_url

def _resource_inlinks(view: ViewCache, site) -> Dict[str, List[dict]]:
    """url -> pages that embed it as an image, stylesheet, script or iframe (Screaming Frog lists these
    as inlinks of the resource, next to normal hyperlinks). Built once per index."""
    if view.res_inlinks is not None:
        return view.res_inlinks
    out: Dict[str, List[dict]] = {}
    for p in site.pages:
        seo = getattr(p, "seo", None) or {}
        alts = {i.get("src"): i.get("alt") for i in seo.get("images") or [] if i.get("src")}
        seen = set()
        for res in seo.get("resources") or []:
            url = res.get("url") or ""
            if not url or url in seen:
                continue
            seen.add(url)
            kind = resource_type_label(res.get("type"))     # the same Type as All Inlinks (JavaScript, CSS...)
            out.setdefault(url, []).append({"source": p.url, "anchor": alts.get(url) or "", "nofollow": False,
                                            "region": kind, "line": int(res.get("line") or 0), "type": kind})
        for bg in seo.get("background_images") or []:      # CSS background images of the page's HTML
            if isinstance(bg, str) and bg and bg not in seen and not bg.startswith("data:"):
                seen.add(bg)
                out.setdefault(bg, []).append({"source": p.url, "anchor": "", "nofollow": False,
                                               "region": "CSS background image", "line": 0, "type": "Background Image"})
    view.res_inlinks = out
    return out

def _url_candidates(u: str) -> List[str]:
    u = (u or "").strip()
    out = [u]
    try:
        n = normalize_url(u)
        if n:
            out.append(n)
    except Exception:
        pass
    for c in list(out):
        out += [c.rstrip("/"), c + "/"] if not c.endswith("/") else [c.rstrip("/")]
    return [c for i, c in enumerate(out) if c and c not in out[:i]]

def _status_of(url: str, site, ext: Dict[str, Any]) -> Tuple[Optional[int], str]:
    """(status code, category) of a linked URL from the crawl, (None, "") when it was not requested.
    The code is the URL's own response (301 for a redirect), as Screaming Frog shows it."""
    target = site.by_url.get(url) or ext.get(url)
    if target is None:
        return None, ""
    return row_status_code(target), row_category(target)

def _size_of(target) -> Optional[int]:
    if target is None:
        return None
    seo = getattr(target, "seo", None) or {}
    if seo.get("size_bytes"):
        return int(seo["size_bytes"])
    length = (getattr(target, "response_headers", None) or {}).get("content-length", "")
    return int(length) if str(length).isdigit() else None

def _empty_url_detail() -> dict:
    return {
        "result": {}, "seo": {}, "headers": {}, "indexable": False, "indexability_status": "",
        "link_score": 0, "crawl_depth": 0, "unique_inlinks": 0, "inlinks": [], "outlinks": [], "images": [],
        "resources": [], "issues": [],
        "serp": {"title": "", "title_px": 0, "description": "", "description_px": 0, "url": ""},
        "evidence": {}, "duplicates": {"title": [], "meta_description": [], "h1": [], "content": [], "near": []},
        "cookies": [], "rendered": None, "source_stored": False, "lighthouse": None,
    }

@app.get("/api/crawl/url")
def get_url_detail(u: str = ""):
    """Everything known about one URL: the lower detail pane (URL details, inlinks, outlinks,
    images, resources, SERP snippet, headers, duplicates, issues and screenshot evidence)."""
    job = session.job
    site, index = _site_index(job)
    detail = _empty_url_detail()
    if index is None or not u:
        return detail
    view = _view(job, index)
    ext = _externals(view, site)
    cands = _url_candidates(u)
    r = next((site.by_url[c] for c in cands if c in site.by_url), None)
    if r is None:
        e = next((ext[c] for c in cands if c in ext), None)
        if e is None:
            return detail
        return _external_detail(e, site, view)

    url = r.url
    seo = getattr(r, "seo", None) or {}
    detail["result"] = r.to_dict()
    detail["status_code"] = row_status_code(r)      # the URL's own code (301 for a redirect, not the 200 it leads to)
    detail["category"] = row_category(r)
    # big lists stay out of the pane's payload: the rendered page has its own endpoint (/api/crawl/render)
    detail["seo"] = json_value({k: v for k, v in seo.items() if k not in DETAIL_HEAVY_KEYS})
    detail["seo"]["ids_count"] = len(seo.get("ids") or [])
    detail["cookies"] = _page_cookies(seo)
    detail["rendered"] = _rendered_summary(job, url, seo)
    detail["source_stored"] = _source_available(job, url)
    detail["lighthouse"] = _safe(lambda: _lighthouse(site, url), None)
    seg = _segment_of_url(job, site, index, url)
    if seg is not None:
        detail["segment"] = seg
    detail["headers"] = safe_headers(getattr(r, "response_headers", None), job.secrets)
    detail["indexable"] = bool(site.indexable.get(url, False))
    detail["indexability_status"] = site.indexability_status.get(url, "")
    detail["link_score"] = int(site.link_score.get(url, 0) or 0)
    detail["crawl_depth"] = int(site.crawl_depth.get(url, getattr(r, "depth", 0)) or 0)
    detail["unique_inlinks"] = int(site.unique_inlinks.get(url, 0))
    detail["inlinks"] = [dict(i) for i in site.inlinks.get(url, [])]
    embeds = _resource_inlinks(view, site).get(url, [])
    if embeds:
        detail["inlinks"] += [dict(i) for i in embeds]   # unique_inlinks stays hyperlinks only, like the grids
    # URLs that redirect here (Screaming Frog lists them as inlinks of type HTTP Redirect)
    detail["inlinks"] += [{"source": e["source"], "anchor": "", "nofollow": False, "region": "HTTP redirect",
                           "line": 0, "type": "HTTP Redirect"}
                          for e in sources_of(site, url) if e.get("type") == "HTTP Redirect"]

    outlinks = []
    for link in seo.get("links") or []:
        code, cat = _status_of(link.get("url") or "", site, ext)
        outlinks.append({"url": link.get("url") or "", "anchor": link.get("anchor") or "",
                         "nofollow": bool(link.get("nofollow")), "internal": bool(link.get("internal")),
                         "status_code": code, "category": cat, "region": link.get("region") or "",
                         "rel": _link_rel(link), "type": "Hyperlink"})
    detail["outlinks"] = outlinks
    for i in detail["inlinks"]:
        i.setdefault("type", "Hyperlink")
    # W5: every other way pages point here / this page points elsewhere (canonical, hreflang, rel next / prev,
    # refreshes, JavaScript redirects...): kept apart so the inlink counts stay hyperlink-based
    detail["other_inlinks"] = _safe(lambda: [dict(e) for e in _other_sources_of(site, url)], [])
    detail["other_outlinks"] = _safe(lambda: [dict(e) for e in _other_outlinks_of(site, r)], [])
    # W5: the details sections every owner registered (structured data, accessibility, spelling, carbon, pdf...)
    detail["sections"] = _safe(lambda: json_value(_detail_sections(site, url)), {})

    images = []
    for img in seo.get("images") or []:
        src = img.get("src") or ""
        target = site.by_url.get(src) or ext.get(src)
        images.append({"src": src, "alt": img.get("alt"), "has_width": bool(img.get("has_width")),
                       "has_height": bool(img.get("has_height")), "size_bytes": _size_of(target),
                       "status_code": row_status_code(target) if target is not None else None})
    detail["images"] = images

    resources = []
    for res in seo.get("resources") or []:
        code, cat = _status_of(res.get("url") or "", site, ext)
        resources.append({"url": res.get("url") or "", "type": res.get("type") or "other",
                          "status_code": code, "category": cat})
    detail["resources"] = resources

    detail["issues"] = [dict(i) for i in _url_issues(view).get(url, [])]

    titles, descs = seo.get("titles") or [], seo.get("meta_descriptions") or []
    # Only a real <title>: the crawler's placeholders ("[Asset] x.pdf", "[Blocked by robots.txt]") are not titles
    fallback = r.page_title or ""
    title = titles[0] if titles else ("" if not seo.get("is_html") or fallback.startswith("[") else fallback)
    desc = descs[0] if descs else ""
    detail["serp"] = {"title": title, "title_px": _px(title, "title"), "description": desc,
                      "description_px": _px(desc, "description"), "url": seo.get("final_url") or url}

    ev = job.evidence.get(url) or {}
    if ev:
        detail["evidence"] = {
            "parent_shot": f"/api/crawl/evidence/{ev['parent_shot']}" if ev.get("parent_shot") else None,
            "error_shot": f"/api/crawl/evidence/{ev['error_shot']}" if ev.get("error_shot") else None,
            "position": ev.get("position"),
        }

    dups = site.duplicates or {}
    first_of = {"title": titles, "meta_description": descs, "h1": seo.get("h1") or []}
    for kind, values in first_of.items():
        value = values[0].strip() if values and isinstance(values[0], str) else ""
        detail["duplicates"][kind] = [x for x in dups.get(kind, {}).get(value, []) if x != url] if value else []
    content_hash = seo.get("content_hash") or ""
    detail["duplicates"]["content"] = (
        [x for x in dups.get("content_hash", {}).get(content_hash, []) if x != url] if content_hash else [])
    detail["duplicates"]["near"] = [[n, p] for n, p in (site.near_duplicates or {}).get(url, [])]
    detail["duplicates"]["near_count"] = int((getattr(site, "near_duplicate_counts", None) or {}).get(url)
                                             or len(detail["duplicates"]["near"]))   # the list keeps the closest
    return detail

def _link_rel(link) -> str:
    from crawler.exporter import link_rel
    return link_rel(link)


def _other_sources_of(site, url: str):
    from crawler.exporter import other_sources_of
    return other_sources_of(site, url)


def _other_outlinks_of(site, r):
    from crawler.exporter import other_outlinks_of
    return other_outlinks_of(site, r)


def _detail_sections(site, url: str) -> Dict[str, Any]:
    """{section: dict | {"not_checked": reason}} from crawler.seo.registry.details (W5): no per-module list."""
    from crawler.seo import registry
    return registry.details(site, url)


def _px(text: str, kind: str) -> int:
    try:
        return int(pixel_width(text, kind)) if text else 0
    except Exception:
        return 0

def _external_detail(e, site, view: ViewCache) -> dict:
    """Detail view for an outbound (external) URL: where it is linked from and its checked status."""
    detail = _empty_url_detail()
    url = getattr(e, "url", "")
    detail["result"] = e.to_dict()
    detail["result"]["category"] = row_category(e)
    detail["indexability_status"] = "External"
    inlinks, sources = [], set()
    for src, links in (site.outlinks or {}).items():
        for link in links or []:
            if link.get("url") == url:
                inlinks.append({"source": src, "anchor": link.get("anchor") or "", "nofollow": bool(link.get("nofollow")),
                                "region": link.get("region") or "", "line": int(link.get("line") or 0)})
                sources.add(src)
    detail["inlinks"] = inlinks
    detail["unique_inlinks"] = len(sources)
    detail["issues"] = [dict(i) for i in _url_issues(view).get(url, [])]
    detail["serp"]["url"] = url
    return detail

def _safe_id(text: str) -> str:
    return re.sub(r"[^\w-]+", "_", text or "")[:60] or "all"

def _xlsx_path(job: Optional[CrawlJob], suffix: str) -> str:
    base = job.base_name if job else "crawl_no_data"
    return os.path.join(get_export_dir(), f"{base}_{suffix}.xlsx")

_NO_JOB_EXCEL_LOCK = threading.RLock()

def _excel_lock(job: Optional[CrawlJob]):
    """The lock that serialises workbook builds of this crawl (two clicks never write one file at once)."""
    return job.excel_lock if job is not None else _NO_JOB_EXCEL_LOCK

def _fmt(format: str) -> str:
    """xlsx (default) or csv; anything else is a 400."""
    fmt = (format or "xlsx").strip().lower()
    if fmt not in ("xlsx", "csv"):
        _bad("Unknown format. Use xlsx or csv.")
    return fmt

def _file_download(path: str, name: str) -> FileResponse:
    """A written export file; the download name keeps the extension the exporter chose (.csv / .zip)."""
    ext = os.path.splitext(path)[1].lower()
    stem = os.path.splitext(name)[0]
    return FileResponse(path=path, filename=stem + ext, media_type=FILE_MEDIA.get(ext, "application/octet-stream"))

def _tab_and_filter(tab: str, filter_id: str, index):
    tab_def = _tab_def(tab, index)
    if tab_def is None:
        raise HTTPException(status_code=404, detail=f"Unknown tab: {tab}")
    if filter_id != "all" and not any(f.id == filter_id for f in tab_def.filters):
        raise HTTPException(status_code=404, detail=f"Unknown filter: {filter_id}")
    fname = tab_def.total_label if filter_id == "all" else next(f.name for f in tab_def.filters if f.id == filter_id)
    return tab_def, fname

@app.get("/api/crawl/export/tab")
def export_seo_tab(tab: str, filter_id: str = Query("all", alias="filter"), format: str = "xlsx",
                   segment: str = "", q: str = "", match: str = "any", sort: str = "",
                   sort_dir: str = Query("asc", alias="dir")):
    """Excel (or CSV) of one tab grid: every row of the chosen filter with the tab's columns; segment=, q= and
    sort= limit and order it like the grid. filter= takes several ids ("a,b") with match=any|all like the grid:
    the file then has exactly the grid's rows (each URL once) and its name lists the filters."""
    fmt = _fmt(format)
    job = session.job
    site, index = _site_index(job)
    ids, m = _filter_sel(tab, filter_id, match, index)
    if len(ids) > 1:
        tab_def = _tab_def(tab, index)
        columns, rows = _registry_columns(tab, index), []
        if index is not None:
            columns, all_rows, order = _grid_view(job, site, index, tab, ids, q, sort, sort_dir, segment, m)
            columns, rows = _with_segments(job, site, index, columns, [all_rows[i] for i in order])
        title = f"{tab_def.name} - {_sel_title(tab_def, ids, m)}" + (f" - {segment}" if segment else "")
        path = _xlsx_path(job, f"{_safe_id(tab)}_{_sel_slug(ids, m)}")
        with _excel_lock(job):
            out = ExcelExporter.export_tab(path, title[:200], columns, rows, fmt=fmt)
        return _file_download(out, os.path.basename(path))
    filter_id = ids[0]
    tab_def, fname = _tab_and_filter(tab, filter_id, index)
    if segment or q or sort:
        return _tab_download(job, site, index, tab, filter_id, fmt, q=q, sort=sort, sort_dir=sort_dir,
                             segment=segment)
    columns, rows = _registry_columns(tab, index), []
    if index is not None:
        columns, rows, _ = _grid(job, site, index, tab, filter_id)
        rows = rows or []
        columns, rows = _with_segments(job, site, index, columns, rows)
    path = _xlsx_path(job, f"{_safe_id(tab)}_{_safe_id(filter_id)}")
    with _excel_lock(job):
        out = ExcelExporter.export_tab(path, f"{tab_def.name} - {fname}", columns, rows, fmt=fmt)
    return _file_download(out, os.path.basename(path))

@app.get("/api/crawl/export/issues")
def export_issues_overview(format: str = "xlsx"):
    """Excel (or CSV) of the issues overview (name, type, priority, URLs, % of total, description, fix)."""
    fmt = _fmt(format)
    job = session.job
    _, index = _site_index(job)
    path = _xlsx_path(job, "issues_overview")
    issues = _safe(lambda: _dismissals(job, index), _empty_dismissals(index))["issues"] if index is not None else []
    with _excel_lock(job):
        out = ExcelExporter.export_issues_overview(path, list(issues), fmt=fmt)
    return _file_download(out, os.path.basename(path))

@app.get("/api/crawl/export/inlinks")
def export_all_inlinks(format: str = "xlsx"):
    """Bulk export of every internal link and embed: source, destination, anchor, nofollow, position,
    line, destination status and type (Hyperlink, Image, CSS, JavaScript, Iframe, Other, Background Image, and
    HTTP Redirect with the redirecting URL as the source)."""
    fmt = _fmt(format)
    job = session.job
    site, _ = _site_index(job)
    path = _xlsx_path(job, "all_inlinks")
    with _excel_lock(job):
        out = ExcelExporter.export_inlinks(path, inlink_rows(site), fmt=fmt)
    return _file_download(out, os.path.basename(path))

@app.get("/api/crawl/export/filter-inlinks")
def export_filter_inlinks(tab: str, filter_id: str = Query("all", alias="filter"), format: str = "xlsx",
                          match: str = "any"):
    """For every URL of one tab filter, every page (or redirect) that links to it or embeds it: the
    pages to edit for broken internal or external links, images without alt text, redirects...
    filter= takes several ids ("a,b") with match=any|all like the grid."""
    fmt = _fmt(format)
    job = session.job
    site, index = _site_index(job)
    ids, m = _filter_sel(tab, filter_id, match, index)
    if len(ids) > 1:
        tab_def = _tab_def(tab, index)
        fname = _sel_title(tab_def, ids, m)
        urls = set()
        if index is not None:
            _c, sel_rows, _t = _grid_sel(job, site, index, tab, ids, m)
            urls = {r["url"] for r in sel_rows or []}
        rows = [r for r in (filter_source_rows(site, index, tab, "all") or []) if r and r[0] in urls]
        slug = _sel_slug(ids, m)
    else:
        filter_id = ids[0]
        tab_def, fname = _tab_and_filter(tab, filter_id, index)
        rows = filter_source_rows(site, index, tab, filter_id) or []
        slug = _safe_id(filter_id)
    path = _xlsx_path(job, f"{_safe_id(tab)}_{slug}_inlinks")
    with _excel_lock(job):
        out = ExcelExporter.export_rows(path, f"{tab_def.name} {fname} inlinks", SOURCE_HEADERS, rows,
                                        widths=(60, 60, 14, 36, 12, 10, 20, 12, 12, 24), fmt=fmt)
    return _file_download(out, os.path.basename(path))

@app.get("/api/crawl/export/all-issues")
def export_all_issue_urls(format: str = "xlsx"):
    """Every URL of every flagged issue: issue name, type, priority, URL, status code, tab and filter."""
    fmt = _fmt(format)
    job = session.job
    _, index = _site_index(job)
    path = _xlsx_path(job, "all_issue_urls")
    d = _safe(lambda: _dismissals(job, index), _empty_dismissals(index))
    rows = issue_url_rows(index, d["issues"] if d["entries"] else None, d["dismissed_urls"] or None)
    with _excel_lock(job):
        out = ExcelExporter.export_rows(path, "All Issue URLs", ISSUE_URL_HEADERS, rows,
                                        widths=(48, 13, 14, 70, 11, 18, 40), fmt=fmt)
    return _file_download(out, os.path.basename(path))

@app.get("/api/crawl/download/{filename}")
async def download_by_filename(filename: str):
    safe_filename = os.path.basename(filename)
    # Only serve report files this app generated, never arbitrary temp files
    if not REPORT_FILE_RE.match(safe_filename):
        raise HTTPException(status_code=404, detail="File not found or expired.")
    filepath = os.path.join(get_export_dir(), safe_filename)
    job = session.job
    if (job is not None and job.crawler is not None and job.phase == "done" and not os.path.exists(filepath)
            and safe_filename == f"{job.base_name}.xlsx"):
        # the workbook is built after "done" (or on request): build / wait for it in the thread pool
        try:
            filepath = await asyncio.get_running_loop().run_in_executor(None, _excel_file, job, "all")
        except Exception as e:
            raise HTTPException(status_code=500, detail=job.mask(f"The Excel report could not be built: {e}"))
    for _ in range(30):
        if os.path.exists(filepath) and os.path.getsize(filepath) > 0:
            break
        await asyncio.sleep(0.5)
    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="File not found or expired.")
    media = FILE_MEDIA.get(os.path.splitext(safe_filename)[1].lower(), XLSX_MEDIA)
    return FileResponse(path=filepath, filename=safe_filename, media_type=media)

@app.get("/api/crawl/export")
def download_export():
    job = session.job
    if job and job.crawler is not None and not job.export_file and job.phase == "done" \
            and job.crawler.status != "auth_failed":
        try:   # built after "done" in the background; a click before it finished builds (or waits) here
            _excel_file(job, "all")
        except Exception as e:
            job.crawler.log(f"Could not build the Excel report: {job.mask(e)}")
    if job and job.export_file:
        filepath = job.export_file
        if job.export_stale and job.phase == "done" and job.crawler is not None:
            try:   # validation, intel or a re-check finished since the workbook was saved
                filepath = _excel_file(job, "all")
            except Exception as e:
                job.crawler.log(f"Could not refresh the Excel report: {e}")
        for _ in range(30):
            if os.path.exists(filepath) and os.path.getsize(filepath) > 0:
                break
            time.sleep(0.5)
        if os.path.exists(filepath):
            return FileResponse(path=filepath, filename=f"{job.base_name}.xlsx", media_type=XLSX_MEDIA)
    raise HTTPException(status_code=404, detail="No export file available.")

# --------------------------------------------------------------------------
# Login: test a login, and the live login window (a remote-controlled Chrome)
# --------------------------------------------------------------------------
class AuthTestRequest(BaseModel):
    start_url: Optional[str] = None
    url: Optional[str] = None                  # same as start_url
    auth: Optional[AuthRequest] = None
    user_agent: Optional[str] = "chrome"
    custom_user_agent: Optional[str] = None
    timeout: Optional[int] = 15
    verify_ssl: Optional[bool] = False

AUTH_TEST_KEYS = {"ok": False, "method": "", "landing_url": "", "cookies_count": 0, "storage_origins": 0,
                  "message": "", "screenshot": ""}

@app.post("/api/auth/test")
def auth_test_endpoint(req: AuthTestRequest):
    """Logs in with the given settings, checks that a protected page opens, then closes everything.
    Passwords, cookie values and header values are never returned. "challenge" says whether a bot check stood
    in the way ({seen, provider, count, first_url, advice, owner_steps}); ARACHNIX never tries to pass one."""
    raw_url = (req.start_url or req.url or "").strip()
    if not raw_url:
        raise HTTPException(status_code=400, detail="Enter the website address to log in to.")
    if not req.auth or (req.auth.auth_type or "none") == "none":
        raise HTTPException(status_code=400, detail="Choose a login method first.")
    _, ua = _user_agent(req)
    config = _auth_config(req.auth, force=True)
    secrets = _auth_secrets(config, _normalize_target(raw_url))
    from crawler.auth import test_login
    try:
        out = test_login(config, _normalize_target(raw_url), user_agent=ua,
                         timeout=max(5, min(int(req.timeout or 15), 60)), verify_ssl=bool(req.verify_ssl),
                         screenshot_dir=LOGIN_TEST_DIR)
    except Exception as e:
        out = {"ok": False, "message": f"The login test could not run: {type(e).__name__}: {str(e)[:200]}"}
    result = {k: (out or {}).get(k, default) for k, default in AUTH_TEST_KEYS.items()}
    result["message"] = mask_secrets(result["message"], secrets)
    result["landing_url"] = mask_secrets(result["landing_url"], secrets)
    result["challenge"] = _masked_challenge((out or {}).get("challenge"), lambda t: mask_secrets(t, secrets))
    return result

# --------------------------------------------------------------------------
# Paste cookies: a masked preview of what was read, and "Test these cookies" (one page that needs sign-in)
# --------------------------------------------------------------------------
MAX_PASTE_CHARS = 2_000_000    # crawler/cookie_import.py reads no more than this

class CookiePasteRequest(BaseModel):
    # a Cookie header, a JSON export (as text, a list or an object), a Netscape cookies.txt, a "Copy as cURL"
    # command or a "Copy as PowerShell" snippet
    cookies: Optional[Union[str, list, dict]] = None
    start_url: Optional[str] = None
    url: Optional[str] = None            # same as start_url
    test_url: Optional[str] = None       # the test: a page (or a path such as /dashboard) that needs sign-in
    user_agent: Optional[str] = None     # the test: a User-Agent string or a preset name (default: the pasted one)
    timeout: Optional[int] = 15
    verify_ssl: Optional[bool] = False
    handoff_id: Optional[str] = None     # W6: read the cookies of a browser-extension hand-off (not consumed)

def _cookie_paste(req: CookiePasteRequest, need_url: bool) -> Tuple[str, str, List[str]]:
    """(paste text, start URL or '', values to mask) of a cookie request; an over-long paste is a 400, and so
    is a missing website address when `need_url`."""
    _extension.apply_to_paste(req)
    raw = _paste_text(req.cookies)
    if len(raw) > MAX_PASTE_CHARS:
        _bad(f"The paste is longer than {MAX_PASTE_CHARS:,} characters. Copy only this site's cookies.")
    start = (req.start_url or req.url or "").strip()
    if need_url and not start:
        _bad("Enter the website address first.")
    start = _normalize_target(start) if start else ""
    from crawler.url_utils import split_url_credentials
    start = split_url_credentials(start)[0] if start else ""
    secrets = _auth_secrets(AuthConfig(enabled=True, auth_type="cookies", cookies=raw), start)
    return raw, start, secrets

@app.post("/api/auth/cookies/preview")
def cookies_preview(req: CookiePasteRequest):
    """What a cookie paste holds, before it is used: {format, count, usable, cookies: [{name, domain, path,
    secure, http_only, expires, expired, other_site, kind: "signin"|"other", value: "***"}], user_agent,
    warnings, error, storage_origins}. No request is made, and no value is ever returned."""
    raw, start, secrets = _cookie_paste(req, need_url=False)
    from crawler.auth import preview_cookies
    try:
        out = preview_cookies(raw, start)
    except Exception as e:
        _bad(f"The paste could not be read ({type(e).__name__}). Copy the cookies again.")
    for row in out.get("cookies") or []:
        row["value"] = "***"
    out["warnings"] = [mask_secrets(w, secrets) for w in out.get("warnings") or []]
    out["error"] = mask_secrets(out.get("error") or "", secrets)
    return out

def _test_user_agent(value: Optional[str]) -> Optional[str]:
    """The cookie test's User-Agent: a preset name becomes its string; a custom string is used as given
    (one line, at most MAX_SIGNIN_UA_CHARS characters, else a 400); None = the pasted one or the default."""
    ua = (value or "").strip()
    if not ua:
        return None
    if ua.lower() in USER_AGENTS:
        return USER_AGENTS[ua.lower()]
    if len(ua) > MAX_SIGNIN_UA_CHARS or _CTRL_CHARS.search(ua):
        _bad(f"The User-Agent must be one line of at most {MAX_SIGNIN_UA_CHARS} characters.")
    return ua

@app.post("/api/auth/cookies/test")
def cookies_test(req: CookiePasteRequest):
    """"Test these cookies": opens one page that needs sign-in with the pasted cookies (and, when it opens, once
    without them) and answers {signed_in: true | false | null (couldn't tell), status, final_url, reason, test_url,
    user_agent_used: "given"|"pasted"|"default", challenge}. Two or three requests; a sign-out or data-changing
    address is never opened, and no cookie value is ever returned."""
    raw, start, secrets = _cookie_paste(req, need_url=True)
    ua = _test_user_agent(req.user_agent)
    test_url = (req.test_url or "").strip() or None
    if test_url and len(test_url) > 4000:
        _bad("The test page address is too long.")
    from crawler.auth import test_cookies
    try:
        out = test_cookies(raw, start, test_url=test_url, user_agent=ua,
                           timeout=max(5, min(int(req.timeout or 15), 60)), verify_ssl=bool(req.verify_ssl))
    except Exception as e:
        out = {"signed_in": None, "reason": f"Couldn't tell: the test could not run ({type(e).__name__})."}
    out = dict(out or {})
    status = out.get("status")
    return {"signed_in": out.get("signed_in") if out.get("signed_in") in (True, False) else None,
            "status": status if isinstance(status, int) else None,
            "final_url": mask_secrets(out.get("final_url") or "", secrets),
            "reason": mask_secrets(out.get("reason") or "", secrets),
            "test_url": mask_secrets(out.get("test_url") or "", secrets),
            "user_agent_used": out.get("user_agent_used") or "default",
            "challenge": _masked_challenge(out.get("challenge"), lambda t: mask_secrets(t, secrets))}

class LiveStartRequest(BaseModel):
    url: str
    width: Optional[int] = 0    # the viewer's screen: below 700 px wide the window opens in a phone layout
    height: Optional[int] = 0

class LiveSessionRequest(BaseModel):
    session_id: str

class LiveActionRequest(BaseModel):
    session_id: str
    type: str
    x: Optional[float] = 0
    y: Optional[float] = 0
    text: Optional[str] = None
    key: Optional[str] = None
    dy: Optional[float] = 0
    url: Optional[str] = None

def _live():
    from crawler.live_browser import LiveBrowserError, get_live_manager
    return get_live_manager(), LiveBrowserError

def _live_challenge(session_id: str) -> Dict[str, Any]:
    """A bot check the live window shows now ({seen: False, ...} when none, or when the window is gone). Only
    reads the page: ARACHNIX never clicks, types into or otherwise tries to pass a bot check."""
    try:
        from crawler.auth import live_window_bot_check
        return _masked_challenge(live_window_bot_check(session_id), str)
    except Exception:
        return _no_challenge()

# live-window inputs after which the page may have changed (a bot check is looked for again)
LIVE_PAGE_ACTIONS = ("click", "navigate", "back", "reload")

@app.post("/api/live/start")
def live_start(req: LiveStartRequest):
    """Opens the login page in a server-side Chrome the person drives through /api/live/frame and /action.
    "challenge" says whether the page shows a bot check ARACHNIX can't complete (then: Paste cookies)."""
    manager, LiveBrowserError = _live()
    try:
        out = manager.start(_normalize_target(req.url), width=max(0, int(req.width or 0)),
                            height=max(0, int(req.height or 0)))
    except LiveBrowserError as e:
        raise HTTPException(status_code=400, detail=str(e))
    out = dict(out or {})
    out["challenge"] = _live_challenge(out.get("session_id") or "")
    return out

@app.get("/api/live/challenge")
def live_challenge(session_id: str = "", x_live_session: str = Header("")):
    """Whether the live window shows a bot check now: the challenge object (send the id in X-Live-Session)."""
    session_id = x_live_session or session_id
    manager, LiveBrowserError = _live()
    try:
        manager.current_url(session_id)        # 404 for a closed window
    except LiveBrowserError as e:
        raise HTTPException(status_code=404, detail=str(e))
    return {"challenge": _live_challenge(session_id)}

@app.get("/api/live/frame")
def live_frame(session_id: str = "", x_live_session: str = Header("")):
    """The live window as a JPEG. Send the id in the X-Live-Session header: a query-string id ends up in proxy
    logs (this server's own access log masks it)."""
    session_id = x_live_session or session_id
    manager, LiveBrowserError = _live()
    try:
        data = manager.frame(session_id)
        current = manager.current_url(session_id)
    except LiveBrowserError as e:
        raise HTTPException(status_code=404, detail=str(e))
    # X-Live-Url keeps the window's address bar right after redirects (e.g. login -> dashboard)
    return Response(content=data, media_type="image/jpeg",
                    headers={"Cache-Control": "no-store", "X-Live-Url": quote(current[:2000], safe=":/?#[]@!$&'()*+,;=%~")})

@app.post("/api/live/action")
def live_action(req: LiveActionRequest):
    manager, LiveBrowserError = _live()
    act = {"type": req.type, "x": req.x or 0, "y": req.y or 0, "text": req.text or "", "key": req.key or "",
           "dy": req.dy or 0, "url": req.url or ""}
    try:
        out = manager.action(req.session_id, act)
    except LiveBrowserError as e:
        raise HTTPException(status_code=404, detail=str(e))
    answer = {"ok": bool(out.get("ok")), "current_url": out.get("current_url") or ""}
    # a click, Enter or a page load can bring up a bot check (typing and scrolling do not)
    if req.type in LIVE_PAGE_ACTIONS or (req.type == "key" and req.key == "Enter"):
        answer["challenge"] = _live_challenge(req.session_id)
    return answer

@app.post("/api/live/finish")
def live_finish(req: LiveSessionRequest):
    """Keeps the window's login for the next crawl (auth_type "live" with this live_session_id). "challenge":
    a bot check the window still showed when the person finished (the login may then not work)."""
    manager, LiveBrowserError = _live()
    challenge = _live_challenge(req.session_id)
    try:
        out = manager.finish(req.session_id)
    except LiveBrowserError as e:
        raise HTTPException(status_code=404, detail=str(e))
    return {"ok": bool(out.get("ok")), "cookies_count": int(out.get("cookies_count") or 0),
            "storage_origins": int(out.get("storage_origins") or 0), "current_url": out.get("current_url") or "",
            "challenge": challenge}

@app.post("/api/live/close")
def live_close(req: LiveSessionRequest):
    manager, _ = _live()
    try:
        return {"ok": bool(manager.close(req.session_id).get("ok"))}
    except Exception:
        return {"ok": False}

# --------------------------------------------------------------------------
# Find subdomains: other hosts of the audited site's domain (links, sitemaps, certificate logs, Search Console,
# typed in), each checked with one request; after a crawl, every host of the domain with a plain reason.
# Hosts outside the audited registrable domain are never probed (crawler/subdomains.py).
# --------------------------------------------------------------------------
class SubdomainDiscoverRequest(BaseModel):
    url: Optional[str] = None                     # the audited site (default: the loaded crawl's start URL)
    sources: Optional[List[str]] = None           # any of links, sitemaps, ct, gsc (default: all)
    extra_hosts: Optional[Union[List[str], str]] = None   # host names the person typed in

def _subdomains():
    from crawler import subdomains
    return subdomains

def _server_proxies() -> Optional[Dict[str, str]]:
    """Settings > Proxy as a requests proxies dict (None when no proxy is on). Never logged."""
    try:
        from crawler.engine import proxy_url
        p = proxy_url(_proxy_config() or {})
    except Exception:
        p = ""
    return {"http": p, "https": p} if p else None

def _loaded_crawler():
    job = session.job
    return job.crawler if job is not None else None

def _crawl_start_url(crawler) -> str:
    return str(getattr(crawler, "start_url", "") or getattr(getattr(crawler, "config", None), "start_url", "") or "")

@app.post("/api/subdomains/discover")
def subdomains_discover(req: Optional[SubdomainDiscoverRequest] = None):
    """Starts Find subdomains in the background: {job_id, domain}; poll GET /api/subdomains/discover/{job_id}.
    The loaded crawl's links and sitemaps are used when it is of the same domain (no request needed). At most two
    searches run at once on this server (429 otherwise); the same search asked again returns the running job."""
    sd = _subdomains()
    req = req or SubdomainDiscoverRequest()
    crawler = _loaded_crawler()
    raw = (req.url or "").strip() or (_crawl_start_url(crawler) if crawler is not None else "")
    if not raw:
        _bad("Enter the website address first.")
    if len(raw) > 2000:
        _bad("The website address is too long.")
    from crawler.url_utils import registrable_domain, split_url_credentials
    url = split_url_credentials(_normalize_target(raw))[0]     # never a user name or password
    typed = _host_values(req.extra_hosts, "Typed subdomains")
    sources = None
    if req.sources is not None:
        sources = list(dict.fromkeys(str(s).strip().lower() for s in req.sources if str(s or "").strip()))
        unknown = [s for s in sources if s not in sd.SOURCES]
        if unknown:
            _bad(f"Unknown place to look: {', '.join(unknown[:5])}. Use any of: {', '.join(sd.SOURCES)}.")
        if not sources and not typed:
            _bad("Choose at least one place to look, or type the subdomains to check.")
    domain = registrable_domain(url)
    same = crawler is not None and registrable_domain(_crawl_start_url(crawler)) == domain
    ua = (getattr(getattr(crawler, "config", None), "user_agent", "") if same else "") or DEFAULT_USER_AGENT
    try:
        job_id = sd.start_job(url, sources, typed, crawler=crawler if same else None, user_agent=ua,
                              proxies=_server_proxies())
    except sd.DiscoveryBusy as e:
        raise HTTPException(status_code=429, detail=str(e), headers={"Retry-After": "10"})
    return {"job_id": job_id, "domain": domain}

@app.get("/api/subdomains/discover/{job_id}")
def subdomains_discover_status(job_id: str):
    """{state: running|done|error, domain, hosts: [{host, sources, resolves, status, title, note, state, url,
    redirect}], errors: {source: message}, sources: {source: {state, hosts, note}}, message, done, total, ...}.
    While it runs, hosts holds the candidates found so far (state not_checked until probed)."""
    out = _subdomains().get_job(job_id)
    if out is None:
        raise HTTPException(status_code=404, detail="This subdomain search has ended or expired. "
                                                    "Run Find subdomains again.")
    return out

@app.post("/api/subdomains/discover/{job_id}/stop")
def subdomains_discover_stop(job_id: str):
    """Stops a running search (the hosts checked so far are kept): {"stopped": bool}."""
    sd = _subdomains()
    if sd.get_job(job_id) is None:
        raise HTTPException(status_code=404, detail="This subdomain search has ended or expired.")
    return {"stopped": sd.stop_job(job_id)}

@app.get("/api/subdomains/known")
def subdomains_known(url: str = ""):
    """What the searches run on this server found for a site's domain, merged (kept in memory until the server
    restarts): {domain, finished_at, hosts: [entry]}; hosts is empty when no search ran."""
    crawler = _loaded_crawler()
    raw = (url or "").strip() or (_crawl_start_url(crawler) if crawler is not None else "")
    from crawler.url_utils import registrable_domain
    domain = registrable_domain(_normalize_target(raw)) if raw else ""
    found = _subdomains().last_discovery(domain) if domain else None
    return found or {"domain": domain, "finished_at": "", "hosts": []}

@app.get("/api/crawl/subdomains")
def crawl_subdomains():
    """Every host of the crawled site's domain that ARACHNIX saw, crawled or not, with a plain reason: {domain,
    start_host, scope, mode, hosts: [{host, crawled, linked_from, in_scope, reason_code, reason, sources, extra,
    home_status, pages_ok}], counts, totals, discovery, complete}. Makes no request; works on saved crawls."""
    crawler = _loaded_crawler()
    try:
        return _subdomains().subdomain_report(crawler)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"The subdomain list could not be built ({type(e).__name__}).")

# --------------------------------------------------------------------------
# AI layer status (free tiers by default; every feature also works with AI off)
# --------------------------------------------------------------------------
@app.get("/api/ai/status")
def get_ai_status():
    try:
        from crawler.ai import client as ai_client
        return ai_client.ai_status()
    except Exception as e:
        return {"enabled": False, "provider": "", "model": "", "bulk_model": "", "private_ok": False,
                "reason": f"AI off - site facts only (the AI layer could not load: {type(e).__name__})",
                "usage": {"calls": 0, "input_tokens": 0, "output_tokens": 0, "provider": "", "est_cost_usd": 0.0}}

# --------------------------------------------------------------------------
# Business insights (with or without a crawl)
# --------------------------------------------------------------------------
class BusinessRequest(BaseModel):
    url: Optional[str] = None
    force: Optional[bool] = False        # True ("Regenerate"): research again even when a saved report exists
    regenerate: Optional[bool] = False   # the same as force

# The last business report per site (lower-case host without "www."), kept next to the crawls so "Generate business
# report" for a company already researched shows that report at once (existing: true, created_at) instead of
# researching it again; "Regenerate" (force) makes a new one, which replaces it.
BUSINESS_CACHE_DIR = os.path.join(REPORTS_DIR, "business_reports")
_business_cache_lock = threading.Lock()

def _business_key(url: str) -> str:
    from crawler.seo.compare import host_key
    return host_key(url)

def _business_cache_path(url: str) -> str:
    return os.path.join(BUSINESS_CACHE_DIR, f"{_business_key(url)}.json")

def _business_cached(url: str) -> Optional[dict]:
    """{"report", "created_at" (ISO 8601, UTC), "url"} of the saved report for this site, None when none."""
    path = _business_cache_path(url)
    try:
        with _business_cache_lock, open(path, encoding="utf-8") as f:
            data = json.load(f)
    except (OSError, ValueError):
        return None
    return data if isinstance(data, dict) and isinstance(data.get("report"), dict) else None

def _business_save(url: str, report: dict) -> str:
    """Saves the finished report as this site's last one; returns its created_at ('' when it could not be saved)."""
    created = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
    try:
        os.makedirs(BUSINESS_CACHE_DIR, exist_ok=True)
        path = _business_cache_path(url)
        tmp = f"{path}.{uuid.uuid4().hex[:6]}.tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump({"url": url, "key": _business_key(url), "created_at": created, "report": json_value(report)}, f)
        with _business_cache_lock:
            os.replace(tmp, path)
        return created
    except (OSError, TypeError, ValueError):
        return ""

def _same_site(a: str, b: str) -> bool:
    from crawler.seo.compare import host_key
    return bool(a and b) and host_key(a) == host_key(b)

@app.post("/api/business/start")
def start_business(req: Optional[BusinessRequest] = None):
    """Company profile, market, contacts, tech stack and SEO snapshot of a website (background job). Uses the
    finished crawl's data (pages, titles, schema, health score, issues) when it is the same site."""
    job = session.job
    crawler = job.crawler if job and job.crawler else None
    url = ((req.url if req else "") or "").strip() or (crawler.start_url if crawler else "")
    if not url:
        raise HTTPException(status_code=400, detail="Enter the website address to research.")
    url = _normalize_target(url)
    force = bool(req and (req.force or req.regenerate))
    use = crawler if (crawler is not None and job.phase == "done" and crawler.status != "auth_failed"
                      and _same_site(url, crawler.start_url)) else None
    mask = job.mask if job is not None else (lambda s: s)
    runner = session.business
    runner.mask = mask
    busy = HTTPException(status_code=409, detail="A business report is already being prepared. Wait for it to finish.")

    if not force:    # this company was researched before: show that report (the dashboard offers "Regenerate")
        shown = runner.result if runner.state == "done" and isinstance(runner.result, dict) else None
        same = shown is not None and _same_site(url, runner.meta.get("url") or shown.get("url") or "")
        if same and runner.meta.get("created_at"):
            runner.meta["existing"] = True
            return {"job_id": runner.job_id, "existing": True, "created_at": runner.meta["created_at"],
                    "url": runner.meta.get("url") or url, "report": shown}
        saved = _business_cached(url)
        if saved is not None:
            created = str(saved.get("created_at") or "")
            try:
                job_id = runner.set_done(saved["report"], "Saved report", url=saved.get("url") or url,
                                         created_at=created, existing=True)
            except JobBusy:
                raise busy
            return {"job_id": job_id, "existing": True, "created_at": created, "url": saved.get("url") or url,
                    "report": saved["report"]}

    def work(runner: JobRunner) -> dict:
        from crawler.ai.business import build_business_report
        report = build_business_report(url, crawler=use, progress=lambda pct, msg: runner.update(pct, msg))
        runner.meta["created_at"] = _business_save(url, report) or time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
        runner.meta["existing"] = False
        runner.update(None, report.get("note") or "Done")
        return report

    try:
        return {"job_id": runner.start(work, "Starting...", url=url), "existing": False}
    except JobBusy:
        raise busy

@app.get("/api/business/status")
def business_status():
    runner = session.business
    out = runner.status()
    out["report"] = runner.result if runner.state == "done" and isinstance(runner.result, dict) else None
    if out["report"] is not None:
        out["created_at"] = runner.meta.get("created_at") or ""
        out["existing"] = bool(runner.meta.get("existing"))
    return out

@app.get("/api/business/pdf")
def business_pdf():
    """The compact business insights PDF of the last finished business report."""
    runner = session.business
    report = runner.result if runner.state == "done" and isinstance(runner.result, dict) else None
    if not report:
        raise HTTPException(status_code=404, detail="No business report yet. Start one first.")
    path = runner.meta.get("pdf")
    if not path or not os.path.exists(path):
        from crawler.business_pdf import build_business_pdf
        domain = _safe_id(get_domain(report.get("url") or "") or "website")
        path = os.path.join(get_export_dir(), f"crawl_business_{domain}_{time.strftime('%Y%m%d_%H%M%S')}.pdf")
        build_business_pdf(report, path)
        runner.meta["pdf"] = path
    return FileResponse(path=path, filename=os.path.basename(path), media_type="application/pdf")

# --------------------------------------------------------------------------
# Validation agent: re-checks reported errors live (and re-checks chosen URLs)
# --------------------------------------------------------------------------
class ValidateRequest(BaseModel):
    scope: Optional[str] = "errors"
    urls: Optional[List[str]] = None

class RecheckRequest(BaseModel):
    urls: List[str] = []

def _finished_job(what: str) -> CrawlJob:
    """The current crawl once it and its reports are finished; 404 / 409 otherwise."""
    job = session.job
    if not job or not job.crawler:
        raise HTTPException(status_code=404, detail="No crawl data available. Run a crawl first.")
    if job.crawler.status == "auth_failed":
        raise HTTPException(status_code=409, detail="The login failed, so nothing was crawled.")
    if job.phase != "done":
        raise HTTPException(status_code=409, detail=f"Wait until the crawl and its reports are finished before {what}.")
    return job

def _clean_urls(urls: Optional[List[str]]) -> List[str]:
    out: List[str] = []
    for u in urls or []:
        u = str(u or "").strip()
        if u.startswith(("http://", "https://")) and u not in out:
            out.append(u)
    return out[:MAX_RECHECK_URLS]

def _empty_validation(state: str = "idle") -> dict:
    return {"state": state, "done": 0, "total": 0, "summary": {k: 0 for k in VERDICT_KEYS}, "results": {},
            "login_guard": [], "job_id": "", "message": "", "scope": ""}

def _start_validation(job: CrawlJob, scope: str, urls: List[str]) -> str:
    def work(runner: JobRunner) -> dict:
        from crawler.validate import Validator, apply_rechecks
        validator = Validator(job.crawler, stop_event=runner.stop_event,
                              progress=lambda done, total: runner.update(done * 100 // max(1, total),
                                                                         f"Re-checked {done} of {total}"))
        runner.meta["validator"] = validator
        out = validator.run(urls=urls or None, scope=scope)
        if scope == "urls":   # a re-check writes the new statuses into the crawl; every count still adds up
            changed = apply_rechecks(job.crawler, out.get("results") or {})
            out["applied"] = changed
            if changed:
                job.results_version += 1
        _merge_validation(job, out)
        _reports_changed(job, resave=scope == "urls")
        s = out.get("summary") or {}
        runner.update(None, f"{s.get('confirmed', 0)} confirmed, {s.get('not_reproduced', 0)} not reproduced, "
                            f"{s.get('intermittent', 0)} intermittent, {s.get('fixed', 0)} fixed"
                            + (f", {s.get('not_checked')} not checked" if s.get("not_checked") else "") + ".")
        return out

    try:
        return job.runners["validate"].start(work, "Starting the live re-check...", scope=scope)
    except JobBusy:
        raise HTTPException(status_code=409, detail="A validation is already running. Wait for it to finish.")

@app.post("/api/validate/start")
def start_validation(req: Optional[ValidateRequest] = None):
    """errors = every 4xx/5xx/no-response/blocked internal URL + broken external links; issues = errors + one
    sample URL per SEO issue; urls = the given URLs (a re-check that updates the crawl)."""
    from crawler.validate import SCOPES
    scope = ((req.scope if req else "") or "errors").strip().lower()
    if scope not in SCOPES:
        raise HTTPException(status_code=400, detail=f"Unknown scope. Use one of: {', '.join(SCOPES)}")
    urls = _clean_urls(req.urls if req else None)
    if scope == "urls" and not urls:
        raise HTTPException(status_code=400, detail="Choose at least one URL to re-check.")
    job = _finished_job("re-checking errors")
    return {"job_id": _start_validation(job, scope, urls)}

@app.post("/api/crawl/recheck")
def recheck_urls(req: RecheckRequest):
    """Re-checks the given URLs live; when finished, their crawl rows (status, category) are updated and the
    SEO index rebuilt. Progress and verdicts: GET /api/validate/status."""
    urls = _clean_urls(req.urls)
    if not urls:
        raise HTTPException(status_code=400, detail="Choose at least one URL to re-check.")
    job = _finished_job("re-checking URLs")
    return {"job_id": _start_validation(job, "urls", urls)}

@app.get("/api/validate/status")
def validation_status(merged: int = 0):
    """The latest validation run (contract shape). merged=1 also returns "merged": every run of this crawl
    combined (latest verdict per URL, last login guard), so a reloaded dashboard shows all verdicts."""
    job = session.job
    if not job:
        return _empty_validation()
    runner = job.runners["validate"]
    out = _empty_validation(runner.state)
    snap = runner.result if runner.state == "done" and isinstance(runner.result, dict) else None
    if snap is None and runner.meta.get("validator") is not None:
        snap = runner.meta["validator"].snapshot()
    if snap:
        out.update(snap)
    out.update(state=runner.state, job_id=runner.job_id, message=runner.message,
               scope=out.get("scope") or runner.meta.get("scope", ""))
    if merged:
        m = job.validation or {}
        out["merged"] = {k: m.get(k, default) for k, default in (
            ("results", {}), ("summary", {k: 0 for k in VERDICT_KEYS}), ("skipped", []), ("login_guard", []),
            ("login_guard_summary", {}))}
    return out

# --------------------------------------------------------------------------
# Crawl intelligence: root causes, content intent, Passed / Failed / Not checked
# --------------------------------------------------------------------------
@app.get("/api/intel/errors")
def intel_errors():
    """Issues grouped by root cause with plain fix steps, in fix-first order (rule-based at once; AI-enriched
    text after POST /api/intel/run)."""
    job = session.job
    site, index = _site_index(job)
    return _intel_errors(job, site, index)

@app.get("/api/intel/content")
def intel_content():
    """Pages whose URL promises one thing and whose content delivers another, thin pages, off-topic pages."""
    job = session.job
    site, index = _site_index(job)
    return _intel_content(job, site, index)

def _intel_work(job: CrawlJob, runner: JobRunner) -> dict:
    """crawler.intel.run_intel with dismissed URLs left out of the groups."""
    from crawler.intel import content, detect_platform, empty_content, empty_errors, root_cause
    crawler = job.crawler
    site, index = _site_index(job)
    if site is None or index is None:
        runner.update(None, "No crawl to analyse.")
        return {"errors": empty_errors(), "content": empty_content()}
    d = _dismissals(job, index)
    runner.update(5, "Detecting the platform")
    plat = detect_platform(site, fetch=False)
    fetched = _safe(lambda: detect_platform(crawler, fetch=True), {})
    if fetched.get("cms") and (not plat.get("cms") or fetched.get("confidence", 0) >= plat.get("confidence", 0)
                               or len(fetched.get("plugins") or []) > len(plat.get("plugins") or [])):
        plat = fetched
    runner.update(20, "Grouping issues by root cause")
    errors = root_cause.analyse(site, index, crawler, platform=plat, validation=_validation_of(job),
                                issues=d["issues"] if d["entries"] else None, dismissed=d["dismissed_urls"] or None)
    runner.update(40, "Checking what each page delivers")
    cont = content.analyse(site)
    if not runner.stop_event.is_set():
        runner.update(55, "AI: explaining root causes")
        errors = _safe(lambda: root_cause.enrich_with_ai(errors, site, crawler), errors)
    if not runner.stop_event.is_set():
        runner.update(75, "AI: reviewing flagged pages")
        cont = _safe(lambda: content.enrich_with_ai(cont, site, crawler, fetch=True), cont)
    runner.meta.update(index=index, dismiss_version=d["version"], results_version=job.results_version)
    note = errors.get("ai_note") or cont.get("ai_note") or ""
    runner.update(None, "Done" + (f" ({note})" if note else "") + ".")
    _reports_changed(job, resave=False)
    return {"errors": errors, "content": cont}

@app.post("/api/intel/run")
def run_intel_job():
    """AI enrichment of the root-cause groups and the content review (background job)."""
    job = _finished_job("running the analysis")
    try:
        return {"job_id": job.runners["intel"].start(lambda r: _intel_work(job, r), "Starting the analysis...")}
    except JobBusy:
        raise HTTPException(status_code=409, detail="The analysis is already running. Wait for it to finish.")

@app.get("/api/intel/status")
def intel_status():
    job = session.job
    runner = job.runners["intel"] if job else JobRunner("intel")
    out = runner.status()
    res = runner.result if runner.state == "done" and isinstance(runner.result, dict) else {}
    out["ai_used"] = bool((res.get("errors") or {}).get("ai_used") or (res.get("content") or {}).get("ai_used"))
    out["ai_note"] = (res.get("errors") or {}).get("ai_note") or (res.get("content") or {}).get("ai_note") or ""
    return out

@app.get("/api/crawl/checks")
def crawl_checks():
    """Every check as passed / failed / not_checked with the reason (a check that did not run is never a pass)."""
    job = session.job
    site, index = _site_index(job)
    checks = _checks(job, site, index)
    from crawler.intel.checks import summary
    return {"checks": checks, "summary": summary(checks)}

# --------------------------------------------------------------------------
# Tracking: crawl history and comparison, dismissed issues, site graph
# --------------------------------------------------------------------------
@app.get("/api/crawls")
def list_crawls(site: str = ""):
    """Crawl history, newest first: every snapshot (one is saved when a crawl finishes; used by comparisons)
    and every saved crawl (File > Crawls; "saved": it can be opened again). site= limits to one site."""
    try:
        from crawler.store import crawl_history
        rows = crawl_history(REPORTS_DIR, site=site.strip() or None)
    except Exception:
        from crawler.seo.compare import list_snapshots
        rows = list_snapshots(REPORTS_DIR, host=site.strip() or None)
    job = session.job
    current = job.saved_meta.get("id") if job is not None and job.saved_meta else ""
    return {"crawls": rows, "current": current or (job.snapshot_id if job is not None else "")}

class CompareRequest(BaseModel):
    a: str = ""
    b: str = "current"
    url_map: Optional[Any] = None                 # URL Mapping: [{"regex", "replace"}] applied to crawl a
    change_detection: Optional[List[str]] = None  # elements to compare (compare.CHANGE_ELEMENTS ids)
    thresholds: Optional[Dict[str, Any]] = None

def _current_snapshot(a: str, b: str) -> Optional[dict]:
    from crawler.seo import compare
    if "current" not in (str(a).strip().lower(), str(b).strip().lower()):
        return None
    job = session.job
    site, index = _site_index(job)
    if index is None:
        return None
    segments = _segment_assignment(job, site, index)[1] or None
    return compare.snapshot_from(job.crawler, index, site=site, segments=segments)

def _check_url_map(url_map: Any):
    if not url_map:
        return
    from crawler.seo.compare import parse_url_map
    _compiled, errors = parse_url_map(url_map)
    if errors:
        _bad("URL Mapping: " + "; ".join(map(str, errors[:5])))

@app.post("/api/crawl/compare")
def compare_crawls(req: CompareRequest):
    """What changed between two crawls (ids from GET /api/crawls; "current" = the crawl on screen): per filter
    the counts with Added / New / Removed / Missing, the change detection elements and the issue changes."""
    from crawler.seo import compare
    _check_url_map(req.url_map)
    if req.change_detection is not None:
        unknown = [e for e in req.change_detection if e not in compare.CHANGE_ELEMENTS]
        if unknown:
            _bad(f"Unknown change detection element: {', '.join(map(str, unknown))}")
    current = _current_snapshot(req.a, req.b)
    return compare.compare(req.a, req.b, REPORTS_DIR, current=current, url_map=req.url_map,
                           change_detection=req.change_detection, thresholds=req.thresholds)

class CompareUrlsRequest(CompareRequest):
    key: str = ""        # "tab:filter" (or an element id for changed:<el> kinds)
    kind: str = "added"
    offset: int = 0
    limit: int = 100

def _compare_urls(req: CompareUrlsRequest) -> dict:
    from crawler.seo import compare
    _check_url_map(req.url_map)
    if not req.kind:
        _bad("Choose which URLs to list (kind).")
    current = _current_snapshot(req.a, req.b)
    try:
        return compare.compare_urls(req.a, req.b, req.key, req.kind, REPORTS_DIR, current=current,
                                    url_map=req.url_map, offset=max(0, req.offset),
                                    limit=max(1, min(req.limit, 5000)), thresholds=req.thresholds)
    except (KeyError, ValueError) as e:
        _bad(str(e) or "Unknown comparison key or kind.")

@app.get("/api/crawl/compare/urls")
def compare_urls_get(a: str = "", b: str = "current", key: str = "", kind: str = "added", offset: int = 0,
                     limit: int = 100, url_map: str = "", tab: str = "", filter_id: str = Query("", alias="filter")):
    """The URLs behind one number of a comparison (paged). key = "tab:filter" (or tab= and filter=). kind: added
    | new | removed | missing | current | previous | fixed | not_rechecked | unchanged | changed:<element> |
    not_changed:<element>. url_map: JSON."""
    if not key and tab:
        key = f"{tab}:{filter_id or 'all'}"
    rules = None
    if url_map.strip():
        try:
            rules = json.loads(url_map)
        except ValueError:
            _bad("url_map must be JSON: [{\"regex\": ..., \"replace\": ...}]")
    return _compare_urls(CompareUrlsRequest(a=a, b=b, key=key, kind=kind, offset=offset, limit=limit,
                                            url_map=rules))

@app.post("/api/crawl/compare/urls")
def compare_urls_post(req: CompareUrlsRequest):
    return _compare_urls(req)

@app.get("/api/crawl/compare/options")
def compare_options():
    """Change detection elements and default thresholds for the Compare options dialog."""
    from crawler.seo import compare
    return {"elements": [{"id": k, "label": v.get("label", k), "kind": v.get("kind", "")}
                         for k, v in compare.CHANGE_ELEMENTS.items()],
            "thresholds": dict(compare.DEFAULT_THRESHOLDS), "buckets": list(compare.BUCKETS)}

class DismissRequest(BaseModel):
    tab: str
    filter: str
    pattern: Optional[str] = ""
    note: Optional[str] = ""
    site: Optional[str] = None      # default: the current crawl's site

@app.post("/api/issues/dismiss")
def dismiss_issue(req: DismissRequest):
    """Marks an issue (or its URLs matching the glob `pattern`) as known and fine for this site. Dismissed URLs
    drop out of the issue counts only; they stay in the tabs and are listed as dismissed."""
    _, index = _site_index(session.job)
    tab = _tab_def(req.tab, index)
    flt = next((f for f in tab.filters if f.id == req.filter), None) if tab else None
    if flt is None:
        raise HTTPException(status_code=400, detail=f"Unknown tab or filter: {req.tab} / {req.filter}")
    store = _dismiss_store(session.job, req.site or "")
    if store is None:
        raise HTTPException(status_code=400, detail="No site: run a crawl first or pass the site address.")
    dismiss_id = store.add(req.tab, req.filter, req.pattern or "", req.note or "", name=f"{tab.name}: {flt.name}")
    if not dismiss_id:
        raise HTTPException(status_code=400, detail="Could not dismiss this issue (pattern too long or too many dismissals).")
    if session.job is not None:
        session.job.export_stale = True
    return {"ok": True, "id": dismiss_id}

@app.delete("/api/issues/dismiss")
def undismiss_issue(id: str = "", site: str = ""):
    store = _dismiss_store(session.job, site)
    ok = bool(store is not None and id and store.remove(id))
    if ok and session.job is not None:
        session.job.export_stale = True
    return {"ok": ok}

@app.get("/api/issues/dismissed")
def list_dismissed(site: str = ""):
    """Every dismissal of the site; for the current crawl also how many URLs each one covered."""
    job = session.job
    store = _dismiss_store(job, site)
    if store is None:
        return {"dismissed": []}
    crawl_site = job.crawler.start_url if job and job.crawler else ""
    if not site or _same_site(site, crawl_site):
        _, index = _site_index(job)
        if index is not None:
            d = _dismissals(job, index)
            return {"dismissed": list(d["dismissed"]), "summary": d.get("summary") or {}}
    return {"dismissed": store.list()}

GRAPH_TYPES = ("force", "crawl_tree", "dir_tree")

@app.get("/api/crawl/graph")
def crawl_graph(limit: int = 500, type: str = "force"):
    """Visualisations. type=force (default): the site graph - page nodes (size = Link Score, colour = HTTP
    status category) and internal link edges, above `limit` nodes the smaller pages grouped by folder.
    type=crawl_tree: every crawled URL under the page it was first found on (its parent link), depth by
    depth. type=dir_tree: URL folders with their URL counts and the worst status inside."""
    from crawler.seo.graph import build_graph
    kind = (type or "force").strip().lower()
    if kind not in GRAPH_TYPES:
        _bad(f"Unknown graph type. Use one of: {', '.join(GRAPH_TYPES)}")
    job = session.job
    site, index = _site_index(job)
    if kind == "crawl_tree":
        if site is None:
            return {"type": kind, "nodes": [], "total": 0, "shown": 0, "root": ""}
        return _memo(job, ("crawl_tree", id(index), int(limit or 0)), lambda: _crawl_tree(site, limit))
    if kind == "dir_tree":
        if site is None:
            return {"type": kind, "total": 0, "tree": [], "depth": []}
        from crawler.seo.structure import site_structure
        return dict(_memo(job, ("dir_tree", id(index)), lambda: site_structure(site)), type=kind)
    if index is None:
        return build_graph(None)
    return _memo(job, ("graph", id(index), int(limit or 500)), lambda: build_graph(site, index, limit))

class GraphTreeRequest(BaseModel):
    """POST /api/crawl/graph/tree: the GET parameters, with long include / expand lists (a search result's
    way down, the branches open on screen) that would not fit in a URL."""
    node: str = ""
    depth: int = 2
    limit: int = 50
    offset: int = 0
    include: List[str] = []
    expand: List[str] = []
    budget: int = 2000
    files: bool = False

def _graph_tree_index(files: bool):
    """(DrillTree of the crawl on screen, or None without an analysed crawl). Built once per analysis."""
    from crawler.seo.graph import build_tree_index
    job = session.job
    site, index = _site_index(job)
    if site is None or index is None:
        return None
    return _memo(job, ("graph_tree", id(index), bool(files)), lambda: build_tree_index(site, bool(files)))

def _graph_tree_answer(req: GraphTreeRequest) -> dict:
    from crawler.seo.graph import build_tree
    tree = _graph_tree_index(req.files)
    if tree is None:
        return dict(build_tree(None, files=req.files), note="No crawl yet: crawl a website first.")
    try:
        return build_tree(None, node=req.node or None, depth=req.depth, limit=req.limit, offset=req.offset,
                          include=req.include, expand=req.expand, budget=req.budget, files=req.files, tree=tree)
    except KeyError:
        raise HTTPException(status_code=404, detail="That part of the tree is not in this crawl (it changed): "
                                                    "reload the tree.")

@app.get("/api/crawl/graph/tree")
def crawl_graph_tree(node: str = "", depth: int = 2, limit: int = 50, offset: int = 0, files: bool = False,
                     budget: int = 2000, include: List[str] = Query(default=[]),
                     expand: List[str] = Query(default=[])):
    """Site graph > Drill-down: the crawled pages as a tree (registrable domain > hosts > folders > pages), one
    slice at a time. node = a node id ("" = the domain); depth = levels listed below it; at most `limit` children
    per node (the asked node's from `offset`; "more" / "more_count" / "more_categories" cover the rest); include
    = node ids listed even past the limit, expand = node ids opened at any level; budget = nodes per answer.
    files=true counts every internal URL (the Internal tab's All) instead of pages. Each node: id, name, kind
    (domain|host|folder|page), url, count, categories (from the HTTP status), worst_category, indexable_count,
    child_count, loaded, children, more; pages add status, category, indexable, title, depth, folder.
    "version" changes when the crawl's analysis does. 404 for a node id this crawl does not have."""
    return _graph_tree_answer(GraphTreeRequest(node=node, depth=depth, limit=limit, offset=offset, files=files,
                                               budget=budget, include=include, expand=expand))

@app.post("/api/crawl/graph/tree")
def crawl_graph_tree_post(req: GraphTreeRequest):
    """The drill-down tree slice of GET /api/crawl/graph/tree with long include / expand lists in the body."""
    return _graph_tree_answer(req)

@app.get("/api/crawl/graph/tree/search")
def crawl_graph_tree_search(q: str = "", limit: int = 200, files: bool = False):
    """Drill-down search: hosts, folders and pages whose name or page title contains q (an address when q has
    "/" or ":"), hosts and folders first. Each match has "path" (the node ids from the domain down to it, to open
    the way with include / expand); "inside" counts the matches below each node; "version" as the tree's."""
    tree = _graph_tree_index(files)
    if tree is None:
        return {"query": q, "total": 0, "matches": [], "inside": {}, "version": ""}
    return dict(tree.search(q, limit), version=tree.version)

CRAWL_TREE_MAX = 20_000

def _crawl_tree(site, limit: int = 0) -> dict:
    """Crawl Tree Graph: one node per internal URL with its parent (the page it was first found on, as the
    Parent link column shows it; the shortest-depth crawled page linking to it otherwise), crawl depth,
    status category and indexability. Nodes beyond `limit` (default CRAWL_TREE_MAX), deepest first, are
    left out and counted."""
    # page-like URLs only (images, scripts and PDFs are files of a page, not branches of the crawl)
    rows = [r for r in site.results if getattr(r, "url", "") and getattr(r, "resource_type", "page") == "page"]
    by_url = {r.url: r for r in rows}
    depth = getattr(site, "crawl_depth", None) or {}
    inlinks = getattr(site, "inlinks", None) or {}

    def parent_of(r) -> str:
        p = getattr(r, "parent_url", "") or ""
        if p in by_url and p != r.url:
            return p
        best, best_d = "", None
        for e in inlinks.get(r.url, ()) or ():
            src = e.get("source") or ""
            if src in by_url and src != r.url:
                d = depth.get(src, getattr(by_url[src], "depth", 0))
                if best_d is None or (d or 0) < best_d:
                    best, best_d = src, d or 0
        return best

    nodes = []
    for r in rows:
        nodes.append({"id": r.url, "parent": parent_of(r), "depth": int(depth.get(r.url, getattr(r, "depth", 0)) or 0),
                      "status_code": row_status_code(r), "category": row_category(r),
                      "indexable": bool((getattr(site, "indexable", None) or {}).get(r.url, False)),
                      "title": (getattr(r, "page_title", "") or "")[:120]})
    # a parent that would make a cycle (redirect loops) is cut
    ids = {n["id"]: n for n in nodes}
    for n in nodes:
        seen, p = {n["id"]}, n["parent"]
        while p:
            if p in seen:
                n["parent"] = ""
                break
            seen.add(p)
            p = (ids.get(p) or {}).get("parent", "")
    cap = max(1, min(int(limit or CRAWL_TREE_MAX), CRAWL_TREE_MAX))
    shown = sorted(nodes, key=lambda n: (n["depth"], n["id"]))[:cap]
    keep = {n["id"] for n in shown}
    for n in shown:
        if n["parent"] not in keep:
            n["parent"] = ""
    root = getattr(site, "start_url", "") or (shown[0]["id"] if shown else "")
    return {"type": "crawl_tree", "root": root, "nodes": shown, "total": len(nodes), "shown": len(shown),
            "more": len(nodes) - len(shown)}

def _json_copy(obj: Any) -> Any:
    """A JSON-safe copy of a report another thread may still be filling."""
    for _ in range(3):
        try:
            return json.loads(json.dumps(obj, default=str))
        except RuntimeError:   # changed size during iteration: try again
            time.sleep(0.05)
    return {}

@app.get("/api/crawl/bots")
def crawl_bots():
    """Can search engines and AI crawlers read the site? A page sample fetched as a browser and as each crawler."""
    from crawler.seo.postcrawl import not_run_reason
    job = session.job
    rep = getattr(job.crawler, "bot_report", None) if job and job.crawler else None
    if not isinstance(rep, dict):
        return {"status": "not_checked", "reason": not_run_reason(None, "bot check"), "sample": [], "pages": {},
                "agents": [], "verdicts": {}, "notes": [], "headline": []}
    out = _json_copy(rep)
    out["job"] = job.runners["postcrawl"].status()
    return out

@app.get("/api/crawl/robots-audit")
def crawl_robots_audit():
    """robots.txt read for search engines and AI crawlers (blocked bots, Content-Signal, Sitemap lines)."""
    from crawler.seo.postcrawl import not_run_reason
    job = session.job
    rep = getattr(job.crawler, "robots_audit", None) if job and job.crawler else None
    if not isinstance(rep, dict):
        return {"status": "not_checked", "reason": not_run_reason(None, "robots.txt audit"), "bots": [],
                "issues": [], "groups": [], "sitemaps": [], "text": ""}
    out = _json_copy(rep)
    text = str(out.get("text") or "")
    out["text"], out["text_truncated"] = text[:ROBOTS_TEXT_SHOWN], len(text) > ROBOTS_TEXT_SHOWN
    return out

@app.get("/api/crawl/crawl-budget")
def crawl_budget_report():
    """URL parameters and redirects that waste crawl budget, with robots.txt and canonical advice."""
    job = session.job
    site, index = _site_index(job)
    if site is None:
        return {"params": [], "total_urls": 0, "param_urls": 0, "wasted_urls": 0, "wasted_share": 0.0, "notes": []}
    from crawler.seo import crawl_budget
    return _json_copy(crawl_budget.get(site))

# --------------------------------------------------------------------------
# Import-ready fix files and developer tickets
# --------------------------------------------------------------------------
def _download(name: str, data: bytes, mime: str, extra: Optional[Dict[str, str]] = None) -> Response:
    safe = re.sub(r'[^\w.-]+', "_", name or "download")
    headers = {"Content-Disposition": f'attachment; filename="{safe}"'}
    headers.update(extra or {})
    return Response(content=data, media_type=mime, headers=headers)

def _ai_enabled() -> bool:
    try:
        from crawler.ai import client as ai_client
        return bool(ai_client.ai_status().get("enabled"))
    except Exception:
        return False

@app.get("/api/export/fixes")
def export_fixes(format: str = "csv", ai: int = 1):
    """Import-ready titles, meta descriptions, alt text and redirects: yoast | rankmath | shopify (.zip) |
    htaccess | nextjs | csv. AI drafts when AI is on (made once per crawl in the background; the request
    waits up to CRAWLER_FIXES_AI_WAIT seconds), otherwise rule-based drafts marked TODO."""
    from crawler import fixfiles
    fmt = (format or "csv").strip().lower()
    if fmt not in fixfiles.FORMATS:
        raise HTTPException(status_code=400, detail=f"Unknown format. Use one of: {', '.join(fixfiles.FORMATS)}")
    job = session.job
    site, index = _site_index(job)
    crawler = job.crawler if job else None
    dismissed = _dismissals(job, index)["dismissed_urls"] if index is not None else None
    use_ai = bool(ai) and index is not None and _ai_enabled()
    if use_ai:
        runner = job.runners["fixes"]
        key = (id(index), json.dumps(sorted((k, len(v)) for k, v in (dismissed or {}).items())))
        if not (runner.busy or (runner.state == "done" and runner.meta.get("key") == key)):
            try:
                runner.start(lambda r: fixfiles.fixes_summary(
                    fixfiles.draft_fixes(site, index, crawler, use_ai=True, dismissed=dismissed)),
                    "Drafting fixes with AI...", key=key)
            except JobBusy:
                pass
        if not runner.wait(FIXES_AI_WAIT) or runner.state != "done":
            use_ai = False    # still drafting (or failed): rule-based file now; AI drafts on a later download
    # the rule-based fallback is built without the crawler so it never replaces the cached AI drafts
    name, data, mime = fixfiles.build_fixes(site, index, crawler if use_ai else None, fmt, use_ai=use_ai,
                                            dismissed=dismissed)
    # "1" only when the drafts really came from the AI (not when it was skipped: private crawl, quota, error)
    summary = job.runners["fixes"].result if use_ai else None
    drafted = isinstance(summary, dict) and bool(summary.get("ai_used"))
    return _download(name, data, mime, {"X-Fixes-AI": "1" if drafted else "0"})

@app.get("/api/export/tickets")
def export_tickets(format: str = "csv"):
    """One developer ticket per root-cause group (title, where, element HTML, steps, how to verify, screenshot
    file names): csv | jira | markdown | zip (all three plus the screenshots/ folder they name)."""
    from crawler import tickets
    fmt = (format or "csv").strip().lower()
    if fmt not in tickets.FORMATS + ("md",):
        raise HTTPException(status_code=400, detail=f"Unknown format. Use one of: {', '.join(tickets.FORMATS)}")
    job = session.job
    site, index = _site_index(job)
    intel = _intel_errors(job, site, index) if index is not None else None
    evidence = {"seo": job.seo_evidence, "links": job.evidence} if job else None
    site_url = job.crawler.start_url if job and job.crawler else ""
    name, data, mime = tickets.build_tickets(intel, evidence, fmt, site_url=site_url,
                                             evidence_dir=job.evidence_dir if job else "")
    return _download(name, data, mime)

# ==========================================================================
# Build 3: shared helpers for the menu-bar features
# ==========================================================================
SEGMENT_NONE = "(No segment)"
_version_cache: Dict[str, str] = {}

def app_version() -> str:
    """APP_VERSION, else the git short commit of the checkout (read once), else "dev"."""
    if "v" in _version_cache:
        return _version_cache["v"]
    v = (os.environ.get("APP_VERSION") or "").strip()
    if not v:
        try:
            root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            out = subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=root, capture_output=True, text=True,
                                 timeout=3)
            v = out.stdout.strip() if out.returncode == 0 else ""
        except Exception:
            v = ""
    _version_cache["v"] = v or "dev"
    return _version_cache["v"]

def _notify():
    from crawler import notify
    return notify

def _integrations():
    import crawler.integrations as integrations
    return integrations

def _store():
    from crawler import store
    return store

def _settings():
    from crawler import settings_store
    return settings_store

def _is_local_request(request: Request) -> bool:
    """The request comes from the server's own computer (loopback, or the machine's own network address).
    Behind a proxy (Render, CRAWLER_TRUST_PROXY, CRAWLER_CLIENT_IP_HEADER) every request looks local, so none
    counts as local there."""
    if os.environ.get("RENDER") or _trust_proxy() or os.environ.get("CRAWLER_CLIENT_IP_HEADER"):
        return False
    peer = str((request.scope.get("client") or ("",))[0] or "")
    server = str((request.scope.get("server") or ("",))[0] or "")
    try:
        ip = ipaddress.ip_address(peer)
        if ip.is_loopback or (ip.version == 6 and ip.ipv4_mapped is not None and ip.ipv4_mapped.is_loopback):
            return True
    except ValueError:
        return False
    return bool(peer) and peer == server

def settings_editable(request: Request) -> Tuple[bool, str]:
    """Server-wide settings (proxy, retention, notifications, browser, API keys) change for everyone, so they
    can be changed when the server has a team access key (the request passed it), from the server's own
    computer, or when CRAWLER_SETTINGS_EDITABLE=1. Never for AI agents (in-process requests, W1)."""
    if _agents.internal_state(request.scope):
        return False, "AI agents cannot change server settings, connections or tokens."
    if access_key():
        return True, ""
    if os.environ.get("CRAWLER_SETTINGS_EDITABLE", "").strip().lower() in ("1", "true", "yes"):
        return True, ""
    if _is_local_request(request):
        return True, ""
    return False, ("Server settings are read-only from other computers while this server has no access key. "
                   "Change them on the server itself (http://localhost:<port>), set CRAWLER_ACCESS_KEY, or set "
                   "CRAWLER_SETTINGS_EDITABLE=1.")

def _require_admin(request: Request):
    ok, why = settings_editable(request)
    if not ok:
        raise HTTPException(status_code=403, detail=why)
    if not access_key():          # admin from this PC or CRAWLER_SETTINGS_EDITABLE: refuse other websites (W1)
        _agents.require_same_origin(request)

def _public_saved_meta(job: Optional[CrawlJob]) -> Optional[dict]:
    m = job.saved_meta if job is not None else None
    if not m:
        return None
    return {k: m.get(k) for k in ("id", "name", "project", "locked", "saved_at", "complete")}

def _autosave_on() -> bool:
    try:
        return bool(_settings().effective("retention", REPORTS_DIR).get("autosave", True))
    except Exception:
        return True

def _apply_retention() -> Optional[dict]:
    """Settings > Crawl Retention: deletes saved crawls past their time (never the crawl on screen) and prunes
    scheduled-run export folders. Never raises."""
    job = session.job
    keep = [str((job.saved_meta or {}).get("id") or "")] if job is not None and job.saved_meta else []
    store = _store()
    return _safe(lambda: store.apply_retention(_settings().get_settings(False, REPORTS_DIR), REPORTS_DIR,
                                               keep_ids=keep), None)

def _autosave(job: CrawlJob, log: bool = False, force: bool = False, restore: bool = False) -> Optional[dict]:
    """Saves the finished crawl (File > Crawls) when Settings > Crawl Retention says so (default on), then
    applies the retention rules. A crawl opened from the saved list is re-saved under its own id. Never raises.
    `force` (File > Save, a scheduled task's "save crawl") also saves a finished crawl that another one replaced
    on screen meanwhile; a crawl replaced before it finished is never saved. A crawl whose saved copy was deleted
    is only saved again with `restore` (File > Save)."""
    crawler = job.crawler
    if crawler is None or job.origin == "direct" or not crawler.results_list:
        return None
    if job.retired and (not force or getattr(crawler, "superseded", False)):
        return None
    if crawler.status in ("running", "paused", "idle", "auth_failed"):
        return None
    if job.saved_deleted and not restore:
        return None    # the team deleted its saved copy: a re-check or fetch does not bring it back
    if not force and not _autosave_on() and not job.saved_meta:
        return None
    store = _store()
    try:
        site, index = _site_index(job)
        meta = store.save_crawl(job, REPORTS_DIR, name=job.save_name or None, project=job.save_project or None,
                                index=(site, index) if index is not None else None,
                                dismissals=_safe(lambda: _dismiss_store(job).export_entries(), None),
                                extras={"segments": _segment_store(job).list() if _segment_store(job) else []},
                                allow_deleted=restore)
        job.saved_meta = meta
        job.saved_deleted = False
        if log:
            crawler.log(f"Crawl saved: File > Crawls > {meta.get('name') or meta.get('id')}")
        _apply_retention()
        return meta
    except Exception as e:
        crawler.log(f"The crawl could not be saved: {job.mask(e)}")
        return None

def _segment_store(job: Optional[CrawlJob], site_url: str = ""):
    url = (site_url or "").strip() or (job.crawler.start_url if job and job.crawler else "")
    if not url:
        return None
    from crawler.segments import store_for
    return store_for(REPORTS_DIR, url)

def _segment_assignment(job: Optional[CrawlJob], site, index) -> Tuple[List[dict], Dict[str, str]]:
    """(the site's segments, {url: segment name or ''}) for the crawl on screen, worked out once per index and
    segment list."""
    if job is None or site is None or index is None:
        return [], {}
    store = _safe(lambda: _segment_store(job), None)
    segs = _safe(lambda: store.list(), []) if store is not None else []
    if not segs:
        return [], {}
    from crawler import segments as seg_mod
    version = seg_mod.segments_version(segs)
    return segs, _memo(job, ("segments", id(index), version),
                       lambda: _safe(lambda: seg_mod.evaluate(site, index, segs), {}, job, "Segments"))

def _segment_of_url(job, site, index, url: str) -> Optional[str]:
    segs, assignment = _segment_assignment(job, site, index)
    return (assignment.get(url) or SEGMENT_NONE) if segs else None

def _page_cookies(seo: dict) -> List[dict]:
    """Cookies the page set (names and attributes; values are never stored): its server answer's Set-Cookie and,
    when it was rendered in Chrome, the ones received while it loaded ("source": request / script, "set_by")."""
    out = []
    for c in (seo or {}).get("cookies") or []:
        if isinstance(c, dict):
            out.append({k: c.get(k) for k in ("name", "domain", "path", "expires", "max_age", "secure", "httponly",
                                              "samesite", "source", "set_by") if k in c})
    return out

def _shot_url(name: str) -> Optional[str]:
    return f"/api/crawl/evidence/{name}" if name else None

def _not_rendered_reason(job: Optional[CrawlJob], url: str) -> str:
    pool = job.render.get("pool") if job is not None else None
    why = (getattr(pool, "failures", None) or {}).get(url) if pool is not None else ""
    if why:
        return f"Not rendered: {job.mask(why)}"
    cfg = getattr(job.crawler, "config", None) if job is not None and job.crawler is not None else None
    if not getattr(cfg, "render_all", False):
        return ("Only available when every page is rendered in Chrome (Configuration > Rendering), or render this "
                "page now.")
    return "This page was not rendered."

def _rendered_summary(job: CrawlJob, url: str, seo: dict) -> dict:
    """The Rendered Page / Chrome Console Log summary for the URL details pane (the full data: /api/crawl/render)."""
    rd = (seo or {}).get("rendered")
    if not isinstance(rd, dict):
        return {"rendered": False, "reason": _not_rendered_reason(job, url)}
    console = rd.get("console") or []
    kept = (seo or {}).get("original_html")
    kept = kept if isinstance(kept, dict) else {}

    def raw(key: str):   # the raw-HTML value (the page's own keys hold the rendered one once promoted)
        return kept[key] if key in kept else (seo or {}).get(key)
    return {
        "rendered": True, "title": rd.get("title") or "", "h1": rd.get("h1") or [], "word_count": rd.get("word_count"),
        "promoted": bool(rd.get("promoted")),   # the tabs read this page's rendered HTML (JavaScript rendering)
        "html": {"title": (raw("titles") or [""])[0], "word_count": raw("word_count"), "h1": raw("h1") or [],
                 "canonical": (raw("canonicals") or [""])[0], "meta_robots": raw("meta_robots") or []},
        "meta_robots": rd.get("meta_robots") or [], "canonical": rd.get("canonical") or "",
        "final_url": rd.get("final_url") or "", "js_redirect": bool(rd.get("js_redirect")),
        "js_links": len(rd.get("js_links") or []), "console": len(console),
        "console_errors": sum(1 for c in console if (c or {}).get("level") == "SEVERE"),
        "blocked_resources": len(rd.get("blocked_resources") or []),
        "failed_resources": len(rd.get("failed_resources") or []),
        "chrome_issues": len(rd.get("chrome_issues") or []) if rd.get("chrome_issues_checked", True) else None,
        "screenshot": _shot_url(rd.get("screenshot") or ""), "render_ms": rd.get("render_ms"),
    }

def _source_available(job: Optional[CrawlJob], url: str) -> bool:
    store = getattr(job.crawler, "source_store", None) if job is not None and job.crawler is not None else None
    return bool(store is not None and _safe(lambda: store.has(url), False))

def _lighthouse(site, url: str) -> Optional[dict]:
    from crawler.integrations import psi
    return psi.details(site, url)

def _report_ctx(job: Optional[CrawlJob], site, index) -> Dict[str, Any]:
    """The extra data reports and bulk exports read: dismissal-adjusted issues, dismissed URLs, segments, the
    crawl's secrets (masked in header exports) and its stats."""
    if job is None or index is None:
        return {}
    d = _safe(lambda: _dismissals(job, index), _empty_dismissals(index))
    ctx: Dict[str, Any] = {"issues": d["issues"], "dismissed_urls": d["dismissed_urls"], "secrets": list(job.secrets),
                           "stats": _safe(lambda: _public_stats(job), {})}
    segs, assignment = _segment_assignment(job, site, index)
    if segs:
        ctx["segments"] = assignment
    return ctx

def _need_index(job: Optional[CrawlJob] = None):
    """(job, site, index) of the crawl on screen; 404 without one."""
    job = job if job is not None else session.job
    site, index = _site_index(job)
    if index is None:
        raise HTTPException(status_code=404, detail="No crawl data available. Run a crawl first.")
    return job, site, index

async def _read_upload(request: Request, limit: int = MAX_UPLOAD_BYTES) -> Tuple[str, bytes]:
    """(file name, bytes) of an upload sent as multipart/form-data (the first file field), as JSON
    {"filename", "data": base64} or {"filename", "text"}, or as the raw body (name from ?filename= or the
    X-Filename header). 413 above `limit`."""
    body = await request.body()
    if len(body) > limit:
        raise HTTPException(status_code=413, detail=f"The file is larger than {limit // (1024 * 1024)} MB.")
    # the boundary is case-sensitive (browsers send "----WebKitFormBoundaryAbC..."): only the type is lowercased
    raw_ctype = (request.headers.get("content-type") or "").strip()
    ctype = raw_ctype.lower()
    name = request.query_params.get("filename") or request.headers.get("x-filename") or ""
    if ctype.startswith("multipart/form-data"):
        from email.parser import BytesParser
        from email.policy import default as email_policy
        msg = BytesParser(policy=email_policy).parsebytes(b"Content-Type: " + raw_ctype.encode("latin-1") + b"\r\n\r\n"
                                                          + body)
        text_fields: Dict[str, str] = {}
        for part in msg.iter_parts():
            fname = part.get_filename()
            data = part.get_payload(decode=True) or b""
            if fname is not None:
                return os.path.basename(fname) or name, data
            field = part.get_param("name", header="content-disposition") or ""
            text_fields[field] = data.decode("utf-8", "replace")
        text = text_fields.get("text") or text_fields.get("urls") or next(iter(text_fields.values()), "")
        return name or "pasted.txt", text.encode("utf-8")
    if ctype.startswith("application/json"):
        try:
            obj = json.loads(body.decode("utf-8") or "{}")
        except ValueError:
            _bad("The request body is not valid JSON.")
        if not isinstance(obj, dict):
            _bad("Send {\"filename\", \"data\"} (base64) or {\"text\"}.")
        if obj.get("data"):
            try:
                data = base64.b64decode(str(obj["data"]), validate=False)
            except Exception:
                _bad("data must be base64.")
            if len(data) > limit:
                raise HTTPException(status_code=413, detail=f"The file is larger than {limit // (1024 * 1024)} MB.")
            return os.path.basename(str(obj.get("filename") or name or "upload.bin")), data
        return os.path.basename(str(obj.get("filename") or name or "pasted.txt")), str(obj.get("text") or "").encode("utf-8")
    return os.path.basename(name or "upload.bin"), body

# --------------------------------------------------------------------------
# URL detail tabs: Cookies, View Source, Rendered Page, Chrome Console Log, detail Export
# --------------------------------------------------------------------------
def _find_crawled(site, u: str):
    return next((site.by_url[c] for c in _url_candidates(u) if c in site.by_url), None) if site is not None else None

@app.get("/api/crawl/cookies")
def get_url_cookies(u: str = ""):
    """Cookies one URL set (Set-Cookie): names and attributes only; values are never stored or shown."""
    job, site, _ = _need_index()
    r = _find_crawled(site, u)
    if r is None:
        raise HTTPException(status_code=404, detail="This URL is not in the crawl.")
    return {"url": r.url, "cookies": _page_cookies(getattr(r, "seo", None) or {})}

@app.get("/api/crawl/source")
def get_url_source(u: str = ""):
    """View Source: the HTML as crawled ("original") and, when the page was rendered in Chrome, the rendered
    HTML. Logins' values are masked. Pages whose HTML was not kept say why ("reason")."""
    job = session.job
    if job is None or job.crawler is None:
        raise HTTPException(status_code=404, detail="No crawl data available.")
    site, _ = _site_index(job)
    r = _find_crawled(site, u) if site is not None else None
    url = r.url if r is not None else (u or "").strip()
    if not url:
        _bad("Choose a URL.")
    out: Dict[str, Any] = {"url": url, "original": None, "rendered": None, "reason": "", "truncated": False}
    store = getattr(job.crawler, "source_store", None)
    if store is None:
        out["reason"] = ("The page HTML was not kept for this crawl (Configuration > Spider > Store HTML is off, "
                         "or this server runs serverless).")
        return out
    html = _safe(lambda: store.get(url), None)
    if html is None:
        out["reason"] = ("Only HTML pages of the crawled site are kept; this URL is not one, or its HTML was "
                         "removed with old crawl files.")
    else:
        out["original"] = job.mask(html[:MAX_SOURCE_CHARS])
        out["truncated"] = len(html) > MAX_SOURCE_CHARS
        out["size"] = len(html)
    rstore = _safe(lambda: _rendered_store(job), None)
    rendered = _safe(lambda: rstore.get(url), None) if rstore is not None else None
    if rendered is not None:
        out["rendered"] = job.mask(rendered[:MAX_SOURCE_CHARS])
    else:
        out["rendered_reason"] = _not_rendered_reason(job, url)
    return out

@app.get("/api/crawl/render")
def get_url_render(u: str = "", refresh: int = 0, stored: int = 0):
    """Rendered Page: the page as Chrome shows it (rendered title, links, console, resources with the blocked
    and failed ones, mobile and accessibility checks, screenshot). The crawl's stored render is used; without
    one (or refresh=1) the page is rendered now in its own Chrome (a few seconds). stored=1: the render kept
    during the crawl as it is, even without a screenshot (the details pane of a "render every page" crawl)."""
    job = session.job
    if job is None or job.crawler is None:
        raise HTTPException(status_code=404, detail="No crawl data available.")
    site, _ = _site_index(job)
    r = _find_crawled(site, u) if site is not None else None
    url = r.url if r is not None else (u or "").strip()
    stored_render = (getattr(r, "seo", None) or {}).get("rendered") if r is not None else None
    if isinstance(stored_render, dict) and (stored_render.get("screenshot") or stored) and not refresh:
        return _render_payload(job, url, stored_render, "crawl")
    if job.phase == "crawling" and not isinstance(stored_render, dict):
        raise HTTPException(status_code=409, detail="Wait until the crawl has finished to render single pages.")
    if not _browser_enabled():
        if isinstance(stored_render, dict):
            return _render_payload(job, url, stored_render, "crawl")
        raise HTTPException(status_code=409, detail="The embedded browser is switched off in Settings.")
    from crawler.render_capture import render_one
    rendered, why = render_one(url, job.crawler, evidence_dir=job.evidence_dir,
                               on_html=lambda page, html: _keep_rendered_html(job, page, html))
    if rendered is None:
        if isinstance(stored_render, dict):
            return _render_payload(job, url, stored_render, "crawl")
        raise HTTPException(status_code=422, detail=job.mask(why or "The page could not be rendered."))
    return _render_payload(job, url, rendered, "now")

@app.post("/api/crawl/render/warm")
def warm_url_render():
    """Starts the Chrome that "Render this page now" uses, in the background (the details pane shows a page
    without a render), so the render does not also wait for Chrome to start. Closes itself after 90 s unused."""
    job = session.job
    if job is None or job.crawler is None or job.phase != "done" or not _browser_enabled():
        return {"warming": False}
    from crawler.render_capture import warm_ondemand
    return {"warming": bool(_safe(lambda: warm_ondemand(job.crawler, job.evidence_dir), False))}

def _rendered_store(job: CrawlJob, create: bool = False):
    """The SourceStore holding the rendered HTML next to the raw HTML (None when no HTML is kept)."""
    rstore = job.render.get("html")
    if rstore is not None:
        return rstore
    store = getattr(job.crawler, "source_store", None)
    if store is not None and callable(getattr(store, "rendered", None)):
        saved = store.rendered()       # an opened saved crawl: the rendered pages packed with it (read-only)
        if saved is not None or not create or not job.evidence_dir:
            return saved
        # "Render this page now" on an opened crawl without rendered pages: kept next to its pictures
        from crawler.source_store import SourceStore
        try:
            rstore = SourceStore(os.path.join(job.evidence_dir, "rendered_now"))
        except Exception:
            return None
        job.render["html"] = rstore
        return rstore
    base = getattr(store, "dir", "") if store is not None else ""
    if not base:
        return None
    path = os.path.join(base, "rendered")
    from crawler.source_store import SourceStore
    if os.path.isdir(path):
        rstore = SourceStore.open(path)
    elif create:
        try:
            rstore = SourceStore(path)
        except Exception:
            return None
    else:
        return None
    job.render["html"] = rstore
    return rstore

def _keep_rendered_html(job: CrawlJob, url: str, html: str):
    """Render now: keeps the rendered HTML so View Source can show it next to the original."""
    if not html:
        return
    rstore = _safe(lambda: _rendered_store(job, create=True), None)
    if rstore is not None:
        _safe(lambda: rstore.put(url, html), None)

def _render_payload(job: CrawlJob, url: str, rendered: dict, source: str) -> dict:
    data = json_value(rendered)
    data.pop("links", None)                   # thousands of URLs; js_links (found only after rendering) stay
    data["screenshot_url"] = _shot_url(rendered.get("screenshot") or "")
    if not data["screenshot_url"]:   # why there is no picture (the dashboard offers "Take a screenshot now")
        why = next((n for n in (rendered.get("notes") or []) if isinstance(n, str) and "screenshot" in n.lower()), "")
        data["screenshot_note"] = job.mask(why) if why else (
            "No screenshot was kept for this page during the crawl (screenshots are capped per crawl)."
            if source == "crawl" else "Chrome could not take a screenshot of this page.")
    data["console"] = [dict(c, text=job.mask(c.get("text"))) for c in data.get("console") or [] if isinstance(c, dict)]
    return {"url": url, "source": source, "rendered": data}

@app.get("/api/crawl/console")
def get_url_console(u: str = ""):
    """Chrome Console Log of one rendered page: level (SEVERE / WARNING / INFO / DEBUG), source, text, line.
    Not checked (with the reason) when the page was not rendered."""
    job, site, _ = _need_index()
    r = _find_crawled(site, u)
    if r is None:
        raise HTTPException(status_code=404, detail="This URL is not in the crawl.")
    rd = (getattr(r, "seo", None) or {}).get("rendered")
    if not isinstance(rd, dict):
        return {"url": r.url, "checked": False, "reason": _not_rendered_reason(job, r.url), "console": []}
    return {"url": r.url, "checked": True, "reason": "",
            "console": [dict(c, text=job.mask(c.get("text"))) for c in rd.get("console") or [] if isinstance(c, dict)]}

@app.get("/api/crawl/export/detail")
def export_url_detail(u: str = "", part: str = "inlinks", format: str = "csv"):
    """The URL details pane's table as CSV / XLSX: part = inlinks | outlinks | images | resources | headers |
    cookies (header values that can be logins are masked), or a registered details section (structured_data,
    accessibility, spelling, carbon, pdf...: W5)."""
    fmt = _fmt(format)
    job, site, _ = _need_index()
    from crawler.exporter import export_detail
    try:
        name, data, mime = export_detail(site, u, part, fmt, secrets=job.secrets)
    except LookupError:
        raise HTTPException(status_code=404, detail="This URL is not in the crawl.")
    except ValueError as e:
        _bad(str(e) or "Unknown part.")
    return _download(name, data, mime)

# --------------------------------------------------------------------------
# Reports, Bulk Export, Multi-Export, Queued URLs, List mode export
# --------------------------------------------------------------------------
@app.get("/api/reports")
def list_reports():
    """Everything the Reports, Bulk Export and Multi-Export menus offer: {"reports": [...], "bulk": [...]}
    (bulk items carry their row count and a not-checked reason once a crawl is loaded)."""
    from crawler import reports
    _, index = _site_index(session.job)
    return reports.catalog(index)

@app.get("/api/crawl/report/{report_id}")
def download_report(report_id: str, format: str = "xlsx"):
    """One report of the Reports menu (Crawl Overview, Redirect Chains, Orphan Pages, SERP Summary, ...) as
    XLSX or CSV (a .zip when it has several tables)."""
    fmt = _fmt(format)
    job, site, index = _need_index()
    from crawler import reports
    try:
        name, data, mime = reports.build_report(report_id, site, index, job.crawler, fmt=fmt,
                                                **_report_ctx(job, site, index))
    except KeyError:
        raise HTTPException(status_code=404, detail=f"Unknown report: {report_id}")
    return _download(name, data, mime)

@app.get("/api/crawl/export/bulk")
def download_bulk(id: str = "", format: str = "xlsx"):
    """One Bulk Export item: queued_urls, links_all_inlinks, ..., or filter:<tab>:<filter>, inlinks:<tab>:<filter>,
    issue:<tab>:<filter> (ids from GET /api/reports)."""
    fmt = _fmt(format)
    job, site, index = _need_index()
    from crawler import reports
    try:
        name, data, mime = reports.build_bulk(id, site, index, job.crawler, fmt=fmt, **_report_ctx(job, site, index))
    except KeyError:
        raise HTTPException(status_code=404, detail=f"Unknown export: {id}")
    return _download(name, data, mime)

@app.get("/api/crawl/export/queued")
def download_queued(format: str = "xlsx"):
    """Queued URLs: found but not crawled yet (paused or stopped crawl) and those left out by the limits."""
    job = session.job
    if job is None or job.crawler is None:
        raise HTTPException(status_code=404, detail="No crawl data available.")
    site, index = _site_index(job)
    from crawler import reports
    name, data, mime = reports.build_bulk("queued_urls", site, index, job.crawler, fmt=_fmt(format))
    return _download(name, data, mime)

LIST_EXPORT_HEADERS = ["Line", "Address (as listed)", "URL Crawled", "Status Code", "Status", "Category",
                       "Redirect URL", "Indexability", "Indexability Status", "Title"]

@app.get("/api/crawl/export/list")
def download_list_export(format: str = "xlsx"):
    """List mode: the listed URLs in upload order with their status, category, redirect, indexability and
    title; lines that were not web addresses are listed on a second sheet."""
    job = session.job
    if job is None or job.crawler is None:
        raise HTTPException(status_code=404, detail="No crawl data available.")
    crawler = job.crawler
    if getattr(crawler, "mode", "spider") != "list":
        raise HTTPException(status_code=409, detail="This crawl is not a List mode crawl.")
    from crawler.exporter import Table, row_status_text, tables_bytes
    site, _ = _site_index(job)
    seeds = list(getattr(crawler.config, "seed_urls", None) or [])
    listed = sorted((r for r in list(crawler.results_list) if getattr(r, "list_index", None) is not None),
                    key=lambda r: r.list_index)
    rows = []
    for r in listed:
        i = r.list_index
        redirect = getattr(r, "redirect_destination", None) or ""   # CrawlResult has no redirect_url
        if redirect == "-":
            redirect = ""
        rows.append([i + 1, seeds[i] if i < len(seeds) else r.url, r.url, row_status_code(r), row_status_text(r),
                     row_category(r), redirect,
                     ("Indexable" if (site.indexable or {}).get(r.url) else "Non-Indexable") if site else "",
                     (site.indexability_status or {}).get(r.url, "") if site else "",
                     # the page's <title> (Title 1), not the crawler's display label (a path or "[Status 404]")
                     ((getattr(r, "seo", None) or {}).get("titles") or [""])[0] or ""])
    tables = [Table("Listed URLs", list(LIST_EXPORT_HEADERS), rows, url_cols=(2, 3, 7))]
    summary = getattr(crawler, "list_summary", None) or {}
    rejected = summary.get("rejected") or []
    if rejected:
        tables.append(Table("Not Web Addresses", ["Line", "Value", "Reason"],
                            [[x.get("index", 0) + 1, x.get("value", ""), x.get("reason", "")] for x in rejected]))
    name, data, mime = tables_bytes(tables, _fmt(format), f"{job.base_name}_list")
    return _download(name, data, mime)

class MultiExportRequest(BaseModel):
    items: List[str] = []
    format: str = "xlsx"
    skip_empty: bool = True

@app.post("/api/crawl/export/multi")
def start_multi_export(req: MultiExportRequest, wait: float = 0.0):
    """Multi-Export: many reports and bulk exports in one workbook (xlsx, an Index sheet first) or a zip of
    CSVs, built as a background job. Poll GET /api/crawl/export/multi, then download .../multi/file. With
    wait=<seconds> (at most 120) the answer is the file itself when it is ready in time."""
    out = _start_multi(req)
    if wait and wait > 0:
        job = session.job
        runner = job.runners["export"]
        runner.wait(min(float(wait), 120.0))
        if runner.state == "done":
            return multi_export_file()
        if runner.state == "error":
            raise HTTPException(status_code=500, detail=runner.message)
    return out

def _start_multi(req: MultiExportRequest) -> dict:
    fmt = _fmt(req.format)
    items = [str(i).strip() for i in req.items if str(i or "").strip()]
    if not items:
        _bad("Choose at least one export.")
    if len(items) > 2000:
        _bad("Choose at most 2,000 exports.")
    job, site, index = _need_index()
    if job.phase != "done":
        raise HTTPException(status_code=409, detail="Wait until the crawl and its reports are finished.")

    def work(runner: JobRunner) -> dict:
        from crawler import reports
        runner.update(5, f"Building {len(items)} export(s)...")
        name, data, mime = reports.multi_export(items, site, index, job.crawler, fmt=fmt, skip_empty=req.skip_empty,
                                                **_report_ctx(job, site, index))
        ext = os.path.splitext(name)[1] or (".xlsx" if fmt == "xlsx" else ".zip")
        path = os.path.join(get_export_dir(), f"{job.base_name}_multi_export{ext}")
        with open(path, "wb") as fh:
            fh.write(data)
        runner.update(None, f"Ready: {name} ({len(data) // 1024:,} KB).")
        return {"name": name, "path": path, "mime": mime, "bytes": len(data), "items": len(items)}

    try:
        return {"job_id": job.runners["export"].start(work, "Starting the Multi-Export...", items=len(items))}
    except JobBusy:
        raise HTTPException(status_code=409, detail="A Multi-Export is already being built. Wait for it to finish.")

@app.get("/api/crawl/export/multi")
def multi_export_status():
    job = session.job
    runner = job.runners["export"] if job is not None else JobRunner("export")
    out = runner.status()
    res = runner.result if runner.state == "done" and isinstance(runner.result, dict) else None
    out["file"] = {k: res[k] for k in ("name", "bytes", "items")} if res else None
    return out

@app.get("/api/crawl/export/multi/file")
def multi_export_file():
    job = session.job
    runner = job.runners["export"] if job is not None else None
    res = runner.result if runner is not None and runner.state == "done" and isinstance(runner.result, dict) else None
    if not res or not os.path.exists(res.get("path") or ""):
        raise HTTPException(status_code=404, detail="No Multi-Export is ready. Start one first.")
    return FileResponse(res["path"], filename=res["name"], media_type=res.get("mime") or "application/octet-stream")

# --------------------------------------------------------------------------
# Sitemaps > XML Sitemap / Images Sitemap
# --------------------------------------------------------------------------
class SitemapRequest(BaseModel):
    kind: str = "xml"            # xml | images
    options: Optional[dict] = None

def _sitemap_kind(kind: str) -> str:
    k = (kind or "xml").strip().lower()
    if k not in ("xml", "images"):
        _bad("Unknown sitemap kind. Use xml or images.")
    return k

def _sitemap_options(options: Optional[dict]) -> dict:
    from crawler import sitemap_gen
    opts = dict(options or {})
    unknown = [k for k in opts if k not in sitemap_gen.DEFAULT_OPTIONS]
    if unknown:
        _bad(f"Unknown sitemap option: {', '.join(map(str, unknown))}")
    return opts

@app.post("/api/crawl/sitemap")
def generate_sitemap(req: SitemapRequest):
    """An XML sitemap (or image sitemap) of the crawl with the export dialog's options: sitemap.xml, or a .zip
    with the files split at 50,000 URLs / 50 MB plus a sitemap index. 400 with the reason when no URL fits."""
    job, site, _ = _need_index()
    from crawler import sitemap_gen
    try:
        name, data, mime = sitemap_gen.build_sitemap(site, job.crawler, _sitemap_kind(req.kind),
                                                     _sitemap_options(req.options))
    except sitemap_gen.SitemapEmpty as e:
        _bad(str(e))
    return _download(name, data, mime)

@app.post("/api/crawl/sitemap/preview")
def preview_sitemap(req: SitemapRequest):
    """How many URLs / images / files the sitemap would hold with these options, and what was left out why."""
    job, site, _ = _need_index()
    from crawler import sitemap_gen
    return sitemap_gen.preview(site, _sitemap_kind(req.kind), _sitemap_options(req.options))

@app.get("/api/crawl/sitemap/options")
def sitemap_options():
    from crawler import sitemap_gen
    return {"defaults": dict(sitemap_gen.DEFAULT_OPTIONS), "changefreqs": list(sitemap_gen.CHANGEFREQS)}

@app.get("/api/crawl/export/sitemap")
def download_sitemap(kind: str = "xml"):
    """The sitemap with the default options, for a plain download link."""
    return generate_sitemap(SitemapRequest(kind=kind))

# --------------------------------------------------------------------------
# Right panel: Site Structure, Response Times; Visualisations: word clouds
# --------------------------------------------------------------------------
@app.get("/api/crawl/structure")
def get_structure(levels: int = 0):
    """Site Structure: folder tree (URLs, indexable / non-indexable, worst status per folder) and the crawl
    depth histogram; the numbers add up to the internal URLs crawled."""
    job = session.job
    site, index = _site_index(job)
    if site is None:
        return {"total": 0, "indexable": 0, "non_indexable": 0, "status": {}, "tree": [], "depth": [],
                "top_folders": []}
    from crawler.seo.structure import site_structure
    return _memo(job, ("structure", id(index), int(levels or 0)), lambda: site_structure(site, max(0, int(levels or 0))))

def _buckets(raw: str) -> Optional[List[float]]:
    if not (raw or "").strip():
        return None
    try:
        vals = sorted({float(x) for x in raw.split(",") if x.strip()})
    except ValueError:
        _bad("buckets must be numbers of seconds, comma separated (e.g. 1,2,3,5,10).")
    if not vals or vals[0] <= 0 or len(vals) > 50:
        _bad("buckets must be 1 to 50 positive numbers of seconds.")
    return vals

def _rt_scope(scope: str) -> str:
    s = (scope or "all").strip().lower()
    if s not in ("all", "html"):
        _bad("scope must be all or html.")
    return s

@app.get("/api/crawl/response-times")
def get_response_times(scope: str = "all", buckets: str = ""):
    """Response Times: URLs per response-time bucket (0-1 s ... 10 s+, or custom `buckets` edges in seconds),
    plus No response, average, median, 90th percentile and slowest; the buckets add up to the total."""
    job = session.job
    site, _ = _site_index(job)
    from crawler.seo.structure import response_times
    if site is None:
        return {"total": 0, "answered": 0, "buckets": [], "no_response": 0, "blocked": 0}
    return response_times(site, _buckets(buckets), _rt_scope(scope))

@app.get("/api/crawl/response-times/urls")
def get_response_time_urls(bucket: Optional[int] = None, scope: str = "all", buckets: str = "", offset: int = 0,
                           limit: int = 500):
    """The URLs of one response-time bucket (bucket omitted = No response)."""
    job = session.job
    site, _ = _site_index(job)
    if site is None:
        return {"total": 0, "urls": []}
    from crawler.seo.structure import urls_in_bucket
    found = urls_in_bucket(site, bucket, _buckets(buckets), _rt_scope(scope))
    urls = list(found.get("urls") if isinstance(found, dict) else found or [])
    offset, limit = max(0, offset), max(1, min(limit, 5000))
    out = dict(found) if isinstance(found, dict) else {}
    out.update(total=len(urls), urls=urls[offset:offset + limit])
    return out

@app.get("/api/crawl/wordcloud")
def get_word_cloud(url: str = "", source: str = "anchors", limit: int = 100):
    """Inlink Anchor Text Word Cloud (source=anchors) or Body Text Word Cloud (source=body) of one URL."""
    src = (source or "anchors").strip().lower()
    if src not in ("anchors", "body"):
        _bad("source must be anchors or body.")
    job, site, _ = _need_index()
    from crawler.seo.structure import word_cloud
    r = _find_crawled(site, url)
    return word_cloud(site, r.url if r is not None else url, src, max(10, min(int(limit or 100), 300)))

# --------------------------------------------------------------------------
# Mode > SERP (stateless: never touches the crawl on screen)
# --------------------------------------------------------------------------
class SerpRequest(BaseModel):
    rows: List[dict] = []
    format: Optional[str] = "csv"

SERP_MAX_ROWS = 50_000

def _serp_rows(rows: List[dict]) -> List[dict]:
    if len(rows) > SERP_MAX_ROWS:
        _bad(f"At most {SERP_MAX_ROWS:,} rows.")
    return [{"url": str((r or {}).get("url") or "")[:4000], "title": str((r or {}).get("title") or ""),
             "description": str((r or {}).get("description") or "")} for r in rows if isinstance(r, dict)]

@app.post("/api/serp/measure")
def serp_measure(req: SerpRequest):
    """SERP mode: each row's title and description length and pixel width, with the Page Titles / Meta
    Description filters it falls under, plus the filter counts (same thresholds as the crawl's tabs)."""
    from crawler import serp
    rows = serp.measure(_serp_rows(req.rows))
    return {"rows": rows, "summary": serp.summary(rows), "filters": serp.filter_names()}

@app.post("/api/serp/upload")
async def serp_upload(request: Request):
    """SERP mode upload (.csv / .txt / .xlsx with URL, Title, Description columns): the rows measured, and the
    lines that could not be read."""
    name, data = await _read_upload(request)
    from crawler import serp
    rows, rejected = await asyncio.get_running_loop().run_in_executor(None, serp.parse_upload, data, name)
    rows = serp.measure(rows[:SERP_MAX_ROWS])
    return {"rows": rows, "rejected": rejected, "summary": serp.summary(rows), "filters": serp.filter_names()}

@app.post("/api/serp/export")
def serp_export(req: SerpRequest):
    """The measured SERP rows as CSV / XLSX (SERP Summary layout, re-uploadable)."""
    from crawler import serp
    from crawler.exporter import Table, tables_bytes
    rows = serp.measure(_serp_rows(req.rows))
    name, data, mime = tables_bytes([Table("SERP Summary", list(serp.SUMMARY_HEADERS), serp.summary_rows(rows),
                                           url_cols=(1,))], _fmt(req.format or "csv"), "serp_summary")
    return _download(name, data, mime)

@app.get("/api/serp/crawl")
def serp_from_crawl():
    """The crawl's titles and descriptions as SERP rows (Reports > SERP Summary, editable in SERP mode)."""
    job = session.job
    site, _ = _site_index(job)
    from crawler import serp
    rows = serp.crawl_rows(site) if site is not None else []
    return {"rows": rows, "summary": serp.summary(rows), "filters": serp.filter_names()}

# --------------------------------------------------------------------------
# Mode > List: read an uploaded / pasted list (or an XML sitemap) into URLs
# --------------------------------------------------------------------------
LIST_SITEMAP_MAX_FILES = 50

def _list_from_table(name: str, data: bytes) -> List[str]:
    """Lines of a .txt / .csv (first column, or a column named URL / Address) / .xlsx / .xml upload."""
    lower = name.lower()
    if lower.endswith(".gz") or data[:2] == b"\x1f\x8b":
        import gzip
        try:
            data = gzip.decompress(data)[:MAX_UPLOAD_BYTES * 4]
        except OSError:
            _bad("The .gz file could not be unpacked.")
        lower = lower[:-3]
    if lower.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            _bad("The Excel file could not be read.")
        ws = wb.worksheets[0]
        values = [list(r) for r in ws.iter_rows(values_only=True, max_row=_list_cap() + 1000)]
        wb.close()
        return _column_values(values)
    text = data.decode("utf-8-sig", "replace")
    if lower.endswith(".xml") or text.lstrip()[:200].lower().startswith(("<?xml", "<urlset", "<sitemapindex")):
        return [html_lib.unescape(m) for m in re.findall(r"<loc>\s*([^<\s]+)\s*</loc>", text, re.I)]
    if lower.endswith(".csv") or ("," in text.split("\n", 1)[0] and lower.endswith((".csv", ".tsv"))):
        import csv
        rows = list(csv.reader(io.StringIO(text)))
        return _column_values(rows)
    return text.splitlines()

def _column_values(rows: List[list]) -> List[str]:
    """The URL column of a table: a column whose header is URL / Address / Loc, else the first column."""
    if not rows:
        return []
    head = [str(c or "").strip().lower() for c in rows[0]]
    col = next((i for i, h in enumerate(head) if h in ("url", "urls", "address", "loc", "page", "link")), None)
    body = rows[1:] if col is not None else rows
    col = col or 0
    return [str(r[col]).strip() for r in body if len(r) > col and r[col] is not None and str(r[col]).strip()]

def _fetch_sitemap_urls(sitemap_url: str, user_agent: str = DEFAULT_USER_AGENT) -> Tuple[List[str], List[str]]:
    """Page URLs of an XML sitemap (or a sitemap index, LIST_SITEMAP_MAX_FILES files at most) fetched now."""
    import gzip
    urls: List[str] = []
    notes: List[str] = []
    todo, seen = [sitemap_url], set()
    sess = requests.Session()
    sess.headers.update({"User-Agent": user_agent})
    proxy = _proxy_config() or {}
    try:
        from crawler.engine import proxy_url
        p = proxy_url(proxy)
        if p:
            sess.proxies.update({"http": p, "https": p})
    except ImportError:
        pass
    while todo and len(seen) < LIST_SITEMAP_MAX_FILES and len(urls) < _list_cap():
        u = todo.pop(0)
        if u in seen:
            continue
        seen.add(u)
        try:
            resp = sess.get(u, timeout=20, verify=False, stream=True)
            raw = resp.raw.read(MAX_UPLOAD_BYTES * 2, decode_content=True)
        except Exception as e:
            notes.append(f"{u}: could not be fetched ({type(e).__name__}).")
            continue
        if resp.status_code != 200:
            notes.append(f"{u}: answered HTTP {resp.status_code}.")
            continue
        if raw[:2] == b"\x1f\x8b":
            try:
                raw = gzip.decompress(raw)
            except OSError:
                notes.append(f"{u}: the .gz file could not be unpacked.")
                continue
        text = raw.decode("utf-8", "replace")
        locs = [html_lib.unescape(m) for m in re.findall(r"<loc>\s*([^<\s]+)\s*</loc>", text, re.I)]
        if "<sitemapindex" in text[:5000].lower():
            todo.extend(locs)
        else:
            urls.extend(locs)
    if todo:
        notes.append(f"Only the first {LIST_SITEMAP_MAX_FILES} sitemap files were read.")
    return urls, notes

def _clean_list(lines: List[str]) -> dict:
    """The upload's lines as web addresses (in order, repeats removed) and the lines that are not addresses."""
    urls: List[str] = []
    rejected: List[dict] = []
    keys = set()
    dup = 0
    for i, line in enumerate(lines):
        addr = _as_web_address(line)
        if not addr:
            rejected.append({"line": i + 1, "value": line[:300], "reason": "Not a web address"})
            continue
        try:
            key = normalize_url(addr) or addr
        except Exception:
            key = addr
        if key in keys:
            dup += 1
            continue
        keys.add(key)
        urls.append(addr)
    cap = _list_cap()
    return {"urls": urls[:cap], "rejected": rejected[:1000], "rejected_count": len(rejected), "lines": len(lines),
            "unique": len(urls), "duplicates": dup, "truncated": len(urls) > cap, "cap": cap}

@app.post("/api/list/parse")
async def parse_list(request: Request):
    """Mode > List: turns a pasted list ({"text"}), an uploaded file (.txt / .csv / .xlsx / sitemap .xml, as
    multipart, base64 JSON or the raw body) or an XML sitemap address ({"sitemap_url"}) into the web addresses
    to crawl, in order and without repeats, plus the lines that are not addresses. Nothing is crawled yet."""
    ctype = (request.headers.get("content-type") or "").lower()
    if ctype.startswith("application/json"):
        body = await request.body()
        try:
            obj = json.loads(body.decode("utf-8") or "{}")
        except ValueError:
            _bad("The request body is not valid JSON.")
        if isinstance(obj, dict) and obj.get("sitemap_url"):
            sm = _as_web_address(str(obj["sitemap_url"]))
            if not sm:
                _bad("Enter the sitemap's full address (https://...).")
            lines, notes = await asyncio.get_running_loop().run_in_executor(None, _fetch_sitemap_urls, sm)
            out = _clean_list(lines)
            out["notes"] = notes
            return out
        if isinstance(obj, dict) and isinstance(obj.get("urls"), list):
            return _clean_list([str(x) for x in obj["urls"]])
    name, data = await _read_upload(request)
    lines = await asyncio.get_running_loop().run_in_executor(None, _list_from_table, name, data)
    return _clean_list(lines)

# --------------------------------------------------------------------------
# Configuration: live testers for URL Rewriting / Include / Exclude and the custom rules
# --------------------------------------------------------------------------
class ConfigTestRequest(BaseModel):
    url: Optional[str] = ""
    urls: Optional[List[str]] = None
    url_rewrite: Optional[dict] = None
    include_patterns: Optional[Union[List[str], str]] = None
    exclude_regex: Optional[Union[List[str], str]] = None
    custom_searches: Optional[List[dict]] = None
    custom_extractors: Optional[List[dict]] = None
    custom_js: Optional[List[dict]] = None
    request_headers: Optional[List[dict]] = None

def _regex_errors(raw) -> List[dict]:
    out = []
    for i, p in enumerate(_lines(raw), start=1):
        try:
            re.compile(p)
        except re.error as e:
            out.append({"line": i, "pattern": p[:200], "error": str(e)})
    return out

@app.post("/api/config/test-url")
def config_test_url(req: ConfigTestRequest):
    """URL Rewriting / Include / Exclude tester: each URL before and after rewriting and whether Include /
    Exclude let it be crawled (Python's regex rules, exactly as the crawl applies them)."""
    from crawler.url_utils import normalize_url as norm, rules_from_config
    rw = default_url_rewrite()
    if req.url_rewrite:
        rw.update({k: v for k, v in req.url_rewrite.items() if k in rw})
        if req.url_rewrite.get("strip_tracking_params") is False:
            rw["remove_params"] = []
    rules, problems = rules_from_config(rw)
    inc_err, exc_err = _regex_errors(req.include_patterns), _regex_errors(req.exclude_regex)
    inc = [re.compile(p) for p in _lines(req.include_patterns) if not any(e["pattern"] == p[:200] for e in inc_err)]
    exc = [re.compile(p) for p in _lines(req.exclude_regex) if not any(e["pattern"] == p[:200] for e in exc_err)]
    tests = []
    for u in ([req.url] if req.url else []) + list(req.urls or [])[:200]:
        raw = (u or "").strip()
        if not raw:
            continue
        try:
            after = norm(raw, rules=rules) or ""
        except Exception as e:
            after, problems = "", problems + [f"{raw[:80]}: {type(e).__name__}"]
        target = after or raw
        excluded_by = next((p.pattern for p in exc if p.search(target)), "")
        included = not inc or any(p.search(target) for p in inc)
        tests.append({"url": raw, "rewritten": after, "changed": bool(after) and after != raw,
                      "included": included, "excluded_by": excluded_by,
                      "crawled": bool(after) and included and not excluded_by})
    return {"results": tests, "problems": [str(p) for p in problems], "include_errors": inc_err,
            "exclude_errors": exc_err}

@app.post("/api/config/validate")
def config_validate(req: ConfigTestRequest):
    """Every crawl option the dialogs edit, checked without starting a crawl: {"ok", "errors": {field: msg}}."""
    errors: Dict[str, Any] = {}
    for field, raw in (("include_patterns", req.include_patterns), ("exclude_regex", req.exclude_regex)):
        bad = _regex_errors(raw)
        if bad:
            errors[field] = bad
    for field, fn in (("url_rewrite", lambda: _url_rewrite_config(req.url_rewrite)),
                      ("request_headers", lambda: _request_headers(req.request_headers)),
                      ("custom", lambda: _custom_rules(CrawlRequest(url="", custom_searches=req.custom_searches,
                                                                    custom_extractors=req.custom_extractors,
                                                                    custom_js=req.custom_js)))):
        try:
            fn()
        except HTTPException as e:
            errors[field] = e.detail
    return {"ok": not errors, "errors": errors}

# --------------------------------------------------------------------------
# File > Crawls: saved crawls (open, close, rename, project, lock, duplicate, delete, import, export)
# --------------------------------------------------------------------------
def _crawl_id(host: str, ts: str) -> str:
    cid = f"{host}/{ts}"
    if not re.fullmatch(r"[a-z0-9._-]{1,200}/\d{8}T\d{6}Z(?:-\d{1,3})?", cid):
        raise HTTPException(status_code=404, detail="No such saved crawl.")
    return cid

def _store_call(fn: Callable[[], Any]) -> Any:
    """A crawler/store.py call with its StoreError as a 400 (404 when the crawl does not exist)."""
    store = _store()
    try:
        return fn()
    except store.StoreError as e:
        msg = str(e)
        raise HTTPException(status_code=404 if "not found" in msg.lower() or "no saved crawl" in msg.lower() else 400,
                            detail=msg)

@app.get("/api/crawls/saved")
def list_saved_crawls(site: str = "", project: Optional[str] = None, q: str = ""):
    """File > Crawls: every saved crawl (newest first) with its meta (name, project, lock, URLs, status
    classes, health score, issues, size); search with q, limit to a site or a project."""
    store = _store()
    rows = store.list_crawls(REPORTS_DIR, site=site.strip() or None, project=project, q=q)
    job = session.job
    current = (job.saved_meta or {}).get("id") if job is not None else ""
    return {"crawls": rows, "projects": store.projects(REPORTS_DIR), "current": current or "",
            "busy": _crawl_busy()}

@app.get("/api/crawls/storage")
def crawls_storage():
    """Settings > Storage (read-only): where crawls are kept, how much space they use, free disk space."""
    info = _store().storage_info(REPORTS_DIR)
    if not os.environ.get("CRAWLER_REPORTS_DIR"):
        info.setdefault("warning", "CRAWLER_REPORTS_DIR is not set: crawls are kept in the system temp folder, "
                                   "which a redeploy or clean-up can empty.")
    return info

class SaveCrawlRequest(BaseModel):
    name: Optional[str] = None
    project: Optional[str] = None

@app.post("/api/crawls/save")
def save_current_crawl(req: Optional[SaveCrawlRequest] = None):
    """Saves the crawl on screen now (File > Save), e.g. when automatic saving is off."""
    job = session.job
    if job is None or job.crawler is None:
        raise HTTPException(status_code=404, detail="No crawl to save.")
    if _crawl_busy(job) or job.phase != "done":
        raise HTTPException(status_code=409, detail="Wait until the crawl and its reports are finished.")
    if req is not None:
        if req.name is not None:
            job.save_name = req.name.strip()[:120]
        if req.project is not None:
            job.save_project = req.project.strip()[:60]
    meta = _autosave(job, force=True, restore=True)
    if meta is None:
        _bad("The crawl could not be saved (see the crawl log).")
    if req is not None and req.name:
        meta = _store_call(lambda: _store().rename(meta["id"], req.name, REPORTS_DIR))
    if req is not None and req.project is not None:
        meta = _store_call(lambda: _store().set_project(meta["id"], req.project or "", REPORTS_DIR))
    job.saved_meta = meta
    return meta

@app.post("/api/crawls/import")
async def import_crawl(request: Request, name: str = "", project: str = ""):
    """Imports a crawl exported from this or another ARACHNIX (.zip; checked for format, version, size
    and unsafe paths). The crawl is added to File > Crawls; open it from there."""
    store = _store()
    name_, data = await _read_upload(request, limit=store.MAX_IMPORT_BYTES)
    if not data:
        _bad("Choose the exported crawl file (.zip).")
    try:
        meta = await asyncio.get_running_loop().run_in_executor(
            None, lambda: store.import_bytes(data, name_, REPORTS_DIR, name=name or None, project=project or None))
    except store.StoreError as e:
        _bad(str(e))
    return meta

@app.post("/api/crawls/projects/lock")
def lock_project(project: str = "", locked: bool = True):
    if not project.strip():
        _bad("Name the project.")
    return {"crawls": _store_call(lambda: _store().lock_project(project.strip(), locked, REPORTS_DIR))}

@app.get("/api/crawls/{host}/{ts}")
def get_saved_crawl(host: str, ts: str):
    meta = _store().get_meta(_crawl_id(host, ts), REPORTS_DIR)
    if meta is None:
        raise HTTPException(status_code=404, detail="No such saved crawl.")
    return meta

@app.post("/api/crawls/{host}/{ts}/open")
def open_saved_crawl(host: str, ts: str):
    """Opens a saved crawl in place of the one on screen: every tab, report and export works on it. Refused
    while a crawl runs (one shared crawl per server)."""
    cid = _crawl_id(host, ts)
    store = _store()
    with _start_lock:
        if _crawl_busy():
            raise HTTPException(status_code=409, detail=_busy_detail() + " Open the saved crawl afterwards.")
        loaded = _store_call(lambda: store.load_crawl(cid, REPORTS_DIR))
        old = session.job
        if old is not None:
            old.stop_jobs()
        job = CrawlJob(loaded.crawler, loaded.base_name, screenshots=False)
        loaded.apply_to_job(job)
        job.origin = "opened"
        job.saved_meta = dict(loaded.meta)
        job.notes = list(loaded.notes or [])
        job.after_done = True
        if loaded.dismissals:
            _safe(lambda: _dismiss_store(job).merge(loaded.dismissals), None)
        integ = (loaded.extras or {}).get("integrations")
        if integ and not getattr(loaded.crawler, "integrations", None):
            _safe(lambda: _integrations().import_state(loaded.crawler, integ), None)
        session.job = job
    site, index = _site_index(job, force=True)
    diffs = _safe(lambda: loaded.count_differences(index), []) if index is not None else []
    return {"status": "opened", "id": cid, "meta": loaded.meta, "notes": loaded.notes,
            "count_differences": diffs[:50], "base_name": job.base_name}

@app.post("/api/crawl/close")
def close_crawl():
    """File > Close: takes the crawl off the screen without deleting it (saved crawls stay in File > Crawls)."""
    job = session.job
    if job is None:
        return {"status": "idle"}
    if _crawl_busy(job):
        raise HTTPException(status_code=409, detail="Stop the running crawl first." if _crawl_running(job)
                            else _busy_detail(job))
    job.stop_jobs()
    session.job = None
    return {"status": "idle", "closed": (job.saved_meta or {}).get("id") or job.snapshot_id}

class RenameRequest(BaseModel):
    name: str = ""

class ProjectRequest(BaseModel):
    project: str = ""

class LockRequest(BaseModel):
    locked: bool = True

def _sync_current(meta: dict) -> dict:
    job = session.job
    if job is not None and job.saved_meta and job.saved_meta.get("id") == meta.get("id"):
        job.saved_meta = meta
        job.save_name = meta.get("name") or job.save_name
        job.save_project = meta.get("project") or ""
    return meta

@app.post("/api/crawls/{host}/{ts}/rename")
def rename_saved_crawl(host: str, ts: str, req: RenameRequest):
    if not req.name.strip():
        _bad("Enter a name.")
    return _sync_current(_store_call(lambda: _store().rename(_crawl_id(host, ts), req.name, REPORTS_DIR)))

@app.post("/api/crawls/{host}/{ts}/project")
def set_saved_crawl_project(host: str, ts: str, req: ProjectRequest):
    return _sync_current(_store_call(lambda: _store().set_project(_crawl_id(host, ts), req.project, REPORTS_DIR)))

@app.post("/api/crawls/{host}/{ts}/lock")
def lock_saved_crawl(host: str, ts: str, req: LockRequest):
    return _sync_current(_store_call(lambda: _store().lock(_crawl_id(host, ts), req.locked, REPORTS_DIR)))

@app.post("/api/crawls/{host}/{ts}/duplicate")
def duplicate_saved_crawl(host: str, ts: str, req: Optional[RenameRequest] = None):
    name = (req.name if req else "") or None
    return _store_call(lambda: _store().duplicate(_crawl_id(host, ts), REPORTS_DIR, name=name))

@app.delete("/api/crawls/{host}/{ts}")
def delete_saved_crawl(host: str, ts: str, force: bool = False, with_snapshot: bool = True):
    """Deletes a saved crawl (and its snapshot). Locked crawls are refused unless force=1; the crawl on screen
    stays on screen."""
    cid = _crawl_id(host, ts)
    out = _store_call(lambda: _store().delete(cid, REPORTS_DIR, with_snapshot=with_snapshot, force=force))
    if not out:
        raise HTTPException(status_code=404, detail="No such saved crawl.")
    job = session.job
    if job is not None and job.saved_meta and job.saved_meta.get("id") == cid:
        job.saved_meta = None
        job.saved_deleted = True     # a later re-check or fetch does not save it again; File > Save brings it
                                     # back under the same id (crawler.saved_crawl keeps it)
    return {"deleted": cid, "result": out}

@app.get("/api/crawls/{host}/{ts}/export")
def export_saved_crawl(host: str, ts: str):
    """The saved crawl as one .zip (results, settings without secrets, page sources, pictures) to import on
    another ARACHNIX."""
    cid = _crawl_id(host, ts)
    store = _store()
    fd, tmp = tempfile.mkstemp(prefix="crawl_export_", suffix=".zip", dir=get_export_dir())
    os.close(fd)
    try:
        with open(tmp, "wb") as fh:
            name = _store_call(lambda: store.export_to(cid, fh, REPORTS_DIR))
    except Exception:
        _safe(lambda: os.remove(tmp), None)
        raise
    from starlette.background import BackgroundTask
    return FileResponse(tmp, filename=name, media_type="application/zip",
                        background=BackgroundTask(lambda: _safe(lambda: os.remove(tmp), None)))

# --------------------------------------------------------------------------
# File > Settings (server-wide) and Configuration > Profiles
# --------------------------------------------------------------------------
@app.get("/api/settings")
def get_settings(request: Request):
    """Every server-wide settings page. Secrets come back only as {"set", "source"}, never their value;
    "editable" says whether this viewer may change them (see settings_editable)."""
    s = _settings()
    out = s.get_settings(True, REPORTS_DIR)
    ok, why = settings_editable(request)
    out["_editable"], out["_editable_reason"] = ok, why
    out["_notify"] = _safe(lambda: _notify().status(REPORTS_DIR), {})
    out["_storage"] = _safe(lambda: crawls_storage(), {})
    out["_pages"] = list(s.PAGES)
    return out

@app.put("/api/settings/{page}")
def put_settings(page: str, request: Request, values: Dict[str, Any]):
    """Saves one settings page after checking every field: 400 {"errors": {field: message}} or the page as
    GET shows it. A secret sent as "", "***" or {"set": ...} is kept; {"clear": true} removes it."""
    _require_admin(request)
    s = _settings()
    if page not in s.PAGES:
        raise HTTPException(status_code=404, detail=f"Unknown settings page: {page}")
    errors = s.update(page, values or {}, REPORTS_DIR)
    if errors:
        return JSONResponse({"detail": "Some fields are not valid.", "errors": errors}, status_code=400)
    if page == "retention":
        _apply_retention()
    return {"page": page, "values": s.get_page(page, True, REPORTS_DIR)}

@app.delete("/api/settings/{page}")
def reset_settings(page: str, request: Request):
    """Restores one settings page to its defaults."""
    _require_admin(request)
    s = _settings()
    if page not in s.PAGES:
        raise HTTPException(status_code=404, detail=f"Unknown settings page: {page}")
    s.reset(page, REPORTS_DIR)
    return {"page": page, "values": s.get_page(page, True, REPORTS_DIR)}

@app.post("/api/settings/proxy/test")
def proxy_test(request: Request, values: Dict[str, Any]):
    """Opens one page through the proxy as the form shows it (saved password kept when the form leaves it
    unchanged): {"ok", "message", "status_code", "url", "ms"}. Messages never contain the proxy password."""
    _require_admin(request)
    s = _settings()
    body = dict(values or {})
    test_url = str(body.pop("test_url", "") or "").strip()
    merged, errors = s.validate("proxy", body, REPORTS_DIR)
    if errors:
        return {"ok": False, "message": "; ".join(f"{k}: {v}" for k, v in errors.items())}
    if not merged.get("enabled") or not str(merged.get("host") or "").strip():
        return {"ok": False, "message": "Turn on Use Proxy Server and enter the proxy host first."}
    creds = bool(merged.get("use_credentials"))
    proxy = {"enabled": True, "host": merged.get("host") or "", "port": int(merged.get("port") or 3128),
             "username": merged.get("username") if creds else "", "password": merged.get("password") if creds else ""}
    from crawler.engine import proxy_url
    purl = proxy_url(proxy)
    if not purl:
        return {"ok": False, "message": "The proxy host is not valid."}
    if not test_url:
        job = session.job
        test_url = (getattr(job.crawler, "start_url", "") if job is not None and job.crawler else "") or \
            "https://example.com/"
    if not re.match(r"^https?://[^\s/]+", test_url, re.I):
        return {"ok": False, "message": "The test address must start with http:// or https://."}
    hidden = [str(proxy.get("password") or ""), str(proxy.get("username") or "")]

    def masked(text: str) -> str:
        out = str(text)
        for h in hidden:
            if h:
                out = out.replace(h, "***")
                out = out.replace(requests.utils.quote(h, safe=""), "***")
        return out[:300]

    started = time.time()
    try:
        r = requests.get(test_url, proxies={"http": purl, "https": purl}, timeout=12, allow_redirects=False,
                         headers={"User-Agent": "Mozilla/5.0 (compatible; SiteAuditor proxy test)"}, stream=True)
        r.close()
    except requests.exceptions.ProxyError as e:
        text = str(e)
        if "407" in text:
            return {"ok": False, "url": test_url, "message": "The proxy asks for a user name and password, or "
                    "refused the ones given (407 Proxy Authentication Required)."}
        return {"ok": False, "url": test_url, "message": "The proxy could not be reached or refused the "
                f"connection ({masked(type(e).__name__)})."}
    except requests.exceptions.RequestException as e:
        return {"ok": False, "url": test_url,
                "message": f"The page did not load through the proxy: {masked(type(e).__name__)}."}
    ms = int((time.time() - started) * 1000)
    if r.status_code == 407:
        return {"ok": False, "url": test_url, "status_code": 407, "ms": ms,
                "message": "The proxy asks for a user name and password, or refused the ones given (407)."}
    return {"ok": True, "url": test_url, "status_code": r.status_code, "ms": ms,
            "message": f"{test_url} answered {r.status_code} through the proxy in {ms:,} ms."}

class NotifyTestRequest(BaseModel):
    to: Optional[Union[str, List[str]]] = None

@app.post("/api/settings/notifications/test")
@app.post("/api/notify/test")
def notify_test(request: Request, req: NotifyTestRequest):
    """Sends "Test email from ARACHNIX" through the saved SMTP settings (errors in plain words, never the
    password)."""
    _require_admin(request)
    ok, bad = _settings().parse_emails(req.to or [])
    if bad:
        _bad(f"Not an email address: {', '.join(bad[:3])}")
    if not ok:
        _bad("Enter at least one email address.")
    notify = _notify()
    if not notify.configured(reports_dir=REPORTS_DIR):
        _bad("Set the SMTP server first (Settings > Notifications).")
    return notify.send_test(ok, reports_dir=REPORTS_DIR)

@app.get("/api/profiles")
def list_profiles():
    s = _settings()
    return {"profiles": s.list_profiles(REPORTS_DIR), "default": s.get_default(REPORTS_DIR)}

class ProfileRequest(BaseModel):
    name: Optional[str] = None
    options: Dict[str, Any] = {}

@app.get("/api/profiles/_default")
def get_default_profile():
    return {"default": _settings().get_default(REPORTS_DIR)}

@app.put("/api/profiles/_default")
def set_default_profile(req: ProfileRequest):
    """Makes these crawl options the default for new crawls (secrets are removed first)."""
    s = _settings()
    errors = s.set_default(req.options or {}, REPORTS_DIR, name=(req.name or "Default"))
    if errors:
        return JSONResponse({"detail": "Some options are not valid.", "errors": errors}, status_code=400)
    return {"default": s.get_default(REPORTS_DIR)}

@app.delete("/api/profiles/_default")
def clear_default_profile():
    return {"cleared": bool(_settings().clear_default(REPORTS_DIR))}

@app.post("/api/profiles/import")
async def import_profile(request: Request, name: str = ""):
    _fname, data = await _read_upload(request, limit=1024 * 1024)
    meta, errors = _settings().import_profile(data, name=name or None, reports_dir=REPORTS_DIR)
    if errors:
        return JSONResponse({"detail": "The profile could not be imported.", "errors": errors}, status_code=400)
    return meta

@app.post("/api/profiles")
def create_profile(req: ProfileRequest):
    if not (req.name or "").strip():
        _bad("Name the profile.")
    return put_profile(req.name, req)

@app.get("/api/profiles/{name}/export")
def export_profile(name: str):
    out = _settings().export_profile(name, REPORTS_DIR)
    if not out:
        raise HTTPException(status_code=404, detail="No such profile.")
    fname, data = out
    return _download(fname, data, "application/json")

@app.get("/api/profiles/{name}")
def get_profile(name: str):
    prof = _settings().load_profile(name, REPORTS_DIR)
    if prof is None:
        raise HTTPException(status_code=404, detail="No such profile.")
    return prof

@app.put("/api/profiles/{name}")
def put_profile(name: str, req: ProfileRequest):
    """Saves the crawl options as a named profile (shared by the team; passwords, cookie and header values,
    tokens and keys are always removed)."""
    meta, errors = _settings().save_profile(name, req.options or {}, REPORTS_DIR)
    if errors:
        return JSONResponse({"detail": "The profile could not be saved.", "errors": errors}, status_code=400)
    return meta

@app.delete("/api/profiles/{name}")
def delete_profile(name: str):
    if not _settings().delete_profile(name, REPORTS_DIR):
        raise HTTPException(status_code=404, detail="No such profile.")
    return {"deleted": name}

# --------------------------------------------------------------------------
# Notifications: the crawl-complete email (Settings > Notifications > Email on Crawl Complete)
# --------------------------------------------------------------------------
def _previous_snapshot(job: CrawlJob) -> str:
    """The newest earlier snapshot of the same site (not this crawl's own, nor a re-check entry)."""
    from crawler.seo.compare import list_snapshots
    own = job.snapshot_id or ""
    for s in list_snapshots(REPORTS_DIR, host=job.crawler.start_url):
        sid = str(s.get("id") or "")
        if sid and sid != own and not sid.endswith("-r") and not (own and sid.startswith(own)):
            return sid
    return ""

def _comparison_with_previous(job: CrawlJob) -> Optional[dict]:
    prev = _previous_snapshot(job)
    if not prev:
        return None
    from crawler.seo import compare
    return compare.compare(prev, "current", REPORTS_DIR, current=_current_snapshot("current", ""))

def _crawl_email(job: CrawlJob, to: List[str], attach: bool, dashboard_url: str = "",
                 compare_result: Optional[dict] = None, extra_files: Optional[List[str]] = None,
                 status_override: str = "", error: str = "") -> dict:
    notify = _notify()
    site, index = _site_index(job)
    d = _safe(lambda: _dismissals(job, index), _empty_dismissals(index)) if index is not None else None
    summary = notify.build_summary(job.crawler, site=site, index=index, compare_result=compare_result,
                                   dashboard_url=dashboard_url, issues=d["issues"] if d else None,
                                   status_override=status_override, error=error)
    files: List[Any] = list(extra_files or [])
    if attach and not files and index is not None:
        path = _safe(lambda: _excel_file(job, "all"), None, job, "Email attachment")
        if path:
            files.append(path)
    return notify.send_crawl_complete(summary, to, attachments=files or None, reports_dir=REPORTS_DIR, mask=job.mask)

def _email_on_complete(job: CrawlJob):
    """Settings > Notifications > Email on Crawl Complete (dashboard crawls; scheduled tasks send their own)."""
    try:
        cfg = _settings().effective("notifications", REPORTS_DIR)
    except Exception:
        return
    if not cfg.get("email_on_complete") or not cfg.get("recipients") or job.retired:
        return
    notify = _notify()
    if not notify.configured(reports_dir=REPORTS_DIR):
        job.crawler.log("[Email] Email on crawl complete is on, but no SMTP server is set.")
        return
    to, _bad_addr = _settings().parse_emails(cfg.get("recipients"))
    try:
        res = _crawl_email(job, to, bool(cfg.get("attach_exports")), cfg.get("dashboard_url") or "",
                           compare_result=_safe(lambda: _comparison_with_previous(job), None))
        job.crawler.log(f"[Email] Crawl summary sent to {len(res.get('sent_to') or [])} recipient(s)."
                        if res.get("ok") else f"[Email] Not sent: {job.mask(res.get('error'))}")
    except Exception as e:
        job.crawler.log(f"[Email] Not sent: {job.mask(e)}")

# --------------------------------------------------------------------------
# File > Scheduling: tasks that start crawls on their own (crawler/scheduler.py decides when)
# --------------------------------------------------------------------------
_scheduler_holder: Dict[str, Any] = {"scheduler": None, "lock_fh": None, "reason": "", "passive": None}
SCHEDULER_LOCK_OFFSET = 1 << 30
_scheduled_jobs: Dict[str, CrawlJob] = {}       # run_ref -> the job a task started, until its after-crawl steps end

# A scheduled login reads a secret from a server environment variable and sends it to the task's start URL. Only
# variables named for that purpose can be used (SITE_SESSION_COOKIE, SITE_TOKEN...): the server's own keys and
# passwords (AI keys, SMTP password, OAuth secret, access key...) are never sent to a website, and a task cannot
# be used to find out which other variables exist.
SCHEDULE_LOGIN_ENV_PREFIX = "SITE_"
LOGIN_ENV_RULE = (f"For safety a scheduled login can only use a server variable whose name starts with "
                  f"{SCHEDULE_LOGIN_ENV_PREFIX} (for example SITE_SESSION_COOKIE). Other server variables, such as "
                  "API keys and passwords, are never sent to a website.")
LOGIN_ADMIN_RULE = ("A task with a login sends a secret of this server to the site it crawls, so it can be added, "
                    "imported or pointed at another site only on the server itself (http://localhost:<port>), with "
                    "the team access key, or when CRAWLER_SETTINGS_EDITABLE=1.")

def _sched_login(task: Any) -> dict:
    """The task's login section ({} when it has none)."""
    st = task.get("start") if isinstance(task, dict) else None
    login = st.get("login") if isinstance(st, dict) else None
    return login if isinstance(login, dict) else {}

def _sched_uses_login(task: Any) -> bool:
    return str(_sched_login(task).get("type") or "none").strip().lower() in ("cookies", "header")

def _login_env_allowed(name: Any) -> bool:
    """Whether a scheduled login may read this server variable: SITE_<something> only."""
    name = str(name or "").strip().upper()
    return name.startswith(SCHEDULE_LOGIN_ENV_PREFIX) and len(name) > len(SCHEDULE_LOGIN_ENV_PREFIX)

def _sched_login_errors(task: Any) -> Dict[str, str]:
    """{field: error} when the task's login names a server variable it may not use."""
    env = str(_sched_login(task).get("env") or "").strip()
    if _sched_uses_login(task) and env and not _login_env_allowed(env):
        return {"start.login.env": LOGIN_ENV_RULE}
    return {}

def _sched_invalid(task: Any, login_errors: Dict[str, str]) -> JSONResponse:
    """400 with every error of the task (the scheduler's own checks too), so the form shows them at once."""
    from crawler.scheduler import validate_task
    errors = _safe(lambda: validate_task(task)[1], {}) or {}
    return JSONResponse({"detail": "The task is not valid.", "errors": {**errors, **login_errors}}, status_code=400)

def _sched_public(task: Any) -> Any:
    """A task as the API shows it. A login naming a variable it may not use (a task saved before that rule) is
    reported as such, never as "variable X is (not) set", which would tell which variables the server has."""
    if not isinstance(task, dict) or not _sched_uses_login(task):
        return task
    env = str(_sched_login(task).get("env") or "").strip()
    if env and not _login_env_allowed(env):
        task["problems"] = [p for p in task.get("problems") or []
                            if env not in str(p) and p != LOGIN_ENV_RULE] + [LOGIN_ENV_RULE]
        task["valid"] = False
    return task

_ENV_NOT_SET = re.compile(r"The environment variable (\S+) is not set on this server\.")

def _sched_runs_public(runs: List[dict]) -> List[dict]:
    """The run history as the API shows it: a skipped run of such an old task does not say whether a variable
    outside the SITE_ names is set either."""
    for r in runs:
        note = r.get("note") if isinstance(r, dict) else None
        if note and "environment variable" in note:
            r["note"] = _ENV_NOT_SET.sub(lambda m: m.group(0) if _login_env_allowed(m.group(1)) else LOGIN_ENV_RULE,
                                         note)
    return runs

def _sched_start_section(task: Any) -> dict:
    """The task's start section as the scheduler stores it (normalised), to tell whether an edit changes it."""
    from crawler.scheduler import validate_task
    return _safe(lambda: validate_task(task)[0].get("start") or {}, {}) or {}

def _login_admin_refusal(request: Request) -> Optional[JSONResponse]:
    """403 (also as the login field's error, which the Scheduling form shows) when this request may not set up a
    task with a login, else None."""
    if settings_editable(request)[0]:
        return None
    return JSONResponse({"detail": LOGIN_ADMIN_RULE, "errors": {"start.login.env": LOGIN_ADMIN_RULE}},
                        status_code=403)

def _scheduler_lock() -> Optional[Any]:
    """An exclusive lock on <reports_dir>/scheduler.lock, so only one server process on this computer runs the
    scheduled tasks of a reports folder (a second server, or a test server, only edits them)."""
    path = os.path.join(REPORTS_DIR, "scheduler.lock")
    try:
        os.makedirs(REPORTS_DIR, exist_ok=True)
        fh = open(path, "a+")
    except OSError:
        return None
    try:
        if os.name == "nt":
            import msvcrt
            # one byte far past the end of the (empty) file: the file itself stays readable for backups and scans
            fh.seek(SCHEDULER_LOCK_OFFSET)
            msvcrt.locking(fh.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl
            fcntl.flock(fh.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        return fh
    except OSError:
        fh.close()
        return None

def _schedule_start(task: dict, request: dict) -> str:
    """Scheduler start_fn: starts the task's crawl through the same code as POST /api/crawl/start. A task whose
    login names a server variable outside the SITE_ names (one saved before that rule, or edited by hand in
    schedules.json) is refused here, so no other secret of the server ever goes to a website."""
    from crawler.scheduler import SchedulerBusy
    if _sched_uses_login(task) and not _login_env_allowed(_sched_login(task).get("env")):
        raise ValueError(LOGIN_ENV_RULE)
    if _crawl_busy():
        raise SchedulerBusy("A crawl is running.")
    try:
        req = CrawlRequest(**{k: v for k, v in (request or {}).items() if k in CrawlRequest.model_fields})
    except Exception as e:
        raise ValueError(f"The task's crawl options are not valid: {str(e)[:200]}")
    req.force = False
    if not req.name:
        req.name = str(task.get("name") or "")[:120] or None
    if not req.project and task.get("project"):
        req.project = str(task.get("project"))[:60]
    ref = uuid.uuid4().hex[:16]
    try:
        _launch_crawl(req, origin="schedule", schedule_ref=ref)
    except HTTPException as e:
        if e.status_code == 409:
            raise SchedulerBusy(str(e.detail))
        raise ValueError(str(e.detail))
    return ref

def _release_scheduled(job: CrawlJob) -> None:
    """The scheduled run is over (after-crawl steps done, or there is nothing to report to): the server no longer
    keeps the crawl for it (it stays on screen until another replaces it) and it no longer counts as busy."""
    job.post_since = 0.0
    if job.schedule_ref and _scheduled_jobs.get(job.schedule_ref) is job:
        _scheduled_jobs.pop(job.schedule_ref, None)

def _schedule_finished(job: CrawlJob):
    """Tells the scheduler that the crawl a task started has ended (its exports and email then run). Also on a
    server that does not run the scheduled tasks itself (CRAWLER_SCHEDULER=0, a second server on the same reports
    folder): its "Run now" crawl is reported to the scheduler it was started from. Until the after-crawl steps
    have finished the crawl counts as busy (_post_busy), so the next task cannot replace it half-way."""
    if not job.schedule_ref:
        return
    sched = _scheduler_holder.get("scheduler") or _scheduler_holder.get("passive")
    if sched is None:
        _release_scheduled(job)
        return
    job.post_since = time.time()   # set while the crawl thread still runs: no moment where the server looks free
    status, error = _crawl_outcome(job)
    stats = _safe(lambda: _public_stats(job), {})
    info = {k: stats.get(k) for k in ("total_crawled", "status_2xx", "status_3xx", "status_4xx", "status_5xx",
                                      "status_err", "status_blocked", "health_score", "elapsed_seconds")}
    info["snapshot_id"] = job.snapshot_id

    def report():
        # a crawl that ends within a second can finish before the scheduler has recorded its run: retry briefly
        for _ in range(20):
            run = _safe(lambda: sched.crawl_finished(job.run_ref, status, error, info), None, job, "Scheduler")
            if run is not None:
                # the after-crawl steps release the crawl when they end; a task deleted meanwhile has none
                if getattr(sched, "on_complete", None) is None or _safe(lambda: sched.get(run.get("task_id") or ""),
                                                                         None) is None:
                    _release_scheduled(job)
                return
            time.sleep(0.5)
        _release_scheduled(job)
    threading.Thread(target=report, daemon=True, name="schedule-finished").start()

def _schedule_complete(task: dict, run: dict) -> dict:
    """After a scheduled crawl: its exports (Multi-Export items, Excel, PDF, sitemap) into
    <reports_dir>/scheduled/<task>/<time>/, the saved crawl, the comparison with the previous crawl and the
    email. Returns what the run history records. The crawl counts as busy until these steps end; then the
    server lets go of it (it stays on screen until another crawl replaces it)."""
    job = _scheduled_jobs.get(str(run.get("run_ref") or ""))
    out: Dict[str, Any] = {"exports": [], "emailed": [], "compare": None, "notes": [], "files": []}
    if job is None or job.crawler is None:
        out["notes"].append("The crawl is no longer on this server (it was replaced or the server restarted).")
        return out
    try:
        _schedule_steps(task, run, job, out)
        if job.retired and not getattr(job.crawler, "superseded", False):
            pdf = bool((task.get("exports") or {}).get("pdf"))
            out["notes"].append("Another crawl replaced this one on screen while its exports were being built"
                                + (": the PDF may miss some pictures." if pdf else "."))
    finally:
        _release_scheduled(job)
        _apply_retention()   # old scheduled-run folders go even when the task does not save its crawl
    return out

def _schedule_apis(task: dict, job: CrawlJob, out: Dict[str, Any]) -> None:
    """W8: the task's API data (Search Console, GA4, PageSpeed, link metrics, grammar) fetched after its crawl and
    waited for, at most apis.wait_minutes, so the exports and the email include it."""
    a = task.get("apis") or {}
    sources = [s for s in a.get("sources") or [] if s in ("gsc", "ga4", "psi", "link_metrics", "grammar")]
    if not sources or job.crawler is None:
        return
    I = _integrations()
    started, failed = [], []
    for s in sources:
        try:
            if s in ("gsc", "ga4", "psi"):
                if not (I.jobs(job.crawler).get(s) or {}).get("running"):
                    I.start_fetch(s, job.crawler)
            elif s == "link_metrics":
                from crawler.integrations import linkmetrics
                linkmetrics.start_fetch(job.crawler)
            elif s == "grammar":
                from crawler.connectors import languagetool
                languagetool.start_check(job.crawler)
            started.append(s)
        except Exception as e:
            failed.append(f"{s}: {job.mask(str(getattr(e, 'message', '') or e))[:120]}")

    def running() -> bool:
        busy = any((I.jobs(job.crawler).get(s) or {}).get("running") for s in started if s in ("gsc", "ga4", "psi"))
        try:
            if "link_metrics" in started:
                from crawler.integrations import linkmetrics
                busy = busy or (linkmetrics.fetch_status(job.crawler) or {}).get("state") == "running"
            if "grammar" in started:
                from crawler.connectors import languagetool
                busy = busy or (languagetool.job_status(job.crawler) or {}).get("state") == "running"
        except Exception:
            pass
        return busy
    until = time.time() + 60 * int(a.get("wait_minutes") if a.get("wait_minutes") is not None else 10)
    while running() and time.time() < until and not job.cancel_event.is_set():
        time.sleep(3)
    if started:
        out["notes"].append("API data fetched before the exports: " + ", ".join(started)
                            + (" (still running when the wait ended)" if running() else "") + ".")
    if failed:
        out["notes"].append("API data not fetched: " + "; ".join(failed))
    _site_index(job, force=True)


def _schedule_steps(task: dict, run: dict, job: CrawlJob, out: Dict[str, Any]) -> None:
    """The after-crawl steps of _schedule_complete; fills `out`."""
    if run.get("status") in ("completed", "stopped"):
        _schedule_apis(task, job, out)
    ex = task.get("exports") or {}
    notify_cfg = task.get("notify") or {}
    emails = list(notify_cfg.get("emails") or [])
    ok_run = run.get("status") in ("completed", "stopped")
    site, index = _site_index(job)
    folder = os.path.join(REPORTS_DIR, "scheduled", re.sub(r"[^\w-]+", "_", str(task.get("id") or "task"))[:60],
                          time.strftime("%Y%m%d_%H%M%S"))
    if ok_run and index is not None:
        os.makedirs(folder, exist_ok=True)
        fmt = ex.get("format") if ex.get("format") in ("xlsx", "csv") else "xlsx"

        def keep(name: str, data: bytes):
            path = os.path.join(folder, re.sub(r'[^\w.-]+', "_", name))
            with open(path, "wb") as fh:
                fh.write(data)
            out["files"].append(path)
            out["exports"].append(os.path.basename(path))

        if ex.get("excel", True):
            path = _safe(lambda: _excel_file(job, "all"), None, job, "Scheduled Excel")
            if path:
                dest = os.path.join(folder, os.path.basename(path))
                _safe(lambda: shutil.copyfile(path, dest), None)
                if os.path.exists(dest):
                    out["files"].append(dest)
                    out["exports"].append(os.path.basename(dest))
        if ex.get("items"):
            from crawler import reports
            res = _safe(lambda: reports.multi_export(list(ex["items"]), site, index, job.crawler, fmt=fmt,
                                                     skip_empty=ex.get("skip_empty", True),
                                                     **_report_ctx(job, site, index)), None, job, "Scheduled exports")
            if res:
                keep(res[0], res[1])
            else:
                out["notes"].append("The selected exports could not be built (see the crawl log).")
        if ex.get("pdf"):
            _ensure_pinpoint(job, wait=PDF_PINPOINT_WAIT)
            path = _safe(lambda: build_pdf(job), None, job, "Scheduled PDF")
            if path:
                dest = os.path.join(folder, os.path.basename(path))
                _safe(lambda: shutil.copyfile(path, dest), None)
                if os.path.exists(dest):
                    out["files"].append(dest)
                    out["exports"].append(os.path.basename(dest))
        if ex.get("sitemap"):
            from crawler import sitemap_gen
            try:
                name, data, _mime = sitemap_gen.build_sitemap(site, job.crawler, "xml")
                keep(name, data)
            except Exception as e:
                out["notes"].append(f"Sitemap not built: {job.mask(e)}")
        if ex.get("save_crawl", True):
            meta = _autosave(job, force=True)
            if meta:
                out["notes"].append(f"Saved as {meta.get('name') or meta.get('id')}.")
            elif getattr(job.crawler, "superseded", False):
                out["notes"].append("Not saved: another crawl replaced this one before it finished.")
            elif not job.crawler.results_list:
                out["notes"].append("Not saved: the crawl has no results.")
            else:
                out["notes"].append("Not saved: the crawl could not be written (see the crawl log).")
        if task.get("auto_compare", True):
            cmp = _safe(lambda: _comparison_with_previous(job), None, job, "Scheduled comparison")
            if cmp:
                out["compare"] = {"a": (cmp.get("a") or {}).get("id") if isinstance(cmp.get("a"), dict) else cmp.get("a"),
                                  "summary": cmp.get("summary")}
            else:
                out["notes"].append("No earlier crawl of this site to compare with.")
    elif not ok_run:
        out["notes"].append(f"The crawl ended as {run.get('status')}: no exports.")
    if emails and (ok_run or notify_cfg.get("on_failure", True)):
        if not _notify().configured(reports_dir=REPORTS_DIR):
            out["notes"].append("No email sent: set the SMTP server under Settings > Notifications.")
        else:
            res = _safe(lambda: _crawl_email(job, emails, bool(notify_cfg.get("attach")),
                                             compare_result=None if not ok_run else _safe(
                                                 lambda: _comparison_with_previous(job), None),
                                             extra_files=out["files"] if notify_cfg.get("attach") else None,
                                             status_override="" if ok_run else str(run.get("status") or ""),
                                             error=str(run.get("error") or "")),
                        {"ok": False, "error": "the email could not be built"}, job, "Scheduled email")
            if res.get("ok"):
                out["emailed"] = list(res.get("sent_to") or emails)
            else:
                out["notes"].append(f"Email not sent: {job.mask(res.get('error'))}")

def _get_scheduler():
    """The scheduler of this server (or a passive one that only edits tasks when another server process owns
    the tick), created on first use."""
    sched = _scheduler_holder.get("scheduler") or _scheduler_holder.get("passive")
    if sched is not None:
        return sched
    from crawler.scheduler import Scheduler

    class _Passive(Scheduler):
        def _mark_interrupted(self):   # the owning process keeps its runs; never mark them here
            return None

    sched = _Passive(_schedule_start, on_complete=_schedule_complete, reports_dir=REPORTS_DIR, busy_fn=_crawl_busy)
    _scheduler_holder["passive"] = sched
    return sched

def _scheduler_state() -> dict:
    sched = _scheduler_holder.get("scheduler")
    running = bool(sched is not None and _safe(lambda: sched.running, False))
    return {"running": running, "reason": "" if running else (_scheduler_holder.get("reason") or "Not started."),
            "sleeps_when_idle": _is_serverless() or bool(os.environ.get("RENDER"))}

def _sched_task(task_id: str) -> dict:
    t = _get_scheduler().get(task_id)
    if t is None:
        raise HTTPException(status_code=404, detail="No such scheduled task.")
    return t

@app.get("/api/schedules")
def list_schedules(q: str = ""):
    """File > Scheduling: every task with its next run, interval and whether it is valid (and why not)."""
    return {"tasks": [_sched_public(t) for t in _get_scheduler().list(q)], "scheduler": _scheduler_state(),
            "env_hint": (f"Logins for scheduled crawls are read from a server environment variable named in the "
                         f"task; its name must start with {SCHEDULE_LOGIN_ENV_PREFIX} (for example "
                         "SITE_SESSION_COOKIE). Tasks with a login are set up on the server itself or with the "
                         "team access key.")}

@app.post("/api/schedules")
def add_schedule(request: Request, task: Dict[str, Any]):
    """Adds a task. A login may only name a SITE_ server variable, and a task with a login is added only by
    whoever may change the server's settings (it sends a server secret to the site it crawls)."""
    task = task or {}
    login_errors = _sched_login_errors(task)
    if login_errors:
        return _sched_invalid(task, login_errors)
    refusal = _login_admin_refusal(request) if _sched_uses_login(task) else None
    if refusal is not None:
        return refusal
    out, errors = _get_scheduler().add(task)
    if errors:
        return JSONResponse({"detail": "The task is not valid.", "errors": errors}, status_code=400)
    return _sched_public(out)

@app.get("/api/schedules/export")
def export_schedules(ids: str = ""):
    data = _get_scheduler().export_json([i for i in ids.split(",") if i.strip()] or None)
    return _download("scheduled_tasks.json", data, "application/json")

@app.post("/api/schedules/import")
async def import_schedules(request: Request):
    """Adds the tasks of an exported file. Tasks whose login names a variable outside the SITE_ names are not
    imported, nor tasks with a login when this request may not change the server's settings (they say why)."""
    _name, data = await _read_upload(request, limit=2 * 1024 * 1024)
    try:
        obj = json.loads(data.decode("utf-8-sig")) if isinstance(data, (bytes, bytearray)) else data
    except (UnicodeDecodeError, ValueError):
        return _get_scheduler().import_json(data)   # its own "not valid JSON" answer
    items = obj.get("tasks") if isinstance(obj, dict) else obj if isinstance(obj, list) else None
    if not isinstance(items, list) or (isinstance(obj, dict) and obj.get("kind") not in (None, "schedules")):
        return _get_scheduler().import_json(obj)
    editable = settings_editable(request)[0]
    kept, refused = [], []
    for item in items:
        if isinstance(item, dict) and _sched_uses_login(item):
            errors = _sched_login_errors(item) or ({} if editable else {"start.login.env": LOGIN_ADMIN_RULE})
            if errors:
                refused.append({"task": str(item.get("name") or "")[:80], "errors": errors})
                continue
        kept.append(item)
    out = _get_scheduler().import_json(dict(obj, tasks=kept) if isinstance(obj, dict) else kept)
    out["errors"] = refused + list(out.get("errors") or [])
    return out

@app.get("/api/schedules/history")
def schedules_history(task_id: str = "", limit: int = 200):
    return {"runs": _sched_runs_public(_get_scheduler().history(task_id or None, max(1, min(limit, 1000))))}

@app.get("/api/schedules/{task_id}")
def get_schedule(task_id: str):
    return _sched_public(_sched_task(task_id))

@app.put("/api/schedules/{task_id}")
def edit_schedule(request: Request, task_id: str, task: Dict[str, Any]):
    """Changes a task (a partial section keeps the rest of the stored one). Changing where a task with a login
    goes (its start section) is refused like adding one; its time, exports and emails can be changed by all."""
    old = _sched_task(task_id)
    task = task or {}
    merged = copy.deepcopy(old)
    for k, v in task.items():   # the scheduler's own merge: whole keys, one level deep for sections
        merged[k] = dict(merged[k], **v) if isinstance(v, dict) and isinstance(merged.get(k), dict) else v
    login_errors = _sched_login_errors(merged)
    if login_errors:
        return _sched_invalid(merged, login_errors)
    if _sched_uses_login(merged) and _sched_start_section(merged) != _sched_start_section(old):
        refusal = _login_admin_refusal(request)
        if refusal is not None:
            return refusal
    out, errors = _get_scheduler().edit(task_id, task)
    if errors:
        return JSONResponse({"detail": "The task is not valid.", "errors": errors}, status_code=400)
    return _sched_public(out)

@app.delete("/api/schedules/{task_id}")
def delete_schedule(task_id: str):
    if not _get_scheduler().delete(task_id):
        raise HTTPException(status_code=404, detail="No such scheduled task.")
    return {"deleted": task_id}

@app.post("/api/schedules/{task_id}/duplicate")
def duplicate_schedule(task_id: str):
    out = _get_scheduler().duplicate(task_id)
    if out is None:
        raise HTTPException(status_code=404, detail="No such scheduled task.")
    return _sched_public(out)

class EnableRequest(BaseModel):
    enabled: bool = True

@app.post("/api/schedules/{task_id}/enable")
def enable_schedule(task_id: str, req: EnableRequest):
    out = _get_scheduler().set_enabled(task_id, req.enabled)
    if out is None:
        raise HTTPException(status_code=404, detail="No such scheduled task.")
    return _sched_public(out)

@app.post("/api/schedules/{task_id}/run")
def run_schedule_now(task_id: str):
    """Starts a task now (its schedule is unchanged); 409 while the server is busy (a crawl, Crawl Analysis, or a
    scheduled task still building its exports), 400 when the task is not valid."""
    from crawler.scheduler import BUSY_SUFFIX
    _sched_task(task_id)
    ref, msg = _get_scheduler().run_now(task_id)
    if not ref:
        raise HTTPException(status_code=409 if msg.endswith(BUSY_SUFFIX) else 400, detail=msg)
    return {"run_ref": ref, "message": msg}

@app.get("/api/schedules/{task_id}/history")
def schedule_history(task_id: str, limit: int = 200):
    return {"runs": _sched_runs_public(_get_scheduler().history(task_id, max(1, min(limit, 1000))))}

# --------------------------------------------------------------------------
# Segments (Configuration > Segments, right panel Segments tab, Reports > Segments Overview)
# --------------------------------------------------------------------------
class SegmentsRequest(BaseModel):
    segments: List[dict] = []
    site: Optional[str] = None

def _segments_store_or_400(site: str = ""):
    store = _segment_store(session.job, site or "")
    if store is None:
        _bad("No site: run a crawl first or pass the site address.")
    return store

@app.get("/api/crawl/segments")
def get_segments(tab: str = "", filter_id: str = Query("all", alias="filter"), site: str = ""):
    """The site's segments and, for the crawl on screen, URLs per segment (with indexable / non-indexable,
    status classes, issues by priority, average depth and response time). The rows add up to the total,
    "(No segment)" included. tab / filter limit the counts to one tab filter."""
    store = _segments_store_or_400(site)
    segs = store.list()
    from crawler import segments as seg_mod
    out: Dict[str, Any] = {"segments": segs, "updated_at": store.updated_at(), "fields": seg_mod.field_options(),
                           "overview": None}
    job = session.job
    site_obj, index = _site_index(job)
    crawl_site = job.crawler.start_url if job and job.crawler else ""
    if index is not None and (not site or _same_site(site, crawl_site)):
        _s, assignment = _segment_assignment(job, site_obj, index)
        if segs:
            out["overview"] = seg_mod.overview(site_obj, index, assignment, segs, tab=tab or None,
                                               filter_id=filter_id or "all")
    return out

@app.put("/api/crawl/segments")
def put_segments(body: Union[SegmentsRequest, List[dict]]):
    """Saves the site's segments (in order: the first matching segment wins) and re-segments the crawl. The body
    is {"segments": [...], "site"?} or just the list."""
    req = body if isinstance(body, SegmentsRequest) else SegmentsRequest(segments=list(body or []))
    store = _segments_store_or_400(req.site or "")
    errors = store.save(req.segments)
    if errors:
        return JSONResponse({"detail": "Some segments are not valid.",
                             "errors": {str(k): v for k, v in errors.items()}}, status_code=400)
    job = session.job
    if job is not None:
        job.export_stale = True
    return get_segments(site=req.site or "")

@app.delete("/api/crawl/segments")
def clear_segments(site: str = ""):
    _segments_store_or_400(site).clear()
    if session.job is not None:
        session.job.export_stale = True
    return {"segments": []}

@app.get("/api/crawl/segments/suggest")
def suggest_segments():
    """Segments suggested from the crawl's page types and top folders (with their URL counts)."""
    job, site, index = _need_index()
    from crawler import segments as seg_mod
    return {"suggestions": seg_mod.suggest(site, index)}

@app.get("/api/crawl/segments/issues")
def segment_issues():
    """Issues x segments: URLs of each issue per segment."""
    job, site, index = _need_index()
    segs, assignment = _segment_assignment(job, site, index)
    if not segs:
        return {"segments": [], "rows": [], "not_checked": "No segments are defined for this site."}
    from crawler import segments as seg_mod
    return seg_mod.issues_by_segment(index, assignment, segs)

@app.get("/api/crawl/segments/counts")
def segment_counts(tab: str = "", filter_id: str = Query("all", alias="filter"), match: str = "any"):
    """URLs of one tab filter (or a selection "a,b" with match=any|all, as the grid) per segment (the grid's
    segment dropdown)."""
    job, site, index = _need_index()
    segs, assignment = _segment_assignment(job, site, index)
    if not segs:
        return {"counts": {}}
    from crawler import segments as seg_mod
    ids, m = _filter_sel(tab, filter_id, match, index) if tab else (["all"], "any")
    if len(ids) > 1:
        _c, rows, _t = _grid_sel(job, site, index, tab, ids, m)
        return {"counts": dict(Counter(assignment.get(r["url"]) or seg_mod.NO_SEGMENT for r in rows or []))}
    return {"counts": seg_mod.counts_for(index, assignment, tab, ids[0])}

# --------------------------------------------------------------------------
# Crawl Analysis > Start / Stop / Configure
# --------------------------------------------------------------------------
class AnalysisConfigRequest(BaseModel):
    analysis: Dict[str, Any] = {}
    save_default: bool = False          # also make these the server's defaults for new crawls

@app.get("/api/crawl/analysis")
def get_analysis():
    """The Crawl Analysis steps of the crawl on screen, the server defaults, which steps were skipped and why,
    and the analysis job."""
    job = session.job
    cfg = dict(getattr(job.crawler.config, "analysis", None) or {}) if job and job.crawler else {}
    site, _ = _site_index(job) if job is not None and job.phase == "done" else (None, None)
    defaults = _safe(lambda: _settings().get_page("analysis", True, REPORTS_DIR), dict(default_analysis()))
    runner = job.runners["analysis"] if job is not None else JobRunner("analysis")
    return {"analysis": cfg, "defaults": defaults, "skipped": dict(getattr(site, "analysis_skipped", None) or {}),
            "job": runner.status(), "can_start": bool(job and job.crawler and job.phase == "done"
                                                     and not _crawl_busy(job))}

@app.put("/api/crawl/analysis/config")
def configure_analysis(request: Request, req: AnalysisConfigRequest):
    """Crawl Analysis > Configure: the steps for the crawl on screen (applied at the next Start), and with
    save_default the server's defaults for new crawls."""
    job = session.job
    base = dict(default_analysis())
    out: Dict[str, Any] = {}
    clean = _bool_map(req.analysis, base, "Crawl Analysis")
    if job is not None and job.crawler is not None:
        if job.busy:
            raise HTTPException(status_code=409, detail="Wait until the crawl has finished.")
        job.crawler.config.analysis = clean
        out["analysis"] = clean
    if req.save_default:
        _require_admin(request)
        errors = _settings().update("analysis", clean, REPORTS_DIR)
        if errors:
            return JSONResponse({"detail": "Some steps are not valid.", "errors": errors}, status_code=400)
        out["defaults"] = _settings().get_page("analysis", True, REPORTS_DIR)
    return out

def _analysis_work(job: CrawlJob, runner: JobRunner) -> dict:
    """Crawl Analysis > Start: the post-crawl checks (robots.txt audit, URL-trap probe), the whole site analysis
    (link score, sitemaps / orphans, duplicates, pagination, hreflang, semantic similarity) and the bot check,
    as configured."""
    from crawler.seo import bot_check, postcrawl
    crawler = job.crawler
    a = dict(getattr(crawler.config, "analysis", None) or {})
    done: Dict[str, str] = {}
    runner.update(5, "robots.txt audit and URL-trap probe")
    opts: Dict[str, Any] = {"bot_check": "deferred", "stop_event": runner.stop_event}   # run below, after the analysis
    if a.get("post_crawl_checks") is False:
        opts.update(robots_audit=False, trap_probe=False)
    done.update(_safe(lambda: postcrawl.run(crawler, opts), {}, job, "Post-crawl checks") or {})
    if runner.stop_event.is_set():
        runner.update(None, "Stopped.")
        return {"steps": done, "stopped": True}
    runner.update(35, "Site analysis: link score, sitemaps, duplicates, pagination, hreflang...")
    crawler._index_key = None
    site, index = _site_index(job, force=True)
    done["site_analysis"] = "done" if index is not None else "error"
    if a.get("bot_check") is not False and not _is_serverless() and not runner.stop_event.is_set():
        runner.update(60, "Checking pages as search and AI crawlers...")
        sample = int(getattr(crawler.config, "bot_check_sample", 0) or BOT_CHECK_SAMPLE)
        try:
            report = bot_check.run(crawler, sample=sample, render=_browser_enabled(), stop_event=runner.stop_event,
                                   on_progress=lambda msg: runner.update(None, msg))
        except Exception as e:
            report = postcrawl.status_report("error", f"{type(e).__name__}: {str(e)[:160]}")
        if isinstance(report, dict):
            if report.get("status") == "running":
                report = postcrawl.status_report("stopped", "the analysis was stopped")
            crawler.bot_report = report
            done["bot_check"] = str(report.get("status"))
        crawler._index_key = None
        _site_index(job, force=True)
    elif a.get("bot_check") is False:
        crawler.bot_report = postcrawl.status_report("skipped", "turned off for this crawl (Crawl Analysis > Configure)")
        crawler._index_key = None
        _site_index(job, force=True)
    _reports_changed(job)
    runner.update(None, "Crawl Analysis finished: " + ", ".join(f"{k.replace('_', ' ')} {v}" for k, v in done.items()))
    return {"steps": done, "stopped": runner.stop_event.is_set()}

@app.post("/api/crawl/analysis/start")
def start_analysis():
    """Re-runs the crawl analysis on the crawl on screen (background job; refused while a crawl runs)."""
    job = _finished_job("running Crawl Analysis")
    if job.busy:
        raise HTTPException(status_code=409, detail=_busy_detail(job))
    for kind in ("postcrawl", "pinpoint"):
        if job.runners[kind].busy:
            raise HTTPException(status_code=409, detail="Wait until the pictures and the bot check are finished.")
    try:
        return {"job_id": job.runners["analysis"].start(lambda r: _analysis_work(job, r), "Starting Crawl Analysis...")}
    except JobBusy:
        raise HTTPException(status_code=409, detail="Crawl Analysis is already running.")

@app.post("/api/crawl/analysis/stop")
def stop_analysis():
    job = session.job
    if job is None or not job.runners["analysis"].busy:
        return {"state": "idle"}
    job.runners["analysis"].stop()
    return {"state": "stopping"}

# --------------------------------------------------------------------------
# Configuration > API Access: optional Google Search Console, GA4 and PageSpeed Insights connections.
# Nothing has to be connected; without a connection their tabs say "Not connected", never 0.
# --------------------------------------------------------------------------
def _public_base(request: Request) -> str:
    """The address the browser uses for this server (CRAWLER_PUBLIC_URL wins, e.g. behind a proxy)."""
    env = (os.environ.get("CRAWLER_PUBLIC_URL") or "").strip().rstrip("/")
    if env:
        return env
    proto = request.headers.get("x-forwarded-proto") if (_trust_proxy() or os.environ.get("RENDER")) else ""
    host = request.headers.get("x-forwarded-host") if (_trust_proxy() or os.environ.get("RENDER")) else ""
    base = str(request.base_url).rstrip("/")
    if proto or host:
        u = urlparse(base)
        base = f"{(proto or u.scheme).split(',')[0].strip()}://{(host or u.netloc).split(',')[0].strip()}"
    return base

def _integration_call(fn: Callable[[], Any]) -> Any:
    """An integrations call with its errors as plain 400 / 401 / 409 / 429 answers (never a secret)."""
    I = _integrations()
    try:
        return fn()
    except I.NotConnected as e:
        raise HTTPException(status_code=409, detail=str(e))
    except I.AuthExpired as e:
        raise HTTPException(status_code=401, detail=str(e))
    except I.QuotaExceeded as e:
        raise HTTPException(status_code=429, detail=str(e))
    except I.IntegrationError as e:
        raise HTTPException(status_code=400, detail=str(e))

def _integration_done(crawler, source: str, job_obj):
    """A fetch finished: the tabs rebuild, the workbook and PDF are refreshed, the saved crawl re-saved."""
    job = session.job
    if job is None or job.crawler is not crawler:
        return
    crawler._index_key = None
    _reports_changed(job)

def _auto_fetch_integrations(job: CrawlJob):
    """"Auto Connect on Start": fetch every connected source marked auto-connect once the crawl is done."""
    if job.crawler is None or not job.crawler.results_list or job.origin in ("opened", "direct"):
        return
    started = _integrations().auto_fetch(job.crawler, on_done=_integration_done)
    if started:
        job.crawler.log("Fetching connected data: " + ", ".join(started))

@app.get("/api/integrations")
def integrations_status(request: Request):
    """Everything the API Access dialog shows, no secrets: per source whether it is connected and why not,
    the accounts, the chosen property, the options and the running fetch; Google's set-up steps (and the
    redirect address to register) when no OAuth client is configured."""
    I = _integrations()
    out = I.status(_public_base(request))
    out["editable"], out["editable_reason"] = settings_editable(request)
    client_set = bool(((out.get("google") or {}).get("client") or {}).get("set"))
    out["google_setup_needed"] = not client_set
    if not client_set:
        out["google_setup_hint"] = ("To connect Search Console or GA4, create a Google Cloud OAuth client (web "
                                    "application) and set GOOGLE_OAUTH_CLIENT_ID and GOOGLE_OAUTH_CLIENT_SECRET on "
                                    "the server, or paste them here. Or upload a service-account key. PageSpeed "
                                    "Insights only needs a free API key.")
    job = session.job
    data = getattr(job.crawler, "integrations", None) if job is not None and job.crawler is not None else None
    out["crawl"] = {s: {"status": (data or {}).get(s, {}).get("status"), "fetched_at": (data or {}).get(s, {}).get("fetched_at")}
                    for s in I.SOURCES} if isinstance(data, dict) else {}
    return out

@app.get("/api/integrations/google/connect")
def google_connect(request: Request, sources: str = "gsc,ga4"):
    """Begins "Connect with Google": the address to open (Google's consent page); Google sends the browser back
    to /api/integrations/google/callback."""
    _require_admin(request)
    I = _integrations()
    wanted = [s.strip() for s in sources.split(",") if s.strip() in ("gsc", "ga4")] or ["gsc", "ga4"]
    st = I.status(_public_base(request))
    if not ((st.get("google") or {}).get("client") or {}).get("set"):
        return JSONResponse({"detail": "No Google OAuth client is set up on this server.",
                             "setup": (st.get("google") or {}).get("setup"),
                             "redirect_uri": (st.get("google") or {}).get("redirect_uri")}, status_code=400)
    return _integration_call(lambda: I.google_begin(_public_base(request), wanted))

@app.get(GOOGLE_CALLBACK_PATH, response_class=HTMLResponse)
def google_callback(state: str = "", code: str = "", error: str = ""):
    """Where Google sends the browser after consent: stores the sign-in (server side) and shows a short page
    that sends the person back to the dashboard."""
    I = _integrations()
    try:
        acc = I.google_callback(state, code, error)
        ok, msg = True, f"Connected {html_lib.escape(str(acc.get('email') or acc.get('name') or 'the Google account'))}."
    except I.IntegrationError as e:
        ok, msg = False, html_lib.escape(str(e))
    color = "#1f7a3a" if ok else "#a12b2b"
    page = (f"<!doctype html><html><head><meta charset='utf-8'><title>Google connection</title>"
            f"<meta name='viewport' content='width=device-width,initial-scale=1'></head>"
            f"<body style='font-family:system-ui,sans-serif;padding:32px;max-width:560px'>"
            f"<h2 style='color:{color}'>{'Connected' if ok else 'Not connected'}</h2><p>{msg}</p>"
            f"<p>Close this tab and go back to the dashboard (Configuration &gt; API Access).</p>"
            f"<script>try{{window.opener&&window.opener.postMessage({{type:'google-connected',ok:{str(ok).lower()}}},'*');}}catch(e){{}}"
            f"setTimeout(function(){{try{{window.close()}}catch(e){{}}}},1500)</script></body></html>")
    return HTMLResponse(page, status_code=200 if ok else 400, headers={"Cache-Control": "no-store"})

class PasteRequest(BaseModel):
    url: str = ""

@app.post("/api/integrations/google/paste")
def google_paste(request: Request, req: PasteRequest):
    """The address the browser landed on after consent, pasted (when the redirect cannot reach this server)."""
    _require_admin(request)
    I = _integrations()
    return _integration_call(lambda: I.google_paste(req.url))

class ServiceAccountRequest(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    key_json: str = Field("", alias="json")     # the key file's text, sent as "json"
    name: Optional[str] = ""

@app.post("/api/integrations/google/service-account")
def google_service_account(request: Request, req: ServiceAccountRequest):
    """Connects with a Google service-account key (JSON). The key is stored server side, never returned."""
    _require_admin(request)
    I = _integrations()
    return _integration_call(lambda: I.add_service_account(req.key_json, req.name or ""))

class GoogleClientRequest(BaseModel):
    client_id: str = ""
    client_secret: str = ""

@app.put("/api/integrations/google/client")
def google_client(request: Request, req: GoogleClientRequest):
    """Saves the Google OAuth client (empty values remove it). The secret is never returned."""
    _require_admin(request)
    I = _integrations()
    return _integration_call(lambda: I.set_google_client(req.client_id, req.client_secret))

class AccountUpdateRequest(BaseModel):
    name: Optional[str] = None
    auto_connect: Optional[bool] = None

@app.patch("/api/integrations/google/accounts/{account_id}")
def google_account_update(request: Request, account_id: str, req: AccountUpdateRequest):
    _require_admin(request)
    I = _integrations()
    return _integration_call(lambda: I.update_account(account_id, name=req.name, auto_connect=req.auto_connect))

@app.delete("/api/integrations/google/accounts/{account_id}")
def google_account_disconnect(request: Request, account_id: str):
    """Disconnect: revokes the sign-in at Google and deletes it here; data already fetched for crawls stays."""
    _require_admin(request)
    I = _integrations()
    return _integration_call(lambda: I.disconnect_account(account_id))

@app.get("/api/integrations/gsc/properties")
def gsc_properties(account_id: str = ""):
    I = _integrations()
    return {"properties": _integration_call(lambda: I.gsc_properties(account_id))}

@app.get("/api/integrations/ga4/accounts")
def ga4_accounts(account_id: str = ""):
    I = _integrations()
    return {"accounts": _integration_call(lambda: I.ga4_accounts(account_id))}

@app.get("/api/integrations/ga4/streams")
def ga4_streams(property: str = "", account_id: str = ""):
    I = _integrations()
    return {"streams": _integration_call(lambda: I.ga4_streams(property, account_id))}

@app.post("/api/integrations/{source}/select")
def integration_select(source: str, values: Dict[str, Any]):
    """Chooses the account / property (/ GA4 data stream) a Google source uses."""
    if source not in ("gsc", "ga4"):
        raise HTTPException(status_code=404, detail="Unknown source.")
    if not str((values or {}).get("property") or "").strip():
        _bad("Choose the " + ("Search Console property." if source == "gsc" else "GA4 property."))
    I = _integrations()
    return _integration_call(lambda: I.select(source, values or {}))

class PsiKeyRequest(BaseModel):
    key: Optional[str] = ""
    auto_connect: Optional[bool] = None

@app.put("/api/integrations/psi/key")
def psi_key(request: Request, req: PsiKeyRequest):
    """Connects PageSpeed Insights with a free Google API key (env PSI_API_KEY wins). Never returned."""
    _require_admin(request)
    I = _integrations()
    return _integration_call(lambda: I.psi_connect(req.key or "", req.auto_connect))

@app.delete("/api/integrations/psi/key")
def psi_disconnect(request: Request):
    _require_admin(request)
    I = _integrations()
    return _integration_call(lambda: I.psi_disconnect())

class PsiTestRequest(BaseModel):
    key: Optional[str] = ""       # test this key before saving it ("" = the saved / env key)
    url: Optional[str] = ""

@app.post("/api/integrations/psi/test")
def psi_test(request: Request, req: PsiTestRequest):
    """Tests a PageSpeed key with one real request (the crawl's start page, or the given URL; about 10-30 s):
    {"ok", "message", "scores"}. Uses one unit of the key's daily quota."""
    _require_admin(request)
    I = _integrations()
    from crawler.integrations import psi as psi_mod
    from crawler.integrations import store as istore
    key = (req.key or "").strip()
    if not key or key == "***":
        key, _where = istore.psi_key()
    job = session.job
    url = _as_web_address(req.url or "") or (job.crawler.start_url if job and job.crawler else "") or "https://www.example.com/"
    sess = requests.Session()
    try:
        res = psi_mod.fetch_one(sess, url, psi_mod.options({"device": "mobile", "categories": ["performance"]}), key)
    except I.QuotaExceeded as e:
        return {"ok": False, "message": str(e), "url": url}
    except I.IntegrationError as e:
        return {"ok": False, "message": str(e), "url": url}
    except Exception as e:
        return {"ok": False, "message": f"The test request failed: {type(e).__name__}", "url": url}
    finally:
        sess.close()
    ok = (res or {}).get("status") == "success"
    return {"ok": ok, "message": "The key works." if ok else str((res or {}).get("message") or "PageSpeed did not test the page."),
            "url": url, "scores": (res or {}).get("scores") or {}}

class FetchRequest(BaseModel):
    options: Optional[Dict[str, Any]] = None
    urls: Optional[List[str]] = None

@app.post("/api/integrations/{source}/fetch")
def integration_fetch(source: str, req: Optional[FetchRequest] = None):
    """Fetches the source's data for the crawl on screen (background job with progress; GET
    /api/integrations/jobs). Refused when not connected or before a crawl."""
    I = _integrations()
    if source not in I.SOURCES:
        raise HTTPException(status_code=404, detail="Unknown source.")
    job = session.job
    if job is None or job.crawler is None or job.phase != "done":
        raise HTTPException(status_code=409, detail="Crawl a site first (and wait until it has finished): the data "
                                                    "is matched to the crawled URLs.")
    urls = _clean_urls(req.urls) if req and req.urls else None
    return _integration_call(lambda: I.start_fetch(source, job.crawler, options=(req.options if req else None),
                                                   urls=urls, on_done=_integration_done))

@app.post("/api/integrations/{source}/stop")
def integration_stop(source: str):
    I = _integrations()
    if source not in I.SOURCES:
        raise HTTPException(status_code=404, detail="Unknown source.")
    return {"job": I.stop_fetch(source)}

@app.get("/api/integrations/jobs")
def integration_jobs():
    """Fetch jobs of the crawl on screen (a fetch still running for a replaced crawl is not shown)."""
    job = session.job
    return {"jobs": _integrations().jobs(crawler=(job.crawler if job is not None else None))}

class PsiRunRequest(BaseModel):
    urls: Optional[List[str]] = None
    device: Optional[str] = None          # mobile | desktop (default: the saved option)

@app.post("/api/psi/run")
def psi_run(req: Optional[PsiRunRequest] = None):
    """PageSpeed Insights for the crawl on screen (all indexable pages up to the URL cap, or the given URLs):
    the same job as POST /api/integrations/psi/fetch."""
    opts = {"device": req.device} if req is not None and req.device else None
    return integration_fetch("psi", FetchRequest(options=opts, urls=req.urls if req is not None else None))

@app.get("/api/psi/status")
def psi_status():
    I = _integrations()
    return {"connection": _safe(lambda: I.connection("psi"), {}), "job": (I.jobs() or {}).get("psi")}

@app.get("/api/integrations/{source}/discovered")
def integration_discovered(source: str):
    """URLs Search Console / GA4 know that the crawl did not find ("Crawl New URLs Discovered"): start a List
    mode crawl with them."""
    if source not in ("gsc", "ga4"):
        raise HTTPException(status_code=404, detail="Unknown source.")
    job = session.job
    if job is None or job.crawler is None:
        return {"urls": []}
    return {"urls": _safe(lambda: _integrations().discovered_urls(job.crawler, source), [])}

# --------------------------------------------------------------------------
# Help: About, Debug, Usage Stats, User Guide, third-party notices
# --------------------------------------------------------------------------
def _repo_file(name: str) -> str:
    return os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), name)

def _read_text(path: str, limit: int = 2_000_000) -> str:
    try:
        with open(path, "r", encoding="utf-8") as fh:
            return fh.read(limit)
    except OSError:
        return ""

@app.get("/api/about")
def about():
    """Help > About: name, version, platform, Python, Chrome status and the documents shipped with the app."""
    env = get_environment()
    return {"name": APP_NAME, "version": app_version(), "platform": env.get("platform"),
            "python": platform_lib.python_version(), "os": f"{platform_lib.system()} {platform_lib.release()}",
            "chrome": env.get("chrome"), "documentation": "/api/help/guide",
            "notices": "/api/help/notices", "notices_present": os.path.exists(_repo_file("THIRD_PARTY_NOTICES.md"))}

@app.get("/api/help/notices")
def help_notices():
    text = _read_text(_repo_file("THIRD_PARTY_NOTICES.md"))
    if not text:
        raise HTTPException(status_code=404, detail="THIRD_PARTY_NOTICES.md is not on this server.")
    return Response(text, media_type="text/markdown; charset=utf-8")

@app.get("/api/help/guide")
def help_guide():
    """Help > User Guide: DOCUMENTATION.md as text (the dashboard renders it)."""
    text = _read_text(_repo_file("DOCUMENTATION.md"))
    if not text:
        raise HTTPException(status_code=404, detail="DOCUMENTATION.md is not on this server.")
    return Response(text, media_type="text/markdown; charset=utf-8")

def _lib_versions() -> Dict[str, str]:
    out = {}
    for mod in ("fastapi", "starlette", "pydantic", "uvicorn", "requests", "bs4", "lxml", "openpyxl", "reportlab",
                "PIL", "selenium"):
        try:
            m = __import__(mod)
            out[mod] = str(getattr(m, "__version__", "") or getattr(m, "VERSION", "") or "installed")
        except Exception:
            out[mod] = "not installed"
    return out

@app.get("/api/debug")
def debug_info(request: Request, download: int = 0):
    """Help > Debug: environment, versions, settings (secrets as set / not set), storage, scheduler,
    connections, the crawl's stats, jobs and recent log lines - never a password, key, token or cookie value.
    download=1 returns it as a JSON file."""
    job = session.job
    data: Dict[str, Any] = {
        "generated_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "app": {"name": APP_NAME, "version": app_version(), "python": sys.version.split()[0],
                "os": f"{platform_lib.system()} {platform_lib.release()}", "libraries": _lib_versions()},
        "environment": get_environment(),
        "settings": _safe(lambda: _settings().get_settings(True, REPORTS_DIR), {}),
        "storage": _safe(crawls_storage, {}),
        "scheduler": _scheduler_state(),
        "integrations": _safe(lambda: {s: {k: v for k, v in c.items() if k in ("connected", "reason")}
                                       for s, c in ((s, _integrations().connection(s)) for s in _integrations().SOURCES)}, {}),
        "ai": _safe(get_ai_status, {}),
        "access": {"key_set": bool(access_key()), "settings_editable": settings_editable(request)[0]},
        "env_vars_set": sorted(k for k in os.environ if k.startswith(("CRAWLER_", "GOOGLE_OAUTH", "PSI_", "APP_",
                                                                       "RENDER", "MAX_CHROME", "CHROME_BIN"))),
        "crawl": None,
    }
    if job is not None and job.crawler is not None:
        stats = _safe(lambda: _public_stats(job), {})
        data["crawl"] = {"start_url": job.crawler.start_url, "phase": job.phase, "origin": job.origin,
                         "stats": stats, "jobs": _jobs_status(job), "notes": list(job.notes),
                         "logs": [job.mask(x) for x in list(job.logs)], "screenshot_note": job.screenshot_note,
                         "config": _safe(lambda: _store().public_config(job.crawler.config), {})}
    text = json.dumps(json_value(data), indent=1, default=str)
    for secret in (job.secrets if job is not None else []):
        text = text.replace(json.dumps(secret)[1:-1], "***") if secret else text
    if download:
        return _download(f"arachnix_debug_{time.strftime('%Y%m%d_%H%M%S')}.json", text.encode("utf-8"),
                         "application/json")
    return JSONResponse(json.loads(text))

@app.get("/api/usage")
def usage_stats(cumulative: int = 0):
    """Help > Usage Stats: crawls and URLs crawled per month and the most crawled sites (from the crawl
    history on this server). cumulative=1 adds the months up."""
    try:
        from crawler.store import crawl_history
        rows = crawl_history(REPORTS_DIR)
    except Exception:
        from crawler.seo.compare import list_snapshots
        rows = list_snapshots(REPORTS_DIR)
    months: "OrderedDict[str, Dict[str, int]]" = OrderedDict()
    sites: Counter = Counter()
    total_urls = 0
    for r in sorted(rows, key=lambda x: str(x.get("started_at") or "")):
        month = str(r.get("started_at") or "")[:7] or "unknown"
        urls = int(r.get("urls") or r.get("pages") or 0)
        m = months.setdefault(month, {"crawls": 0, "urls": 0})
        m["crawls"] += 1
        m["urls"] += urls
        sites[str(r.get("site") or r.get("host") or "")] += 1
        total_urls += urls
    series = [{"month": k, **v} for k, v in months.items()]
    if cumulative:
        c, u = 0, 0
        for s in series:
            c, u = c + s["crawls"], u + s["urls"]
            s["crawls"], s["urls"] = c, u
    return {"months": series, "total_crawls": len(rows), "total_urls": total_urls, "cumulative": bool(cumulative),
            "top_sites": [{"site": s, "crawls": n} for s, n in sites.most_common(10) if s]}

# --------------------------------------------------------------------------
# Server start / stop
# --------------------------------------------------------------------------
def _on_startup():
    """Starts the scheduler (not on serverless hosts, and only in the one server process that holds the reports
    folder's scheduler lock; CRAWLER_SCHEDULER=0 turns it off) and lets connection changes refresh the tabs."""
    _safe(lambda: _integrations().add_listener(_on_integrations_changed), None)
    if _is_serverless():
        _scheduler_holder["reason"] = "Serverless hosts cannot run scheduled crawls."
        return
    if os.environ.get("CRAWLER_SCHEDULER", "").strip().lower() in ("0", "off", "false", "no"):
        _scheduler_holder["reason"] = "Turned off on this server (CRAWLER_SCHEDULER=0)."
        return
    fh = _scheduler_lock()
    if fh is None:
        _scheduler_holder["reason"] = ("Another ARACHNIX server on this computer runs the scheduled tasks of "
                                       "this reports folder.")
        return
    try:
        from crawler.scheduler import Scheduler
        sched = Scheduler(_schedule_start, on_complete=_schedule_complete, reports_dir=REPORTS_DIR,
                          busy_fn=_crawl_busy, logger=lambda m: logging.getLogger("uvicorn.error").info(m),
                          busy_detail=lambda: _busy_detail().split(" Wait")[0])
        sched.start()
        _scheduler_holder.update(scheduler=sched, lock_fh=fh, reason="", passive=None)
    except Exception as e:
        fh.close()
        _scheduler_holder["reason"] = f"The scheduler could not start: {type(e).__name__}: {str(e)[:160]}"

def _on_shutdown():
    sched = _scheduler_holder.get("scheduler")
    if sched is not None:
        _safe(lambda: sched.stop(), None)
    fh = _scheduler_holder.get("lock_fh")
    if fh is not None:
        _safe(fh.close, None)
    _scheduler_holder.update(scheduler=None, lock_fh=None)

def _on_integrations_changed():
    job = session.job
    if job is not None and job.crawler is not None:
        _integrations().refresh(job.crawler)
        job.crawler._index_key = None

if os.path.exists("web/static"):
    app.mount("/static", StaticFiles(directory="web/static"), name="static")
elif os.path.exists("static"):
    app.mount("/static", StaticFiles(directory="static"), name="static")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("web.app:app", host="0.0.0.0", port=8000, reload=True)


# Build 4 W1: /mcp, MCP settings and tokens, export links, the crawl log of agent calls
_agents.install(app)
_connections.install(app)   # Build 4 W2
_ai_routes.install(app)     # Build 4 W3 / W4
_extension.install(app)     # Build 4 W6
_v2_misc.install(app)       # Build 4 W10 / W11
