"""
Reports menu, Bulk Export menu and Multi-Export (Screaming Frog parity, built our own way).

    REPORTS                                   [{"id", "name", "menu_path", "description"}] in menu order
    build_report(report_id, site, index, crawler, fmt="xlsx", **ctx) -> (file name, bytes, media type)
    report_tables(report_id, site, index, crawler, **ctx) -> [Table]
    BULK_EXPORTS                              the fixed Bulk Export items (same shape as REPORTS)
    register_bulk_export(export_id, title, headers, rows | files, not_checked, group)
                                              adds a Bulk Export item from another module (table rows, or a .zip
                                              of files), so bulk_catalog, build_bulk and multi_export list and
                                              build it with no special case
    bulk_catalog(index=None) -> [items]       BULK_EXPORTS + one item per tab filter ("filter:<tab>:<filter>")
                                              and per filter's linking pages ("inlinks:<tab>:<filter>")
    build_bulk(export_id, site, index, crawler, fmt="xlsx", **ctx) -> (file name, bytes, media type)
    multi_export(items, site, index, crawler, fmt="xlsx", skip_empty=True, **ctx) -> (file name, bytes, mime)
        xlsx: one workbook (an "Index" sheet, then one sheet per export); csv: a .zip of CSV files
        (index.csv + one folder per menu); skip_empty leaves out exports with no rows or that could not run,
        and the index says so.

ctx (all optional): issues = the Issues overview rows after dismissals (default index.issues);
dismissed_urls = {"tab:filter": [urls]}; segments = {url: segment name} (Segments Overview);
secrets = [str] masked in header exports; stats = crawler.get_stats(); comparison = a compare.compare() result, or
compare = {"a", "b" ("current" = this crawl), "url_map", "change_detection", "thresholds"} with reports_dir
(Comparison Overview).

Every report is built from the same TabIndex the tabs show, so its rows reconcile with the tab and filter
counts: e.g. Redirect Chains lists the URLs of Response Codes > Internal Redirect Chain + Internal Redirect
Loop, Non-Indexable Canonicals the Canonicals filter of that name, and a pair report (hreflang, pagination,
insecure content) lists exactly the pages of its filter, one row per failing link. A report whose check did
not run is a one-row "Not checked" table with the reason, never a table of zeros. Status classes come from
the HTTP status (a 404 is always 4xx). AMP is deprecated: no AMP report or export exists.
"""
import time
from collections import OrderedDict
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

from .exporter import (DETAIL_COOKIE_HEADERS, INLINK_HEADERS, ISSUE_URL_HEADERS, ISSUES_OVERVIEW_HEADERS,
                       PARENT_LINK, SOURCE_HEADERS, Table, XLSX_MIME, ZIP_MIME, check_format, cookie_set_by,
                       csv_bytes, externals_by_url, filter_name, filter_source_rows, inlink_rows, issue_url_rows,
                       issues_overview_rows, not_checked_table, path_type, row_category, row_status_code,
                       row_status_text, safe_header_value, safe_name, tab_table, tables_bytes, write_table,
                       zip_bytes, _RESOURCE_TYPES)

CATEGORY_NAMES = OrderedDict([("2xx", "Success (2xx)"), ("3xx", "Redirection (3xx)"),
                              ("4xx", "Client Error (4xx)"), ("5xx", "Server Error (5xx)"),
                              ("error", "No Response"), ("blocked", "Blocked by Robots.txt")])
TEMP_REDIRECTS = (302, 303, 307)
MAX_HOPS = 20
NO_SEGMENT = "(No segment)"
PAGE_SOURCE_CAP = 20_000            # pages in one "All Page Source" zip


def _r(id_: str, name: str, *path: str, description: str = "") -> dict:
    return {"id": id_, "name": name, "menu_path": list(path) or [name], "description": description}


REPORTS: List[dict] = [
    _r("crawl_overview", "Crawl Overview", "Crawl Overview",
       description="Every tab and filter with its URL count and % of the tab, plus the crawl summary."),
    _r("issues_overview", "Issues Overview", "Issues Overview",
       description="Every flagged issue: type, priority, URLs, % of total, what it means and how to fix it."),
    _r("segments_overview", "Segments Overview", "Segments Overview",
       description="URLs, status classes and issue counts per segment."),
    _r("redirects_all", "All Redirects", "Redirects", "All Redirects",
       description="Every redirect with each hop, the final URL and whether it can be indexed."),
    _r("redirect_chains", "Redirect Chains", "Redirects", "Redirect Chains",
       description="Internal URLs that redirect two or more times, or in a loop."),
    _r("redirects_to_error", "Redirects to Error", "Redirects", "Redirects to Error",
       description="Redirects whose chain ends on a broken URL (4xx), a server error (5xx) or no answer, with "
                   "every hop and the page that links to the first URL."),
    _r("redirect_canonical_chains", "Redirect & Canonical Chains", "Redirects", "Redirect & Canonical Chains",
       description="Chains of two or more redirects and canonicals mixed."),
    _r("canonical_chains", "Canonical Chains", "Canonicals", "Canonical Chains",
       description="Pages whose canonical points to a page that is canonicalised again."),
    _r("non_indexable_canonicals", "Non-Indexable Canonicals", "Canonicals", "Non-Indexable Canonicals",
       description="Pages whose canonical points to a URL that cannot be indexed."),
    _r("pagination_non_200", "Non-200 Pagination URLs", "Pagination", "Non-200 Pagination URLs",
       description="rel=next / prev links to URLs that do not answer 200."),
    _r("pagination_unlinked", "Unlinked Pagination URLs", "Pagination", "Unlinked Pagination URLs",
       description="rel=next / prev URLs that no crawled page links to."),
    _r("hreflang_all", "All hreflang URLs", "Hreflang", "All hreflang URLs",
       description="Every hreflang annotation of every page."),
    _r("hreflang_non_200", "Non-200 hreflang URLs", "Hreflang", "Non-200 hreflang URLs"),
    _r("hreflang_unlinked", "Unlinked hreflang URLs", "Hreflang", "Unlinked hreflang URLs"),
    _r("hreflang_missing_return", "Missing Return Links", "Hreflang", "Missing Return Links"),
    _r("hreflang_inconsistent_return", "Inconsistent Language & Region Return Links", "Hreflang",
       "Inconsistent Language & Region Return Links"),
    _r("hreflang_non_canonical_return", "Non Canonical Return Links", "Hreflang", "Non Canonical Return Links"),
    _r("hreflang_noindex_return", "Noindex Return Links", "Hreflang", "Noindex Return Links"),
    _r("insecure_content", "Insecure Content", "Insecure Content",
       description="HTTPS pages that load or link to HTTP addresses."),
    _r("serp_summary", "SERP Summary", "SERP Summary",
       description="URL, title and description with their lengths and pixel widths (re-uploadable in SERP mode)."),
    _r("orphan_pages", "Orphan Pages", "Orphan Pages",
       description="URLs found in XML sitemaps that no crawled page links to."),
    _r("sd_validation_summary", "Validation Errors & Warnings Summary", "Structured Data",
       "Validation Errors & Warnings Summary"),
    _r("sd_validation", "Validation Errors & Warnings", "Structured Data", "Validation Errors & Warnings"),
    _r("sd_rich_summary", "Google Rich Result Features Summary", "Structured Data",
       "Google Rich Result Features Summary"),
    _r("sd_rich", "Google Rich Result Features", "Structured Data", "Google Rich Result Features"),
    _r("js_console_summary", "JavaScript Console Log Summary", "Javascript", "JavaScript Console Log Summary",
       description="Console errors and warnings of the pages rendered in Chrome."),
    _r("pagespeed_summary", "PageSpeed Summary", "PageSpeed", "PageSpeed Summary",
       description="PageSpeed Insights filters (needs a PageSpeed connection)."),
    _r("mobile_summary", "Mobile Summary", "Mobile", "Mobile Summary"),
    _r("accessibility_summary", "Accessibility Violations Summary", "Accessibility",
       "Accessibility Violations Summary"),
    _r("http_header_summary", "HTTP Header Summary", "HTTP Headers", "HTTP Header Summary",
       description="Each response header and how many URLs sent it."),
    _r("cookie_summary", "Cookie Summary", "Cookies", "Cookie Summary",
       description="Every cookie the site sets: server answers and, with Render every page in Chrome, what scripts "
                   "and requests set while pages load (names and attributes; values are never kept)."),
    _r("compare_overview", "Comparison Overview", "Compare", "Comparison Overview",
       description="The crawl comparison on screen (Mode > Compare): every filter's current, previous, added, new, "
                   "removed and missing URLs, the summary rows, change detection and the issue changes."),
]
REPORT_IDS = {r["id"] for r in REPORTS}

BULK_EXPORTS: List[dict] = [
    _r("queued_urls", "Queued URLs", "Queued URLs",
       description="URLs found but not crawled yet (paused or stopped crawl), and those left out by the limits."),
    _r("links_all_inlinks", "All Inlinks", "Links", "All Inlinks"),
    _r("links_all_outlinks", "All Outlinks", "Links", "All Outlinks"),
    _r("links_all_anchor_text", "All Anchor Text", "Links", "All Anchor Text"),
    _r("links_external", "External Links", "Links", "External Links"),
    _r("web_http_headers", "All HTTP Headers", "Web", "All HTTP Headers"),
    _r("web_cookies", "All Cookies", "Web", "All Cookies"),
    _r("web_page_source", "All Page Source", "Web", "All Page Source",
       description="The raw HTML of every crawled page, one file each (always a .zip)."),
    _r("path_absolute", "Absolute", "Path Type", "Absolute"),
    _r("path_protocol_relative", "Protocol-Relative", "Path Type", "Protocol-Relative"),
    _r("path_root_relative", "Root-Relative", "Path Type", "Root-Relative"),
    _r("path_path_relative", "Path-Relative", "Path Type", "Path-Relative"),
    _r("issues_all", "All Issues", "Issues", "All Issues",
       description="Every URL of every flagged issue, with its tab and filter."),
    _r("issues_overview", "Issues Overview", "Issues", "Issues Overview"),
]
PATH_EXPORTS = {"path_absolute": "Absolute", "path_protocol_relative": "Protocol-Relative",
                "path_root_relative": "Root-Relative", "path_path_relative": "Path-Relative"}


# ---------------------------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------------------------
def _tr(index, tab_id: str):
    return index.tabs.get(tab_id) if index is not None else None


def _filter_rows(index, tab_id: str, filter_id: str) -> Tuple[Optional[List[Any]], str]:
    """(rows of a tab filter, why it was not checked); rows None when the tab is not in the index."""
    tr = _tr(index, tab_id)
    if tr is None:
        return None, ""
    why = (getattr(tr, "not_checked", None) or {}).get(filter_id, "")
    return list(tr.filter_rows.get(filter_id) or []), why


def _union(index, tab_id: str, filter_ids: Sequence[str]) -> List[Any]:
    out, seen = [], set()
    tr = _tr(index, tab_id)
    if tr is None:
        return out
    for fid in filter_ids:
        for r in tr.filter_rows.get(fid) or ():
            u = getattr(r, "url", "")
            if u not in seen:
                seen.add(u)
                out.append(r)
    return out


def _missing_tab(title: str, tab_id: str) -> Table:
    return not_checked_table(title, f"The {tab_id.replace('_', ' ').title()} tab is not available in this build.")


def _code(r) -> Any:
    c = row_status_code(r) if r is not None else None
    return "" if c is None else c


def _status(r) -> str:
    return row_status_text(r) if r is not None else "Not crawled"


def _find(site, url: str):
    from .seo.tabs.canonicals import find_result
    return find_result(site, url) if url else None


def _norm(url: str) -> str:
    from .seo.tabs.canonicals import norm
    return norm(url)


def _indexability(site, r) -> Tuple[str, str]:
    if r is None:
        return "", ""
    from .seo.tabs.internal import indexability_label, indexability_status_label
    try:
        return indexability_label(site, r), indexability_status_label(site, r)
    except Exception:
        return "", ""


def _parent_link(site, r) -> Tuple[str, int]:
    """(first page linking to the URL, number of linking pages); the URL that led the crawler there when no
    crawled page links to it."""
    from .exporter import linking_pages
    url = getattr(r, "url", "")
    pages = [p for p in linking_pages(site, url) if p != url]
    return (pages[0] if pages else (getattr(r, "parent_url", "") or "")), len(pages)


def _when(ts: Any) -> str:
    try:
        ts = float(ts or 0)
    except (TypeError, ValueError):
        return ""
    return time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(ts)) if ts > 0 else ""


def _pct(n: int, total: int) -> float:
    return round(n / total * 100, 2) if total else 0.0


def _issues(index, ctx: dict) -> List[dict]:
    issues = ctx.get("issues")
    return list(issues if issues is not None else (index.issues if index is not None else []))


# ---------------------------------------------------------------------------------------------
# Crawl Overview / Issues Overview / Segments Overview / tab summaries
# ---------------------------------------------------------------------------------------------
OVERVIEW_HEADERS = ["Tab", "Filter", "Type", "Priority", "URLs", "% of Total", "Checked"]


def _filter_base(tr, f) -> int:
    """The "% of Total" base of one filter: the rows its check covers (Filter.of_rows, e.g. the Performance
    cache checks cover files only), else every row of the tab. The same base as the dashboard Overview and
    the Issues overview percentage (registry.filter_base), so a filter shows one percentage everywhere."""
    if getattr(f, "of_rows", None) is None:
        return len(tr.rows)
    from .seo.registry import filter_base
    return filter_base(tr, f)


def crawl_overview(site, index, crawler, ctx: dict) -> List[Table]:
    """Summary (Screaming Frog's Summary rows, the same counts and formula as the dashboard Overview:
    Total URLs Crawled = internal + external URLs minus the ones never requested) and every tab filter with
    its count and its own "% of Total" base."""
    from .exporter import crawl_totals, crawled_note
    results = list((getattr(site, "by_url", None) or {}).values()) if site is not None else []
    total = len(results)
    counts = {k: 0 for k in CATEGORY_NAMES}
    for r in results:
        c = row_category(r) or "error"
        counts[c] = counts.get(c, 0) + 1
    indexable = sum(1 for r in results if (getattr(site, "indexable", None) or {}).get(r.url))
    totals = crawl_totals(results, getattr(site, "externals", None) or [])
    cfg = getattr(site, "config", None)
    summary = [["Site", getattr(site, "start_url", "") or ""],
               ["Mode", (getattr(cfg, "mode", "") or "spider").title()],
               ["Crawl Started", _when(getattr(crawler, "start_time", 0))],
               ["Crawl Finished", _when(getattr(crawler, "end_time", 0))],
               ["Crawl Status", str(getattr(crawler, "status", "") or "")],
               ["Total URLs Encountered", f"{totals['encountered']:,} ({total:,} internal + "
                                          f"{totals['external']:,} external)"],
               ["Total Internal Blocked by robots.txt", totals["internal_blocked"]],
               ["Total External Blocked by robots.txt", totals["external_blocked"]],
               ["Total URLs Crawled", f"{totals['crawled']:,} ({crawled_note(totals)})"],
               ["Total Internal URLs", total],
               ["Total External URLs", totals["external"]],
               ["Total Internal Indexable URLs", indexable],
               ["Total Internal Non-Indexable URLs", total - indexable]]
    for k, name in CATEGORY_NAMES.items():   # the internal URLs by status class: they add up to Total Internal
        summary.append([f"{name} URLs", f"{counts.get(k, 0):,} ({_pct(counts.get(k, 0), total)}%)"])
    summary.append(["Issues Flagged", len(_issues(index, ctx))])
    stats = ctx.get("stats") or {}
    if stats.get("limit_reached"):
        summary.append(["Note", "The crawl stopped at the Max URLs limit: these counts cover part of the site."])

    rows = []
    for tab_id, tr in sorted((index.tabs if index is not None else {}).items(), key=lambda kv: kv[1].tab.order):
        tab = tr.tab
        if getattr(tab, "hidden", False):
            continue
        n_all = len(tr.rows)
        nc = getattr(tr, "not_checked", None) or {}
        rows.append([tab.name, tab.total_label, "", "", n_all, 100.0 if n_all else 0.0,
                     f"Not checked: {nc['all']}" if nc.get("all") else "Checked"])
        for f in tab.filters:
            why = nc.get(f.id, "")
            if why:
                rows.append([tab.name, f.name, f.issue_type or "Stat", f.priority or "", "", "",
                             f"Not checked: {why}"])
                continue
            n = len(tr.filter_rows.get(f.id) or ())
            rows.append([tab.name, f.name, f.issue_type or "Stat", f.priority or "", n,
                         _pct(n, _filter_base(tr, f)), "Checked"])
    return [Table("Summary", ["Item", "Value"], summary, widths=(30, 70)),
            Table("Tabs and Filters", OVERVIEW_HEADERS, rows, widths=(22, 52, 13, 10, 10, 11, 60))]


def issues_overview(site, index, crawler, ctx: dict) -> List[Table]:
    return [Table("Issues Overview", ISSUES_OVERVIEW_HEADERS, issues_overview_rows(_issues(index, ctx)),
                  widths=(48, 13, 14, 9, 11, 70, 70))]


def _segment_of(segments: dict, url: str) -> str:
    v = segments.get(url)
    if isinstance(v, (list, tuple)):
        v = v[0] if v else None
    if isinstance(v, dict):
        v = v.get("name")
    return str(v) if v else NO_SEGMENT


SEGMENT_HEADERS = ["Segment", "URLs", "% of URLs", "Indexable", "Non-Indexable"] + \
                  [CATEGORY_NAMES[k] for k in CATEGORY_NAMES] + ["URLs With Issues"]
SEGMENT_ISSUE_HEADERS = ["Segment", "Issue Name", "Issue Type", "Issue Priority", "URLs", "% of Segment"]


def segments_overview(site, index, crawler, ctx: dict) -> List[Table]:
    segments = ctx.get("segments")
    if not segments:
        return [not_checked_table("Segments Overview", "No segments are set up (Configuration > Segments).")]
    results = list(site.by_url.values())
    order: List[str] = []
    stats: Dict[str, dict] = {}
    for r in results:
        seg = _segment_of(segments, r.url)
        if seg not in stats:
            order.append(seg)
            stats[seg] = {"urls": 0, "indexable": 0, "cats": {k: 0 for k in CATEGORY_NAMES}, "issue_urls": set()}
        s = stats[seg]
        s["urls"] += 1
        s["indexable"] += 1 if site.indexable.get(r.url) else 0
        c = row_category(r) or "error"
        s["cats"][c] = s["cats"].get(c, 0) + 1
    issue_rows = []
    for issue in _issues(index, ctx):
        tr = _tr(index, issue.get("tab"))
        if tr is None:
            continue
        hidden = set((ctx.get("dismissed_urls") or {}).get(f"{issue.get('tab')}:{issue.get('filter')}") or ())
        per: Dict[str, int] = {}
        for r in tr.filter_rows.get(issue.get("filter")) or ():
            u = getattr(r, "url", "")
            if u in hidden:
                continue
            seg = _segment_of(segments, u)
            per[seg] = per.get(seg, 0) + 1
            if seg in stats:
                stats[seg]["issue_urls"].add(u)
        for seg in order + sorted(set(per) - set(order)):
            if per.get(seg):
                base = stats[seg]["urls"] if seg in stats else 0
                issue_rows.append([seg, issue.get("name", ""), issue.get("issue_type") or "",
                                   issue.get("priority") or "", per[seg], _pct(per[seg], base)])
    total = len(results)
    seg_rows = []
    for seg in sorted(order, key=lambda s: (s == NO_SEGMENT, order.index(s))):
        s = stats[seg]
        seg_rows.append([seg, s["urls"], _pct(s["urls"], total), s["indexable"], s["urls"] - s["indexable"]]
                        + [s["cats"].get(k, 0) for k in CATEGORY_NAMES] + [len(s["issue_urls"])])
    return [Table("Segments", SEGMENT_HEADERS, seg_rows),
            Table("Issues by Segment", SEGMENT_ISSUE_HEADERS, issue_rows, widths=(24, 52, 13, 14, 10, 13))]


TAB_SUMMARY_HEADERS = ["Filter", "Issue Type", "Issue Priority", "URLs", "% of Total", "Checked", "Description"]


def tab_summary(index, tab_id: str, title: str, missing_reason: str = "") -> List[Table]:
    """Every filter of one tab with its count (a report-shaped copy of the Overview tree for that tab)."""
    tr = _tr(index, tab_id)
    if tr is None:
        if missing_reason:
            return [not_checked_table(title, missing_reason)]
        return [_missing_tab(title, tab_id)]
    nc = getattr(tr, "not_checked", None) or {}
    if nc.get("all"):
        # the whole tab could not run (an integration not connected, no rendering...): one honest row, no zeros
        return [not_checked_table(title, nc["all"])]
    total = len(tr.rows)
    rows = [[tr.tab.total_label, "", "", total, 100.0 if total else 0.0, "Checked", tr.tab.description or ""]]
    for f in tr.tab.filters:
        why = nc.get(f.id, "")
        n = len(tr.filter_rows.get(f.id) or ())
        rows.append([f.name, f.issue_type or "", f.priority or "", "" if why else n,
                     "" if why else _pct(n, _filter_base(tr, f)),
                     f"Not checked: {why}" if why else "Checked", f.description or ""])
    return [Table(title, TAB_SUMMARY_HEADERS, rows, widths=(44, 12, 14, 9, 11, 50, 70))]


# ---------------------------------------------------------------------------------------------
# Redirect / canonical chains
# ---------------------------------------------------------------------------------------------
def _hop(site, r, follow: str) -> Optional[Tuple[str, str]]:
    """(next URL, hop type) when the URL sends the visitor or search engines elsewhere, else None.
    follow: "redirects" (HTTP, meta refresh, HTTP refresh, JavaScript: internal.redirect_hop, the same hops the
    Redirect Chain / Loop filters walk), "canonicals" or "both"."""
    from .seo.tabs.internal import is_external, redirect_hop
    if follow in ("redirects", "both"):
        target, kind = redirect_hop(r)
        if kind:
            return target, kind
    if follow in ("canonicals", "both") and not is_external(r) and (row_category(r) == "2xx"):
        from .seo.tabs.canonicals import canonical_target, is_canonicalised
        if is_canonicalised(r):
            t = canonical_target(r)
            return (_norm(t) or t), "Canonical"
    return None


REDIRECT_LIMIT = "Redirect limit reached"      # crawler/engine.py: the crawl stopped this chain (Max redirects)


def walk_chain(site, start, follow: str = "redirects") -> dict:
    """Follows one URL's redirects (and / or canonicals) through the crawl: {"hops": [{"url", "code",
    "status", "type", "target"}], "loop", "final" (result or None), "final_url", "stopped" (why the crawl did not
    request the last target, e.g. the Max redirects limit; "" otherwise)}."""
    hops: List[dict] = []
    seen = {getattr(start, "url", "")}
    cur, loop, final, final_url = start, False, None, ""
    while True:
        h = _hop(site, cur, follow)
        if h is None:
            final, final_url = cur, getattr(cur, "url", "")
            break
        nxt, kind = h
        hops.append({"url": cur.url, "code": _code(cur), "status": _status(cur), "type": kind, "target": nxt})
        if not nxt:
            break
        if nxt in seen or len(hops) >= MAX_HOPS:
            loop = nxt in seen
            final_url = nxt
            break
        seen.add(nxt)
        t = _find(site, nxt)
        if t is None:
            final_url = nxt
            break
        cur = t
    stopped = ""
    if final is None and not loop and hops:
        last = _find(site, hops[-1]["url"])
        if REDIRECT_LIMIT in str(getattr(last, "error_message", "") or ""):
            stopped = REDIRECT_LIMIT
    return {"hops": hops, "loop": loop, "final": final, "final_url": final_url, "stopped": stopped}


def _chain_table(site, title: str, starts: Sequence[Any], follow: str, min_hops: int = 1,
                 keep: Optional[Callable[[dict], bool]] = None) -> Table:
    chains = []
    for r in starts:
        c = walk_chain(site, r, follow)
        if (len(c["hops"]) >= min_hops or c["loop"]) and (keep is None or keep(c)):
            chains.append((r, c))
    n = max([len(c["hops"]) for _, c in chains] + [1])
    headers = ["Address", PARENT_LINK, "Linking Pages", "Number of Redirects" if follow == "redirects" else
               "Number of Hops", "Redirect Loop" if follow != "canonicals" else "Canonical Loop",
               "Temp Redirect in Chain", "Chain Type"]
    for i in range(1, n + 1):
        headers += [f"Address {i}", f"Status Code {i}", f"Status {i}", f"Type {i}"]
    headers += ["Final Address", "Final Status Code", "Final Status", "Final Indexability",
                "Final Indexability Status"]
    rows = []
    for r, c in chains:
        hops = c["hops"]
        parent, linking = _parent_link(site, r)
        types = []
        for h in hops:
            if h["type"] not in types:
                types.append(h["type"])
        row = [r.url, parent, linking, len(hops), c["loop"], any(h["code"] in TEMP_REDIRECTS for h in hops),
               " + ".join(types)]
        for i in range(n):
            if i < len(hops):
                h = hops[i]
                row += [h["url"], h["code"], h["status"], h["type"]]
            else:
                row += ["", "", "", ""]
        final = c["final"]
        if c["loop"]:
            row += [c["final_url"], "", "Loop: comes back to a URL already in the chain", "Non-Indexable",
                    "Redirect loop" if follow != "canonicals" else "Canonical loop"]
        elif final is not None:
            label, why = _indexability(site, final)
            row += [final.url, _code(final), _status(final), label, why]
        else:
            row += [c["final_url"], "", f"Not crawled: {c['stopped']}" if c.get("stopped") else "Not crawled", "", ""]
        rows.append(row)
    url_cols = (1, 2) + tuple(8 + 4 * i for i in range(n)) + (8 + 4 * n,)
    return Table(title, headers, rows, url_cols=url_cols)


REDIRECT_FILTERS = ("internal_redirection", "internal_meta_refresh", "internal_http_refresh", "internal_js_redirect",
                    "external_redirection", "external_meta_refresh", "external_http_refresh", "external_js_redirect")
ERROR_CLASSES = ("4xx", "5xx", "error")


def redirects_all(site, index, crawler, ctx: dict) -> List[Table]:
    if _tr(index, "response_codes") is None:
        return [_missing_tab("All Redirects", "response_codes")]
    starts = _union(index, "response_codes", REDIRECT_FILTERS)
    return [_chain_table(site, "All Redirects", starts, "redirects")]


def _ends_in_error(chain: dict) -> bool:
    """The redirect chain ends on a URL the crawl requested that answered 4xx / 5xx or not at all."""
    final = chain.get("final")
    return not chain.get("loop") and final is not None and row_category(final) in ERROR_CLASSES


def redirects_to_error(site, index, crawler, ctx: dict) -> List[Table]:
    """Every redirect (HTTP, meta refresh, HTTP refresh, JavaScript; internal and external) whose chain ends on a
    4xx, a 5xx or no response. Chains whose last URL was not requested are left out (their end is unknown);
    loops are in Redirect Chains."""
    if _tr(index, "response_codes") is None:
        return [_missing_tab("Redirects to Error", "response_codes")]
    starts = _union(index, "response_codes", REDIRECT_FILTERS)
    return [_chain_table(site, "Redirects to Error", starts, "redirects", keep=_ends_in_error)]


def redirect_chains(site, index, crawler, ctx: dict) -> List[Table]:
    if _tr(index, "response_codes") is None:
        return [_missing_tab("Redirect Chains", "response_codes")]
    starts = _union(index, "response_codes", ("internal_redirect_chain", "internal_redirect_loop"))
    return [_chain_table(site, "Redirect Chains", starts, "redirects")]


def redirect_canonical_chains(site, index, crawler, ctx: dict) -> List[Table]:
    starts = list(site.by_url.values()) if site is not None else []
    return [_chain_table(site, "Redirect & Canonical Chains", starts, "both", min_hops=2)]


def canonical_chains(site, index, crawler, ctx: dict) -> List[Table]:
    rows, why = _filter_rows(index, "canonicals", "canonicalised")
    if rows is None:
        return [_missing_tab("Canonical Chains", "canonicals")]
    if why:
        return [not_checked_table("Canonical Chains", why)]
    return [_chain_table(site, "Canonical Chains", rows, "canonicals", min_hops=2)]


NON_INDEXABLE_CANONICAL_HEADERS = ["Address", "Indexability Status", "Canonical Link Element",
                                   "Canonical Status Code", "Canonical Status", "Canonical Indexability",
                                   "Canonical Indexability Status"]


def non_indexable_canonicals(site, index, crawler, ctx: dict) -> List[Table]:
    rows, why = _filter_rows(index, "canonicals", "non_indexable_canonical")
    if rows is None:
        return [_missing_tab("Non-Indexable Canonicals", "canonicals")]
    if why:
        return [not_checked_table("Non-Indexable Canonicals", why)]
    from .seo.tabs.canonicals import canonical_target
    out = []
    for r in rows:
        target = canonical_target(r)
        t = _find(site, target)
        label, status = _indexability(site, t)
        out.append([r.url, site.indexability_status.get(r.url, ""), target, _code(t), _status(t), label, status])
    return [Table("Non-Indexable Canonicals", NON_INDEXABLE_CANONICAL_HEADERS, out, url_cols=(1, 3))]


# ---------------------------------------------------------------------------------------------
# Pagination / hreflang pair reports
# ---------------------------------------------------------------------------------------------
PAGINATION_HEADERS = ["Source", "Pagination URL", "Type", "Status Code", "Status", "Indexability",
                      "Indexability Status"]


def _pagination(site, index, filter_id: str, title: str, test: Callable[[str], bool]) -> List[Table]:
    rows, why = _filter_rows(index, "pagination", filter_id)
    if rows is None:
        return [_missing_tab(title, "pagination")]
    if why:
        return [not_checked_table(title, why)]
    from .seo.tabs.pagination import next_urls, prev_urls
    out = []
    for r in rows:
        before = len(out)
        for kind, urls in (("Next", next_urls(r)), ("Prev", prev_urls(r))):
            for u in urls:
                if test(u):
                    t = _find(site, u)
                    label, status = _indexability(site, t)
                    out.append([r.url, u, kind, _code(t), _status(t), label, status])
        if len(out) == before:          # keep every page of the filter, even if no single link is singled out
            out.append([r.url, "", "", "", "See the Pagination tab", "", ""])
    return [Table(title, PAGINATION_HEADERS, out, url_cols=(1, 2))]


def pagination_non_200(site, index, crawler, ctx: dict) -> List[Table]:
    from .seo.tabs.canonicals import target_non_200
    return _pagination(site, index, "non_200", "Non-200 Pagination URLs", lambda u: target_non_200(site, u))


def pagination_unlinked(site, index, crawler, ctx: dict) -> List[Table]:
    from .seo.tabs.canonicals import is_linked
    return _pagination(site, index, "unlinked", "Unlinked Pagination URLs", lambda u: not is_linked(site, u))


HREFLANG_HEADERS = ["Source", "Hreflang Language Code", "Hreflang URL", "Annotation In", "Status Code", "Status",
                    "Indexability", "Indexability Status", "Note"]


def _hl_points_to(e: dict, r) -> bool:
    from .seo.tabs.canonicals import same_url
    p = getattr(r, "seo", None) or {}
    u = e.get("url") or ""
    return same_url(u, r.url) or same_url(u, p.get("final_url") or "")


def _hl_page(site, url: str):
    """The crawled internal URL behind an hreflang URL when it answered 200 (else None; the Hreflang tab's rule)."""
    from .seo.tabs.hreflang import _alt_page
    return _alt_page(site, url)


def _hl_lang(e: dict) -> str:
    return str(e.get("lang") or "").strip().lower()


def _hl_entries(r, site=None) -> List[Tuple[dict, str]]:
    """(annotation, where it is written) for the page: HTML, HTTP Header, and the XML sitemaps (P8)."""
    from .seo.tabs.hreflang import sitemap_entries
    p = getattr(r, "seo", None) or {}
    return ([(e, "HTML") for e in p.get("hreflang") or () if e.get("lang") or e.get("url")]
            + [(e, "HTTP Header") for e in p.get("http_hreflang") or () if e.get("lang") or e.get("url")]
            + [(e, "Sitemap") for e in sitemap_entries(r, site)])


def _hl_tests(site) -> Dict[str, Tuple[str, str, Callable[[Any, dict], str]]]:
    """report id -> (filter id, title, test(page, entry) -> "" (not listed) or the note for the row)."""
    from .seo.tabs.canonicals import has_noindex, is_linked, is_start_url, target_non_200
    from .seo.tabs.hreflang import unjudged_alternates
    not_requested = "Not checked: the crawl did not request this alternate"

    def alternate(r, e):
        return bool(e.get("url")) and not _hl_points_to(e, r)

    def unjudged(r, e, filter_id):
        return alternate(r, e) and e["url"] in unjudged_alternates(site, r, filter_id)

    def all_(r, e):
        return "Self reference" if _hl_points_to(e, r) else "Alternate"

    def non_200(r, e):
        if unjudged(r, e, "non_200"):
            return not_requested
        return "Does not answer 200" if e.get("url") and target_non_200(site, e["url"]) else ""

    def unlinked(r, e):
        return ("No crawled page links to it" if alternate(r, e) and not is_start_url(site, e["url"])
                and not is_linked(site, e["url"]) else "")

    def missing_return(r, e):
        if not alternate(r, e):
            return ""
        if unjudged(r, e, "missing_return_links"):
            return not_requested
        t = _hl_page(site, e["url"])
        return ("The alternate does not link back" if t is not None
                and not any(_hl_points_to(x, r) for x, _ in _hl_entries(t, site)) else "")

    def inconsistent(r, e):
        own = {_hl_lang(x) for x, _ in _hl_entries(r, site) if _hl_points_to(x, r)}
        if not own or not alternate(r, e):
            return ""
        if unjudged(r, e, "inconsistent_return_links"):
            return not_requested
        t = _hl_page(site, e["url"])
        if t is None:
            return ""
        back = {_hl_lang(x) for x, _ in _hl_entries(t, site) if _hl_points_to(x, r)}
        extra = back - own
        return f"Links back as {', '.join(sorted(extra))}; this page calls itself {', '.join(sorted(own))}" \
            if extra else ""

    def non_canonical(r, e):
        if not alternate(r, e):
            return ""
        if unjudged(r, e, "non_canonical_return_links"):
            return not_requested
        from .seo.tabs.internal import is_external
        t = _find(site, e["url"])
        return ("The alternate is canonicalised to another URL" if t is not None and not is_external(t)
                and site.indexability_status.get(t.url) == "Canonicalised" else "")

    def noindex(r, e):
        if unjudged(r, e, "noindex_return_links"):
            return not_requested
        return "The alternate is noindex" if alternate(r, e) and has_noindex(site, e["url"]) else ""

    return {
        "hreflang_all": ("contains_hreflang", "All hreflang URLs", all_),
        "hreflang_non_200": ("non_200", "Non-200 hreflang URLs", non_200),
        "hreflang_unlinked": ("unlinked", "Unlinked hreflang URLs", unlinked),
        "hreflang_missing_return": ("missing_return_links", "Missing Return Links", missing_return),
        "hreflang_inconsistent_return": ("inconsistent_return_links",
                                         "Inconsistent Language & Region Return Links", inconsistent),
        "hreflang_non_canonical_return": ("non_canonical_return_links", "Non Canonical Return Links",
                                          non_canonical),
        "hreflang_noindex_return": ("noindex_return_links", "Noindex Return Links", noindex),
    }


def hreflang_report(report_id: str, site, index) -> List[Table]:
    filter_id, title, test = _hl_tests(site)[report_id]
    rows, why = _filter_rows(index, "hreflang", filter_id)
    if rows is None:
        return [_missing_tab(title, "hreflang")]
    if why:
        return [not_checked_table(title, why)]
    out = []
    for r in rows:
        before = len(out)
        for e, where in _hl_entries(r, site):
            note = test(r, e)
            if not note:
                continue
            url = e.get("url") or ""
            t = _find(site, url)
            label, status = _indexability(site, t)
            out.append([r.url, e.get("lang") or "", url, where, _code(t), _status(t) if url else "", label, status,
                        note])
        if len(out) == before:
            out.append([r.url, "", "", "", "", "", "", "", "See the Hreflang tab"])
    return [Table(title, HREFLANG_HEADERS, out, url_cols=(1, 3))]


# ---------------------------------------------------------------------------------------------
# Insecure content, SERP summary, orphan pages
# ---------------------------------------------------------------------------------------------
INSECURE_HEADERS = ["Page", "Insecure URL", "Type", "Internal", "Status Code", "Status"]


def insecure_content(site, index, crawler, ctx: dict) -> List[Table]:
    """HTTPS pages loading HTTP resources (the Security > Mixed Content pages), plus their insecure form
    actions and hyperlinks to HTTP addresses. Pages with a resource row == the Mixed Content count."""
    from .seo.tabs.internal import is_html_doc
    out = []
    for p in getattr(site, "pages", None) or ():
        if not p.url.lower().startswith("https://") or not is_html_doc(p):
            continue
        seo = getattr(p, "seo", None) or {}
        types = {res.get("url"): res for res in seo.get("resources") or () if res.get("url")}
        for u in seo.get("mixed_content") or ():
            res = types.get(u) or {}
            t = _find(site, u)
            out.append([p.url, u, _RESOURCE_TYPES.get(res.get("type") or "", "Other"), bool(res.get("internal")),
                        _code(t), _status(t)])
        for f in seo.get("forms") or ():
            if f.get("insecure") and f.get("action"):
                t = _find(site, f["action"])
                out.append([p.url, f["action"], "Form", t is not None and not hasattr(t, "domain"), _code(t),
                            _status(t)])
        for link in seo.get("links") or ():
            u = link.get("url") or ""
            if u.lower().startswith("http://"):
                t = _find(site, u)
                out.append([p.url, u, "Hyperlink", bool(link.get("internal")), _code(t), _status(t)])
    return [Table("Insecure Content", INSECURE_HEADERS, out, url_cols=(1, 2))]


def serp_summary(site, index, crawler, ctx: dict) -> List[Table]:
    from . import serp
    return [Table("SERP Summary", serp.SUMMARY_HEADERS, serp.summary_rows(serp.crawl_rows(site)),
                  widths=(60, 60, 12, 16, 80, 18, 22), url_cols=(1,))]


ORPHAN_HEADERS = ["Address", "Source", "Sitemaps", "Status Code", "Status", "Indexability", "Indexability Status"]


# Orphan sources besides XML sitemaps: (tab id, source label) of the optional integrations. Their "Orphan URLs"
# filter is used when the tab exists and the integration is connected (a tab that is not connected reports
# "Not connected" itself; its URLs simply do not appear here).
ORPHAN_SOURCES = (("search_console", "Google Search Console"), ("analytics", "Google Analytics"))


def _orphan_filter(tr) -> Optional[str]:
    return next((f.id for f in tr.tab.filters if "orphan" in f.id.lower() or "orphan" in f.name.lower()), None)


def orphan_pages(site, index, crawler, ctx: dict) -> List[Table]:
    """URLs a source knows (XML sitemaps; Search Console and Analytics when connected) that no crawled page
    links to: the union of the Sitemaps > Orphan URLs filter and the integrations' Orphan URLs filters."""
    rows, why = _filter_rows(index, "sitemaps", "orphan")
    sources: "OrderedDict[str, List[str]]" = OrderedDict()
    reasons = []
    if rows is None:
        reasons.append("the Sitemaps tab is not available")
    elif why:
        reasons.append(f"XML Sitemap: {why}")
    else:
        for r in rows:
            sources.setdefault(getattr(r, "url", ""), []).append("XML Sitemap")
    for tab_id, label in ORPHAN_SOURCES:
        tr = _tr(index, tab_id)
        fid = _orphan_filter(tr) if tr is not None else None
        if fid is None or (getattr(tr, "not_checked", None) or {}).get(fid):
            continue
        for r in tr.filter_rows.get(fid) or ():
            src = sources.setdefault(getattr(r, "url", ""), [])
            if label not in src:
                src.append(label)
    if not sources and reasons:
        return [not_checked_table("Orphan Pages", "; ".join(reasons))]
    from .seo.tabs.sitemaps import sitemaps_of
    out = []
    for url, found_in in sources.items():
        crawled = site.by_url.get(url)
        label, status = _indexability(site, crawled)
        out.append([url, ", ".join(found_in), sitemaps_of(site, url), _code(crawled) if crawled is not None else "",
                    _status(crawled) if crawled is not None else "Not crawled", label, status])
    return [Table("Orphan Pages", ORPHAN_HEADERS, out, url_cols=(1,))]


# ---------------------------------------------------------------------------------------------
# Structured data
# ---------------------------------------------------------------------------------------------
def _sd_split(msg: str) -> Tuple[str, str]:
    """'Product (JSON-LD): missing offers' -> ('Product (JSON-LD)', 'missing offers')."""
    head, sep, tail = str(msg).partition(": ")
    return (head, tail) if sep else ("", str(msg))


def _sd_rows(site, index) -> Tuple[Optional[List[Any]], str]:
    tr = _tr(index, "structured_data")
    if tr is None:
        return None, ""
    return [r for r in tr.rows if not hasattr(r, "domain")], ""


def _sd_messages(r) -> List[Tuple[str, str, str]]:
    """[(severity, schema type, issue)] of a page: parse errors, then validation errors and warnings."""
    from .seo.tabs.structured_data import parse_errors, validation_errors, validation_warnings
    out = []
    n = parse_errors(r)
    if n:
        out.append(("Error", "JSON-LD", f"{n} block{'s' if n != 1 else ''} could not be parsed"))
    out += [("Error",) + _sd_split(m) for m in validation_errors(r)]
    out += [("Warning",) + _sd_split(m) for m in validation_warnings(r)]
    return out


SD_SUMMARY_HEADERS = ["Severity", "Schema Type", "Issue", "URLs", "% of Pages With Structured Data"]
SD_DETAIL_HEADERS = ["Address", "Severity", "Schema Type", "Issue"]


def sd_validation_summary(site, index, crawler, ctx: dict) -> List[Table]:
    rows, _ = _sd_rows(site, index)
    if rows is None:
        return [_missing_tab("Validation Errors & Warnings Summary", "structured_data")]
    from .seo.tabs.structured_data import contains
    with_sd = [r for r in rows if contains(r)]
    agg: Dict[Tuple[str, str, str], set] = OrderedDict()
    for r in with_sd:
        for key in _sd_messages(r):
            agg.setdefault(key, set()).add(r.url)
    out = [[sev, typ, issue, len(urls), _pct(len(urls), len(with_sd))]
           for (sev, typ, issue), urls in sorted(agg.items(), key=lambda kv: (kv[0][0] != "Error", -len(kv[1])))]
    return [Table("Validation Errors & Warnings Summary", SD_SUMMARY_HEADERS, out, widths=(10, 30, 60, 9, 18))]


def sd_validation(site, index, crawler, ctx: dict) -> List[Table]:
    rows, _ = _sd_rows(site, index)
    if rows is None:
        return [_missing_tab("Validation Errors & Warnings", "structured_data")]
    out = [[r.url, sev, typ, issue] for r in rows for sev, typ, issue in _sd_messages(r)]
    return [Table("Validation Errors & Warnings", SD_DETAIL_HEADERS, out, url_cols=(1,), widths=(60, 10, 30, 60))]


SD_RICH_SUMMARY_HEADERS = ["Rich Result Feature", "URLs", "URLs With Errors", "URLs With Warnings", "Note"]
SD_RICH_HEADERS = ["Address", "Rich Result Feature", "Result", "Errors", "Warnings"]


def _rich(site, r) -> dict:
    from .seo.tabs.structured_data import rich_results
    return rich_results(r, site)


def sd_rich_summary(site, index, crawler, ctx: dict) -> List[Table]:
    rows, _ = _sd_rows(site, index)
    if rows is None:
        return [_missing_tab("Google Rich Result Features Summary", "structured_data")]
    agg: Dict[str, dict] = OrderedDict()
    for r in rows:
        rr = _rich(site, r)
        for f in rr["features"]:
            a = agg.setdefault(f, {"urls": set(), "err": set(), "warn": set(), "note": ""})
            a["urls"].add(r.url)
            if any(e.startswith(f + " - ") for e in rr["errors"]):
                a["err"].add(r.url)
            if any(w.startswith(f + " - ") for w in rr["warnings"]):
                a["warn"].add(r.url)
        for f in rr["retired"]:
            a = agg.setdefault(f, {"urls": set(), "err": set(), "warn": set(),
                                   "note": "No longer shown by Google"})
            a["urls"].add(r.url)
    out = [[f, len(a["urls"]), len(a["err"]), len(a["warn"]), a["note"]]
           for f, a in sorted(agg.items(), key=lambda kv: -len(kv[1]["urls"]))]
    return [Table("Google Rich Result Features Summary", SD_RICH_SUMMARY_HEADERS, out, widths=(36, 9, 16, 18, 30))]


def sd_rich(site, index, crawler, ctx: dict) -> List[Table]:
    rows, _ = _sd_rows(site, index)
    if rows is None:
        return [_missing_tab("Google Rich Result Features", "structured_data")]
    out = []
    for r in rows:
        rr = _rich(site, r)
        for f in rr["features"]:
            errs = [e.split(" - ", 1)[1] for e in rr["errors"] if e.startswith(f + " - ")]
            warns = [w.split(" - ", 1)[1] for w in rr["warnings"] if w.startswith(f + " - ")]
            result = "Invalid (errors)" if errs else ("Valid with warnings" if warns else "Valid")
            out.append([r.url, f, result, errs, warns])
        for f in rr["retired"]:
            out.append([r.url, f, "No longer shown by Google", [], []])
    return [Table("Google Rich Result Features", SD_RICH_HEADERS, out, url_cols=(1,), widths=(60, 30, 20, 60, 60))]


# ---------------------------------------------------------------------------------------------
# JavaScript console, headers, cookies
# ---------------------------------------------------------------------------------------------
CONSOLE_HEADERS = ["Level", "Message", "Source", "Pages", "First Page"]


def js_console_summary(site, index, crawler, ctx: dict) -> List[Table]:
    from .seo.registry import RENDER_REASON, rendered
    pages = [p for p in getattr(site, "pages", None) or () if rendered(p) is not None]
    if not pages:
        return [not_checked_table("JavaScript Console Log Summary", RENDER_REASON)]
    agg: Dict[Tuple[str, str, str], List[str]] = OrderedDict()
    for p in pages:
        for c in rendered(p).get("console") or ():
            if not isinstance(c, dict):
                continue
            key = (str(c.get("level") or "").lower() or "log", " ".join(str(c.get("text") or "").split())[:500],
                   str(c.get("source") or ""))
            urls = agg.setdefault(key, [])
            if p.url not in urls:
                urls.append(p.url)
    order = {"severe": 0, "error": 0, "warning": 1, "info": 2, "log": 3}
    out = [[lvl, msg, src, len(urls), urls[0]]
           for (lvl, msg, src), urls in sorted(agg.items(), key=lambda kv: (order.get(kv[0][0], 4), -len(kv[1])))]
    return [Table("JavaScript Console Log Summary", CONSOLE_HEADERS, out, url_cols=(5,), widths=(10, 80, 40, 9, 60))]


HEADER_SUMMARY_HEADERS = ["Header", "URLs", "% of URLs With Headers"]


def http_header_summary(site, index, crawler, ctx: dict) -> List[Table]:
    counts: Dict[str, int] = {}
    with_headers = 0
    for r in (getattr(site, "by_url", None) or {}).values():
        h = getattr(r, "response_headers", None) or {}
        if not h:
            continue
        with_headers += 1
        for name in {str(k).lower() for k in h}:
            counts[name] = counts.get(name, 0) + 1
    out = [[k, n, _pct(n, with_headers)] for k, n in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))]
    return [Table("HTTP Header Summary", HEADER_SUMMARY_HEADERS, out, widths=(44, 10, 22))]


COOKIE_SUMMARY_HEADERS = ["Cookie Name", "Domain", "Path", "Secure", "HttpOnly", "SameSite", "Expires",
                          "Max-Age", "URLs", "First URL"]


def _cookies(site):
    for r in (getattr(site, "by_url", None) or {}).values():
        for c in (getattr(r, "seo", None) or {}).get("cookies") or ():
            if isinstance(c, dict) and c.get("name"):
                yield r, c


def cookie_summary(site, index, crawler, ctx: dict) -> List[Table]:
    agg: Dict[Tuple[str, str, str], dict] = OrderedDict()
    for r, c in _cookies(site):
        key = (str(c.get("name")), str(c.get("domain") or ""), str(c.get("path") or ""))
        a = agg.setdefault(key, {"c": c, "urls": []})
        if r.url not in a["urls"]:
            a["urls"].append(r.url)
    out = []
    for (name, domain, path), a in agg.items():
        c = a["c"]
        out.append([name, domain, path, bool(c.get("secure")), bool(c.get("httponly")), c.get("samesite") or "",
                    c.get("expires") or "", c.get("max_age") if c.get("max_age") not in (None, "") else "",
                    len(a["urls"]), a["urls"][0]])
    return [Table("Cookie Summary", COOKIE_SUMMARY_HEADERS, out, url_cols=(10,))]


# ---------------------------------------------------------------------------------------------
# Report registry
# ---------------------------------------------------------------------------------------------
# ---------------------------------------------------------------------------------------------
# Comparison Overview (P10)
# ---------------------------------------------------------------------------------------------
COMPARE_META_HEADERS = ["", "Previous Crawl", "Current Crawl"]
COMPARE_OVERVIEW_HEADERS = ["Tab", "Filter", "Issue Type", "Priority", "Current", "Previous", "Change", "Change %",
                            "Added", "New", "Removed", "Missing", "Note"]
COMPARE_TOTAL_HEADERS = ["Summary", "Current", "Previous", "Change", "Change %", "Added", "New", "Removed", "Missing"]
COMPARE_CHANGE_HEADERS = ["Element", "Compared", "Changed", "Not Changed", "Not Checked"]
COMPARE_ISSUE_HEADERS = ["Issue", "Issue Type", "Priority", "Previous", "Current", "New", "Regressed", "Fixed",
                         "Not Rechecked", "Unchanged", "Note"]


def _comparison(site, index, crawler, ctx: dict) -> Tuple[Optional[dict], str]:
    """(compare.compare() result, "" ) for the report, or (None, why it cannot be built)."""
    given = ctx.get("comparison")
    if isinstance(given, dict):
        return (None, str(given["error"])) if given.get("error") else (given, "")
    req = ctx.get("compare")
    if not isinstance(req, dict) or not str(req.get("a") or "").strip():
        return None, ("No comparison was chosen: open Mode > Compare, pick the crawl to compare with, then download "
                      "this report.")
    from .seo import compare
    a, b = str(req.get("a") or ""), str(req.get("b") or "current")
    current = None
    if "current" in (a.strip().lower(), b.strip().lower()):
        current = compare.snapshot_from(crawler, (site, index), site=site, segments=ctx.get("segments") or None)
    result = compare.compare(a, b, ctx.get("reports_dir") or None, current=current, url_map=req.get("url_map"),
                             change_detection=req.get("change_detection"), thresholds=req.get("thresholds"))
    if result.get("error"):
        return None, str(result["error"])
    return result, ""


def _cmp_note(r: dict) -> str:
    """Why a row of the comparison cannot show added / removed URLs ("" when it can)."""
    if r.get("not_comparable"):
        return r["not_comparable"]
    if r.get("check_current") == "not_checked":
        return "Not checked in the current crawl"
    if r.get("check_previous") == "not_checked":
        return "Not checked in the previous crawl"
    if not r.get("exact"):
        return "The older crawl kept only this filter's count, not its URLs"
    return ""


def _blank(v: Any) -> Any:
    return "" if v is None else v


def compare_overview(site, index, crawler, ctx: dict) -> List[Table]:
    """The comparison of Mode > Compare as tables: the two crawls (and settings that changed), every filter's
    Current / Previous / Change / Added / New / Removed / Missing (Added + New - Removed - Missing = Change on every
    comparable row), the summary rows, change detection and the issue changes. A filter whose deciding setting
    changed between the crawls shows "Not comparable" with the change instead of buckets."""
    result, why = _comparison(site, index, crawler, ctx)
    if result is None:
        return [not_checked_table("Comparison Overview", why)]
    a, b = result.get("a") or {}, result.get("b") or {}
    meta_rows = [[label, _blank(a.get(k)), _blank(b.get(k))] for label, k in (
        ("Crawl", "id"), ("Site", "site"), ("Started", "started_at"), ("Internal URLs", "urls"), ("Pages", "pages"),
        ("Pages that load", "pages_ok"), ("Issue types", "issues"), ("Issue URLs", "issue_urls"),
        ("Health score", "health_score"), ("Snapshot version", "version"))]
    for s in result.get("settings_changed") or []:
        meta_rows.append([f"Setting changed: {s.get('setting')}", str(s.get("previous")), str(s.get("current"))])
    if (result.get("url_map") or {}).get("rules"):
        meta_rows.append(["URL Mapping rules", f"{result['url_map']['rules']} rule(s), "
                                               f"{result['url_map'].get('mapped', 0)} URLs rewritten", ""])
    overview = []
    for r in result.get("overview") or []:
        nc = _cmp_note(r)
        side_nc = r.get("check_current") == "not_checked" or r.get("check_previous") == "not_checked"
        overview.append([r.get("tab_name") or r.get("tab"), r.get("filter_name") or r.get("filter"),
                         r.get("issue_type") or "", r.get("priority") or "", r.get("current"), r.get("previous"),
                         "" if side_nc else _blank(r.get("change")),
                         "" if side_nc else _blank(r.get("change_pct")),
                         "" if side_nc else _blank(r.get("added")), "" if side_nc else _blank(r.get("new")),
                         "" if side_nc else _blank(r.get("removed")), "" if side_nc else _blank(r.get("missing")),
                         nc])
    totals = [[t.get("label"), t.get("current"), t.get("previous"), t.get("change"), _blank(t.get("change_pct")),
               t.get("added"), t.get("new"), t.get("removed"), t.get("missing")] for t in result.get("totals") or []]
    changes = [[c.get("label") or c.get("element"), c.get("compared"), c.get("changed"), c.get("not_changed"),
                c.get("not_checked") or ""] for c in result.get("changes") or []]
    issues = [[i.get("name"), i.get("issue_type") or "", i.get("priority") or "", i.get("a"), i.get("b"),
               i.get("new_count"), i.get("regressed_count"), i.get("fixed_count"), i.get("not_rechecked_count"),
               i.get("unchanged_count"), i.get("not_comparable") or ""] for i in result.get("issues") or []]
    return [Table("Comparison", COMPARE_META_HEADERS, meta_rows, widths=(36, 40, 40)),
            Table("Comparison - Every Filter", COMPARE_OVERVIEW_HEADERS, overview,
                  widths=(22, 36, 12, 10, 10, 10, 10, 10, 10, 10, 10, 10, 50)),
            Table("Comparison - Summary", COMPARE_TOTAL_HEADERS, totals),
            Table("Comparison - Change Detection", COMPARE_CHANGE_HEADERS, changes),
            Table("Comparison - Issues", COMPARE_ISSUE_HEADERS, issues,
                  widths=(50, 12, 10, 10, 10, 8, 10, 8, 14, 10, 50))]


def _hl(report_id):
    return lambda site, index, crawler, ctx: hreflang_report(report_id, site, index)


_BUILDERS: Dict[str, Callable[[Any, Any, Any, dict], List[Table]]] = {
    "crawl_overview": crawl_overview,
    "issues_overview": issues_overview,
    "segments_overview": segments_overview,
    "redirects_all": redirects_all,
    "redirect_chains": redirect_chains,
    "redirects_to_error": redirects_to_error,
    "redirect_canonical_chains": redirect_canonical_chains,
    "canonical_chains": canonical_chains,
    "non_indexable_canonicals": non_indexable_canonicals,
    "pagination_non_200": pagination_non_200,
    "pagination_unlinked": pagination_unlinked,
    "insecure_content": insecure_content,
    "serp_summary": serp_summary,
    "orphan_pages": orphan_pages,
    "sd_validation_summary": sd_validation_summary,
    "sd_validation": sd_validation,
    "sd_rich_summary": sd_rich_summary,
    "sd_rich": sd_rich,
    "js_console_summary": js_console_summary,
    "pagespeed_summary": lambda s, i, c, ctx: tab_summary(i, "pagespeed", "PageSpeed Summary",
                                                          _not_connected()),
    "mobile_summary": lambda s, i, c, ctx: tab_summary(i, "mobile", "Mobile Summary"),
    "accessibility_summary": lambda s, i, c, ctx: tab_summary(i, "accessibility", "Accessibility Violations Summary"),
    "http_header_summary": http_header_summary,
    "cookie_summary": cookie_summary,
    "compare_overview": compare_overview,
}
for _id in ("hreflang_all", "hreflang_non_200", "hreflang_unlinked", "hreflang_missing_return",
            "hreflang_inconsistent_return", "hreflang_non_canonical_return", "hreflang_noindex_return"):
    _BUILDERS[_id] = _hl(_id)


def _not_connected() -> str:
    try:
        from .seo.registry import NOT_CONNECTED
        return NOT_CONNECTED
    except Exception:
        return "Not connected - connect under Configuration > API Access."


def report_info(report_id: str) -> Optional[dict]:
    return next((r for r in REPORTS if r["id"] == report_id), None)


def report_tables(report_id: str, site, index, crawler=None, **ctx) -> List[Table]:
    """The tables of one report (KeyError for an unknown id). Without a crawl every report is "Not checked"."""
    if report_id not in _BUILDERS:
        raise KeyError(f"Unknown report: {report_id}")
    info = report_info(report_id) or {"name": report_id}
    if site is None or index is None:
        return [not_checked_table(info["name"], "There is no finished crawl to report on yet.")]
    return _BUILDERS[report_id](site, index, crawler, ctx)


def build_report(report_id: str, site, index, crawler=None, fmt: str = "xlsx", **ctx) -> Tuple[str, bytes, str]:
    """(file name, bytes, media type) of one Reports menu item: a workbook (one sheet per table), a .csv, or
    a .zip of CSVs when the report has more than one table (Crawl Overview, Segments Overview)."""
    check_format(fmt)
    tables = report_tables(report_id, site, index, crawler, **ctx)
    return tables_bytes(tables, fmt, _file_stem(site, report_id))


def _file_stem(site, what: str) -> str:
    host = ""
    try:
        from urllib.parse import urlsplit
        host = urlsplit(getattr(site, "start_url", "") or "").netloc
    except Exception:
        pass
    return safe_name(f"{host}_{what}" if host else what)


# ---------------------------------------------------------------------------------------------
# Bulk exports
# ---------------------------------------------------------------------------------------------
def bulk_catalog(index=None) -> List[dict]:
    """Every Bulk Export item: the fixed ones, then per visible tab one item per filter ("filter:tab:id", the
    filter's URLs with the tab's columns) and one for the pages pointing at them ("inlinks:tab:id"), then
    per flagged issue an Issues item ("issue:tab:id", the same rows as its filter item). Ids are unique.
    kind: "bulk" | "filter" | "inlinks" | "issue"; count is None before a crawl; not_checked holds the reason
    a filter could not run ("" when it ran)."""
    out = [dict(b, kind="bulk") for b in BULK_EXPORTS]
    if index is None:
        try:
            from .seo import registry
            tabs = [(t, None) for t in registry.tabs()]
        except Exception:
            tabs = []
    else:
        tabs = [(tr.tab, tr) for _, tr in sorted(index.tabs.items(), key=lambda kv: kv[1].tab.order)]
    for tab, tr in tabs:
        if getattr(tab, "hidden", False):
            continue
        for fid, fname in [("all", tab.total_label)] + [(f.id, f.name) for f in tab.filters]:
            count = None
            if tr is not None:
                count = len(tr.rows) if fid == "all" else len(tr.filter_rows.get(fid) or ())
            nc = (getattr(tr, "not_checked", None) or {}).get(fid, "") if tr is not None else ""
            out.append({"id": f"filter:{tab.id}:{fid}", "name": fname, "menu_path": [tab.name, fname],
                        "kind": "filter", "tab": tab.id, "filter": fid, "count": count, "not_checked": nc})
            out.append({"id": f"inlinks:{tab.id}:{fid}", "name": f"{fname} Inlinks",
                        "menu_path": [tab.name, f"{fname} Inlinks"], "kind": "inlinks", "tab": tab.id,
                        "filter": fid, "count": count, "not_checked": nc})
    for issue in (index.issues if index is not None else []) or []:
        out.append({"id": f"issue:{issue['tab']}:{issue['filter']}", "name": issue["name"],
                    "menu_path": ["Issues", issue["name"]], "kind": "issue", "tab": issue["tab"],
                    "filter": issue["filter"], "count": issue.get("urls"), "not_checked": ""})
    return out


QUEUED_HEADERS = ["URL", PARENT_LINK, "Depth", "Source", "Reason"]


def queued_urls(site, index, crawler, ctx: dict) -> List[Table]:
    """URLs waiting in the crawl queue (paused / stopped crawl) and those found but left out by the Max URLs,
    Max depth, memory, Max redirects or Max sitemap-only URLs limits."""
    if crawler is None or not hasattr(crawler, "queue"):
        return [not_checked_table("Queued URLs", "No crawl is loaded.")]
    lock = getattr(crawler, "_lock", None)
    rows = []
    for _ in range(3):              # the queue can change while it is copied: retry, then take what we have
        try:
            if lock is not None:
                with lock:
                    queue = list(crawler.queue)
                    enqueued = getattr(crawler, "enqueued_urls", None) or set()
                    limits = [(u, "Not crawled: Max URLs limit") for u in getattr(crawler, "_over_limit", ())] + \
                             [(u, "Not crawled: Max depth limit") for u in getattr(crawler, "_over_depth", ())] + \
                             [(u, "Not crawled: memory limit") for u in getattr(crawler, "_over_memory", ())] + \
                             [(u, f"Not crawled: {REDIRECT_LIMIT} (after {p})")
                              for u, p in (getattr(crawler, "_over_redirects", None) or {}).items()] + \
                             [(u, "Not crawled: Max sitemap-only URLs limit")
                              for u in getattr(crawler, "_over_sitemap_only", ()) if u not in enqueued]
            else:
                queue, limits = list(crawler.queue), []
            break
        except RuntimeError:
            queue, limits = [], []
    for item in queue:
        try:
            url, parent, depth, source = (list(item) + ["", "", 0, ""])[:4]
        except Exception:
            continue
        rows.append([url, parent or "", depth, source or "", "Waiting in the queue"])
    for url, why in sorted(limits):
        rows.append([url, "", "", "", why])
    return [Table("Queued URLs", QUEUED_HEADERS, rows, url_cols=(1, 2))]


def all_inlinks(site, index, crawler, ctx: dict) -> List[Table]:
    return [Table("All Inlinks", INLINK_HEADERS, inlink_rows(site), url_cols=(1, 2),
                  widths=(60, 60, 36, 10, 20, 12, 12, 24, 16, 16, 20))]


OUTLINK_HEADERS = ["Source", "Destination", "Type", "Anchor Text", "Alt Text", "Follow", "Target", "Internal",
                   "Path Type", "Link Position", "Source Line", "Destination Status Code", "Destination Status",
                   "Rel"]


def outlink_rows(site, external_only: bool = False, hyperlinks_only: bool = False) -> List[list]:
    """OUTLINK_HEADERS rows: every hyperlink, then every resource (image, script, stylesheet, frame, CSS
    background image) of every crawled page, then every HTTP redirect (its source is the redirecting URL), then
    every other link type of every crawled URL (canonical, hreflang, rel next / prev, meta refresh, Refresh header,
    JavaScript redirect, mobile alternate, form action, other link elements; P6), internal and external. Rel holds
    the rel values. The internal ones are the rows of All Inlinks (exporter.inlink_rows)."""
    from .exporter import (background_links, is_internal_target, link_rel, other_outlinks_of, redirect_links,
                           resource_rel)
    rows = []

    def status(u):
        t = site.by_url.get(u) or externals_by_url(site).get(u)
        return (_code(t), _status(t)) if t is not None else ("", "Not crawled")

    for p in getattr(site, "pages", None) or ():
        seo = getattr(p, "seo", None) or {}
        for link in seo.get("links") or ():
            u = link.get("url") or ""
            if not u or (external_only and link.get("internal")):
                continue
            code, st = status(u)
            rows.append([p.url, u, "Hyperlink", link.get("anchor") or "", "", not link.get("nofollow"),
                         "_blank" if link.get("target_blank") else "", bool(link.get("internal")),
                         path_type(link.get("raw")), link.get("region") or "", int(link.get("line") or 0) or "",
                         code, st, link_rel(link)])
        if hyperlinks_only:
            continue
        for res in seo.get("resources") or ():
            u = res.get("url") or ""
            if not u or (external_only and res.get("internal")):
                continue
            code, st = status(u)
            kind = _RESOURCE_TYPES.get(res.get("type") or "", "Other")
            rows.append([p.url, u, kind, "", (res.get("alt") or "") if kind == "Image" else "", True, "",
                         bool(res.get("internal")), path_type(res.get("raw")), res.get("region") or "",
                         int(res.get("line") or 0) or "", code, st, resource_rel(res)])
        for bg in background_links(p):
            u = bg["url"]
            internal = is_internal_target(site, u)
            if external_only and internal:
                continue
            code, st = status(u)
            rows.append([p.url, u, "Background Image", "", "", True, "", internal, "", "CSS background image", "",
                         code, st, ""])
    if not hyperlinks_only:
        for rd in redirect_links(site):
            u = rd["url"]
            internal = is_internal_target(site, u)
            if external_only and internal:
                continue
            code, st = status(u)
            rows.append([rd["source"], u, "HTTP Redirect", "", "", True, "", internal, path_type(rd["raw"]),
                         "HTTP redirect", "", code, st, ""])
        for r in (getattr(site, "by_url", None) or {}).values():
            for e in other_outlinks_of(site, r):
                u = e["target"]
                internal = is_internal_target(site, u)
                if external_only and internal:
                    continue
                code, st = status(u)
                rows.append([r.url, u, e["type"], e["anchor"], "", True, "", internal, path_type(e.get("raw")),
                             e["region"], int(e.get("line") or 0) or "", code, st, e["rel"]])
    return rows


def all_outlinks(site, index, crawler, ctx: dict) -> List[Table]:
    return [Table("All Outlinks", OUTLINK_HEADERS, outlink_rows(site), url_cols=(1, 2))]


def external_links(site, index, crawler, ctx: dict) -> List[Table]:
    return [Table("External Links", OUTLINK_HEADERS, outlink_rows(site, external_only=True, hyperlinks_only=True),
                  url_cols=(1, 2))]


ANCHOR_HEADERS = ["Anchor Text", "Links", "Unique Destinations", "Unique Sources", "Nofollow Links",
                  "Top Destination"]


def all_anchor_text(site, index, crawler, ctx: dict) -> List[Table]:
    """Every distinct anchor text of the internal hyperlinks ("" = links with no anchor text)."""
    agg: Dict[str, dict] = {}
    for dest, links in (getattr(site, "inlinks", None) or {}).items():
        for link in links or ():
            a = " ".join(str(link.get("anchor") or "").split())
            e = agg.setdefault(a, {"n": 0, "dest": {}, "src": set(), "nofollow": 0})
            e["n"] += 1
            e["dest"][dest] = e["dest"].get(dest, 0) + 1
            e["src"].add(link.get("source") or "")
            e["nofollow"] += 1 if link.get("nofollow") else 0
    out = []
    for a, e in sorted(agg.items(), key=lambda kv: (-kv[1]["n"], kv[0])):
        top = max(e["dest"].items(), key=lambda kv: kv[1])[0] if e["dest"] else ""
        out.append([a or "(no anchor text)", e["n"], len(e["dest"]), len(e["src"]), e["nofollow"], top])
    return [Table("All Anchor Text", ANCHOR_HEADERS, out, url_cols=(6,), widths=(50, 9, 20, 16, 15, 60))]


def path_type_export(export_id: str):
    wanted = PATH_EXPORTS[export_id]

    def build(site, index, crawler, ctx):
        rows = [r for r in outlink_rows(site) if r[8] == wanted]
        return [Table(f"Path Type - {wanted}", OUTLINK_HEADERS, rows, url_cols=(1, 2))]
    return build


HEADER_EXPORT_HEADERS = ["Address", "Status Code", "Status", "Header", "Value"]


def all_http_headers(site, index, crawler, ctx: dict) -> List[Table]:
    """Every response header of every internal URL; Set-Cookie values and login-like headers are masked."""
    secrets = ctx.get("secrets") or ()
    out = []
    for r in (getattr(site, "by_url", None) or {}).values():
        for k, v in (getattr(r, "response_headers", None) or {}).items():
            out.append([r.url, _code(r), _status(r), str(k), safe_header_value(k, v, secrets)])
    return [Table("All HTTP Headers", HEADER_EXPORT_HEADERS, out, url_cols=(1,), widths=(60, 11, 20, 30, 80))]


def all_cookies(site, index, crawler, ctx: dict) -> List[Table]:
    out = []
    for r, c in _cookies(site):
        out.append([r.url, c.get("name") or "", c.get("domain") or "", c.get("path") or "", c.get("expires") or "",
                    c.get("max_age") if c.get("max_age") not in (None, "") else "", bool(c.get("secure")),
                    bool(c.get("httponly")), c.get("samesite") or "", cookie_set_by(c)])
    return [Table("All Cookies", ["Address"] + DETAIL_COOKIE_HEADERS, out, url_cols=(1,))]


def all_issues(site, index, crawler, ctx: dict) -> List[Table]:
    return [Table("All Issues", ISSUE_URL_HEADERS,
                  issue_url_rows(index, ctx.get("issues"), ctx.get("dismissed_urls")), url_cols=(4,),
                  widths=(48, 13, 14, 70, 11, 18, 40))]


def page_source_files(site, crawler, cap: int = PAGE_SOURCE_CAP) -> Tuple[List[Tuple[str, bytes]], str]:
    """([(file name, html bytes)], why none) for every crawled HTML page whose source was kept."""
    store = getattr(crawler, "source_store", None) if crawler is not None else None
    if store is None:
        return [], ("Page sources were not kept for this crawl (Configuration > Spider: store page source; not "
                    "available on serverless hosts).")
    files: List[Tuple[str, bytes]] = []
    index_rows = []
    for p in getattr(site, "pages", None) or ():
        if len(files) >= cap:
            break
        try:
            html = store.get(p.url)
        except Exception:
            html = None
        if html is None:
            continue
        name = f"{len(files) + 1:05d}_{safe_name(p.url.split('://', 1)[-1], 90)}.html"
        files.append((name, html.encode("utf-8", "replace")))
        index_rows.append([name, p.url])
    if not files:
        return [], "No page source was kept for this crawl."
    files.insert(0, ("index.csv", csv_bytes(["File", "URL"], index_rows)))
    return files, ""


def _page_source_zip(site, crawler) -> Tuple[List[Tuple[str, bytes]], str]:
    return page_source_files(site, crawler)


_BULK_BUILDERS: Dict[str, Callable[[Any, Any, Any, dict], List[Table]]] = {
    "queued_urls": queued_urls,
    "links_all_inlinks": all_inlinks,
    "links_all_outlinks": all_outlinks,
    "links_all_anchor_text": all_anchor_text,
    "links_external": external_links,
    "web_http_headers": all_http_headers,
    "web_cookies": all_cookies,
    "issues_all": all_issues,
    "issues_overview": issues_overview,
}
for _id in PATH_EXPORTS:
    _BULK_BUILDERS[_id] = path_type_export(_id)
# Bulk exports that are a .zip of files (never a sheet): id -> (site, crawler) -> ([(file name, bytes)], why none)
_BULK_FILES: Dict[str, Callable[[Any, Any], Tuple[List[Tuple[str, bytes]], str]]] = {
    "web_page_source": _page_source_zip,
}


def register_bulk_export(export_id: str, title: str, headers: Optional[List[str]] = None,
                         rows: Optional[Callable[[Any, Any], Any]] = None,
                         files: Optional[Callable[[Any, Any], Tuple[List[Tuple[str, bytes]], str]]] = None,
                         not_checked: Optional[Callable[[Any], str]] = None, group: str = "Bulk export",
                         description: str = "") -> None:
    """Adds (or replaces) a Bulk Export item built by another module (PLAN4_build.md 10.3).
    rows(site, crawler) -> iterable of table rows under `headers`; or files(site, crawler) -> ([(name, bytes)],
    why none), always delivered as a .zip. not_checked(site) -> "" or the reason the data is missing (the export
    is then a one-row "Not checked" table / file, never an empty one). group: the Bulk Export sub-menu ("Bulk
    export" = the top level)."""
    if (rows is None) == (files is None):
        raise ValueError("register_bulk_export needs exactly one of rows= or files=")
    try:                       # ARACHNIX 1.0: 2.0-only exports stay hidden (crawler/seo/registry.py V2_FEATURES)
        from .seo.registry import v2_hidden
        if v2_hidden("bulk_exports", export_id):
            return
    except ImportError:
        pass
    path = [title] if not group or group == "Bulk export" else [group, title]
    item = _r(export_id, title, *path, description=description)
    if files is not None:
        item["files"] = True
    for i, b in enumerate(BULK_EXPORTS):
        if b["id"] == export_id:
            BULK_EXPORTS[i] = item
            break
    else:
        BULK_EXPORTS.append(item)
    _BULK_BUILDERS.pop(export_id, None)
    _BULK_FILES.pop(export_id, None)

    def reason(site) -> str:
        if not_checked is None:
            return ""
        try:
            return not_checked(site) or ""
        except Exception as e:
            return f"The data for this export could not be read ({type(e).__name__})."

    if files is not None:
        def build_files(site, crawler):
            why = reason(site)
            if why:
                return [], why
            return files(site, crawler)
        _BULK_FILES[export_id] = build_files
        return

    def build(site, index, crawler, ctx):
        why = reason(site)
        if why:
            return [not_checked_table(title, why)]
        return [Table(title, list(headers or []), [list(r) for r in rows(site, crawler) or ()])]
    _BULK_BUILDERS[export_id] = build


def is_file_export(export_id: str) -> bool:
    """True for a Bulk Export item delivered as a .zip of files (All Page Source, screenshots...)."""
    return export_id in _BULK_FILES


def bulk_files(export_id: str, site, crawler) -> Tuple[List[Tuple[str, bytes]], str]:
    """([(file name, bytes)], why none) of a files Bulk Export item (KeyError for another id)."""
    fn = _BULK_FILES[export_id]
    if site is None:
        return [], "There is no finished crawl to export yet."
    try:
        return fn(site, crawler)
    except Exception as e:
        return [], f"The files could not be collected ({type(e).__name__})."


def _tab_filter(export_id: str) -> Tuple[str, str, str]:
    kind, _, rest = export_id.partition(":")
    tab_id, _, filter_id = rest.partition(":")
    return kind, tab_id, filter_id or "all"


def bulk_tables(export_id: str, site, index, crawler=None, **ctx) -> List[Table]:
    """The tables of one Bulk Export item (KeyError for an unknown id or a tab / filter not in this crawl)."""
    if export_id in _BULK_FILES:
        raise KeyError(f"{export_id} is a folder of files: use build_bulk")
    if export_id.startswith(("filter:", "inlinks:", "issue:")):
        kind, tab_id, filter_id = _tab_filter(export_id)
        if kind in ("filter", "issue"):
            t = tab_table(site, index, tab_id, filter_id) if index is not None else None
            if t is None:
                raise KeyError(f"Unknown tab or filter: {tab_id}/{filter_id}")
            return [t]
        tr = _tr(index, tab_id)
        if tr is None or (filter_id != "all" and filter_id not in {f.id for f in tr.tab.filters}):
            raise KeyError(f"Unknown tab or filter: {tab_id}/{filter_id}")
        title = f"{tr.tab.name} - {filter_name(tr.tab, filter_id)} Inlinks"
        why = (getattr(tr, "not_checked", None) or {}).get(filter_id, "")
        if why:
            return [not_checked_table(title, why)]
        return [Table(title, SOURCE_HEADERS, filter_source_rows(site, index, tab_id, filter_id) or [],
                      url_cols=(1, 2), widths=(60, 60, 14, 36, 12, 10, 20, 12, 12, 24))]
    if export_id not in _BULK_BUILDERS:
        raise KeyError(f"Unknown bulk export: {export_id}")
    if export_id != "queued_urls" and (site is None or index is None):
        name = next((b["name"] for b in BULK_EXPORTS if b["id"] == export_id), export_id)
        return [not_checked_table(name, "There is no finished crawl to export yet.")]
    return _BULK_BUILDERS[export_id](site, index, crawler, ctx)


def build_bulk(export_id: str, site, index, crawler=None, fmt: str = "xlsx", **ctx) -> Tuple[str, bytes, str]:
    """(file name, bytes, media type) of one Bulk Export item. A files item ("web_page_source" and the ones
    registered with files=) is always a .zip (with a not_checked.csv when there is nothing to put in it); every
    other item follows fmt."""
    check_format(fmt)
    if export_id in _BULK_FILES:
        files, why = bulk_files(export_id, site, crawler)
        if not files:
            files = [("not_checked.csv", csv_bytes(["Status", "Reason"], [["Not checked", why]]))]
        stem = "page_source" if export_id == "web_page_source" else export_id
        return f"{_file_stem(site, stem)}.zip", zip_bytes(files), ZIP_MIME
    tables = bulk_tables(export_id, site, index, crawler, **ctx)
    return tables_bytes(tables, fmt, _file_stem(site, export_id.replace(":", "_")))


# ---------------------------------------------------------------------------------------------
# Multi-Export
# ---------------------------------------------------------------------------------------------
MULTI_INDEX_HEADERS = ["Export", "Menu", "Rows", "Result", "Sheet / File"]


def _item_info(item_id: str, index) -> dict:
    info = report_info(item_id)
    if info:
        return dict(info, kind="report")
    b = next((b for b in BULK_EXPORTS if b["id"] == item_id), None)
    if b:
        return dict(b, kind="bulk")
    if item_id.startswith(("filter:", "inlinks:", "issue:")):
        kind, tab_id, filter_id = _tab_filter(item_id)
        tr = _tr(index, tab_id)
        tname = tr.tab.name if tr is not None else tab_id
        fname = filter_name(tr.tab, filter_id) if tr is not None else filter_id
        if kind == "issue":
            return {"id": item_id, "name": f"{tname}: {fname}", "menu_path": ["Issues", f"{tname}: {fname}"],
                    "kind": kind}
        name = fname + (" Inlinks" if kind == "inlinks" else "")
        return {"id": item_id, "name": f"{tname} - {name}", "menu_path": [tname, name], "kind": kind}
    return {"id": item_id, "name": item_id, "menu_path": [item_id], "kind": "unknown"}


class _Names:
    """Unique sheet titles (Excel: 31 characters, no []:*?/\\) or zip paths, decided before writing so the
    index can say where each export went."""

    def __init__(self):
        self.sheetnames: List[str] = []

    def sheet(self, title: str) -> str:
        from .exporter import _sheet_title
        name = _sheet_title(self, title)
        self.sheetnames.append(name)
        return name

    def path(self, folder: str, stem: str) -> str:
        base, n = f"{folder}/{stem}", 2
        name = f"{base}.csv"
        while name.lower() in {s.lower() for s in self.sheetnames}:
            name, n = f"{base} ({n}).csv", n + 1
        self.sheetnames.append(name)
        return name


def multi_export(items: Sequence[str], site, index, crawler=None, fmt: str = "xlsx", skip_empty: bool = True,
                 **ctx) -> Tuple[str, bytes, str]:
    """Many reports and bulk exports in one download. xlsx: one workbook with an "Index" sheet and one sheet
    per table; csv: a .zip with index.csv and one CSV per table in a folder per menu ("Reports/...",
    "Response Codes/..."). skip_empty leaves out tables with no rows and checks that did not run; the index
    lists every requested item with its rows, what happened and the sheet / file it went to, so nothing
    disappears silently."""
    check_format(fmt)
    seen = set()
    names = _Names()
    if fmt == "xlsx":
        names.sheet("Index")
    index_rows: List[list] = []
    sheets: List[Tuple[str, Table]] = []
    extra_files: List[Tuple[str, bytes]] = []
    for item_id in items or ():
        if item_id in seen:
            continue
        seen.add(item_id)
        info = _item_info(item_id, index)
        menu = " > ".join((["Reports"] if info["kind"] == "report" else ["Bulk Export"]) + info["menu_path"])
        if item_id in _BULK_FILES:
            files, why = bulk_files(item_id, site, crawler)
            n_files = len([n for n, _ in files if n != "index.csv"])
            folder = "/".join(str(part).replace("/", "-")[:80] for part in info["menu_path"]) + "/"
            if not files:
                index_rows.append([info["name"], menu, 0, f"Not checked: {why}", ""])
            elif fmt == "xlsx":
                index_rows.append([info["name"], menu, n_files,
                                   "Not in the workbook: files cannot go in a sheet (download it on its own)", ""])
            else:
                extra_files += [(f"{folder}{n}", data) for n, data in files]
                index_rows.append([info["name"], menu, n_files, "Exported", folder])
            continue
        try:
            tables = (report_tables(item_id, site, index, crawler, **ctx) if info["kind"] == "report"
                      else bulk_tables(item_id, site, index, crawler, **ctx))
        except KeyError as e:
            index_rows.append([info["name"], menu, 0, f"Not exported: {e.args[0] if e.args else e}", ""])
            continue
        for t in tables:
            label = info["name"] if len(tables) == 1 else f"{info['name']} - {t.title}"
            rows = 0 if t.not_checked else len(t.rows)
            if t.not_checked:
                result = f"Not checked: {t.not_checked}"
            elif not t.rows:
                result = "Empty"
            else:
                result = "Exported"
            if skip_empty and not t.has_data:
                index_rows.append([label, menu, rows, result + " (skipped)", ""])
                continue
            if fmt == "xlsx":
                where = names.sheet(label)
            else:
                folder = "Reports" if info["kind"] == "report" else safe_name(info["menu_path"][0], 60)
                where = names.path(folder, safe_name(label, 120))
            index_rows.append([label, menu, rows, result, where])
            sheets.append((where, t))
    index_table = Table("Index", MULTI_INDEX_HEADERS, index_rows, widths=(52, 60, 10, 70, 40))
    stem = _file_stem(site, "multi_export")
    if fmt == "xlsx":
        from openpyxl import Workbook
        import io
        from .exporter import _discard
        wb = Workbook(write_only=True)
        try:
            write_table(wb, index_table)
            for where, t in sheets:     # plain-text URLs: hyperlinked cells double openpyxl's time on big exports
                write_table(wb, Table(where, t.headers, t.rows, widths=t.widths, not_checked=t.not_checked))
            buf = io.BytesIO()
            wb.save(buf)
            return f"{stem}.xlsx", buf.getvalue(), XLSX_MIME
        except BaseException:
            _discard(wb)
            raise
    files = [("index.csv", csv_bytes(index_table.headers, index_table.rows))]
    files += [(where, csv_bytes(t.headers, t.rows)) for where, t in sheets]
    return f"{stem}.zip", zip_bytes(files + extra_files), ZIP_MIME


def catalog(index=None) -> dict:
    """Everything the Reports, Bulk Export and Multi-Export menus can offer: {"reports", "bulk"}."""
    return {"reports": [dict(r) for r in REPORTS], "bulk": bulk_catalog(index)}


# P20: Crawl Path, Small Tap Targets, PageSpeed Opportunities, Screenshots and All Page Text (crawler/reports_extra.py)
try:
    from . import reports_extra as _extra
    _extra.install(__import__(__name__, fromlist=["REPORTS"]))
except Exception as _e:  # pragma: no cover - a broken extra must never hide the other reports
    import sys as _sys
    print(f"[reports] the extra reports failed to load: {type(_e).__name__}: {_e}", file=_sys.stderr)
