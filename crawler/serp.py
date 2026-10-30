"""
SERP mode (Mode > SERP) and the Reports > SERP Summary report: the length in characters and the width in
pixels of page titles and meta descriptions, measured with the same font metrics
(seo/text_metrics.pixel_width) and the same thresholds as the Page Titles and Meta Description tabs. Nothing
is crawled, so the pages do not need to be live (check proposed titles before a release, or re-upload an
edited SERP Summary).

    rows = measure([{"url", "title", "description"}, ...])   each row + lengths, pixel widths and flags
    counts = summary(rows)                                   filter counts per tab, like the Overview tree
    rows, rejected = parse_upload(data, filename)            .csv / .txt / .xlsx with URL, Title, Description
    rows = crawl_rows(site)                                  the SERP Summary of a crawl (Page Titles "All")

The flag ids are the Page Titles / Meta Description filter ids ("over_60_characters", ...), so a SERP row
lands in the same filters a crawled page would. As in the tabs, the length and pixel filters look at
non-empty values only, and "Duplicate" means the same text on more than one row.
"""
import csv
import io
import os
from collections import Counter
from typing import Any, Dict, List, Optional, Sequence, Tuple

from .seo.tabs.meta_description import (DESCRIPTION_MAX_CHARS, DESCRIPTION_MAX_PX, DESCRIPTION_MIN_CHARS,
                                        DESCRIPTION_MIN_PX)
from .seo.tabs.page_titles import TITLE_MAX_CHARS, TITLE_MAX_PX, TITLE_MIN_CHARS, TITLE_MIN_PX
from .seo.text_metrics import pixel_width

MAX_ROWS = 50_000                 # rows read from one upload
MAX_TEXT = 2_000                  # characters kept per title / description (nothing real is longer)

SUMMARY_HEADERS = ["URL", "Title", "Title Length", "Title Pixel Width", "Description", "Description Length",
                   "Description Pixel Width"]

# (filter id, label, test(length, pixels)) per field; ids match the Page Titles / Meta Description tabs
TITLE_CHECKS: List[Tuple[str, str, Any]] = [
    ("over_60_characters", f"Over {TITLE_MAX_CHARS} Characters", lambda n, px: n > TITLE_MAX_CHARS),
    ("below_30_characters", f"Below {TITLE_MIN_CHARS} Characters", lambda n, px: n < TITLE_MIN_CHARS),
    ("over_561_pixels", f"Over {TITLE_MAX_PX} Pixels", lambda n, px: px > TITLE_MAX_PX),
    ("below_200_pixels", f"Below {TITLE_MIN_PX} Pixels", lambda n, px: px < TITLE_MIN_PX),
]
DESCRIPTION_CHECKS: List[Tuple[str, str, Any]] = [
    ("over_155_characters", f"Over {DESCRIPTION_MAX_CHARS} Characters", lambda n, px: n > DESCRIPTION_MAX_CHARS),
    ("below_70_characters", f"Below {DESCRIPTION_MIN_CHARS} Characters", lambda n, px: n < DESCRIPTION_MIN_CHARS),
    ("over_985_pixels", f"Over {DESCRIPTION_MAX_PX} Pixels", lambda n, px: px > DESCRIPTION_MAX_PX),
    ("below_400_pixels", f"Below {DESCRIPTION_MIN_PX} Pixels", lambda n, px: px < DESCRIPTION_MIN_PX),
]
FIELDS = (("title", "page_titles", "title", TITLE_CHECKS),
          ("description", "meta_description", "description", DESCRIPTION_CHECKS))

# Upload header names (lowercased) -> field; Screaming Frog's own exports use "Address" and "Title 1"
_HEADER_ALIASES = {
    "url": "url", "address": "url", "page": "url", "page url": "url",
    "title": "title", "title 1": "title", "page title": "title", "seo title": "title", "meta title": "title",
    "description": "description", "meta description": "description", "meta description 1": "description",
    "description 1": "description", "seo description": "description",
}


def clean(text: Any) -> str:
    """Whitespace collapsed the way the crawler stores titles and descriptions."""
    return " ".join(str(text or "").split())[:MAX_TEXT]


def filter_names() -> Dict[str, Dict[str, str]]:
    """{"page_titles": {id: label}, "meta_description": {id: label}} including Missing and Duplicate."""
    out = {}
    for _, tab, _, checks in FIELDS:
        names = {"missing": "Missing", "duplicate": "Duplicate"}
        names.update({fid: label for fid, label, _ in checks})
        out[tab] = names
    return out


def measure(rows: Sequence[dict]) -> List[dict]:
    """Each row ({"url", "title", "description"}; other keys are kept) with title_length,
    title_pixel_width, description_length, description_pixel_width and title_flags / description_flags
    (the Page Titles / Meta Description filter ids the row falls under)."""
    base = []
    for r in rows or ():
        r = dict(r or {})
        r["url"] = str(r.get("url") or "").strip()
        r["title"] = clean(r.get("title"))
        r["description"] = clean(r.get("description"))
        base.append(r)
    seen = {key: Counter(r[key] for r in base if r[key]) for key, _, _, _ in FIELDS}
    out = []
    for r in base:
        for key, _, kind, checks in FIELDS:
            text = r[key]
            n = len(text)
            px = pixel_width(text, kind) if text else 0
            r[f"{key}_length"] = n
            r[f"{key}_pixel_width"] = px
            flags = []
            if not text:
                flags.append("missing")
            else:
                if seen[key][text] > 1:
                    flags.append("duplicate")
                flags += [fid for fid, _, test in checks if test(n, px)]
            r[f"{key}_flags"] = flags
        out.append(r)
    return out


def summary(rows: Sequence[dict]) -> Dict[str, Dict[str, int]]:
    """Filter counts of measured rows: {"page_titles": {"all": n, "missing": n, ...}, "meta_description": ...}
    (every filter listed, 0 included: SERP mode checks every row, so 0 means none)."""
    out: Dict[str, Dict[str, int]] = {}
    names = filter_names()
    for key, tab, _, _ in FIELDS:
        counts = {"all": len(rows)}
        counts.update({fid: 0 for fid in names[tab]})
        for r in rows:
            for fid in r.get(f"{key}_flags") or ():
                counts[fid] = counts.get(fid, 0) + 1
        out[tab] = counts
    return out


def summary_rows(rows: Sequence[dict]) -> List[list]:
    """SUMMARY_HEADERS rows of measured rows (the SERP Summary layout, re-uploadable in SERP mode)."""
    return [[r.get("url", ""), r.get("title", ""), r.get("title_length", 0), r.get("title_pixel_width", 0),
             r.get("description", ""), r.get("description_length", 0), r.get("description_pixel_width", 0)]
            for r in rows]


def crawl_rows(site) -> List[dict]:
    """The crawl's SERP data, measured: one row per Page Titles "All" row (internal HTML pages that loaded),
    with the first title and the first meta description, in crawl order."""
    from .seo.tabs.page_titles import content_rows, text_list
    rows = []
    for r in content_rows(site) if site is not None else ():
        titles = text_list(r, "titles")
        descs = text_list(r, "meta_descriptions")
        rows.append({"url": r.url, "title": titles[0] if titles else "",
                     "description": descs[0] if descs else ""})
    return measure(rows)


# ---------------------------------------------------------------------------------------------
# Upload (.csv / .txt / .xlsx)
# ---------------------------------------------------------------------------------------------
def _header_map(cells: Sequence[Any]) -> Optional[Dict[str, int]]:
    """{field: column} when the row is a header naming at least one known column, else None."""
    found: Dict[str, int] = {}
    for i, c in enumerate(cells):
        key = _HEADER_ALIASES.get(" ".join(str(c or "").split()).strip().lower())
        if key and key not in found:
            found[key] = i
    return found or None


def _table_rows(data: bytes, filename: str) -> List[List[Any]]:
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            out = []
            for row in ws.iter_rows(values_only=True):
                out.append(list(row))
                if len(out) > MAX_ROWS + 1:
                    break
            return out
        finally:
            wb.close()
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t") if sample.strip() else csv.excel
    except csv.Error:
        dialect = csv.excel_tab if ("\t" in sample and "," not in sample) else csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)][:MAX_ROWS + 1]


def parse_upload(data: bytes, filename: str = "") -> Tuple[List[dict], List[dict]]:
    """Rows ({"url", "title", "description"}) of an uploaded file and the lines left out
    ([{"line", "reason"}]). The header row may name the columns in any order (URL / Address, Title /
    Title 1, Description / Meta Description 1); without a header the columns are URL, Title, Description."""
    table = _table_rows(data or b"", filename)
    rows: List[dict] = []
    rejected: List[dict] = []
    if not table:
        return rows, rejected
    cols = _header_map(table[0])
    start = 1 if cols else 0
    cols = cols or {"url": 0, "title": 1, "description": 2}
    for n, cells in enumerate(table[start:], start=start + 1):
        cells = list(cells or [])

        def cell(key):
            i = cols.get(key)
            return "" if i is None or i >= len(cells) or cells[i] is None else str(cells[i])
        row = {"url": cell("url").strip(), "title": clean(cell("title")), "description": clean(cell("description"))}
        if not (row["url"] or row["title"] or row["description"]):
            if any(str(c or "").strip() for c in cells):
                rejected.append({"line": n, "reason": "No URL, title or description in this row."})
            continue
        if len(rows) >= MAX_ROWS:
            rejected.append({"line": n, "reason": f"Over {MAX_ROWS:,} rows: the rest were not read."})
            break
        rows.append(row)
    return rows, rejected
